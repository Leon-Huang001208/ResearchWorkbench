"""Wind index structure probe workbook builder and reader."""

from __future__ import annotations

from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
import time
from typing import Any, Iterable
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook

from core.observability import get_logger
from services.wind_realtime_workbook import _rows_from_matrix

logger = get_logger(__name__)
SHANGHAI_TZ = ZoneInfo("Asia/Shanghai")

DEFAULT_PROBE_WORKBOOK_PATH = (
    Path.home()
    / "Library"
    / "Application Support"
    / "AlphaFoundry"
    / "wind"
    / "AlphaFoundry_Wind_Index_Structure_Probe.xlsx"
)

PROBE_SHEETS = [
    "README",
    "Config",
    "FormulaCatalog",
    "ProbeTargets",
    "ProbeResults",
    "Health",
    "FormulaLog",
]

FORMULA_CATALOG_HEADERS = [
    "formula_key",
    "domain",
    "label",
    "template",
    "expected_shape",
    "notes",
]

PROBE_TARGET_HEADERS = ["domain", "target_code", "trade_date", "notes"]

PROBE_RESULT_HEADERS = [
    "probe_id",
    "domain",
    "target_code",
    "formula_key",
    "label",
    "formula_text",
    "value",
    "status",
    "updated_at",
    "notes",
]

HEALTH_HEADERS = ["metric", "value", "updated_at", "notes"]
LOG_HEADERS = ["timestamp", "level", "component", "message", "details"]
ISO_NOW_FORMULA = '=TEXT(NOW(),"yyyy-mm-ddThh:mm:ss")'
WIND_FAILURE_TEXTS = frozenset(
    {
        "无法读取数据！",
        "无法读取数据!",
        "no data",
        "n/a",
        "#n/a",
        "#value!",
        "#name?",
    }
)


@dataclass(frozen=True)
class ProbeFormulaDefinition:
    formula_key: str
    domain: str
    label: str
    template: str
    expected_shape: str
    notes: str = ""


@dataclass(frozen=True)
class ProbeResultRow:
    probe_id: str
    domain: str
    target_code: str
    formula_key: str
    label: str
    formula_text: str
    value: Any
    updated_at: datetime | None
    notes: str = ""


@dataclass(frozen=True)
class ProbeSnapshot:
    rows: tuple[ProbeResultRow, ...]
    status: str
    updated_at: datetime | None
    error_count: int = 0
    message: str = ""


PROBE_FORMULAS: tuple[ProbeFormulaDefinition, ...] = (
    ProbeFormulaDefinition(
        formula_key="index_name_wss",
        domain="index",
        label="指数名称 WSS",
        template='=@wss("{code}","sec_name")',
        expected_shape="scalar",
        notes="确认指数代码是否能被 Wind 识别。",
    ),
    ProbeFormulaDefinition(
        formula_key="index_close_wss",
        domain="index",
        label="指数收盘价 WSS",
        template='=@wss("{code}","close","tradeDate={trade_date}")',
        expected_shape="scalar",
        notes="验证指数行情日期参数。",
    ),
    ProbeFormulaDefinition(
        formula_key="index_constituent_wset",
        domain="index",
        label="指数成分 WSET",
        template=(
            '=@wset("indexconstituent","date={trade_date};windcode={code};'
            'field=wind_code,sec_name,i_weight")'
        ),
        expected_shape="table",
        notes="候选全成分公式；以实际 Wind Excel 返回为准。",
    ),
    ProbeFormulaDefinition(
        formula_key="index_constituent_wset_plain",
        domain="index",
        label="指数成分 WSET 无@",
        template=(
            '=wset("indexconstituent","date={trade_date};windcode={code};'
            'field=wind_code,sec_name,i_weight")'
        ),
        expected_shape="table",
        notes="固定模板中测试无 @ 的 WSET 写法。",
    ),
    ProbeFormulaDefinition(
        formula_key="index_weight_wss",
        domain="index",
        label="成分权重点查 WSS",
        template='=@s_info_indexweight("600519.SH","{trade_date}","{code}")',
        expected_shape="scalar",
        notes="用 600519.SH 做样本点查，验证权重函数可用性。",
    ),
    ProbeFormulaDefinition(
        formula_key="index_weight_plain",
        domain="index",
        label="成分权重点查无@",
        template='=s_info_indexweight("600519.SH","{trade_date}","{code}")',
        expected_shape="scalar",
        notes="固定模板中测试无 @ 的指数权重函数写法。",
    ),
    ProbeFormulaDefinition(
        formula_key="etf_name_wss",
        domain="etf",
        label="ETF 名称 WSS",
        template='=@wss("{code}","sec_name")',
        expected_shape="scalar",
        notes="确认 ETF 代码是否能被 Wind 识别。",
    ),
    ProbeFormulaDefinition(
        formula_key="etf_tracking_index_wss",
        domain="etf",
        label="ETF 跟踪指数 WSS",
        template='=@wss("{code}","fund_trackindexcode")',
        expected_shape="scalar",
        notes="候选跟踪指数代码字段；需以探针结果确认。",
    ),
    ProbeFormulaDefinition(
        formula_key="etf_nav_wss",
        domain="etf",
        label="ETF 单位净值 WSS",
        template='=@wss("{code}","nav","tradeDate={trade_date}")',
        expected_shape="scalar",
        notes="ETF 日度净值候选字段。",
    ),
    ProbeFormulaDefinition(
        formula_key="etf_shares_wss",
        domain="etf",
        label="ETF 份额 WSS",
        template='=@wss("{code}","fund_share_total","tradeDate={trade_date}")',
        expected_shape="scalar",
        notes="用于推导申赎净流入的份额字段候选。",
    ),
    ProbeFormulaDefinition(
        formula_key="etf_shares_unit_total_wss",
        domain="etf",
        label="ETF 总份额 WSS",
        template='=@wss("{code}","unit_total","tradeDate={trade_date}")',
        expected_shape="scalar",
        notes="ETF 份额候选字段。",
    ),
    ProbeFormulaDefinition(
        formula_key="etf_shares_fund_share_wss",
        domain="etf",
        label="ETF 基金份额 WSS",
        template='=@wss("{code}","fund_share","tradeDate={trade_date}")',
        expected_shape="scalar",
        notes="ETF 份额候选字段。",
    ),
    ProbeFormulaDefinition(
        formula_key="etf_aum_wss",
        domain="etf",
        label="ETF 规模 WSS",
        template='=@wss("{code}","fund_assetnetvalue","tradeDate={trade_date}")',
        expected_shape="scalar",
        notes="ETF AUM 候选字段；需与 Wind 返回口径确认。",
    ),
    ProbeFormulaDefinition(
        formula_key="etf_aum_netasset_total_wss",
        domain="etf",
        label="ETF 净资产 WSS",
        template='=@wss("{code}","netasset_total","tradeDate={trade_date}")',
        expected_shape="scalar",
        notes="ETF AUM 候选字段。",
    ),
    ProbeFormulaDefinition(
        formula_key="etf_aum_fund_scale_wss",
        domain="etf",
        label="ETF 基金规模 WSS",
        template='=@wss("{code}","fund_fundscale","tradeDate={trade_date}")',
        expected_shape="scalar",
        notes="ETF AUM 候选字段。",
    ),
)


def build_index_structure_probe_workbook(
    *,
    workbook_path: str | Path | None = None,
    trade_date: str | None = None,
    index_codes: Iterable[str] | None = None,
    etf_codes: Iterable[str] | None = None,
) -> Path:
    """Build a fixed Excel probe workbook for Wind index/ETF field discovery."""
    output_path = Path(workbook_path or DEFAULT_PROBE_WORKBOOK_PATH).expanduser()
    output_path.parent.mkdir(parents=True, exist_ok=True)
    trade_date_value = trade_date or datetime.now(SHANGHAI_TZ).strftime("%Y-%m-%d")
    generated_at = datetime.now(UTC).isoformat()
    index_targets = tuple(_dedupe_codes(index_codes or ("000300.SH", "399006.SZ", "HSI.HI")))
    etf_targets = tuple(_dedupe_codes(etf_codes or ("510300.SH", "159919.SZ", "2800.HK")))

    try:
        workbook = Workbook()
        workbook.remove(workbook.active)
        sheets = {name: workbook.create_sheet(name) for name in PROBE_SHEETS}

        sheets["README"]["A1"] = "AlphaFoundry Wind Index Structure Probe"
        sheets["README"]["A2"] = (
            "Open this workbook with Wind Excel logged in, or run the prime script in hidden mode."
        )

        sheets["Config"].append(["key", "value", "description"])
        sheets["Config"].append(["trade_date", trade_date_value, "探针交易日"])
        sheets["Config"].append(["formula_version", "1", "指数结构探针公式版本"])
        sheets["Config"].append(["last_generated_at", generated_at, "工作簿生成时间"])
        sheets["Config"].append(["data_owner", "AlphaFoundry", "数据维护方"])

        sheets["FormulaCatalog"].append(FORMULA_CATALOG_HEADERS)
        for definition in PROBE_FORMULAS:
            sheets["FormulaCatalog"].append(
                [
                    definition.formula_key,
                    definition.domain,
                    definition.label,
                    definition.template,
                    definition.expected_shape,
                    definition.notes,
                ]
            )

        sheets["ProbeTargets"].append(PROBE_TARGET_HEADERS)
        for code in index_targets:
            sheets["ProbeTargets"].append(["index", code, trade_date_value, "指数样本"])
        for code in etf_targets:
            sheets["ProbeTargets"].append(["etf", code, trade_date_value, "ETF样本"])

        sheets["ProbeResults"].append(PROBE_RESULT_HEADERS)
        result_row_number = 1
        for domain, code in [("index", code) for code in index_targets] + [
            ("etf", code) for code in etf_targets
        ]:
            for definition in PROBE_FORMULAS:
                if definition.domain != domain:
                    continue
                result_row_number += 1
                formula_text = definition.template.format(code=code, trade_date=trade_date_value)
                sheets["ProbeResults"].append(
                    [
                        f"{domain}:{code}:{definition.formula_key}",
                        domain,
                        code,
                        definition.formula_key,
                        definition.label,
                        formula_text,
                        formula_text,
                        _status_formula(result_row_number),
                        ISO_NOW_FORMULA,
                        definition.notes,
                    ]
                )

        sheets["Health"].append(HEALTH_HEADERS)
        sheets["Health"].append(["workbook_open", "true", generated_at, "文件已生成"])
        sheets["Health"].append(
            [
                "probe_formula_count",
                len(PROBE_FORMULAS),
                generated_at,
                "候选公式数量，需以 Wind 返回结果确认",
            ]
        )
        sheets["Health"].append(
            [
                "probe_target_count",
                len(index_targets) + len(etf_targets),
                generated_at,
                "探针目标数量",
            ]
        )
        sheets["FormulaLog"].append(LOG_HEADERS)

        for sheet in sheets.values():
            sheet.freeze_panes = "A2"

        workbook.save(output_path)
    except Exception as exc:
        logger.error("Failed to build Wind index structure probe workbook: %s", exc)
        raise RuntimeError(f"Failed to build Wind index structure probe workbook: {exc}") from exc

    logger.info("Built Wind index structure probe workbook: %s", output_path)
    return output_path


def prime_index_structure_probe_workbook(
    workbook_path: str | Path | None = None,
    *,
    visible: bool = False,
    reopen: bool = True,
    wait_seconds: float = 10.0,
    save: bool = True,
) -> Path:
    """Open the probe workbook in Excel hidden mode, calculate Wind formulas, and save."""
    path = Path(workbook_path or DEFAULT_PROBE_WORKBOOK_PATH).expanduser().resolve()
    if not path.exists():
        raise FileNotFoundError(f"Wind index structure probe workbook does not exist: {path}")

    try:
        import xlwings as xw
    except ImportError as exc:
        raise RuntimeError("xlwings is required to prime Wind index structure probe") from exc

    book = _open_probe_workbook(xw, path, reopen=reopen)
    try:
        book.app.visible = visible
    except Exception as exc:
        logger.debug("Unable to set Excel visibility to %s: %s", visible, exc)

    try:
        book.app.calculate()
    except Exception as exc:
        logger.warning("Unable to trigger Excel calculation for probe workbook: %s", exc)
    if wait_seconds > 0:
        time.sleep(wait_seconds)
    if save:
        book.save()

    logger.info("Primed Wind index structure probe workbook: %s", path)
    return path


def _open_probe_workbook(xw: Any, path: Path, *, reopen: bool) -> Any:
    if reopen:
        for app in xw.apps:
            for candidate in list(app.books):
                fullname = str(getattr(candidate, "fullname", "") or "")
                if not fullname:
                    continue
                try:
                    matches = Path(fullname).expanduser().resolve() == path
                except OSError:
                    matches = Path(fullname).expanduser() == path
                if matches:
                    try:
                        candidate.close()
                        logger.info("Closed stale open Wind index probe workbook before priming")
                    except Exception as exc:
                        logger.debug("Unable to close stale probe workbook: %s", exc)

    app = xw.apps.active or xw.App(visible=False)
    return app.books.open(str(path), update_links=False, read_only=False)


def read_index_structure_probe_snapshot(
    workbook_path: str | Path | None = None,
) -> ProbeSnapshot:
    """Read saved probe values from the workbook cache."""
    path = Path(workbook_path or DEFAULT_PROBE_WORKBOOK_PATH).expanduser()
    if not path.exists():
        return ProbeSnapshot(
            rows=(),
            status="workbook_missing",
            updated_at=None,
            message=f"Wind指数结构探针工作簿不存在: {path}",
        )

    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
        rows = _rows_from_matrix(list(workbook["ProbeResults"].iter_rows(values_only=True)))
        return parse_probe_result_rows(rows)
    except Exception as exc:
        logger.error("Failed to read Wind index structure probe workbook %s: %s", path, exc)
        return ProbeSnapshot(
            rows=(),
            status="workbook_read_error",
            updated_at=None,
            error_count=1,
            message=f"读取Wind指数结构探针失败: {exc}",
        )


def parse_probe_result_rows(
    rows: Iterable[dict[str, Any]],
    *,
    now: datetime | None = None,
) -> ProbeSnapshot:
    """Parse cached probe result rows into valid values and error counts."""
    _reference_time = now or datetime.now(UTC)
    parsed: list[ProbeResultRow] = []
    error_count = 0
    for row in rows:
        status = str(row.get("status") or "").strip()
        value = row.get("value")
        if status != "ok" or _is_failed_probe_value(value):
            error_count += 1
            continue
        parsed.append(
            ProbeResultRow(
                probe_id=str(row.get("probe_id") or ""),
                domain=str(row.get("domain") or ""),
                target_code=str(row.get("target_code") or ""),
                formula_key=str(row.get("formula_key") or ""),
                label=str(row.get("label") or ""),
                formula_text=str(row.get("formula_text") or ""),
                value=value,
                updated_at=_parse_datetime(row.get("updated_at")),
                notes=str(row.get("notes") or ""),
            )
        )

    updated_at = max((row.updated_at for row in parsed if row.updated_at), default=None)
    if parsed:
        status = "ok"
        message = ""
    elif error_count:
        status = "all_failed"
        message = "Wind指数结构探针暂无成功字段，请检查Wind登录、公式字段或刷新状态"
    else:
        status = "empty"
        message = "Wind指数结构探针暂无结果"

    return ProbeSnapshot(
        rows=tuple(parsed),
        status=status,
        updated_at=updated_at,
        error_count=error_count,
        message=message,
    )


def _dedupe_codes(codes: Iterable[str]) -> list[str]:
    seen: set[str] = set()
    result: list[str] = []
    for raw_code in codes:
        code = str(raw_code or "").strip().upper()
        if not code or code in seen:
            continue
        seen.add(code)
        result.append(code)
    return result


def _status_formula(row_number: int) -> str:
    return (
        f'=IF(OR(ISERROR(G{row_number}),G{row_number}="",'
        f'G{row_number}="Fetching...",G{row_number}="fetching...",'
        f'G{row_number}="无法读取数据！",G{row_number}="无法读取数据!"),'
        f'"formula_error","ok")'
    )


def _is_failed_probe_value(value: Any) -> bool:
    if value in (None, ""):
        return True
    if not isinstance(value, str):
        return False
    normalized = value.strip().lower()
    return normalized in WIND_FAILURE_TEXTS


def _parse_datetime(value: Any) -> datetime | None:
    if value in (None, ""):
        return None
    if isinstance(value, datetime):
        return value if value.tzinfo else value.replace(tzinfo=SHANGHAI_TZ)
    text = str(value).strip()
    if not text:
        return None
    try:
        return datetime.fromisoformat(text)
    except ValueError:
        try:
            return datetime.strptime(text[:19], "%Y-%m-%dT%H:%M:%S").replace(tzinfo=SHANGHAI_TZ)
        except ValueError:
            return None
