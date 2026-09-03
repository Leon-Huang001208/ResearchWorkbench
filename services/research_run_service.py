"""Application service for durable, evidence-first research runs."""

from __future__ import annotations

import hashlib
import json
from collections.abc import Callable, Collection
from datetime import UTC, datetime
from io import BytesIO
from typing import TYPE_CHECKING, Any
from uuid import uuid4

from docx import Document

from core.contracts.research import (
    QualityGateResult,
    ResearchArtifact,
    ResearchClaim,
    ResearchDecisionCard,
    ResearchEvidenceInput,
    ResearchRun,
    ResearchRunCreateRequest,
    ResearchRunOutputs,
    ResearchRunStatus,
    ResearchTask,
)
from core.observability import get_logger
from data_layer.repositories.research_run_repository import ResearchRunRepository
from services.research_templates import (
    ResearchTemplateRegistry,
    get_default_research_template_registry,
)
from services.runtime_provider_service import (
    InvocationUsage,
    RuntimeBlockedError,
    RuntimeExecutionResult,
    RuntimeFailedError,
    RuntimeProviderService,
    stable_runtime_request_hash,
)

if TYPE_CHECKING:
    from core.contracts.research_workspace import (
        ProviderExecutionResult,
        RuntimeProvider,
        SkillManifest,
    )
    from services.agent_team_service import (
        AgentAssignment,
        AgentSupervisorContext,
        AgentTeamService,
    )
    from services.research_workspace_service import ResearchWorkspaceService

logger = get_logger(__name__)


class ResearchRunService:
    """Persist all research facts around a pure LangGraph execution plan."""

    def __init__(
        self,
        repository: ResearchRunRepository,
        registry: ResearchTemplateRegistry | None = None,
        workspace_service: ResearchWorkspaceService | None = None,
        runtime_service: RuntimeProviderService | None = None,
        agent_team_service: AgentTeamService | None = None,
        provider_invoker: Callable[[RuntimeProvider, dict[str, Any], str], Any] | None = None,
        agent_worker: Callable[[AgentAssignment], Any] | None = None,
        agent_supervisor: Callable[[AgentSupervisorContext], Any] | None = None,
        skill_invoker: (
            Callable[[SkillManifest, dict[str, Any], dict[str, Any]], dict[str, Any]] | None
        ) = None,
        authorized_tool_ids: Collection[str] = (),
    ):
        self._repository = repository
        self._registry = registry or get_default_research_template_registry()
        self._workspace_service = workspace_service
        self._runtime_service = runtime_service
        self._agent_team_service = agent_team_service
        self._provider_invoker = provider_invoker
        self._agent_worker = agent_worker
        self._agent_supervisor = agent_supervisor
        self._skill_invoker = skill_invoker
        self._authorized_tool_ids = frozenset(authorized_tool_ids)

    @property
    def authorized_tool_ids(self) -> frozenset[str]:
        """Expose the effective executable registry for readiness checks and tests."""

        return self._authorized_tool_ids

    def create(
        self,
        request: ResearchRunCreateRequest,
        *,
        idempotency_key: str | None = None,
    ) -> ResearchRun:
        if request.subject is None or request.target_id is None:
            raise ValueError("Research subject normalization failed")
        template = self._registry.require_executable(request.template_key, request.subject)
        self._registry.validate_evidence_kinds(
            request.template_key, [item.evidence_kind for item in request.evidence_inputs]
        )
        now = _utc_now()
        run_id = _new_id("research_run")
        if idempotency_key and self._workspace_service is not None:
            reserved_id = self._workspace_service.reserve_operation(
                operation="create",
                target_id="research_run",
                idempotency_key=idempotency_key,
                request_hash=_request_hash(request.model_dump(mode="json")),
                resource_id=run_id,
            )
            if reserved_id != run_id:
                return self._get_run(reserved_id)
        run = ResearchRun(
            run_id=run_id,
            template_key=request.template_key,
            target_id=request.target_id,
            subject=request.subject,
            as_of=request.as_of,
            question=request.question,
            attachment_refs=request.attachment_refs,
            source_plan={
                "_runtime": {
                    "mode": request.mode,
                    "agent_team_id": request.agent_team_id,
                    "skill_keys": request.skill_keys,
                }
            },
            status=ResearchRunStatus.DRAFT,
            created_at=now,
            updated_at=now,
        )
        created = self._repository.create_run(
            run, [evidence.model_dump(mode="json") for evidence in request.evidence_inputs]
        )
        self._repository.create_task(
            ResearchTask(
                task_id=_new_id("research_task"),
                run_id=created.run_id,
                task_key=template.task_key,
                status=ResearchRunStatus.DRAFT,
                created_at=now,
                updated_at=now,
            )
        )
        if self._workspace_service is not None:
            self._workspace_service.record_run_event(
                created.run_id,
                status=ResearchRunStatus.DRAFT.value,
                stage=ResearchRunStatus.DRAFT.value,
                idempotency_key=f"{created.run_id}:draft",
            )
        logger.info(
            "research run created",
            run_id=created.run_id,
            template_key=created.template_key,
            subject_type=created.subject.subject_type,
            target_id=created.target_id,
        )
        return created

    def execute(
        self,
        run_id: str,
        *,
        idempotency_key: str | None = None,
        project_id: str | None = None,
        workspace_id: str | None = None,
    ) -> ResearchRun:
        run = self._get_run(run_id)
        self._require_run_scope(run_id, project_id=project_id, workspace_id=workspace_id)
        if idempotency_key and self._workspace_service is not None:
            reserved_id = self._workspace_service.reserve_operation(
                operation="execute",
                target_id=run_id,
                idempotency_key=idempotency_key,
                request_hash=_request_hash({"operation": "execute", "run_id": run_id}),
                resource_id=run_id,
            )
            if reserved_id != run_id:
                raise ValueError("idempotency key belongs to another research run")
        if run.status is ResearchRunStatus.COMPLETED:
            return run
        if run.status not in {
            ResearchRunStatus.DRAFT,
            ResearchRunStatus.BLOCKED,
            ResearchRunStatus.FAILED,
        }:
            raise ValueError(f"Research run {run_id} cannot execute from {run.status.value}")
        if self._workspace_service is not None:
            self._workspace_service.claim_run_execution(
                run_id,
                allowed_statuses={
                    ResearchRunStatus.DRAFT.value,
                    ResearchRunStatus.BLOCKED.value,
                    ResearchRunStatus.FAILED.value,
                },
            )
        return self._execute(run, increment_retry=run.status is not ResearchRunStatus.DRAFT)

    def resume(
        self,
        run_id: str,
        *,
        idempotency_key: str | None = None,
        project_id: str | None = None,
        workspace_id: str | None = None,
    ) -> ResearchRun:
        run = self._get_run(run_id)
        self._require_run_scope(run_id, project_id=project_id, workspace_id=workspace_id)
        if idempotency_key and self._workspace_service is not None:
            reserved_id = self._workspace_service.reserve_operation(
                operation="resume",
                target_id=run_id,
                idempotency_key=idempotency_key,
                request_hash=_request_hash({"operation": "resume", "run_id": run_id}),
                resource_id=run_id,
            )
            if reserved_id != run_id:
                raise ValueError("idempotency key belongs to another research run")
        if run.status is ResearchRunStatus.COMPLETED and idempotency_key:
            return run
        if run.status is not ResearchRunStatus.BLOCKED:
            raise ValueError(f"Research run {run_id} is not blocked")
        if self._workspace_service is not None:
            self._workspace_service.claim_run_execution(
                run_id, allowed_statuses={ResearchRunStatus.BLOCKED.value}
            )
        return self._execute(run, increment_retry=True)

    def add_evidence(
        self,
        run_id: str,
        evidence: ResearchEvidenceInput,
        *,
        project_id: str | None = None,
        workspace_id: str | None = None,
    ) -> ResearchRun:
        run = self._get_run(run_id)
        self._require_run_scope(run_id, project_id=project_id, workspace_id=workspace_id)
        if run.status not in {ResearchRunStatus.DRAFT, ResearchRunStatus.BLOCKED}:
            raise ValueError("Evidence can only be added to draft or blocked research runs")
        self._registry.validate_evidence_kinds(run.template_key, [evidence.evidence_kind])
        self._repository.append_evidence_input(run_id, evidence.model_dump(mode="json"))
        logger.info("research evidence added", run_id=run_id, evidence_id=evidence.evidence_id)
        return self._get_run(run_id)

    def get(
        self,
        run_id: str,
        *,
        project_id: str | None = None,
        workspace_id: str | None = None,
    ) -> ResearchRun:
        self._require_run_scope(run_id, project_id=project_id, workspace_id=workspace_id)
        return self._get_run(run_id)

    def list_recent(
        self,
        *,
        limit: int = 20,
        project_id: str | None = None,
        workspace_id: str | None = None,
    ) -> list[ResearchRun]:
        if bool(project_id) != bool(workspace_id):
            raise ValueError("project and workspace scope must be supplied together")
        runs = self._repository.list_runs(limit=max(limit * 4, limit))
        if self._workspace_service is None:
            return runs[:limit]
        visible: list[ResearchRun] = []
        for run in runs:
            scope = self._workspace_service.repository.get_run_scope(run.run_id)
            if (project_id is None and scope is None) or (
                project_id is not None
                and workspace_id is not None
                and scope == (project_id, workspace_id)
            ):
                visible.append(run)
            if len(visible) == limit:
                break
        return visible

    def list_events(
        self,
        run_id: str,
        after_event_id: str | None = None,
        *,
        project_id: str | None = None,
        workspace_id: str | None = None,
    ) -> list[dict]:
        self._get_run(run_id)
        self._require_run_scope(run_id, project_id=project_id, workspace_id=workspace_id)
        if self._workspace_service is None:
            return []
        return self._workspace_service.list_run_events(run_id, after_event_id)

    def accept_provider_result(
        self,
        result: ProviderExecutionResult,
        *,
        request_hash: str,
    ) -> ResearchRun:
        """Correlate a durable DSH callback and apply it to its authoritative Run."""

        from core.contracts.research_workspace import ProviderTerminalStatus

        if self._runtime_service is None:
            raise RuntimeError("runtime provider service is required for callbacks")
        if result.request_hash != request_hash:
            raise ValueError("provider callback request hash mismatch")
        if self._workspace_service is not None:
            self._workspace_service.repository.db.scalar(
                self._workspace_service.repository.run_lock_statement(result.run_id)
            )
        run = self._get_run(result.run_id)
        task = self._repository.get_task_or_raise(result.run_id)
        snapshot = dict(task.state_snapshot or {})
        expected = {
            "provider_request_id": result.request_id,
            "provider_request_hash": request_hash,
            "provider_id": result.provider_id,
        }
        if any(snapshot.get(key) != value for key, value in expected.items()):
            raise ValueError("provider callback does not match the persisted Run request")
        if snapshot.get("provider_result_id") == result.provider_result_id and run.status in {
            ResearchRunStatus.COMPLETED,
            ResearchRunStatus.BLOCKED,
            ResearchRunStatus.FAILED,
        }:
            return run
        self._runtime_service.accept_provider_result(result, request_hash=request_hash)
        attempt = task.attempt
        metadata = {
            **snapshot,
            "provider_result_id": result.provider_result_id,
            "provider_terminal_status": result.terminal_status.value,
        }
        if result.terminal_status is ProviderTerminalStatus.COMPLETED:
            self._persist_execution_result(
                run.run_id,
                result.output,
                attempt=attempt,
                template_key=run.template_key,
            )
            blocked_reasons = list(result.output.get("blocked_reasons") or [])
            final_status = (
                ResearchRunStatus.BLOCKED if blocked_reasons else ResearchRunStatus.COMPLETED
            )
            updated = self._repository.update_run(
                run.run_id,
                status=final_status,
                source_plan=result.output.get("source_plan", {}),
                blocked_reasons=blocked_reasons,
                resume_from="validate" if blocked_reasons else None,
                completed_at=_utc_now() if final_status is ResearchRunStatus.COMPLETED else None,
            )
        else:
            final_status = (
                ResearchRunStatus.BLOCKED
                if result.terminal_status is ProviderTerminalStatus.BLOCKED
                else ResearchRunStatus.FAILED
            )
            updated = self._repository.update_run(
                run.run_id,
                status=final_status,
                blocked_reasons=[result.error_code or "blocked_runtime"],
                resume_from="planning",
            )
        self._repository.update_task(
            run.run_id,
            status=final_status,
            attempt=attempt,
            resume_from=updated.resume_from,
            state_snapshot={**metadata, "status": final_status.value},
            error_message=(
                result.error_code
                if result.terminal_status is not ProviderTerminalStatus.COMPLETED
                else None
            ),
        )
        if self._workspace_service is not None:
            self._workspace_service.record_run_event(
                run.run_id,
                status=final_status.value,
                stage="provider_callback",
                idempotency_key=f"{run.run_id}:provider:{result.provider_result_id}",
            )
            if final_status is ResearchRunStatus.COMPLETED:
                self._workspace_service.archive_completed_run_safely(run.run_id)
        logger.info(
            "DSH callback applied to research Run",
            run_id=run.run_id,
            provider_id=result.provider_id,
            provider_result_id=result.provider_result_id,
            status=final_status.value,
        )
        return updated

    def get_task(self, run_id: str) -> ResearchTask:
        """Expose the durable task snapshot for diagnostics and resume audits."""
        try:
            return self._repository.get_task_or_raise(run_id)
        except KeyError as exc:
            logger.warning("research task not found", run_id=run_id)
            raise ValueError(f"Research task for {run_id} not found") from exc

    def get_outputs(
        self,
        run_id: str,
        *,
        project_id: str | None = None,
        workspace_id: str | None = None,
    ) -> ResearchRunOutputs:
        run = self._get_run(run_id)
        self._require_run_scope(run_id, project_id=project_id, workspace_id=workspace_id)
        artifacts = self._repository.list_artifacts(run_id)
        card_artifact = _latest_artifact(artifacts, "decision_card")
        report_artifact = _latest_artifact(artifacts, "report_markdown")
        notes_artifact = _latest_artifact(artifacts, "research_notes")
        return ResearchRunOutputs(
            run_id=run_id,
            status=run.status,
            decision_card=(
                ResearchDecisionCard.model_validate(card_artifact.payload)
                if card_artifact is not None
                else None
            ),
            research_notes=(
                notes_artifact.payload.get("notes", []) if notes_artifact is not None else []
            ),
            claims=self._repository.list_claims(run_id),
            quality_gates=self._repository.list_quality_gates(run_id),
            artifacts=artifacts,
            report_markdown=(
                report_artifact.payload.get("content") if report_artifact is not None else None
            ),
        )

    def export_markdown(
        self,
        run_id: str,
        *,
        project_id: str | None = None,
        workspace_id: str | None = None,
    ) -> str:
        """Return a completed run's immutable Markdown report for download."""
        run = self._get_run(run_id)
        self._require_run_scope(run_id, project_id=project_id, workspace_id=workspace_id)
        if run.status is not ResearchRunStatus.COMPLETED:
            raise ValueError("Only completed research runs can be exported")
        report = self.get_outputs(
            run_id, project_id=project_id, workspace_id=workspace_id
        ).report_markdown
        if not report:
            raise ValueError("Completed research run has no Markdown report")
        return report

    def export_word(
        self,
        run_id: str,
        *,
        project_id: str | None = None,
        workspace_id: str | None = None,
    ) -> bytes:
        """Render the completed Markdown projection into an in-memory DOCX file."""
        markdown = self.export_markdown(run_id, project_id=project_id, workspace_id=workspace_id)
        try:
            document = Document()
            for line in markdown.splitlines():
                if line.startswith("# "):
                    document.add_heading(line[2:], level=1)
                elif line.startswith("## "):
                    document.add_heading(line[3:], level=2)
                elif line.startswith("- "):
                    document.add_paragraph(line[2:], style="List Bullet")
                elif line:
                    document.add_paragraph(line)
            buffer = BytesIO()
            document.save(buffer)
            return buffer.getvalue()
        except Exception as exc:
            logger.exception("research Word export failed", run_id=run_id, error=str(exc))
            raise ValueError("无法生成 Word 研究报告。") from exc

    def _execute(self, run: ResearchRun, *, increment_retry: bool) -> ResearchRun:
        attempt = run.retry_count + (1 if increment_retry else 0)
        try:
            template = self._registry.require_executable(run.template_key, run.subject)
            if template.graph_factory is None:
                raise ValueError(f"Research template {run.template_key} has no executor")
            graph = template.graph_factory()
            self._publish_progress(
                run.run_id,
                status=ResearchRunStatus.PLANNING,
                attempt=attempt,
                resume_from=run.resume_from,
                increment_retry=increment_retry,
            )
            runtime_input = {
                "run_id": run.run_id,
                "target_id": run.target_id,
                "subject": run.subject.model_dump(mode="json"),
                "as_of": run.as_of.isoformat(),
                "question": run.question,
                "evidence": self._repository.get_evidence_inputs(run.run_id),
            }
            if self._workspace_service is not None:
                workspace_context = self._workspace_service.build_run_runtime_context(run.run_id)
                if workspace_context:
                    runtime_input["workspace_context"] = workspace_context
            self._publish_progress(
                run.run_id,
                status=ResearchRunStatus.COLLECTING,
                attempt=attempt,
                resume_from="collecting",
            )
            skill_outputs = self._execute_skills(run, runtime_input)
            if skill_outputs:
                runtime_input["skill_outputs"] = skill_outputs
                runtime_input["evidence"] = self._repository.get_evidence_inputs(run.run_id)
            self._publish_progress(
                run.run_id,
                status=ResearchRunStatus.ANALYZING,
                attempt=attempt,
                resume_from="analyzing",
            )
            result, runtime_metadata = self._invoke_research_runtime(
                run,
                graph=graph,
                runtime_input=runtime_input,
                attempt=attempt,
            )
            if not isinstance(result, dict):
                raise TypeError("research runtime must return an object result")
            self._publish_progress(
                run.run_id,
                status=ResearchRunStatus.VALIDATING,
                attempt=attempt,
                resume_from="validating",
                state_metadata=runtime_metadata,
            )
            result = dict(result)
            source_plan = dict(result.get("source_plan") or {})
            source_plan["_runtime"] = {
                **dict((run.source_plan or {}).get("_runtime") or {}),
                **runtime_metadata,
            }
            result["source_plan"] = source_plan
            self._persist_execution_result(
                run.run_id, result, attempt=attempt, template_key=run.template_key
            )
            self._publish_progress(
                run.run_id,
                status=ResearchRunStatus.PUBLISHING,
                attempt=attempt,
                resume_from="publishing",
                state_metadata=runtime_metadata,
            )
            blocked_reasons = list(result.get("blocked_reasons", []))
            status = ResearchRunStatus.BLOCKED if blocked_reasons else ResearchRunStatus.COMPLETED
            completed_at = _utc_now() if status is ResearchRunStatus.COMPLETED else None
            updated = self._repository.update_run(
                run.run_id,
                status=status,
                source_plan=result.get("source_plan", {}),
                blocked_reasons=blocked_reasons,
                resume_from="validate" if blocked_reasons else None,
                completed_at=completed_at,
            )
            self._repository.update_task(
                run.run_id,
                status=status,
                attempt=attempt,
                resume_from="validate" if blocked_reasons else None,
                state_snapshot={
                    "status": status.value,
                    "attempt": attempt,
                    "source_plan": result.get("source_plan", {}),
                    "blocked_reasons": blocked_reasons,
                    **runtime_metadata,
                },
            )
            if self._workspace_service is not None:
                self._workspace_service.record_run_event(
                    run.run_id,
                    status=status.value,
                    stage=status.value,
                    idempotency_key=f"{run.run_id}:{status.value}:{attempt}",
                )
                if status is ResearchRunStatus.COMPLETED:
                    self._repository.db.commit()
                    self._workspace_service.archive_completed_run_safely(run.run_id)
            logger.info("research run executed", run_id=run.run_id, status=updated.status.value)
            return updated
        except RuntimeBlockedError as exc:
            logger.warning(
                "research run blocked by runtime",
                run_id=run.run_id,
                attempt=attempt,
                error_code=exc.code,
            )
            self._repository.update_run(
                run.run_id,
                status=ResearchRunStatus.BLOCKED,
                blocked_reasons=[exc.code],
                resume_from="planning",
            )
            self._repository.update_task(
                run.run_id,
                status=ResearchRunStatus.BLOCKED,
                attempt=attempt,
                resume_from="planning",
                state_snapshot={
                    "status": ResearchRunStatus.BLOCKED.value,
                    "attempt": attempt,
                    "error_code": exc.code,
                },
                error_message=exc.code,
            )
            if self._workspace_service is not None:
                self._workspace_service.record_run_event(
                    run.run_id,
                    status=ResearchRunStatus.BLOCKED.value,
                    stage="runtime",
                    idempotency_key=f"{run.run_id}:blocked:{attempt}",
                )
                self._repository.db.commit()
            raise
        except RuntimeFailedError as exc:
            logger.error(
                "research runtime returned a typed failure",
                run_id=run.run_id,
                attempt=attempt,
                error_code=exc.code,
            )
            self._repository.update_run(
                run.run_id,
                status=ResearchRunStatus.FAILED,
                blocked_reasons=[exc.code],
                resume_from="planning",
            )
            self._repository.update_task(
                run.run_id,
                status=ResearchRunStatus.FAILED,
                attempt=attempt,
                resume_from="planning",
                state_snapshot={
                    "status": ResearchRunStatus.FAILED.value,
                    "attempt": attempt,
                    "error_code": exc.code,
                },
                error_message=exc.code,
            )
            if self._workspace_service is not None:
                self._workspace_service.record_run_event(
                    run.run_id,
                    status=ResearchRunStatus.FAILED.value,
                    stage="runtime",
                    idempotency_key=f"{run.run_id}:failed:{attempt}",
                )
                self._repository.db.commit()
            raise
        except Exception as exc:
            logger.exception("research run failed", run_id=run.run_id, error=str(exc))
            self._repository.update_run(
                run.run_id,
                status=ResearchRunStatus.FAILED,
                blocked_reasons=["研究运行时发生内部错误，请检查运行日志后重试。"],
                resume_from="planning",
            )
            self._repository.update_task(
                run.run_id,
                status=ResearchRunStatus.FAILED,
                attempt=attempt,
                resume_from="planning",
                state_snapshot={"status": ResearchRunStatus.FAILED.value, "attempt": attempt},
                error_message="研究运行时发生内部错误，请检查运行日志后重试。",
            )
            if self._workspace_service is not None:
                self._workspace_service.record_run_event(
                    run.run_id,
                    status=ResearchRunStatus.FAILED.value,
                    stage="planning",
                    idempotency_key=f"{run.run_id}:failed:{attempt}",
                )
                self._repository.db.commit()
            raise

    def _publish_progress(
        self,
        run_id: str,
        *,
        status: ResearchRunStatus,
        attempt: int,
        resume_from: str | None,
        increment_retry: bool = False,
        state_metadata: dict[str, Any] | None = None,
    ) -> None:
        """Commit one status checkpoint so a concurrent SSE reader can observe it."""

        snapshot = {"status": status.value, "attempt": attempt, **(state_metadata or {})}
        self._repository.update_task(
            run_id,
            status=status,
            attempt=attempt,
            resume_from=resume_from,
            state_snapshot=snapshot,
        )
        self._repository.update_run(
            run_id,
            status=status,
            increment_retry=increment_retry,
        )
        if self._workspace_service is not None:
            self._workspace_service.record_run_event(
                run_id,
                status=status.value,
                stage=status.value,
                idempotency_key=f"{run_id}:{status.value}:{attempt}",
            )
            self._repository.db.commit()

    def _execute_skills(
        self,
        run: ResearchRun,
        runtime_input: dict[str, Any],
    ) -> dict[str, dict[str, Any]]:
        config = dict((run.source_plan or {}).get("_runtime") or {})
        skill_keys = list(config.get("skill_keys") or [])
        if not skill_keys:
            return {}
        if self._runtime_service is None or self._skill_invoker is None:
            raise RuntimeBlockedError(
                "Configured Skills have no authorized execution adapter",
                code="blocked_runtime",
            )
        outputs: dict[str, dict[str, Any]] = {}
        remaining_tokens = 4096
        remaining_cost = 1.0
        for skill_key in skill_keys:
            try:
                manifest = self._runtime_service.get_skill(skill_key)
            except KeyError as exc:
                raise RuntimeBlockedError(
                    f"Enabled Skill not found: {skill_key}", code="blocked_runtime"
                ) from exc
            usage: list[InvocationUsage] = []
            try:
                from functools import partial

                outputs[skill_key] = self._runtime_service.execute_skill(
                    manifest,
                    runtime_input,
                    partial(self._skill_invoker, manifest),
                    authorized_tool_ids=self._authorized_tool_ids,
                    deadline_seconds=30,
                    reserve_tokens=remaining_tokens,
                    remaining_tokens=remaining_tokens,
                    reserve_cost=remaining_cost,
                    remaining_cost=remaining_cost,
                    usage_callback=usage.append,
                )
            except ValueError as exc:
                raise RuntimeBlockedError(
                    f"Skill {skill_key} failed its authorization/schema boundary",
                    code="skill_boundary_rejected",
                ) from exc
            if not usage:
                raise RuntimeFailedError(
                    f"Skill {skill_key} omitted usage accounting",
                    code="missing_skill_usage",
                )
            remaining_tokens -= usage[0].tokens_used
            remaining_cost -= usage[0].cost_used
        return outputs

    def _invoke_research_runtime(
        self,
        run: ResearchRun,
        *,
        graph: Any,
        runtime_input: dict[str, Any],
        attempt: int,
    ) -> tuple[dict[str, Any], dict[str, Any]]:
        if self._runtime_service is None:
            return graph.invoke(runtime_input), {}
        config = dict((run.source_plan or {}).get("_runtime") or {})
        mode = str(config.get("mode") or "fingpt")
        required = {"agent_team"} if mode == "claw" else {"single_agent"}
        request_id = f"{run.run_id}:{attempt}"

        def invoke(provider: RuntimeProvider) -> Any:
            if provider.provider_type == "langgraph":
                return graph.invoke(runtime_input)
            if self._provider_invoker is None:
                raise RuntimeError("DSH invocation adapter is unavailable")
            return self._provider_invoker(provider, runtime_input, request_id)

        def record_provider_request(provider: RuntimeProvider) -> None:
            if provider.provider_type != "dsh":
                return
            self._repository.update_task(
                run.run_id,
                status=ResearchRunStatus.PLANNING,
                attempt=attempt,
                resume_from="planning",
                state_snapshot={
                    "status": ResearchRunStatus.PLANNING.value,
                    "attempt": attempt,
                    "provider_request_id": request_id,
                    "provider_request_hash": stable_runtime_request_hash(runtime_input),
                    "provider_id": provider.provider_id,
                },
            )
            self._repository.db.commit()

        routed: RuntimeExecutionResult = self._runtime_service.execute(
            mode,
            required,
            invoke,
            before_invoke=record_provider_request,
        )
        runtime_output = routed.output
        if mode == "claw":
            team_run = self._execute_claw_team(run, runtime_output)
            runtime_output = self._merge_agent_team_result(runtime_output, team_run)
        return runtime_output, {
            "provider_id": routed.provider_id,
            "provider_result_id": routed.provider_result_id,
            "provider_terminal_status": routed.terminal_status.value,
            "fallback_from": routed.fallback_from,
        }

    def _execute_claw_team(self, run: ResearchRun, provider_output: Any):
        config = dict((run.source_plan or {}).get("_runtime") or {})
        team_id = config.get("agent_team_id")
        if not team_id or self._agent_team_service is None or self._agent_worker is None:
            raise RuntimeBlockedError(
                "Claw requires an executable Agent Team", code="blocked_runtime"
            )
        try:
            team = self._agent_team_service.get_team(str(team_id))
        except KeyError as exc:
            raise RuntimeBlockedError("Agent Team not found", code="blocked_runtime") from exc
        output = provider_output if isinstance(provider_output, dict) else {}
        steps = list(output.get("agent_steps") or [])
        if not steps:
            steps = [
                {
                    "role": role,
                    "task": {"run_id": run.run_id, "question": run.question},
                }
                for role in team.roles
                if role != team.supervisor_role
            ]
        try:
            return self._agent_team_service.execute_plan(
                team,
                steps,
                self._agent_worker,
                supervisor=self._agent_supervisor,
            )
        except Exception as exc:
            from services.agent_team_service import AgentBudgetError

            if isinstance(exc, AgentBudgetError):
                raise RuntimeBlockedError(str(exc), code=exc.code) from exc
            if isinstance(exc, ValueError) and "sensitive" in str(exc).lower():
                raise RuntimeBlockedError(
                    "Agent output failed the safety boundary",
                    code="unsafe_agent_output",
                ) from exc
            raise

    @staticmethod
    def _merge_agent_team_result(provider_output: Any, team_run) -> dict[str, Any]:
        """Merge typed worker projections into the Run result before persistence."""

        result = dict(provider_output) if isinstance(provider_output, dict) else {}
        notes = list(result.get("research_notes") or [])
        claims = list(result.get("claims") or [])
        report_sections: list[str] = []
        for entry in team_run.blackboard:
            if entry.entry_type != "result":
                continue
            payload = dict(entry.payload or {})
            notes.extend(list(payload.get("research_notes") or []))
            claims.extend(list(payload.get("claims") or []))
            if payload.get("report_markdown"):
                report_sections.append(str(payload["report_markdown"]))
        result["research_notes"] = notes
        result["claims"] = claims
        result["agent_blackboard"] = [item.model_dump(mode="json") for item in team_run.blackboard]
        if report_sections:
            original = str(result.get("report_markdown") or "").rstrip()
            result["report_markdown"] = (
                original + "\n\n## Claw Agent Team\n" + "\n\n".join(report_sections)
            ).strip()
        return result

    def _persist_execution_result(
        self, run_id: str, result: dict, *, attempt: int, template_key: str
    ) -> None:
        now = _utc_now()
        self._registry.validate_evidence_kinds(
            template_key, [claim["category"] for claim in result.get("claims", [])]
        )
        artifacts = {
            "source_plan": result.get("source_plan", {}),
            "research_notes": {"notes": result.get("research_notes", [])},
            "challenge": {"quality_checks": result.get("quality_checks", [])},
        }
        if result.get("agent_blackboard"):
            artifacts["agent_blackboard"] = {"entries": result["agent_blackboard"]}
        if result.get("decision_card"):
            artifacts["decision_card"] = result["decision_card"]
        if result.get("report_markdown"):
            artifacts["report_markdown"] = {"content": result["report_markdown"]}
        for artifact_type, payload in artifacts.items():
            self._repository.append_artifact(
                ResearchArtifact(
                    artifact_id=_new_id("research_artifact"),
                    run_id=run_id,
                    artifact_type=f"{artifact_type}:{attempt}",
                    payload=payload,
                    created_at=now,
                )
            )

        claims = [
            ResearchClaim(
                claim_id=_new_id("research_claim"),
                run_id=run_id,
                category=claim["category"],
                text=claim["text"],
                evidence_refs=claim["evidence_refs"],
                numeric_context=claim["numeric_context"],
                conflict_status=claim["conflict_status"],
                created_at=now,
            )
            for claim in result.get("claims", [])
        ]
        self._repository.replace_claims(run_id, claims)
        gates = [
            QualityGateResult(
                gate_id=_new_id("research_gate"),
                run_id=run_id,
                gate_key=check["gate_key"],
                passed=check["passed"],
                message=check["message"],
                details={},
                checked_at=now,
            )
            for check in result.get("quality_checks", [])
        ]
        self._repository.replace_quality_gates(run_id, gates)

    def _get_run(self, run_id: str) -> ResearchRun:
        try:
            return self._repository.get_run_or_raise(run_id)
        except KeyError as exc:
            logger.warning("research run not found", run_id=run_id)
            raise ValueError(f"Research run {run_id} not found") from exc

    def _require_run_scope(
        self,
        run_id: str,
        *,
        project_id: str | None,
        workspace_id: str | None,
    ) -> None:
        if self._workspace_service is not None:
            self._workspace_service.require_run_scope(
                run_id,
                project_id=project_id,
                workspace_id=workspace_id,
            )


def _new_id(prefix: str) -> str:
    return f"{prefix}_{uuid4().hex}"


def _utc_now() -> datetime:
    return datetime.now(UTC)


def _request_hash(payload: dict[str, Any]) -> str:
    encoded = json.dumps(payload, sort_keys=True, separators=(",", ":"), ensure_ascii=False)
    return "sha256:" + hashlib.sha256(encoded.encode("utf-8")).hexdigest()


def _latest_artifact(
    artifacts: list[ResearchArtifact], artifact_type: str
) -> ResearchArtifact | None:
    """Return the newest immutable artifact revision for a logical output type."""
    revisions: list[tuple[int, ResearchArtifact]] = []
    prefix = f"{artifact_type}:"
    for artifact in artifacts:
        if not artifact.artifact_type.startswith(prefix):
            continue
        try:
            revisions.append((int(artifact.artifact_type.removeprefix(prefix)), artifact))
        except ValueError:
            continue
    return max(revisions, key=lambda item: item[0])[1] if revisions else None
