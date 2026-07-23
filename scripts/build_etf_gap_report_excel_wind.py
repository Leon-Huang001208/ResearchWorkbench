"""Build Huaan ETF gap dataset through Excel Wind add-in.

This script intentionally uses xlwings + Excel Wind formulas only. It does not
call Wind MCP, WindPy, or HTTP APIs. Run it with Anaconda Python because that is
where xlwings is installed on this machine:

    /Users/leon/opt/anaconda3/bin/python scripts/build_etf_gap_report_excel_wind.py
"""

from __future__ import annotations

import argparse
import datetime as dt
import json
import logging
import math
import re
import sys
import time
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

PROJECT_ROOT = Path(__file__).resolve().parents[1]
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

from data_layer.adapters.wind.client import WindExcelClient  # noqa: E402

DEFAULT_INPUT = Path("/Users/leon/Desktop/公募基金_概况.xlsx")
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "outputs" / "etf_gap_report_20260702"
LOG_DIR = PROJECT_ROOT / "logs"
EXCEL_ERRORS = {"#N/A", "#VALUE!", "#REF!", "#DIV/0!", "#NAME?", "#NUM!", "#NULL!"}
WIND_LOADING = {"fetch...", "fetching...", "loading...", "calculating...", "connecting..."}
WIND_CODE_RE = re.compile(r"^[A-Z0-9]+\.[A-Z]{2,5}$")
PE_FETCH_SUFFIXES = {".SH", ".SZ", ".CSI", ".CNI", ".MI"}


def _setup_logging() -> Path:
    LOG_DIR.mkdir(parents=True, exist_ok=True)
    log_path = LOG_DIR / f"etf_gap_excel_wind_{dt.datetime.now():%Y%m%dT%H%M%S}.log"
    logging.basicConfig(
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        handlers=[
            logging.FileHandler(log_path, encoding="utf-8"),
            logging.StreamHandler(sys.stdout),
        ],
    )
    return log_path


def _as_text(value: Any) -> str:
    if value is None:
        return ""
    return str(value).strip()


def _as_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return float(value)
    try:
        text = str(value).replace(",", "").strip()
        if not text or text == "--":
            return None
        parsed = float(text)
        return parsed if math.isfinite(parsed) else None
    except Exception:
        return None


def _jsonable(value: Any) -> Any:
    if isinstance(value, (dt.date, dt.datetime)):
        return value.isoformat()
    if isinstance(value, float) and not math.isfinite(value):
        return None
    return value


def _excel_serial_to_date(value: Any) -> str | None:
    if isinstance(value, dt.datetime):
        return value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    num = _as_float(value)
    if num is None:
        return None
    try:
        # Excel serial date system used by xlwings on Mac/Windows for workbook values.
        return (dt.date(1899, 12, 30) + dt.timedelta(days=int(num))).isoformat()
    except Exception:
        return None


def _five_year_start(end_date: dt.date) -> dt.date:
    try:
        return end_date.replace(year=end_date.year - 5)
    except ValueError:
        return end_date.replace(year=end_date.year - 5, day=28)


def _is_loading(value: Any) -> bool:
    if value is None:
        return True
    if isinstance(value, str):
        return value.strip().lower() in WIND_LOADING
    return False


def _is_error(value: Any) -> bool:
    if isinstance(value, str):
        return value.strip().upper() in EXCEL_ERRORS
    return False


def read_etf_rows(input_path: Path) -> list[dict[str, Any]]:
    if not input_path.exists():
        raise FileNotFoundError(input_path)
    workbook = load_workbook(input_path, read_only=True, data_only=True)
    sheet = workbook[workbook.sheetnames[0]]
    rows = list(sheet.iter_rows(values_only=True))
    if not rows:
        raise ValueError(f"Empty workbook: {input_path}")
    headers = [_as_text(cell) for cell in rows[0]]
    required = {"基金代码", "基金名称", "管理人"}
    missing = sorted(required - set(headers))
    if missing:
        raise ValueError(f"Missing required columns: {missing}")

    out: list[dict[str, Any]] = []
    for row in rows[1:]:
        record = {
            headers[i]: _jsonable(row[i] if i < len(row) else None) for i in range(len(headers))
        }
        if record.get("基金代码") and record.get("基金名称"):
            out.append(record)
    workbook.close()
    return out


def _wind_result_to_value(result: Any) -> tuple[Any, str]:
    if isinstance(result, Exception):
        return None, f"error:{result}"
    if _is_error(result):
        return None, f"error:{result}"
    if _is_loading(result):
        return None, "timeout_or_empty"
    return result, "ok"


def normalize_wind_index_code(value: Any) -> tuple[str, str]:
    """Return a valid Wind index code and status.

    Some ETF rows return numeric placeholders such as 0/0.0 instead of a real
    tracking index. Keep those rows auditable, but exclude them from WSD calls.
    """
    text = _as_text(value).upper()
    if not text or text in {"--", "0", "0.0", "NONE", "NULL"}:
        return "", "invalid_or_empty"
    tokens = re.split(r"[,;；\\s]+", text)
    for token in tokens:
        cleaned = token.strip()
        if WIND_CODE_RE.match(cleaned) and not cleaned.endswith(".OF"):
            return cleaned, "ok"
    return "", f"invalid_code:{text}"


def is_pe_fetch_eligible(code: str) -> bool:
    """Limit WSD PE pulls to A-share/China equity index namespaces.

    Some cross-border/global index namespaces can crash or disconnect the Mac
    Excel automation bridge when asked for pe_ttm histories. They remain in the
    coverage table, but are not considered PE-screen candidates.
    """
    return any(code.endswith(suffix) for suffix in PE_FETCH_SUFFIXES)


def _cache_is_reusable(payload: dict[str, Any]) -> bool:
    return payload.get("status") == "ok" and payload.get("sample_count")


def fetch_wss_batch(
    client: WindExcelClient,
    formulas: list[str],
    *,
    chunk_size: int,
    timeout: float,
) -> list[dict[str, Any]]:
    results: list[dict[str, Any]] = []
    for start in range(0, len(formulas), chunk_size):
        chunk = formulas[start : start + chunk_size]
        logging.info("Fetching WSS batch %s-%s of %s", start + 1, start + len(chunk), len(formulas))
        try:
            values = client.execute_batch(chunk, timeout=timeout)
        except Exception:
            logging.exception("WSS batch failed; falling back to per-formula execution")
            values = []
            for formula in chunk:
                try:
                    values.append(client.execute(formula, timeout=timeout))
                except Exception as inner:
                    values.append(inner)
                    logging.warning("WSS formula failed: %s %s", formula, inner)
        for formula, raw in zip(chunk, values):
            value, status = _wind_result_to_value(raw)
            results.append({"formula": formula, "value": _as_text(value), "status": status})
    return results


def _get_formula_sheet(client: WindExcelClient) -> Any:
    client._ensure_connected()  # Existing project adapter owns connection/session setup.
    return client._sheet


def fetch_wsd_weekly_pe(
    client: WindExcelClient,
    index_code: str,
    *,
    start_date: dt.date,
    end_date: dt.date,
    timeout: float,
    max_rows: int = 290,
) -> dict[str, Any]:
    sheet = _get_formula_sheet(client)
    app = client._app
    start_row = client._allocate_helper_rows(max_rows + 2)
    cell = sheet.range(f"ZX{start_row}")
    block = sheet.range(f"ZX{start_row}:ZY{start_row + max_rows - 1}")
    formula = (
        f'=@wsd("{index_code}","pe_ttm","{start_date:%Y-%m-%d}",'
        f'"{end_date:%Y-%m-%d}","Period=W")'
    )
    try:
        block.value = None
        time.sleep(0.05)
        cell.value = formula
        try:
            app.calculate()
        except Exception:
            pass

        started = time.monotonic()
        values: list[Any] = []
        while time.monotonic() - started < timeout:
            time.sleep(0.8)
            first = cell.value
            if _is_error(first):
                return {"status": f"error:{first}", "formula": formula, "series": []}
            raw = block.value
            if raw and isinstance(raw, list):
                values = raw
                rows_with_pe = 0
                for row in values:
                    if isinstance(row, list) and len(row) >= 2 and _as_float(row[1]) is not None:
                        rows_with_pe += 1
                if rows_with_pe >= 4:
                    break
            elif not _is_loading(first):
                values = [[first, None]]

        series: list[dict[str, Any]] = []
        for row in values if isinstance(values, list) else []:
            if not isinstance(row, list) or len(row) < 2:
                continue
            pe = _as_float(row[1])
            date_text = _excel_serial_to_date(row[0])
            if date_text and pe is not None and pe > 0:
                series.append({"date": date_text, "pe_ttm": pe})

        if not series:
            final = cell.value
            status = f"error:{final}" if _is_error(final) else "timeout_or_empty"
            return {"status": status, "formula": formula, "series": []}

        return {"status": "ok", "formula": formula, "series": series}
    except Exception as exc:
        logging.exception("WSD weekly PE failed for %s", index_code)
        return {"status": f"error:{exc}", "formula": formula, "series": []}
    finally:
        try:
            block.value = None
        except Exception as exc:
            logging.debug("Unable to clear WSD helper block: %s", exc)


def summarize_pe(series: list[dict[str, Any]]) -> dict[str, Any]:
    values = [_as_float(item.get("pe_ttm")) for item in series]
    clean = [value for value in values if value is not None and value > 0]
    if not clean:
        return {"current_pe": None, "pe_percentile_5y": None, "sample_count": 0}
    current = clean[-1]
    percentile = sum(1 for value in clean if value <= current) / len(clean)
    return {
        "current_pe": current,
        "pe_percentile_5y": percentile,
        "sample_count": len(clean),
        "first_pe_date": series[0].get("date"),
        "last_pe_date": series[-1].get("date"),
    }


def build_dataset(args: argparse.Namespace) -> dict[str, Any]:
    etfs = read_etf_rows(args.input)
    if args.limit:
        etfs = etfs[: args.limit]
    logging.info("Loaded %s ETF rows from %s", len(etfs), args.input)

    end_date = args.end_date or dt.date.today()
    start_date = _five_year_start(end_date)
    cache_dir = args.output_dir / "wind_excel_plugin_cache"
    cache_dir.mkdir(parents=True, exist_ok=True)
    wsd_cache_dir = cache_dir / f"wsd_pe_ttm_weekly_{start_date:%Y%m%d}_{end_date:%Y%m%d}"
    wsd_cache_dir.mkdir(parents=True, exist_ok=True)

    client = WindExcelClient(visible=args.visible, timeout=args.timeout)
    client._connect()
    if not client.heartbeat():
        raise RuntimeError("Excel Wind 插件心跳失败：请确认 Excel 已登录 Wind。")

    tracking_formulas = [f'=@wss("{row["基金代码"]}","fund_trackindexcode")' for row in etfs]
    tracking_results = fetch_wss_batch(
        client,
        tracking_formulas,
        chunk_size=args.wss_chunk_size,
        timeout=args.timeout,
    )

    for row, result in zip(etfs, tracking_results):
        code, code_status = normalize_wind_index_code(result["value"])
        row["跟踪指数代码原始值"] = result["value"]
        row["跟踪指数代码"] = code
        row["跟踪指数代码状态"] = result["status"] if result["status"] != "ok" else code_status
        row["是否华安"] = "是" if "华安" in _as_text(row.get("管理人")) else "否"

    index_codes = sorted(
        {_as_text(row.get("跟踪指数代码")) for row in etfs if row.get("跟踪指数代码")}
    )
    if args.index_limit:
        index_codes = index_codes[: args.index_limit]
    logging.info("Unique tracking indices to fetch: %s", len(index_codes))

    name_formulas = [f'=@wss("{code}","sec_name")' for code in index_codes]
    name_results = fetch_wss_batch(
        client,
        name_formulas,
        chunk_size=args.wss_chunk_size,
        timeout=args.timeout,
    )
    index_name_by_code = {
        code: result["value"] if result["status"] == "ok" else ""
        for code, result in zip(index_codes, name_results)
    }
    index_name_status_by_code = {
        code: result["status"] for code, result in zip(index_codes, name_results)
    }

    pe_by_code: dict[str, dict[str, Any]] = {
        code: {
            "status": "not_pe_eligible",
            "series": [],
            "current_pe": None,
            "pe_percentile_5y": None,
            "sample_count": 0,
        }
        for code in index_codes
        if not is_pe_fetch_eligible(code)
    }
    pe_fetch_codes = [code for code in index_codes if is_pe_fetch_eligible(code)]
    logging.info(
        "PE-fetch eligible indices: %s; skipped non-eligible: %s",
        len(pe_fetch_codes),
        len(index_codes) - len(pe_fetch_codes),
    )
    for offset, code in enumerate(pe_fetch_codes, start=1):
        cache_path = wsd_cache_dir / f"{code.replace('.', '_')}.json"
        if args.use_cache and cache_path.exists():
            try:
                cached = json.loads(cache_path.read_text(encoding="utf-8"))
                if _cache_is_reusable(cached):
                    pe_by_code[code] = cached
                    logging.info("Using cached WSD %s/%s %s", offset, len(pe_fetch_codes), code)
                    continue
            except Exception:
                logging.warning("Ignoring unreadable cache: %s", cache_path)

        logging.info("Fetching WSD weekly PE %s/%s %s", offset, len(pe_fetch_codes), code)
        result = fetch_wsd_weekly_pe(
            client,
            code,
            start_date=start_date,
            end_date=end_date,
            timeout=args.wsd_timeout,
        )
        result.update(summarize_pe(result.get("series", [])))
        pe_by_code[code] = result
        cache_path.write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding="utf-8")
        time.sleep(args.wsd_pause)

    coverage_rows: list[dict[str, Any]] = []
    for code in index_codes:
        related = [row for row in etfs if row.get("跟踪指数代码") == code]
        huaan = [row for row in related if row.get("是否华安") == "是"]
        total_scale = sum(_as_float(row.get("规模(亿)")) or 0 for row in related)
        largest = sorted(
            related, key=lambda item: _as_float(item.get("规模(亿)")) or 0, reverse=True
        )
        managers = sorted({_as_text(row.get("管理人")) for row in related if row.get("管理人")})
        pe = pe_by_code.get(code, {})
        percentile = _as_float(pe.get("pe_percentile_5y"))
        sample_count = int(pe.get("sample_count") or 0)
        is_candidate = bool(
            not huaan
            and percentile is not None
            and percentile < args.threshold
            and sample_count >= args.min_samples
        )
        coverage_rows.append(
            {
                "跟踪指数代码": code,
                "跟踪指数名称": index_name_by_code.get(code, ""),
                "指数名称状态": index_name_status_by_code.get(code, ""),
                "华安是否覆盖": "是" if huaan else "否",
                "华安产品": "；".join(
                    f"{row.get('基金代码')} {row.get('基金名称')}" for row in huaan
                ),
                "现有ETF数量": len(related),
                "现有ETF总规模(亿)": round(total_scale, 4),
                "竞品管理人": "；".join(managers),
                "代表ETF产品": "；".join(
                    f"{row.get('基金代码')} {row.get('基金名称')}" for row in largest[:5]
                ),
                "PE(TTM)": pe.get("current_pe"),
                "近五年PE分位": percentile,
                "近五年样本数": sample_count,
                "PE样本起始日": pe.get("first_pe_date"),
                "PE样本截至日": pe.get("last_pe_date"),
                "Wind取数状态": pe.get("status", "missing"),
                "候选": "是" if is_candidate else "否",
            }
        )

    coverage_rows.sort(
        key=lambda item: (
            item["候选"] != "是",
            item.get("近五年PE分位") if item.get("近五年PE分位") is not None else 999,
            -(item.get("现有ETF总规模(亿)") or 0),
        )
    )
    candidates = [row for row in coverage_rows if row["候选"] == "是"]

    return {
        "metadata": {
            "generated_at": dt.datetime.now().isoformat(timespec="seconds"),
            "input_path": str(args.input),
            "source": "Excel Wind 插件（xlwings 写入 WSS/WSD 公式）",
            "wind_mcp_used": False,
            "valuation_date": end_date.isoformat(),
            "five_year_start": start_date.isoformat(),
            "threshold": args.threshold,
            "min_samples": args.min_samples,
            "etf_count": len(etfs),
            "index_count": len(index_codes),
            "candidate_count": len(candidates),
        },
        "etfs": etfs,
        "coverage": coverage_rows,
        "candidates": candidates,
        "pe_series": pe_by_code,
    }


def _parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Build Huaan ETF gap dataset via Excel Wind add-in."
    )
    parser.add_argument("--input", type=Path, default=DEFAULT_INPUT)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--output-json", type=Path)
    parser.add_argument("--end-date", type=lambda text: dt.date.fromisoformat(text), default=None)
    parser.add_argument("--threshold", type=float, default=0.90)
    parser.add_argument("--min-samples", type=int, default=180)
    parser.add_argument("--visible", action="store_true")
    parser.add_argument("--timeout", type=float, default=45.0)
    parser.add_argument("--wsd-timeout", type=float, default=45.0)
    parser.add_argument("--wsd-pause", type=float, default=0.1)
    parser.add_argument("--wss-chunk-size", type=int, default=80)
    parser.add_argument("--limit", type=int, help="Debug: limit ETF rows.")
    parser.add_argument("--index-limit", type=int, help="Debug: limit unique indices.")
    parser.add_argument("--no-cache", action="store_false", dest="use_cache")
    parser.set_defaults(use_cache=True)
    return parser.parse_args()


def main() -> None:
    log_path = _setup_logging()
    args = _parse_args()
    args.input = args.input.expanduser().resolve()
    args.output_dir = args.output_dir.expanduser().resolve()
    args.output_dir.mkdir(parents=True, exist_ok=True)
    output_json = args.output_json or (args.output_dir / "huaan_etf_gap_excel_wind_data.json")

    logging.info("Log path: %s", log_path)
    logging.info("Starting Excel Wind ETF gap dataset build")
    dataset = build_dataset(args)
    output_json.write_text(
        json.dumps(dataset, ensure_ascii=False, indent=2, default=_jsonable), encoding="utf-8"
    )
    logging.info("Wrote dataset JSON: %s", output_json)
    print(json.dumps({"output_json": str(output_json), "log": str(log_path)}, ensure_ascii=False))


if __name__ == "__main__":
    main()
