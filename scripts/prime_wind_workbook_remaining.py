"""Prime remaining Wind formula rows that are still missing G-column data.

Runs synchronously in a single Python process to avoid COM session conflicts.
Safe to re-run — skips rows where G-column already has a value.
"""

import os
import sys
import time

sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
from dotenv import load_dotenv

load_dotenv()

import xlwings as xw
from openpyxl import load_workbook

from services.wind_realtime_workbook import resolve_workbook_path

wb_path = resolve_workbook_path(None)

# Load formula map from disk
print("Reading formulas from disk...", flush=True)
wb_disk = load_workbook(str(wb_path), data_only=False, read_only=True)
formulas = [
    (cell.row, cell.value)
    for row in wb_disk["RealtimeRaw"].iter_rows()
    for cell in row
    if cell.value and isinstance(cell.value, str) and "=wss(" in cell.value
]
wb_disk.close()
print(f"Total formula rows: {len(formulas)}", flush=True)

# Find already-open workbook
book = None
for app in xw.apps:
    for b in app.books:
        if "AlphaFoundry_Wind_Realtime" in b.name:
            book = b
            break

if book is None:
    print("ERROR: workbook not open in Excel. Open it first.", flush=True)
    sys.exit(1)

raw = book.sheets["RealtimeRaw"]

skipped = 0
written = 0
failed = 0

for i, (row_num, formula) in enumerate(formulas):
    # Skip if G-column already has a valid value
    try:
        g = raw.range((row_num, 7)).value
        if g is not None and float(g) > 0:
            skipped += 1
            print(f"[{i+1}/{len(formulas)}] row {row_num}: skip (G={g})", flush=True)
            continue
    except (TypeError, ValueError):
        pass

    # Write formula with retry on COM-busy
    wrote = False
    for attempt in range(10):
        try:
            raw.range((row_num, 6)).formula = formula
            wrote = True
            print(
                f"[{i+1}/{len(formulas)}] row {row_num}: written (attempt {attempt+1})", flush=True
            )
            break
        except Exception as e:
            wait = 10 + attempt * 5
            print(f"  busy attempt {attempt+1}/10, wait {wait}s: {str(e)[:60]}", flush=True)
            time.sleep(wait)

    if not wrote:
        print(f"  FAILED row {row_num} after 10 attempts", flush=True)
        failed += 1
        continue

    # Poll until Wind fills G-column (up to 90s)
    t0 = time.monotonic()
    got = False
    while time.monotonic() - t0 < 90:
        try:
            g = raw.range((row_num, 7)).value
            if g is not None:
                try:
                    if float(g) > 0:
                        print(f"  Wind done in {time.monotonic()-t0:.0f}s: G={g}", flush=True)
                        got = True
                        break
                except (TypeError, ValueError):
                    pass
        except Exception:
            pass
        time.sleep(3)

    if not got:
        print(f"  timeout (90s) waiting for Wind at row {row_num} — continuing", flush=True)

    written += 1

print(f"\nDone. written={written} skipped={skipped} failed={failed}", flush=True)
print("Saving workbook...", flush=True)
try:
    book.save()
    print("Saved OK", flush=True)
except Exception as e:
    print(f"Save error: {e}", flush=True)

# Final snapshot check
from services.wind_realtime_workbook import WindRealtimeWorkbookReader

result = WindRealtimeWorkbookReader(stale_after_seconds=300).get_view("wind_hot_concept", limit=10)
print(
    f"\nSnapshot wind_hot_concept: up={len(result.get('up',[]))} down={len(result.get('down',[]))} status={result.get('status')}",
    flush=True,
)
if result.get("up"):
    print("Top gainer:", result["up"][0].get("name"), result["up"][0].get("change"), flush=True)
if result.get("down"):
    print("Top loser:", result["down"][0].get("name"), result["down"][0].get("change"), flush=True)
