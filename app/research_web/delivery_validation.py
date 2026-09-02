"""Trusted parser source, executed ONLY via sandbox.run_script, never on the host."""

import hashlib
import io
import json
import os
import stat
from html.parser import HTMLParser
from pathlib import Path

METADATA_SHEETS = {
    "sources",
    "source",
    "references",
    "reference",
    "notes",
    "readme",
    "metadata",
    "来源",
    "数据来源",
    "说明",
    "参考资料",
    "参考文献",
}


class VisibleHTML(HTMLParser):
    def __init__(self):
        super().__init__()
        self.hidden = 0
        self.content = False

    def handle_starttag(self, tag, attrs):
        if tag in {"script", "style", "head"}:
            self.hidden += 1

    def handle_endtag(self, tag):
        if tag in {"script", "style", "head"}:
            self.hidden = max(0, self.hidden - 1)

    def handle_data(self, data):
        if not self.hidden and data.strip():
            self.content = True


def meaningful(value):
    return value is not None and (not isinstance(value, str) or bool(value.strip()))


def validate_content(raw, extension):
    if not raw:
        return "文件为空"
    stream = io.BytesIO(raw)
    if extension == "docx":
        from docx import Document

        doc = Document(stream)
        if not any(p.text.strip() for p in doc.paragraphs) and not any(
            cell.text.strip() for table in doc.tables for row in table.rows for cell in row.cells
        ):
            return "Word 可打开，但没有正文或表格内容"
    elif extension == "xlsx":
        from openpyxl import load_workbook

        book = load_workbook(stream, read_only=True, data_only=True)
        try:
            for sheet in book:
                if sheet.title.strip().casefold() in METADATA_SHEETS:
                    continue
                rows = 0
                # Do not trust dimension metadata to hide data or force huge allocation.
                sheet.reset_dimensions()
                for row in sheet.iter_rows(values_only=True):
                    if any(meaningful(value) for value in row):
                        rows += 1
                    if rows >= 2:
                        return None
            return "Excel 可打开，但非元数据工作表没有至少两行有效内容（表头及数据）"
        finally:
            book.close()
    elif extension == "png":
        from PIL import Image

        with Image.open(stream) as image:
            if image.format != "PNG" or min(image.size) < 1:
                return "不是有效的 PNG 图像"
            image.verify()
    elif extension == "html":
        html = VisibleHTML()
        html.feed(raw.decode("utf-8-sig"))
        html.close()
        if not html.content:
            return "HTML 没有可见文本内容"
    elif extension == "md":
        if not raw.decode("utf-8-sig").strip():
            return "Markdown 没有文本内容"
    else:
        return "不支持的交付格式"
    return None


def validate_file(item):
    result = {"id": item["id"], "sha256": item["sha256"], "valid": False, "reason": None}
    descriptor = None
    try:
        relative = Path(item["path"])
        if relative.is_absolute() or ".." in relative.parts or relative.parts[0] != "outputs":
            raise ValueError("invalid path")
        descriptor = os.open("outputs", os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW)
        for part in relative.parts[1:-1]:
            child = os.open(part, os.O_RDONLY | os.O_DIRECTORY | os.O_NOFOLLOW, dir_fd=descriptor)
            os.close(descriptor)
            descriptor = child
        fd = os.open(relative.name, os.O_RDONLY | os.O_NOFOLLOW | os.O_NONBLOCK, dir_fd=descriptor)
        with os.fdopen(fd, "rb") as stream:
            info = os.fstat(stream.fileno())
            if not stat.S_ISREG(info.st_mode) or info.st_nlink != 1:
                raise ValueError("not an independent regular file")
            raw = stream.read(16 * 1024 * 1024 + 1)
        if len(raw) > 16 * 1024 * 1024:
            result["reason"] = "文件超过 16 MiB 验证上限"
        elif hashlib.sha256(raw).hexdigest() != item["sha256"]:
            result["reason"] = "检查期间文件发生变化，未确认交付"
        else:
            result["reason"] = validate_content(raw, item["format"])
            result["valid"] = result["reason"] is None
    except Exception as exc:  # noqa: BLE001 - untrusted parser boundary; type-only metadata
        # Fixed type only: parser errors may contain untrusted content or paths.
        result["reason"] = "文件无法安全解析或已损坏：" + type(exc).__name__
    finally:
        if descriptor is not None:
            os.close(descriptor)
    return result


if __name__ == "__main__":
    print(json.dumps([validate_file(item) for item in json.loads(globals()["CANDIDATES"])]))
