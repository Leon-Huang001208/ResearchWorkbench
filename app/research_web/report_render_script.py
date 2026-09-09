"""Trusted sandbox program for report-project file projection.

This file is read by ``report_rendering.py`` and executed with network and host
filesystem access denied.  It never performs research or model calls.
"""

import html
import json
import re
import shutil
import zipfile
from pathlib import Path

FORMATS = list(globals().get("FORMATS", []))


def clean_text(value, limit=100_000):
    if value is None:
        return ""
    if not isinstance(value, str):
        value = json.dumps(value, ensure_ascii=False)
    return value.strip()[:limit]


def load_context():
    path = Path("inputs/report-workflow/run-context.json")
    if path.is_symlink() or not path.is_file() or path.stat().st_size > 1_000_000:
        return {}
    value = json.loads(path.read_text(encoding="utf-8"))
    return value if isinstance(value, dict) else {}


def payload_sections(value, context):
    sections = value.get("sections")
    if isinstance(sections, list):
        return sections, []
    if isinstance(sections, dict):
        return [
            {"id": key, "title": key, "content": content}
            for key, content in sections.items()
        ], []

    indexed = {}
    for key, family in value.items():
        if key != "blocks" and not key.endswith("_blocks"):
            continue
        if not isinstance(family, dict):
            continue
        for block_id, item in family.items():
            if not isinstance(item, dict):
                item = {"content": item}
            content = (
                item.get("content")
                or item.get("text")
                or item.get("summary")
                or item.get("value")
            )
            if not clean_text(content) and str(item.get("status", "")).startswith(
                "supported"
            ):
                content = {
                    name: field
                    for name, field in item.items()
                    if name not in {"title", "status"}
                }
            if clean_text(content):
                indexed[block_id] = {
                    "id": block_id,
                    "title": item.get("title") or block_id,
                    "content": content,
                }

    missing_by_id = {}
    for item in value.get("missing") or []:
        if not isinstance(item, dict):
            continue
        block_id = clean_text(item.get("block") or item.get("id"), 200)
        if block_id:
            missing_by_id[block_id] = item

    declared = context.get("blocks") if isinstance(context.get("blocks"), list) else []
    missing_blocks = []
    for block in declared:
        if not isinstance(block, dict):
            continue
        block_id = clean_text(block.get("id"), 200)
        if not block_id or not block.get("required", True):
            continue
        issue = missing_by_id.get(block_id)
        if block_id in indexed and issue is None:
            continue
        missing_blocks.append(block_id)
        issue = issue or {}
        reason = clean_text(issue.get("reason") or "模型未提供该必需区块")
        indexed[block_id] = {
            "id": block_id,
            "title": issue.get("title") or block.get("title") or block_id,
            "content": "数据缺失：" + reason,
        }

    order = {
        clean_text(block.get("id"), 200): index
        for index, block in enumerate(declared)
        if isinstance(block, dict) and block.get("id")
    }
    return sorted(
        indexed.values(), key=lambda item: order.get(item["id"], len(order))
    ), missing_blocks


def load_payload():
    path = Path("outputs/report_payload.json")
    if path.is_symlink() or not path.is_file() or path.stat().st_size > 2_000_000:
        raise ValueError("invalid report payload")
    value = json.loads(path.read_text(encoding="utf-8"))
    if not isinstance(value, dict):
        raise TypeError("payload must be an object")
    context = load_context()
    title = clean_text(value.get("title") or context.get("name"), 500)
    if not title:
        manifest = Path("inputs/report-workflow/manifest.yaml")
        if (
            manifest.is_file()
            and not manifest.is_symlink()
            and manifest.stat().st_size <= 1_000_000
        ):
            match = re.search(
                r"(?m)^name:\s*['\"]?([^'\"\r\n]+)",
                manifest.read_text(encoding="utf-8"),
            )
            title = clean_text(match.group(1) if match else None, 500)
    if not title:
        title = clean_text(value.get("workflow_id"), 500)
    if not title:
        raise ValueError("payload title is required")
    sections, missing_blocks = payload_sections(value, context)
    if not isinstance(sections, list) or not sections:
        raise ValueError("payload sections are required")
    normalized = []
    for index, item in enumerate(sections[:80]):
        if not isinstance(item, dict):
            continue
        content = clean_text(item.get("content"))
        if not content:
            continue
        normalized.append(
            {
                "id": clean_text(item.get("id") or f"section-{index + 1}", 200),
                "title": clean_text(
                    item.get("title") or item.get("id") or f"章节 {index + 1}", 500
                ),
                "content": content,
            }
        )
    if not normalized:
        raise ValueError("payload has no report content")
    value["title"] = title
    value["summary"] = clean_text(
        value.get("summary")
        or "基于锁定底稿与共享数据快照生成；数据缺失和口径限制已在正文中列示。"
    )
    period = (
        value.get("report_period")
        if isinstance(value.get("report_period"), dict)
        else {}
    )
    source_notes = (
        value.get("report_data_source_notes")
        if isinstance(value.get("report_data_source_notes"), dict)
        else {}
    )
    value["as_of"] = clean_text(
        value.get("as_of")
        or period.get("end_date_trading")
        or source_notes.get("last_available_trading_date")
        or source_notes.get("refreshed"),
        200,
    )
    value["sections"] = normalized
    value["sources"] = (
        value.get("sources") if isinstance(value.get("sources"), list) else []
    )
    missing = value.get("missing") if isinstance(value.get("missing"), list) else []
    value["missing"] = [
        clean_text(item.get("reason") or item.get("title") or item)
        if isinstance(item, dict)
        else clean_text(item)
        for item in missing
    ]
    value["missing_blocks"] = missing_blocks
    return value


def replacements(payload):
    rows = {
        "标题": payload["title"],
        "报告标题": payload["title"],
        "数据截止时间": payload["as_of"],
        "摘要": payload["summary"],
    }
    for item in payload["sections"]:
        rows[item["id"]] = item["content"]
        rows[item["title"]] = item["content"]
    return rows


def replace_paragraph(paragraph, values, replaced):
    text = paragraph.text
    if "{{" not in text:
        return
    updated = text
    for key, value in values.items():
        token = "{{" + key + "}}"
        if token in updated:
            updated = updated.replace(token, value)
            replaced.add(key)
    if updated != text:
        paragraph.text = updated


def render_docx(payload, values):
    from docx import Document

    templates = sorted(Path("inputs/report-workflow/templates").rglob("*.docx"))
    if not templates:
        templates = sorted(Path("inputs/report-project").rglob("*.docx"))
    document = Document(str(templates[0])) if templates else Document()
    replaced = set()
    for paragraph in document.paragraphs:
        replace_paragraph(paragraph, values, replaced)
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    replace_paragraph(paragraph, values, replaced)
    if not replaced:
        document.add_heading(payload["title"], level=0)
        if payload["as_of"]:
            document.add_paragraph("数据截止时间：" + payload["as_of"])
        if payload["summary"]:
            document.add_paragraph(payload["summary"])
        for item in payload["sections"]:
            document.add_heading(item["title"], level=1)
            document.add_paragraph(item["content"])
    if payload["missing"]:
        document.add_heading("数据缺失与限制", level=1)
        for item in payload["missing"]:
            document.add_paragraph(clean_text(item), style="List Bullet")
    output = Path("outputs/report.docx")
    document.save(output)
    return output


def render_xlsx(payload):
    context_path = Path("inputs/report-workflow/run-context.json")
    if context_path.is_file() and not context_path.is_symlink():
        context = json.loads(context_path.read_text(encoding="utf-8"))
        primary = clean_text(context.get("primary_workbook"), 240)
        if primary:
            logical = Path(primary)
            if (
                logical.is_absolute()
                or ".." in logical.parts
                or logical.suffix.lower() != ".xlsx"
            ):
                raise ValueError("invalid primary workbook")
            source = Path("inputs/report-workflow") / logical
            if source.is_symlink() or not source.is_file():
                raise ValueError("primary workbook is missing")
            output = Path("outputs/report.xlsx")
            shutil.copyfile(source, output)
            return output
    from openpyxl import Workbook

    book = Workbook()
    overview = book.active
    overview.title = "报告摘要"
    overview.append(["字段", "内容"])
    overview.append(["标题", payload["title"]])
    overview.append(["数据截止时间", payload["as_of"]])
    overview.append(["摘要", payload["summary"]])
    sections = book.create_sheet("章节底稿")
    sections.append(["章节ID", "章节", "正文"])
    for item in payload["sections"]:
        sections.append([item["id"], item["title"], item["content"]])
    sources = book.create_sheet("来源")
    sources.append(["来源", "链接", "截止时间"])
    for item in payload["sources"]:
        if isinstance(item, dict):
            sources.append(
                [
                    clean_text(item.get("title") or item.get("source")),
                    clean_text(item.get("url")),
                    clean_text(item.get("as_of")),
                ]
            )
        else:
            sources.append([clean_text(item), "", ""])
    missing = book.create_sheet("缺失项")
    missing.append(["缺失或限制"])
    for item in payload["missing"]:
        missing.append([clean_text(item)])
    output = Path("outputs/report.xlsx")
    book.save(output)
    return output


def render_html(payload):
    sections = "".join(
        f"<section><h2>{html.escape(item['title'])}</h2><p>{html.escape(item['content']).replace(chr(10), '<br>')}</p></section>"
        for item in payload["sections"]
    )
    sources = "".join(
        f"<li>{html.escape(clean_text(item.get('title') or item.get('source'))) if isinstance(item, dict) else html.escape(clean_text(item))}</li>"
        for item in payload["sources"]
    )
    missing = "".join(
        f"<li>{html.escape(clean_text(item))}</li>" for item in payload["missing"]
    )
    text = f"""<!doctype html><html lang=\"zh-CN\"><meta charset=\"utf-8\"><meta name=\"viewport\" content=\"width=device-width\"><title>{html.escape(payload["title"])}</title><style>body{{max-width:920px;margin:40px auto;padding:0 24px;font:16px/1.75 system-ui;color:#202124}}h1,h2{{line-height:1.3}}.meta{{color:#666}}section{{margin:32px 0}}@media print{{body{{margin:0}}}}</style><body><h1>{html.escape(payload["title"])}</h1><p class=\"meta\">数据截止时间：{html.escape(payload["as_of"] or "未提供")}</p><p>{html.escape(payload["summary"])}</p>{sections}<section><h2>来源</h2><ul>{sources or "<li>未提供</li>"}</ul></section><section><h2>数据缺失与限制</h2><ul>{missing or "<li>无已声明缺失项</li>"}</ul></section></body></html>"""
    output = Path("outputs/report.html")
    output.write_text(text, encoding="utf-8")
    return output


PPTX_PARAGRAPH = re.compile(r"(<a:p(?:\s[^>]*)?>)(.*?)(</a:p>)", re.DOTALL)
PPTX_TEXT = re.compile(r"(<a:t(?:\s[^>]*)?>)(.*?)(</a:t>)", re.DOTALL)
PPTX_PLACEHOLDER = re.compile(r"\{\{[^{}\r\n]{1,200}\}\}")


def replace_pptx_paragraph(body, values):
    nodes = list(PPTX_TEXT.finditer(body))
    if not nodes:
        return body
    texts = [html.unescape(node.group(2)) for node in nodes]
    replacements = {
        "{{" + key + "}}": value for key, value in values.items() if key
    }
    if not replacements:
        return body
    pattern = re.compile(
        "|".join(
            re.escape(token)
            for token in sorted(replacements, key=len, reverse=True)
        )
    )
    matches = list(pattern.finditer("".join(texts)))
    for match in reversed(matches):
        start = match.start()
        end = match.end()
        offset = 0
        start_index = end_index = None
        start_offset = end_offset = 0
        for index, text in enumerate(texts):
            boundary = offset + len(text)
            if start_index is None and start < boundary:
                start_index = index
                start_offset = start - offset
            if end_index is None and end <= boundary:
                end_index = index
                end_offset = end - offset
                break
            offset = boundary
        if start_index is None or end_index is None:
            continue
        replacement = replacements[match.group(0)]
        if start_index == end_index:
            text = texts[start_index]
            texts[start_index] = (
                text[:start_offset] + replacement + text[end_offset:]
            )
            continue
        texts[start_index] = texts[start_index][:start_offset] + replacement
        for index in range(start_index + 1, end_index):
            texts[index] = ""
        texts[end_index] = texts[end_index][end_offset:]

    pieces = []
    cursor = 0
    for node, text in zip(nodes, texts, strict=True):
        pieces.append(body[cursor : node.start()])
        pieces.append(node.group(1))
        pieces.append(html.escape(text, quote=False))
        pieces.append(node.group(3))
        cursor = node.end()
    pieces.append(body[cursor:])
    return "".join(pieces)


def replace_pptx_placeholders(xml, values):
    def replace_paragraph(match):
        return (
            match.group(1)
            + replace_pptx_paragraph(match.group(2), values)
            + match.group(3)
        )

    updated = PPTX_PARAGRAPH.sub(replace_paragraph, xml)
    unresolved = []
    for match in PPTX_PARAGRAPH.finditer(updated):
        visible = "".join(
            html.unescape(node.group(2)) for node in PPTX_TEXT.finditer(match.group(2))
        )
        unresolved.extend(PPTX_PLACEHOLDER.findall(visible))
    if unresolved:
        raise ValueError(
            "unresolved PPTX placeholders: " + ", ".join(sorted(set(unresolved)))
        )
    return updated


def render_pptx(payload, values):
    templates = sorted(Path("inputs/report-workflow/templates").rglob("*.pptx"))
    if not templates:
        templates = sorted(Path("inputs/report-project").rglob("*.pptx"))
    if not templates:
        raise ValueError("PPTX template is required")
    source = templates[0]
    output = Path("outputs/report.pptx")
    with zipfile.ZipFile(source) as package, zipfile.ZipFile(output, "w") as target:
        for info in package.infolist():
            raw = package.read(info.filename)
            if info.filename.startswith("ppt/slides/slide") and info.filename.endswith(
                ".xml"
            ):
                text = raw.decode("utf-8")
                text = replace_pptx_placeholders(text, values)
                raw = text.encode("utf-8")
            target.writestr(info, raw)
    return output


payload = load_payload()
values = replacements(payload)
writers = {
    "docx": lambda: render_docx(payload, values),
    "html": lambda: render_html(payload),
    "xlsx": lambda: render_xlsx(payload),
    "pptx": lambda: render_pptx(payload, values),
}
created = []
for requested in FORMATS:
    output = writers[requested]()
    if output.is_symlink() or not output.is_file() or output.stat().st_size == 0:
        raise ValueError("renderer did not produce a regular file")
    created.append(output.as_posix())
print(
    json.dumps(
        {
            "status": "completed",
            "files": created,
            "missing_blocks": payload.get("missing_blocks", []),
        },
        ensure_ascii=False,
    )
)
