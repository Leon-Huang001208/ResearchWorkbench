"""Offline research artifacts sharing one payload and an embedded Chinese font."""

import hashlib
import html
import json
import logging
import os
import re
import shutil
import stat
from pathlib import Path
from urllib.parse import urlsplit
from uuid import uuid4

log = logging.getLogger("research_web.report")
_REPORT_FORMATS = ("html", "docx", "pdf")


def _report_formats(formats):
    if (
        not isinstance(formats, (list, tuple))
        or not 1 <= len(formats) <= 3
        or any(not isinstance(kind, str) or kind not in _REPORT_FORMATS for kind in formats)
        or len(set(formats)) != len(formats)
    ):
        raise ValueError("research_report_formats_invalid")
    return tuple(kind for kind in _REPORT_FORMATS if kind in formats)


def required_report_formats(contract_or_question):
    """Recognize only an explicit, bounded format list from the frozen question."""
    question = (
        contract_or_question.get("question")
        if isinstance(contract_or_question, dict)
        else contract_or_question
    )
    if not isinstance(question, str) or len(question) > 100000:
        return _REPORT_FORMATS
    declarations = []
    pattern = re.compile(
        r"(?:只|仅)生成\s*(?P<formats>(?:HTML|DOCX|PDF)"
        r"(?:\s*(?:和|与|及|、|,|，|/|\+)\s*(?:HTML|DOCX|PDF)){0,2})"
        r"\s*(?=$|[。；;！!\n\r]|[，,](?!\s*[A-Za-z]))",
        re.IGNORECASE,
    )
    for match in pattern.finditer(question):
        if re.search(
            r"(?:不|不要|不能|不得|不应|并非|不是|无需|无须)\s*$", question[: match.start()]
        ):
            continue
        try:
            declarations.append(
                _report_formats(re.findall(r"html|docx|pdf", match["formats"].lower()))
            )
            if len(declarations) > 1:
                return _REPORT_FORMATS
        except ValueError:
            log.warning("research_report_format_declaration_rejected")
            return _REPORT_FORMATS
    if len(declarations) == 1:
        log.info("research_report_formats_frozen count=%d", len(declarations[0]))
        return declarations[0]
    return _REPORT_FORMATS


def source_id(url):
    return "web:" + hashlib.sha256(url.encode()).hexdigest()[:24]


def load_business_evidence(task_id, contract_sha256):
    """Read this turn's convenience snapshot; it is never a Host verdict."""
    try:
        if (
            not isinstance(task_id, str)
            or not re.fullmatch(r"[a-f0-9]{24}", task_id)
            or not isinstance(contract_sha256, str)
            or not re.fullmatch(r"[a-f0-9]{64}", contract_sha256)
        ):
            raise ValueError("research_evidence_snapshot_binding")
        descriptor = os.open(
            "resources/research_evidence.json",
            os.O_RDONLY | getattr(os, "O_NOFOLLOW", 0) | getattr(os, "O_NONBLOCK", 0),
        )
        with os.fdopen(descriptor, "rb") as stream:
            info = os.fstat(stream.fileno())
            if not stat.S_ISREG(info.st_mode) or info.st_nlink != 1 or info.st_size > 256000:
                raise ValueError("research_evidence_snapshot_file")
            raw = stream.read(256001)
        if len(raw) > 256000:
            raise ValueError("research_evidence_snapshot_limit")
        value = json.loads(raw)
        if (
            not isinstance(value, dict)
            or set(value)
            != {"schema_version", "task_id", "contract_sha256", "sources", "observations"}
            or value["schema_version"] != 1
            or value["task_id"] != task_id
            or value["contract_sha256"] != contract_sha256
            or not isinstance(value["sources"], list)
            or len(value["sources"]) > 40
            or not isinstance(value["observations"], list)
            or len(value["observations"]) > 200
        ):
            raise ValueError("research_evidence_snapshot_binding")
        log.info("research_evidence_snapshot_loaded observations=%d", len(value["observations"]))
        return value
    except (OSError, ValueError, TypeError, KeyError) as exc:
        log.warning("research_evidence_snapshot_rejected type=%s", type(exc).__name__)
        raise ValueError("research_evidence_snapshot_invalid") from exc


def calculation_evidence(observations, requirement_ids, *, methods_requirement_ids=()):
    """Serialize one calculation's evidence; only the independent Host can verify it.

    Pass the actual four calculation_observations from research_check_quality.
    This creates sections/claims/notes/coverage, never missing or quality status.
    Requirement mappings express the author's intent, not verified coverage.
    """
    metrics = ("total_return", "max_drawdown", "terminal_scenario_return", "row_count")

    def ids(value):
        return (
            isinstance(value, (list, tuple))
            and len(value) <= 40
            and all(isinstance(v, str) and re.fullmatch(r"[A-Za-z0-9_-]{1,80}", v) for v in value)
            and len(set(value)) == len(value)
        )

    try:
        if (
            not isinstance(observations, list)
            or len(observations) != 4
            or not isinstance(requirement_ids, dict)
            or not set(requirement_ids).issubset(metrics)
            or not all(ids(value) for value in requirement_ids.values())
            or not ids(methods_requirement_ids)
        ):
            raise ValueError("research_calculation_evidence_invalid")
        rows, scope, methods = {}, None, None
        for row in observations:
            if not isinstance(row, dict):
                raise TypeError("research_calculation_evidence_invalid")
            for key in (
                "id",
                "source_id",
                "calculation_id",
                "metric",
                "value",
                "period",
                "unit",
                "content",
                "input_sha256",
                "recipe_sha256",
            ):
                if not isinstance(row.get(key), str) or not 1 <= len(row[key]) <= (
                    8192 if key == "content" else 256
                ):
                    raise ValueError("research_calculation_evidence_invalid")
            current_scope = tuple(
                row[k] for k in ("source_id", "calculation_id", "input_sha256", "recipe_sha256")
            )
            lines = row["content"].split("\n", 1)
            if (
                row["metric"] not in metrics
                or row["metric"] in rows
                or not re.fullmatch(r"[A-Za-z][A-Za-z0-9_-]{0,31}", row["calculation_id"])
                or len(lines) != 2
                or not all(line.strip() for line in lines)
                or (scope is not None and scope != current_scope)
                or (methods is not None and methods != lines[1])
            ):
                raise ValueError("research_calculation_evidence_invalid")
            rows[row["metric"]] = row
            scope, methods = current_scope, lines[1]
        result = {"sections": [], "notes": [], "numeric_claims": [], "coverage": []}

        def section(identifier, heading, text, refs, requirements):
            result["sections"].append(
                {"id": identifier, "heading": heading, "text": text, "citations": list(refs)}
            )
            result["notes"].append({"stage_id": identifier, "text": text, "citations": list(refs)})
            result["coverage"].extend(
                {"requirement_id": rid, "section_ids": [identifier], "excerpt": text}
                for rid in requirements
            )

        prefix = "calc_" + rows[metrics[0]]["calculation_id"]
        for metric in metrics:
            row = rows[metric]
            text = row["content"].split("\n", 1)[0]
            identifier = prefix + "_" + metric
            section(identifier, metric, text, [row["source_id"]], requirement_ids.get(metric, []))
            result["numeric_claims"].append(
                {
                    "id": "claim_" + rows[metric]["calculation_id"] + "_" + metric,
                    "section_id": identifier,
                    "text": text,
                    "value": row["value"],
                    "period": row["period"],
                    "unit": row["unit"],
                    "operation": "identity",
                    "operands": [
                        {
                            "observation_id": row["id"],
                            "value": row["value"],
                            "period": row["period"],
                            "unit": row["unit"],
                            "quote": text,
                        }
                    ],
                }
            )
        section(
            prefix + "_methods",
            "计算方法与验证范围",
            methods,
            [rows[metrics[0]]["source_id"]],
            methods_requirement_ids,
        )
        log.info("research_calculation_evidence_assembled metrics=%d", len(metrics))
        return result
    except (ValueError, TypeError, KeyError) as exc:
        log.warning("research_calculation_evidence_failed type=%s", type(exc).__name__)
        raise ValueError("research_calculation_evidence_invalid") from exc


def _evidence_decimal(value):
    from decimal import Decimal

    if (
        isinstance(value, bool)
        or not isinstance(value, (str, int, float, Decimal))
        or len(str(value)) > 128
    ):
        raise ValueError("research_evidence_number_invalid")
    number = Decimal(str(value).replace(",", ""))
    if not number.is_finite() or abs(number.adjusted()) > 100:
        raise ValueError("research_evidence_number_invalid")
    return number


def source_references(value, sources):
    """Resolve IDs or an explicit ID/title/locator, never infer identity from prose."""
    try:
        if not isinstance(value, str) or not 1 <= len(value) <= 2000:
            raise ValueError("research_source_reference_invalid")
        value = value.strip()
        if " · " in value:
            parts = value.split(" · ")
            if (
                len(parts) != 3
                or not parts[2].strip()
                or len(parts[2]) > 256
                or "file:" in parts[2]
            ):
                raise ValueError("research_source_reference_invalid")
            row = canonical_source(sources, parts[0])
            if row["id"] != parts[0] or row["title"] != parts[1]:
                raise ValueError("research_source_reference_invalid")
            refs = [row["id"]]
        else:
            if re.fullmatch(r"file:[a-f0-9]{24}(?:\s*/\s*file:[a-f0-9]{24})+", value):
                refs = re.split(r"\s*/\s*", value)
            else:
                refs = (
                    json.loads(value) if value.startswith("[") else re.split(r"[;,，；\s]+", value)
                )
            if not isinstance(refs, list) or not refs or len(refs) > 20:
                raise ValueError("research_source_reference_invalid")
            for ref in refs:
                if not isinstance(ref, str) or canonical_source(sources, ref)["id"] != ref:
                    raise ValueError("research_source_reference_invalid")
        log.info("research_source_references_resolved count=%d", len(refs))
        return refs
    except (ValueError, TypeError, KeyError) as exc:
        log.warning("research_source_reference_rejected type=%s", type(exc).__name__)
        raise ValueError("research_source_reference_invalid") from exc


def canonical_source(sources, identity):
    """Resolve an exact Host ID/title; uploaded filenames are never source IDs."""
    try:
        if (
            not isinstance(sources, list)
            or not 1 <= len(sources) <= 200
            or not isinstance(identity, str)
        ):
            raise ValueError("research_evidence_sources_invalid")
        if any(
            not isinstance(s, dict)
            or not isinstance(s.get("id"), str)
            or not re.fullmatch(r"(?:file|web|dataset):[A-Za-z0-9_-]{1,80}", s["id"])
            or not isinstance(s.get("title"), str)
            for s in sources
        ):
            raise ValueError("research_evidence_sources_invalid")
        if len({s["id"] for s in sources}) != len(sources):
            raise ValueError("research_evidence_source_ambiguous")
        matches = [
            s
            for s in sources
            if (
                s["id"] == identity
                if identity.startswith(("file:", "web:", "dataset:"))
                else s["title"] == identity
            )
        ]
        if len(matches) != 1:
            raise ValueError("research_evidence_source_missing_or_ambiguous")
        result = {k: matches[0][k] for k in ("id", "title", "url", "as_of") if k in matches[0]}
        log.info("research_evidence_source_resolved")
        return result
    except (ValueError, TypeError, KeyError):
        log.warning("research_evidence_source_rejected")
        raise


def evidence_section(identifier, heading, text, citations, requirement_ids):
    """Copy author-nominated text/lineage, never certify it or create observations."""
    try:
        if isinstance(citations, list) and any(
            isinstance(v, str) and v.startswith("business_") for v in citations
        ):
            raise ValueError("research_evidence_observation_use_source_ids")
        if (
            not isinstance(identifier, str)
            or not re.fullmatch(r"[A-Za-z0-9_-]{1,80}", identifier)
            or not isinstance(heading, str)
            or not 1 <= len(heading) <= 200
            or not isinstance(text, str)
            or not 1 <= len(text) <= 10000
        ):
            raise ValueError("research_evidence_section_invalid")
        for values, pattern in (
            (citations, r"(?:file|web|dataset):[A-Za-z0-9_-]{1,80}"),
            (requirement_ids, r"[A-Za-z0-9_-]{1,80}"),
        ):
            if (
                not isinstance(values, list)
                or not 1 <= len(values) <= 100
                or any(not isinstance(v, str) or not re.fullmatch(pattern, v) for v in values)
            ):
                raise ValueError("research_evidence_lineage_invalid")
        citations = list(dict.fromkeys(citations))
        requirement_ids = list(dict.fromkeys(requirement_ids))
        return {
            "numeric_claims": [],
            "sections": [
                {"id": identifier, "heading": heading, "text": text, "citations": list(citations)}
            ],
            "notes": [{"stage_id": identifier, "text": text, "citations": list(citations)}],
            "coverage": [
                {"requirement_id": rid, "section_ids": [identifier], "excerpt": text}
                for rid in requirement_ids
            ],
        }
    except (ValueError, TypeError):
        log.warning("research_evidence_section_rejected")
        raise


def build_numeric_claim(identifier, section_id, text, period, unit, operation, operands):
    """Preflight exact arithmetic and serialization; Host still verifies source bytes."""
    from decimal import Context, InvalidOperation, localcontext

    try:
        if not all(
            isinstance(v, str) and re.fullmatch(r"[A-Za-z0-9_-]{1,80}", v)
            for v in (identifier, section_id)
        ) or not all(isinstance(v, str) and 1 <= len(v) <= 10000 for v in (text, period, unit)):
            raise ValueError("research_evidence_claim_invalid")
        arities = {
            "identity": (1,),
            "difference": (2,),
            "ratio": (2,),
            "percent_change": (2,),
            "sum": tuple(range(1, 51)),
        }
        if (
            operation not in arities
            or not isinstance(operands, list)
            or len(operands) not in arities[operation]
        ):
            raise ValueError("research_evidence_claim_arity")
        number_pattern = re.compile(r"(?<![\w])[-+]?\d+(?:,\d{3})*(?:\.\d+)?|[-+]?\d+(?:\.\d+)?")
        values = []
        for operand in operands:
            if (
                not isinstance(operand, dict)
                or set(operand) != {"observation_id", "value", "quote", "period", "unit"}
                or not isinstance(operand["observation_id"], str)
                or not re.fullmatch(
                    r"(?:[A-Za-z0-9_-]{1,80}|(?:file|web|dataset):[A-Za-z0-9_-]{1,80})",
                    operand["observation_id"],
                )
            ):
                raise ValueError("research_evidence_operand_invalid")
            if (
                not all(
                    isinstance(operand[k], str) and 1 <= len(operand[k]) <= 10000
                    for k in ("quote", "period", "unit")
                )
                or operand["period"] not in operand["quote"]
            ):
                raise ValueError("research_evidence_operand_context")
            quantity = list(
                number_pattern.finditer(operand["quote"].replace(operand["period"], "", 1))
            )
            value = _evidence_decimal(operand["value"])
            context = operand["quote"].replace(operand["period"], "", 1)
            if (
                len(quantity) != 1
                or _evidence_decimal(quantity[0].group()) != value
                or not re.match(
                    r"\s*" + re.escape(operand["unit"]) + r"(?=$|[\s,，;；。.!！？?、)\]】])",
                    context[quantity[0].end() :],
                )
            ):
                raise ValueError("research_evidence_operand_quantity")
            values.append(value)
        units = {o["unit"] for o in operands}
        if len(units) != 1 or period != operands[0]["period"]:
            raise ValueError("research_evidence_operand_units_or_period")
        if operation in {"identity", "difference", "sum"} and (
            unit not in units or any(o["period"] != period for o in operands)
        ):
            raise ValueError("research_evidence_operand_units_or_period")
        if (
            operation == "ratio"
            and unit not in {"ratio", "倍", "%"}
            or operation == "percent_change"
            and unit != "%"
        ):
            raise ValueError("research_evidence_claim_unit")
        with localcontext(Context(prec=28)):
            if operation == "identity":
                value = values[0]
            elif operation == "sum":
                value = sum(values)
            elif operation == "difference":
                value = values[0] - values[1]
            elif operation == "ratio":
                value = values[0] / values[1] * (100 if unit == "%" else 1)
            else:
                value = (values[0] - values[1]) / values[1] * 100
        if unit not in text or [
            _evidence_decimal(m.group())
            for m in number_pattern.finditer(text.replace(period, "", 1))
        ] != [value]:
            raise ValueError("research_evidence_claim_text_or_rounding")
        log.info("research_evidence_claim_built operation=%s", operation)
        return {
            "id": identifier,
            "section_id": section_id,
            "text": text,
            "value": format(value, "f"),
            "period": period,
            "unit": unit,
            "operation": operation,
            "operands": [
                {**operand, "value": format(number, "f")}
                for operand, number in zip(operands, values)
            ],
        }
    except (ValueError, TypeError, KeyError, ArithmeticError, InvalidOperation) as exc:
        log.warning("research_evidence_claim_rejected type=%s", type(exc).__name__)
        raise ValueError("research_evidence_claim_invalid") from exc


def reconcile_amounts(reported_total, components):
    """Exact reported-total minus all provided signed components, not inferred gaps."""
    from decimal import Context, Inexact, InvalidOperation, localcontext

    try:
        if not isinstance(components, list) or not 1 <= len(components) <= 1000:
            raise ValueError("research_evidence_components_invalid")
        values = [_evidence_decimal(v) for v in components]
        with localcontext(Context(prec=512)) as context:
            context.traps[Inexact] = True
            total = sum(values)
            gap = _evidence_decimal(reported_total) - total
        log.info("research_evidence_reconciliation_calculated components=%d", len(values))
        return {"component_sum": format(total, "f"), "difference": format(gap, "f")}
    except (ValueError, TypeError, ArithmeticError, InvalidOperation):
        log.warning("research_evidence_reconciliation_rejected")
        raise


def rate_difference(before_pct, after_pct):
    """Signed cut magnitude: percentage-point change and basis-point conversion."""
    from decimal import Context, Inexact, InvalidOperation, localcontext

    try:
        with localcontext(Context(prec=512)) as context:
            context.traps[Inexact] = True
            difference = _evidence_decimal(before_pct) - _evidence_decimal(after_pct)
            bps = difference * 100
        log.info("research_evidence_rate_difference_calculated")
        return {"percentage_points": format(difference, "f"), "basis_points": format(bps, "f")}
    except (ValueError, TypeError, ArithmeticError, InvalidOperation):
        log.warning("research_evidence_rate_difference_rejected")
        raise


def numeric_summary(rows, expected, *, decimal_places=None):
    """Independent acceptance oracle; explicit display rounding never changes Host gates."""
    from decimal import ROUND_HALF_UP, InvalidOperation

    try:
        if not isinstance(expected, dict) or not 1 <= len(expected) <= 256:
            raise ValueError("research_summary_expected_invalid")
        places = {} if decimal_places is None else decimal_places
        if (
            not isinstance(places, dict)
            or not set(places).issubset(expected)
            or any(
                isinstance(p, bool) or not isinstance(p, int) or not 0 <= p <= 12
                for p in places.values()
            )
        ):
            raise ValueError("research_summary_precision_invalid")
        iterator = iter(rows)
        header = tuple(next(iterator, ()))
        if header[:2] != ("metric", "value") or any(v is not None for v in header[2:]):
            raise ValueError("research_summary_header_invalid")
        actual = {}
        for row in iterator:
            if len(row) < 2 or any(v is not None for v in row[2:]):
                raise ValueError("research_summary_columns_invalid")
            key, value = row[:2]
            if key is None and value is None:
                break
            if (
                not isinstance(key, str)
                or not 1 <= len(key) <= 80
                or key in actual
                or len(actual) >= 256
                or isinstance(value, bool)
                or not isinstance(value, (float, int))
            ):
                raise ValueError("research_summary_value_invalid")
            _evidence_decimal(value)
            actual[key] = value
        if not set(expected).issubset(actual):
            raise ValueError("research_summary_metric_missing")
        for key, value in expected.items():
            wanted = _evidence_decimal(value)
            if _evidence_decimal(actual[key]) == wanted:
                continue
            if key in places:
                from decimal import Decimal

                wanted = wanted.quantize(Decimal(1).scaleb(-places[key]), rounding=ROUND_HALF_UP)
            if _evidence_decimal(actual[key]) != wanted:
                raise ValueError("research_summary_numeric_mismatch")
        log.info("research_summary_verified metrics=%d", len(expected))
        return actual
    except (ValueError, TypeError, KeyError, ArithmeticError, InvalidOperation) as exc:
        log.warning("research_summary_rejected type=%s", type(exc).__name__)
        raise ValueError("research_summary_invalid") from exc


def write_summary_workbook(path, metrics, *, decimal_places=None, sheets=None):
    """Write the fixed numeric table and separate details atomically; never certify facts."""
    from decimal import ROUND_HALF_UP, Context, Decimal, localcontext

    from openpyxl import Workbook, load_workbook

    try:
        from research_helpers import safe_cell
    except ImportError:
        from .research_helpers import safe_cell

    temporary = None
    book = None
    try:
        places = {} if decimal_places is None else decimal_places
        details = {} if sheets is None else sheets
        if (
            not isinstance(metrics, dict)
            or not 1 <= len(metrics) <= 256
            or not isinstance(places, dict)
            or not set(places).issubset(metrics)
            or any(
                isinstance(p, bool) or not isinstance(p, int) or not 0 <= p <= 12
                for p in places.values()
            )
            or not isinstance(details, dict)
            or len(details) > 12
        ):
            raise ValueError("research_summary_workbook_invalid")
        path = Path(path)
        if path.suffix.lower() != ".xlsx":
            raise ValueError("research_summary_workbook_path")
        if path.is_symlink() or (
            path.exists() and (not path.is_file() or path.stat().st_nlink != 1)
        ):
            raise ValueError("research_summary_workbook_path")
        book = Workbook()
        summary = book.active
        summary.title = "Summary"
        summary.append(["metric", "value"])
        precision = []
        for key, value in metrics.items():
            if not isinstance(key, str) or not re.fullmatch(r"[A-Za-z0-9_-]{1,80}", key):
                raise ValueError("research_summary_workbook_metric")
            number = _evidence_decimal(value)
            with localcontext(Context(prec=512)):
                rounded = (
                    number.quantize(Decimal(1).scaleb(-places[key]), rounding=ROUND_HALF_UP)
                    if key in places
                    else number
                )
            numeric = int(rounded) if rounded == rounded.to_integral_value() else float(rounded)
            _evidence_decimal(numeric)
            summary.append([key, numeric])
            if key in places:
                summary.cell(summary.max_row, 2).number_format = "0" + (
                    "." + "0" * places[key] if places[key] else ""
                )
            precision.append([key, format(number, "f")])
        # Materialize a fully empty separator, followed by a second header so it survives save.
        summary.append([None, None])
        summary.append(["details", "See separate worksheets"])
        book.create_sheet("ExactValues").append(["metric", "exact_decimal"])
        for row in precision:
            book["ExactValues"].append(row)
        names = {"summary", "exactvalues"}
        count = 0
        for name, rows in details.items():
            if (
                not isinstance(name, str)
                or not 1 <= len(name) <= 31
                or re.search(r"[\\/*?:\[\]]", name)
                or name.lower() in names
                or not isinstance(rows, (list, tuple))
                or len(rows) > 5000
            ):
                raise ValueError("research_summary_workbook_sheet")
            names.add(name.lower())
            sheet = book.create_sheet(name)
            for row in rows:
                if not isinstance(row, (list, tuple)) or len(row) > 40:
                    raise ValueError("research_summary_workbook_row")
                count += len(row)
                if count > 40000:
                    raise ValueError("research_summary_workbook_limit")
                cells = []
                for value in row:
                    if isinstance(value, (dict, list, tuple)):
                        value = json.dumps(value, ensure_ascii=False, allow_nan=False)
                    elif isinstance(value, Decimal):
                        value = format(_evidence_decimal(value), "f")
                    elif isinstance(value, (int, float)) and not isinstance(value, bool):
                        _evidence_decimal(value)
                    if not isinstance(value, (str, int, float, bool, type(None))) or (
                        isinstance(value, str) and len(value) > 32000
                    ):
                        raise ValueError("research_summary_workbook_cell")
                    cells.append(safe_cell(value))
                sheet.append(cells)
        temporary = path.with_name("." + path.name + "." + uuid4().hex + ".xlsx")
        book.save(temporary)
        book.close()
        book = load_workbook(temporary, read_only=True, data_only=False)
        numeric_summary(book["Summary"].values, metrics, decimal_places=places)
        book.close()
        os.replace(temporary, path)
        log.info("research_summary_workbook_written metrics=%d", len(metrics))
        return path
    except (OSError, ValueError, TypeError, ArithmeticError):
        log.warning("research_summary_workbook_rejected")
        raise
    finally:
        if book is not None:
            book.close()
        if temporary is not None and temporary.exists():
            temporary.unlink()


def select_business_evidence(bundle, requirement_ids):
    """Select exact records and their mappings together, rejecting ambiguous metrics."""
    try:
        if not isinstance(bundle, dict) or not isinstance(bundle.get("observations"), list):
            raise TypeError("research_business_selection_invalid")
        if not isinstance(requirement_ids, dict) or not 1 <= len(requirement_ids) <= 40:
            raise ValueError("research_business_selection_invalid")
        selected, mapping = [], {}
        for key, requirements in requirement_ids.items():
            matches = [
                row
                for row in bundle["observations"]
                if isinstance(row, dict)
                and key
                in (
                    row.get("id"),
                    row.get("metric"),
                    str(row.get("metric")) + "@" + str(row.get("period")),
                )
            ]
            if len(matches) != 1 or matches[0].get("id") in mapping:
                raise ValueError("research_business_selection_missing_or_ambiguous")
            row = matches[0]
            selected.append(row)
            mapping[row["id"]] = requirements
        return business_evidence(selected, mapping)
    except (ValueError, TypeError, KeyError):
        log.warning("research_business_selection_rejected")
        raise


def business_evidence(observations, requirement_ids):
    """Serialize selected actual Host business observations, never grant them trust."""
    try:
        if (
            not isinstance(observations, list)
            or not 1 <= len(observations) <= 40
            or not isinstance(requirement_ids, dict)
        ):
            raise ValueError("research_business_evidence_invalid")
        result = {"sections": [], "notes": [], "coverage": [], "numeric_claims": []}
        seen = set()
        used = set()
        for row in observations:
            if (
                not isinstance(row, dict)
                or not isinstance(row.get("id"), str)
                or not re.fullmatch(r"business_[a-f0-9]{32}", row["id"])
                or row["id"] in seen
                or row.get("kind") not in {"business_numeric", "business_quote"}
                or not isinstance(row.get("content"), str)
                or hashlib.sha256(row["content"].encode()).hexdigest() != row.get("sha256")
            ):
                raise ValueError("research_business_evidence_invalid")
            seen.add(row["id"])
            key = row["id"] if row["id"] in requirement_ids else row.get("metric")
            used.add(key)
            requirements = requirement_ids.get(key)
            refs = row.get("source_ids")
            if not isinstance(refs, list) or row.get("source_id") not in refs:
                raise ValueError("research_business_evidence_invalid")
            heading = row.get("heading", row.get("metric", ""))
            if row.get("period"):
                heading += " (" + row["period"] + ")"
            blocks = evidence_section(row["id"], heading, row["content"], refs, requirements)
            for name in blocks:
                result[name].extend(blocks[name])
            if row["kind"] == "business_numeric":
                quote = row["content"].split("\n", 1)[0]
                operand = {
                    "observation_id": row["id"],
                    "value": row["value"],
                    "period": row["period"],
                    "unit": row["unit"],
                    "quote": quote,
                }
                result["numeric_claims"].append(
                    build_numeric_claim(
                        "claim_" + row["id"],
                        row["id"],
                        quote,
                        row["period"],
                        row["unit"],
                        "identity",
                        [operand],
                    )
                )
            elif re.search(r"\d", row["content"]):
                raise ValueError("research_business_quote_needs_numeric_evidence")
        if set(requirement_ids) != used:
            raise ValueError("research_business_evidence_mapping_unused")
        log.info("research_business_evidence_assembled observations=%d", len(seen))
        return result
    except (ValueError, TypeError, KeyError, ArithmeticError) as exc:
        log.warning("research_business_evidence_rejected type=%s", type(exc).__name__)
        raise ValueError("research_business_evidence_invalid") from exc


def validate_calculation_recipes(recipes):
    """Validate data-only calculation requests; never accept paths or executable code."""
    from decimal import Context, Decimal, InvalidOperation, localcontext

    if not isinstance(recipes, list) or len(recipes) > 4:
        raise ValueError("research_calculation_recipes_invalid")
    ids = set()
    for row in recipes:
        if (
            not isinstance(row, dict)
            or set(row)
            != {"id", "source_id", "source_sha256", "operation", "weights", "terminal_shocks"}
            or not isinstance(row["id"], str)
            or not re.fullmatch(r"[a-zA-Z][a-zA-Z0-9_-]{0,31}", row["id"])
            or row["id"] in ids
            or row["operation"] != "buy_and_hold"
            or not isinstance(row["source_id"], str)
            or not re.fullmatch(r"file:[a-f0-9]{24}", row["source_id"])
            or not isinstance(row["source_sha256"], str)
            or not re.fullmatch(r"[a-f0-9]{64}", row["source_sha256"])
            or not isinstance(row["weights"], dict)
            or not 1 <= len(row["weights"]) <= 8
            or not isinstance(row["terminal_shocks"], dict)
            or not set(row["terminal_shocks"]).issubset(row["weights"])
        ):
            raise ValueError("research_calculation_recipe_invalid")
        ids.add(row["id"])
        try:
            weights = []
            for kind in ("weights", "terminal_shocks"):
                for name, value in row[kind].items():
                    if (
                        not isinstance(name, str)
                        or not 1 <= len(name) <= 80
                        or isinstance(value, bool)
                        or not isinstance(value, (str, int, float))
                        or len(str(value)) > 64
                    ):
                        raise ValueError("research_calculation_number_invalid")
                    if not re.fullmatch(
                        r"[+-]?(?:[0-9]+(?:\.[0-9]*)?|\.[0-9]+)(?:[eE][+-]?[0-9]{1,3})?", str(value)
                    ):
                        raise ValueError("research_calculation_number_invalid")
                    number = Decimal(str(value))
                    if not number.is_finite() or abs(number.as_tuple().exponent) > 30:
                        raise ValueError("research_calculation_number_invalid")
                    if kind == "weights":
                        if not 0 <= number <= 1:
                            raise ValueError("research_calculation_weight_invalid")
                        weights.append(number)
                    elif not -1 <= number <= 1:
                        raise ValueError("research_calculation_shock_invalid")
            with localcontext(Context(prec=128)):
                if sum(weights) != Decimal(1):
                    raise ValueError("research_calculation_weights_invalid")
        except (InvalidOperation, ArithmeticError) as exc:
            raise ValueError("research_calculation_number_invalid") from exc
    return recipes


def validate_payload(value):
    if not isinstance(value, dict) or not re.fullmatch(
        r"[a-f0-9]{24}", str(value.get("task_id", ""))
    ):
        raise ValueError("research_report_task_invalid")
    if not isinstance(value.get("title"), str) or not 1 <= len(value["title"]) <= 200:
        raise ValueError("research_report_title_invalid")
    if "report_formats" in value:
        _report_formats(value["report_formats"])
    sections, sources = value.get("sections"), value.get("sources", [])
    if not isinstance(sections, list) or not 1 <= len(sections) <= 50:
        raise ValueError("research_report_sections_invalid")
    if not isinstance(sources, list) or len(sources) > 200:
        raise ValueError("research_report_sources_invalid")
    ids = set()
    for item in sources:
        if not isinstance(item, dict) or not re.fullmatch(
            r"(?:web|file|dataset):[a-zA-Z0-9_-]{1,80}", str(item.get("id", ""))
        ):
            raise ValueError("research_report_source_invalid")
        if item["id"] in ids or not isinstance(item.get("title"), str) or len(item["title"]) > 1000:
            raise ValueError("research_report_source_duplicate_or_title")
        ids.add(item["id"])
        url = item.get("url")
        if url:
            if not isinstance(url, str) or len(url) > 4000 or any(ord(c) < 33 for c in url):
                raise ValueError("research_report_url_invalid")
            parts = urlsplit(url)
            if (
                parts.scheme not in {"http", "https"}
                or not parts.hostname
                or parts.username
                or parts.password
            ):
                raise ValueError("research_report_url_invalid")
    stage_ids = set()
    for item in sections:
        if not isinstance(item, dict) or any(
            not isinstance(item.get(k), str) for k in ("heading", "text")
        ):
            raise ValueError("research_report_section_invalid")
        if not item["heading"].strip() or len(item["heading"]) > 200 or len(item["text"]) > 100000:
            raise ValueError("research_report_section_size")
        stage = item.get("id")
        if stage is not None:
            if (
                not isinstance(stage, str)
                or not re.fullmatch(r"[a-zA-Z0-9_-]{1,80}", stage)
                or stage in stage_ids
            ):
                raise ValueError("research_report_stage_invalid")
            stage_ids.add(stage)
        refs = item.get("citations", [])
        if (
            not isinstance(refs, list)
            or len(refs) > 200
            or any(not isinstance(ref, str) or ref not in ids for ref in refs)
        ):
            raise ValueError("research_report_citation_unknown")
    for key in ("required_sections", "missing"):
        values = value.get(key, [])
        if (
            not isinstance(values, list)
            or len(values) > 100
            or any(not isinstance(s, str) or len(s) > 10000 for s in values)
        ):
            raise ValueError("research_report_list_invalid")
    notes = value.get("notes", [])
    if not isinstance(notes, list) or len(notes) > 100:
        raise ValueError("research_report_notes_invalid")
    for note in notes:
        if (
            not isinstance(note, dict)
            or not isinstance(note.get("text"), str)
            or len(note["text"]) > 100000
        ):
            raise ValueError("research_report_note_invalid")
        refs = note.get("citations", [])
        if note.get("stage_id") is not None and note["stage_id"] not in stage_ids:
            raise ValueError("research_report_note_stage_unknown")
        if not isinstance(refs, list) or any(
            not isinstance(ref, str) or ref not in ids for ref in refs
        ):
            raise ValueError("research_report_note_citation_unknown")
    coverage = value.get("coverage", [])
    if not isinstance(coverage, list) or len(coverage) > 100:
        raise ValueError("research_report_coverage_invalid")
    for row in coverage:
        if (
            not isinstance(row, dict)
            or not isinstance(row.get("requirement_id"), str)
            or not re.fullmatch(r"[a-zA-Z0-9_-]{1,80}", row["requirement_id"])
            or not isinstance(row.get("section_ids"), list)
            or len(row["section_ids"]) > 50
            or any(not isinstance(s, str) or s not in stage_ids for s in row["section_ids"])
            or not isinstance(row.get("excerpt"), str)
            or not 1 <= len(row["excerpt"]) <= 10000
        ):
            raise ValueError("research_report_coverage_invalid")
    validate_calculation_recipes(value.get("calculations", []))
    claims = value.get("numeric_claims", [])
    if not isinstance(claims, list) or len(claims) > 200:
        raise ValueError("research_report_numeric_claim_invalid")
    claim_ids = set()
    for claim in claims:
        if (
            not isinstance(claim, dict)
            or not isinstance(claim.get("id"), str)
            or not re.fullmatch(r"[a-zA-Z0-9_-]{1,80}", claim["id"])
            or claim["id"] in claim_ids
            or claim.get("section_id") not in stage_ids
            or any(
                not isinstance(claim.get(key), str) or not 1 <= len(claim[key]) <= 10000
                for key in ("text", "unit", "period")
            )
            or isinstance(claim.get("value"), bool)
            or not isinstance(claim.get("value"), (str, int, float))
            or claim.get("operation", "identity")
            not in {"identity", "sum", "difference", "ratio", "percent_change"}
            or not isinstance(claim.get("operands"), list)
            or not 1 <= len(claim["operands"]) <= 50
        ):
            raise ValueError("research_report_numeric_claim_invalid")
        claim_ids.add(claim["id"])
        for operand in claim["operands"]:
            if (
                not isinstance(operand, dict)
                or any(
                    not isinstance(operand.get(key), str) or not 1 <= len(operand[key]) <= 10000
                    for key in ("observation_id", "quote", "period", "unit")
                )
                or isinstance(operand.get("value"), bool)
                or not isinstance(operand.get("value"), (str, int, float))
            ):
                raise ValueError("research_report_numeric_claim_operand_invalid")
    if len(json.dumps(value, ensure_ascii=False).encode()) > 2_000_000:
        raise ValueError("research_report_size")
    return value


def _bounded_bytes(path, maximum=16 * 1024 * 1024):
    if path.is_symlink() or not path.is_file():
        raise ValueError("research_report_archive_alias_or_missing")
    flags = os.O_RDONLY | getattr(os, "O_NOFOLLOW", 0)
    fd = os.open(path, flags)
    with os.fdopen(fd, "rb") as stream:
        opened = os.fstat(stream.fileno())
        original = path.lstat()
        if not stat.S_ISREG(opened.st_mode) or (opened.st_dev, opened.st_ino) != (
            original.st_dev,
            original.st_ino,
        ):
            raise ValueError("research_report_archive_alias")
        body = stream.read(maximum + 1)
    if len(body) > maximum:
        raise ValueError("research_report_archive_size")
    return body


def _version_read(directory):
    if directory.is_symlink() or not directory.is_dir():
        raise ValueError("research_report_archive_alias")
    version = json.loads(_bounded_bytes(directory / "version.json", 2_000_000))
    if not isinstance(version, dict):
        raise TypeError("research_report_archive_changed")
    rows = version.get("files", [])
    if (
        not isinstance(rows, list)
        or not 1 <= len(rows) <= 4
        or any(not isinstance(row, dict) for row in rows)
    ):
        raise ValueError("research_report_archive_changed")
    for row in rows:
        name = row.get("stored_name")
        if not isinstance(name, str) or not re.fullmatch(r"[a-zA-Z0-9._-]{1,240}", name):
            raise ValueError("research_report_archive_changed")
        if hashlib.sha256(_bounded_bytes(directory / name)).hexdigest() != row.get("sha256"):
            raise ValueError("research_report_archive_changed")
    manifest = next((row for row in rows if row.get("name") == "research-report.json"), None)
    if not manifest:
        raise ValueError("research_report_archive_changed")
    raw = _bounded_bytes(directory / manifest["stored_name"], 2_000_000)
    report = validate_payload(json.loads(raw))
    archived_hashes = {row.get("name"): row.get("sha256") for row in rows}
    artifacts = report.get("artifacts", [])
    if not isinstance(artifacts, list) or any(not isinstance(row, dict) for row in artifacts):
        raise ValueError("research_report_archive_changed")
    missing = [row.get("name") for row in artifacts if row.get("name") not in archived_hashes]
    missing_formats = sorted(
        set(report.get("report_formats", _REPORT_FORMATS))
        - {name.rsplit(".", 1)[-1] for name in archived_hashes if name != "research-report.json"}
    )
    snapshot = _snapshot_digest(archived_hashes, missing, missing_formats)
    identity = report["task_id"] + "-" + hashlib.sha256(raw).hexdigest()[:24] + "-" + snapshot[:12]
    if (
        report["task_id"] != version.get("task_id")
        or version.get("version_id") != identity
        or version.get("directory") != "report-versions/" + identity
        or (not directory.name.startswith(".pending-") and directory.name != identity)
        or version.get("missing_artifacts") != missing
        or version.get("missing_formats") != missing_formats
        or any(
            archived_hashes.get(row.get("name")) != row.get("sha256")
            for row in artifacts
            if row.get("name") not in missing
        )
    ):
        raise ValueError("research_report_archive_changed")
    return version


def _snapshot_digest(hashes, missing, formats):
    return hashlib.sha256(
        json.dumps(
            {"hashes": hashes, "missing_artifacts": missing, "missing_formats": formats},
            sort_keys=True,
            separators=(",", ":"),
        ).encode()
    ).hexdigest()


def archive_report_version(output, expected_task_id=None):
    """Verify and atomically retain the current payload and all published bytes."""
    output = Path(output)
    if output.is_symlink():
        raise ValueError("research_report_archive_alias")
    current = output / "research-report.json"
    if not current.exists() and not current.is_symlink():
        return None
    try:
        raw = _bounded_bytes(current, 2_000_000)
        payload = validate_payload(json.loads(raw))
        if expected_task_id is not None and payload["task_id"] != expected_task_id:
            raise ValueError("research_report_archive_task_mismatch")
        artifacts = payload.get("artifacts", [])
        if (
            not isinstance(artifacts, list)
            or len(artifacts) > 3
            or any(
                not isinstance(row, dict) or not isinstance(row.get("name"), str)
                for row in artifacts
            )
            or len({row["name"].rsplit(".", 1)[-1] for row in artifacts}) != len(artifacts)
            or not {row["name"].rsplit(".", 1)[-1] for row in artifacts}.issubset(
                {"html", "docx", "pdf"}
            )
        ):
            raise ValueError("research_report_archive_artifact_invalid")
        bodies = {"research-report.json": raw}
        missing = []
        for artifact in artifacts:
            name = artifact.get("name")
            if not isinstance(name, str) or not re.fullmatch(
                r"[a-zA-Z0-9][a-zA-Z0-9._-]{0,149}", name
            ):
                raise ValueError("research_report_archive_artifact_invalid")
            if not (output / name).exists() and not (output / name).is_symlink():
                missing.append(name)
                continue
            body = _bounded_bytes(output / name)
            if hashlib.sha256(body).hexdigest() != artifact.get("sha256"):
                raise ValueError("research_report_archive_artifact_changed")
            bodies[name] = body
        digest = hashlib.sha256(raw).hexdigest()
        hashes = {name: hashlib.sha256(body).hexdigest() for name, body in bodies.items()}
        missing_formats = sorted(
            set(payload.get("report_formats", _REPORT_FORMATS))
            - {name.rsplit(".", 1)[-1] for name in bodies if name != "research-report.json"}
        )
        snapshot = _snapshot_digest(hashes, missing, missing_formats)
        version_id = payload["task_id"] + "-" + digest[:24] + "-" + snapshot[:12]
        archive = output / "report-versions"
        if archive.is_symlink():
            raise ValueError("research_report_archive_alias")
        archive.mkdir(exist_ok=True)
        target = archive / version_id
        version = {
            "task_id": payload["task_id"],
            "version_id": version_id,
            "directory": "report-versions/" + version_id,
            "missing_artifacts": missing,
            "missing_formats": missing_formats,
            "files": [
                {
                    "name": name,
                    "stored_name": version_id + "-" + name,
                    "sha256": hashlib.sha256(body).hexdigest(),
                }
                for name, body in bodies.items()
            ],
        }
        if target.exists() or target.is_symlink():
            existing = _version_read(target)
            if existing != version:
                raise ValueError("research_report_archive_changed")
            return existing
        # mkdtemp resolves relative directories through getcwd on Python 3.12.
        # Keep the existing session-relative capability without requiring it.
        pending = archive / (".pending-" + uuid4().hex)
        pending.mkdir(mode=0o700)
        try:
            for row in version["files"]:
                destination = pending / row["stored_name"]
                with destination.open("xb") as stream:
                    stream.write(bodies[row["name"]])
                    stream.flush()
                    os.fsync(stream.fileno())
                destination.chmod(0o444)
            (pending / "version.json").write_text(json.dumps(version), encoding="utf-8")
            (pending / "version.json").chmod(0o444)
            _version_read(pending)
            try:
                pending.rename(target)
            except OSError:
                if not target.exists() or _version_read(target) != version:
                    raise
        finally:
            if pending.exists():
                shutil.rmtree(pending)
        log.info("research_report_version_archived files=%d", len(version["files"]))
        return version
    except OSError as exc:
        log.warning("research_report_archive_failed type=%s", type(exc).__name__)
        raise ValueError("research_report_archive_io_failed") from None
    except (ValueError, TypeError, KeyError, RecursionError) as exc:
        log.warning("research_report_archive_failed type=%s", type(exc).__name__)
        raise


def list_report_versions(output, limit=20):
    """Project bounded verified versions; never treat old quality claims as trusted."""
    output = Path(output)
    archive = output / "report-versions"
    if archive.is_symlink() or output.is_symlink():
        raise ValueError("research_report_archive_alias")
    if not archive.exists():
        return []
    if not isinstance(limit, int) or not 1 <= limit <= 20:
        raise ValueError("research_report_archive_limit_invalid")
    versions = []
    for path in sorted(archive.iterdir(), key=lambda item: item.name, reverse=True):
        if path.name.startswith("."):
            continue
        try:
            version = _version_read(path)
            manifest = next(
                row for row in version["files"] if row["name"] == "research-report.json"
            )
            versions.append(
                {
                    **version,
                    "report": json.loads(_bounded_bytes(path / manifest["stored_name"], 2_000_000)),
                    "quality": "incomplete",
                }
            )
        except (OSError, ValueError, TypeError, KeyError, RecursionError) as exc:
            log.warning("research_report_version_invalid type=%s", type(exc).__name__)
            versions.append(
                {
                    "version_id": path.name,
                    "quality": "invalid",
                    "issues": ["report_archive_invalid"],
                }
            )
        if len(versions) >= limit:
            break
    return versions


def report_blocks(payload):
    """The same text/version is consumed by all three renderers."""
    blocks = [(payload["title"], True)]
    if payload.get("as_of"):
        blocks.append(("研究截止时间：" + str(payload["as_of"]), False))
    for section in payload["sections"]:
        blocks.extend([(section["heading"], True), (section["text"], False)])
        if section.get("citations"):
            blocks.append(("引用：" + ", ".join(section["citations"]), False))
    traces = payload.get("business_calculation_trace", [])
    if traces:
        if not isinstance(traces, list) or len(traces) > 16:
            raise ValueError("research_report_calculation_trace_invalid")
        blocks.append(("计算底稿（系统生成）", True))
        for trace in traces:
            if (
                not isinstance(trace, dict)
                or not isinstance(trace.get("formula"), str)
                or len(trace["formula"]) > 4096
            ):
                raise ValueError("research_report_calculation_trace_invalid")
            blocks.append((trace["formula"], False))
            for operand in trace.get("operands", []):
                blocks.append(
                    (
                        " · ".join(
                            str(operand.get(key, ""))
                            for key in ("metric", "period", "value", "unit")
                        ),
                        False,
                    )
                )
        blocks.append(("仅核对所列输入的计算与展示；来源真实性及跨版本可比性仍须核验。", False))
    blocks.append(("来源与证据", True))
    for source in payload.get("sources", []):
        blocks.append(
            (source["id"] + " · " + source["title"] + " " + str(source.get("url", "")), False)
        )
        if source.get("as_of") is not None:
            blocks.append(("来源记录发布日期：" + str(source["as_of"]), False))
    missing = list(payload.get("missing", []))
    headings = {s["heading"] for s in payload["sections"]}
    missing.extend(
        "未覆盖要求：" + s for s in payload.get("required_sections", []) if s not in headings
    )
    if missing:
        blocks.append(("缺口与限制（研究质量未通过）", True))
        blocks.extend((text, False) for text in missing)
    return blocks


def _numeric_table_rows(payload):
    """Project only quantities already present in their canonical section text.

    The table duplicates nominated claims; independent Host checks still decide
    whether their evidence and arithmetic are valid. No display-only facts enter.
    """
    from decimal import InvalidOperation

    sections = {section.get("id"): section for section in payload["sections"]}
    number = re.compile(r"(?<![\w])[-+]?\d+(?:,\d{3})*(?:\.\d+)?|[-+]?\d+(?:\.\d+)?")
    rows = []
    for claim in payload.get("numeric_claims", []):
        try:
            section = sections[claim["section_id"]]
            text, period, unit = claim["text"], claim["period"], claim["unit"]
            context = text.replace(period, "", 1)
            quantities = list(number.finditer(context))
            if (
                text not in section["text"]
                or period not in text
                or len(quantities) != 1
                or _evidence_decimal(quantities[0].group()) != _evidence_decimal(claim["value"])
                or not re.match(
                    r"\s*" + re.escape(unit) + r"(?=$|[\s,，;；。.!！？?、)\]】])",
                    context[quantities[0].end() :],
                )
            ):
                raise ValueError("research_report_table_claim_inconsistent")
            rows.append((text, period, str(claim["value"]), unit))
        except (ValueError, TypeError, KeyError, InvalidOperation) as exc:
            log.warning("research_report_table_claim_omitted type=%s", type(exc).__name__)
    return rows


def render_markdown(payload):
    """Render the complete canonical payload, including genuine gaps, as text."""
    validate_payload(payload)
    log.info("research_report_markdown_rendered sections=%d", len(payload["sections"]))
    return (
        "\n\n".join(("## " if heading else "") + text for text, heading in report_blocks(payload))
        + "\n"
    )


def render_pdf(path, payload):
    path = Path(path)
    session_root = Path(globals().get("_SESSION_ROOT", Path(__file__).resolve().parent.parent))
    if not path.is_absolute():
        path = session_root / path
    try:
        previous_directory = os.getcwd()
    except PermissionError:
        previous_directory = str(session_root)
    try:
        # ReportLab initializes paths from CWD. Only the already writable output
        # directory is readable under Seatbelt; session-root grants stay closed.
        os.chdir(path.parent)
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

        font = Path(__file__).with_name("fonts") / "NotoSansSC-Regular.ttf"
        if font.is_symlink() or not font.is_file():
            raise ValueError("research_pdf_font_missing")
        pdfmetrics.registerFont(TTFont("ResearchCJK", str(font)))
        styles = {
            True: ParagraphStyle(
                "heading",
                fontName="ResearchCJK",
                fontSize=15,
                leading=22,
                wordWrap="CJK",
                spaceBefore=12,
                keepWithNext=True,
            ),
            False: ParagraphStyle(
                "body",
                fontName="ResearchCJK",
                fontSize=10,
                leading=17,
                wordWrap="CJK",
                splitLongWords=True,
            ),
        }
        story = []
        for text, heading in report_blocks(payload):
            if any(
                ord(char) not in pdfmetrics.getFont("ResearchCJK").face.charToGlyph
                for char in text
                if char not in "\n\r\t"
            ):
                raise ValueError("research_pdf_glyph_missing")
            story.extend(
                [Paragraph(html.escape(text).replace("\n", "<br/>"), styles[heading]), Spacer(1, 6)]
            )
        SimpleDocTemplate(
            str(path),
            pagesize=A4,
            title=payload["title"],
            leftMargin=42,
            rightMargin=42,
            topMargin=42,
            bottomMargin=42,
        ).build(story)
        log.info("research_pdf_rendered")
        return Path(path)
    except Exception as exc:
        log.error("research_pdf_render_failed type=%s", type(exc).__name__)
        raise
    finally:
        os.chdir(previous_directory)


def write_research_report(output, payload, *, formats=None):
    try:
        validate_payload(payload)
        selected = _REPORT_FORMATS if formats is None else _report_formats(formats)
        output = Path(output)
        output.mkdir(parents=True, exist_ok=True)
        stem = "research-report-" + uuid4().hex[:12]
        paths = {kind: output / (stem + "." + kind) for kind in selected}
        paths["json"] = output / "research-report.json"
        if any(p.is_symlink() for p in paths.values()):
            raise ValueError("research_report_output_alias")
        archive_report_version(output)
        blocks = report_blocks(payload)
        rows = _numeric_table_rows(payload)
        headings = ("声明原文", "期间", "数值", "单位")
        if "docx" in selected:
            from docx import Document

            doc = Document()
            for text, heading in blocks:
                doc.add_heading(text, level=1) if heading else doc.add_paragraph(text)
            if rows:
                table = doc.add_table(rows=1, cols=4)
                for cell, text in zip(table.rows[0].cells, headings):
                    cell.text = text
                for row in rows:
                    for cell, text in zip(table.add_row().cells, row):
                        cell.text = text
            doc.save(paths["docx"])
            Document(paths["docx"])
        if "html" in selected:
            body = "".join(
                f'<{"h2" if heading else "p"}>{html.escape(text)}</{"h2" if heading else "p"}>'
                for text, heading in blocks
            )
            if rows:
                body += (
                    "<table><thead><tr>"
                    + "".join("<th>" + html.escape(text) + "</th>" for text in headings)
                    + "</tr></thead><tbody>"
                    + "".join(
                        "<tr>"
                        + "".join("<td>" + html.escape(text) + "</td>" for text in row)
                        + "</tr>"
                        for row in rows
                    )
                    + "</tbody></table>"
                )
            paths["html"].write_text(
                '<!doctype html><html lang="zh"><meta charset="utf-8"><title>'
                + html.escape(payload["title"])
                + "</title><body>"
                + body
                + "</body></html>",
                encoding="utf-8",
            )
        if "pdf" in selected:
            import pdfplumber

            render_pdf(paths["pdf"], payload)
            with pdfplumber.open(paths["pdf"]) as pdf:
                if not pdf.pages or not any((p.extract_text() or "").strip() for p in pdf.pages):
                    raise ValueError("research_pdf_empty")
        published = {
            **{key: value for key, value in payload.items() if key != "report_formats"},
            "artifacts": [
                {"name": paths[k].name, "sha256": hashlib.sha256(paths[k].read_bytes()).hexdigest()}
                for k in selected
            ],
        }
        if formats is not None:
            published["report_formats"] = list(selected)
        manifest = output / (stem + ".json")
        manifest.write_text(json.dumps(published, ensure_ascii=False), encoding="utf-8")
        manifest.replace(paths["json"])
        log.info("research_report_rendered formats=%d", len(selected))
        return paths
    except Exception as exc:
        log.error("research_report_failed type=%s", type(exc).__name__)
        raise
