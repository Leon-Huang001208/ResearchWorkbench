"""Deterministic document extraction/delivery helpers for the isolated runner."""

import html
import logging
from pathlib import Path
from uuid import uuid4

log = logging.getLogger("research_web.documents")


def read_hot_evidence() -> dict:
    """Verify the admitted frozen copy and return bounded selected evidence only."""
    import hashlib
    import json
    import re
    from uuid import UUID

    try:
        context = json.loads(Path("inputs/hot-context.json").read_bytes())
        did = str(UUID(context["dataset_id"]))
        root = Path("inputs/datasets") / did
        expected_path = str(root / "evidence-manifest.json")
        if context["evidence_manifest"] != expected_path:
            raise ValueError("hot_evidence_path_invalid")
        canonical = (root / "evidence-manifest.json").read_bytes()
        digest = hashlib.sha256(canonical).hexdigest()
        if len(canonical) > 700000 or digest != context["evidence_manifest_hash"]:
            raise ValueError("hot_evidence_manifest_hash_invalid")
        manifest = json.loads(canonical)
        selected = context["selected_material_ids"]
        if (manifest["instanceId"] != context["logical_instance_id"] or manifest["materialIds"] != selected
                or not isinstance(selected, list) or not 1 <= len(selected) <= 20
                or not isinstance(manifest["objects"], list) or len(manifest["objects"]) > 40):
            raise ValueError("hot_evidence_binding_invalid")
        objects = {}
        total = 0
        for item in manifest["objects"]:
            name = item["filename"]
            if not isinstance(name, str) or not re.fullmatch(r"(?:raw-[a-f0-9-]{36}\.bin|material-[a-f0-9-]{36}-r[1-9][0-9]{0,9}\.json)", name):
                raise ValueError("hot_evidence_object_name_invalid")
            raw = (root / name).read_bytes()
            total += len(raw)
            if (len(raw) > 8 * 1024 * 1024 or total > 32 * 1024 * 1024
                    or len(raw) != item["bytes"] or hashlib.sha256(raw).hexdigest() != item["contentHash"]):
                raise ValueError("hot_evidence_object_hash_invalid")
            objects[item["key"]] = raw
        materials = []
        for item in manifest["materials"]:
            material = json.loads(objects[item["normalizedObject"]])
            if material["materialId"] not in selected or material["revision"] != item["revision"]:
                raise ValueError("hot_evidence_material_invalid")
            summary = {"revision": material["revision"]}
            truncated = []
            for name, limit in (("materialId", 64), ("title", 256), ("originalUrl", 512),
                                ("publishedAt", 64), ("acquiredAt", 64), ("contentHash", 64), ("excerpt", 512)):
                value = material.get(name)
                if value is None:
                    summary[name] = None
                elif not isinstance(value, str):
                    raise ValueError("hot_evidence_field_type_invalid")
                elif len(value.encode("utf-8")) > limit:
                    truncated.append(name)
                    # Do not return a broken URL or identity/hash as if it were complete.
                    summary[name] = value.encode("utf-8")[:limit].decode("utf-8", errors="ignore") if name in {"title", "excerpt"} else None
                else:
                    summary[name] = value
            summary["truncatedFields"] = truncated
            summary["rawHash"] = hashlib.sha256(objects[item["rawObject"]]).hexdigest()
            materials.append(summary)
        log.info("hot_evidence_read materials=%d", len(materials))
        limitations = manifest.get("limitations", [])
        if not isinstance(limitations, list) or any(not isinstance(value, str) for value in limitations):
            raise ValueError("hot_evidence_limitations_invalid")
        result = {"manifestHash": digest, "cutoff": manifest["cutoff"], "timezone": manifest["timezone"],
                  "materials": materials,
                  "limitations": [value.encode("utf-8")[:256].decode("utf-8", errors="ignore") for value in limitations[:10]],
                  "limitationsTruncated": len(limitations) > 10 or any(len(value.encode("utf-8")) > 256 for value in limitations[:10]),
                  "scope": "仅所选已存材料；原始RSS已核验，不是原链接网页全文或完整覆盖"}
        if len(json.dumps(result, ensure_ascii=False).encode("utf-8")) > 32768:
            raise ValueError("hot_evidence_summary_limit")
        return result
    except (OSError, ValueError, KeyError, TypeError) as exc:
        log.error("hot_evidence_read_failed type=%s", type(exc).__name__)
        raise ValueError("hot_evidence_read_failed") from exc


def read_pdf(path: str | Path, max_pages: int = 150) -> list[dict]:
    import pdfplumber

    try:
        with pdfplumber.open(path) as pdf:
            if len(pdf.pages) > max_pages:
                raise ValueError(f"PDF 超过 {max_pages} 页；请明确拆分研究范围")
            pages = [
                {"page": index, "text": page.extract_text() or "", "tables": page.extract_tables()}
                for index, page in enumerate(pdf.pages, 1)
            ]
        log.info("pdf_extracted pages=%d", len(pages))
        return pages
    except (OSError, ValueError) as exc:
        log.error("pdf_extraction_failed type=%s", type(exc).__name__)
        raise


def safe_cell(value):
    # Text remains text, never an Excel formula supplied by an external document.
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value


def write_deliverables(
    output: str | Path, title: str, sections: list[dict], rows: list[dict], sources: list[str]
) -> dict[str, Path]:
    from docx import Document
    from openpyxl import Workbook, load_workbook

    output = Path(output)
    output.mkdir(parents=True, exist_ok=True)
    stem = f"report-{uuid4().hex[:10]}"
    paths = {
        key: output / f"{stem}.{ext}"
        for key, ext in {"markdown": "md", "html": "html", "docx": "docx", "xlsx": "xlsx"}.items()
    }
    try:
        markdown = f"# {title}\n\n" + "\n\n".join(
            f"## {section['heading']}\n\n{section['text']}" for section in sections
        )
        markdown += "\n\n## 来源与缺失项\n\n" + "\n".join(f"- {source}" for source in sources)
        paths["markdown"].write_text(markdown, encoding="utf-8")
        body = "".join(
            f"<section><h2>{html.escape(str(section['heading']))}</h2><p>{html.escape(str(section['text'])).replace(chr(10), '<br>')}</p></section>"
            for section in sections
        )
        paths["html"].write_text(
            f'<!doctype html><html lang="zh"><meta charset="utf-8"><title>{html.escape(title)}</title><style>body{{max-width:900px;margin:40px auto;font:16px/1.8 sans-serif;padding:20px;color:#16304a}}h1,h2{{color:#164e89}}</style><h1>{html.escape(title)}</h1>{body}<h2>来源</h2><pre>{html.escape(chr(10).join(sources))}</pre></html>',
            encoding="utf-8",
        )
        doc = Document()
        doc.add_heading(title, 0)
        for section in sections:
            doc.add_heading(str(section["heading"]), 1)
            doc.add_paragraph(str(section["text"]))
        doc.add_heading("来源与缺失项", 1)
        for source in sources:
            doc.add_paragraph(source)
        doc.save(paths["docx"])
        book = Workbook()
        sheet = book.active
        sheet.title = "analysis"
        columns = list(dict.fromkeys(key for row in rows for key in row))
        if columns:
            sheet.append([safe_cell(column) for column in columns])
            for row in rows:
                sheet.append([safe_cell(row.get(column)) for column in columns])
        references = book.create_sheet("sources")
        for source in sources:
            references.append([safe_cell(source)])
        book.save(paths["xlsx"])
        # Reopen to detect broken containers before reporting delivery.
        Document(paths["docx"])
        check = load_workbook(paths["xlsx"], read_only=True)
        check.close()
        log.info("deliverables_validated count=%d", len(paths))
        return paths
    except (OSError, ValueError, KeyError, TypeError) as exc:
        log.error("deliverable_generation_failed type=%s", type(exc).__name__)
        raise
