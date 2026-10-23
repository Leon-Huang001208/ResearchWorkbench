"""Private Excel-primary weekly inputs; no SDK import or alternate runtime.

Only the Host supplies immutable published Run resources. Excel calls share the
ordinary Wind worker and ownership slot, including cleanup failure poisoning.
"""

from __future__ import annotations

import asyncio
import hashlib
import importlib.util
import json
import math
import re
import zipfile
from datetime import UTC, date, datetime, time, timedelta
from io import BytesIO
from typing import Any
from xml.etree import ElementTree
from zoneinfo import ZoneInfo

from core.observability import get_logger

from ..report_workflows.weekly_data import (
    assess_weekly_snapshot,
    load_weekly_mapping,
    sample_snapshot,
)
from ..report_workflows.weekly_sources import _resource
from . import providers_wind as shared
from .providers import ProviderError

log = get_logger(__name__)
_DAY = timedelta(days=1)
_ECONOMIC_HEADER = ["日期", "时间", "国家/地区", "指标名称", "重要性", "前值", "预测值", "今值"]
COLLECTION_POLICY = "wind-excel-weekly-v1"


def _host_now():
    return datetime.now(UTC)


def _native_date_in_excerpt(excerpt, value):
    day = date.fromisoformat(value)
    return (
        re.search(rf"(?<!\d){day.year}[-/.年]0?{day.month}[-/.月]0?{day.day}日?(?!\d)", excerpt)
        is not None
    )


def _export_observation(run_root, manifest, export, contract, sheets):
    """Consume only operator-captured, reviewed, immutable native UI evidence."""
    try:
        raw = _resource(run_root, manifest, export["observation_path"], immutable=True)
        if len(raw) > 512 * 1024 or hashlib.sha256(raw).hexdigest() != export["observation_sha256"]:
            raise ValueError("observation_hash")
        observation = json.loads(raw)
        if set(observation) != {
            "schema_version",
            "source",
            "kind",
            "application",
            "source_url",
            "export_path",
            "export_sha256",
            "coverage_start",
            "coverage_end",
            "observed_at",
            "pages",
        }:
            raise ValueError("observation_fields")
        if (
            observation["schema_version"] != 1
            or observation["source"] != "wind_excel"
            or observation["kind"] != "native_export_observation"
            or observation["application"] != "com.wind.mac.Windotd"
            or observation["source_url"] != "http://114.80.154.45/eco.web/index.html?lan=cn#/"
            or observation["export_path"] != export["path"]
            or observation["export_sha256"] != export["sha256"]
        ):
            raise ValueError("observation_identity")
        start, end = observation["coverage_start"], observation["coverage_end"]
        if not 0 <= (date.fromisoformat(end) - date.fromisoformat(start)).days <= 7:
            raise ValueError("observation_dates")
        observed, cutoff = datetime.fromisoformat(
            observation["observed_at"]
        ), datetime.fromisoformat(contract["cutoff"])
        if (
            observed.tzinfo is None
            or not cutoff - timedelta(days=1) <= observed <= cutoff
            or observed > datetime.now(UTC)
        ):
            raise ValueError("observation_time")
        if set(observation["pages"]) != {"economic_data", "financial_events"}:
            raise ValueError("observation_pages")
        used = set()
        for key, label in (("economic_data", "经济数据"), ("financial_events", "财经大事")):
            page = observation["pages"][key]
            if set(page) != {"sheet", "result", "filters", "message", "raw_accessibility_excerpt"}:
                raise ValueError("observation_page_fields")
            excerpt, sheet = page["raw_accessibility_excerpt"], page["sheet"]
            if (
                page["filters"]
                != {"countries": "all", "event_types": "all", "importance": ["重要", "中等", "低"]}
                or not isinstance(excerpt, str)
                or label not in excerpt
                or not _native_date_in_excerpt(excerpt, start)
                or not _native_date_in_excerpt(excerpt, end)
                or not re.search(r"(?:选择国家|国家(?:/地区)?)\s*[:：]?\s*全部", excerpt)
                or not re.search(r"(?:数据/事件|事件(?:类型)?)\s*[:：]?\s*全部", excerpt)
                or any(item not in excerpt for item in ("重要", "中等", "低"))
            ):
                raise ValueError("observation_filters_or_excerpt")
            if sheet is not None:
                if sheet not in sheets or sheet in used:
                    raise ValueError("observation_sheet")
                used.add(sheet)
            if key == "economic_data" and sheet != export["sheet"]:
                raise ValueError("observation_primary_sheet")
            if page["result"] == "exported":
                if sheet is None or page["message"] is not None:
                    raise ValueError("observation_exported_page")
            elif page["result"] == "empty":
                message = "当前日期暂无公布日程，请切换至其他日期。"
                if page["message"] != message or message not in excerpt:
                    raise ValueError("observation_empty_page")
                if sheet is not None and any(
                    row
                    and row[0] not in (None, "", "日期")
                    and not str(row[0]).startswith("数据来源")
                    for row in sheets[sheet][1:]
                ):
                    raise ValueError("observation_empty_conflict")
            else:
                raise ValueError("observation_result")
        if used != set(sheets):
            raise ValueError("observation_sheet_coverage")
        return observation
    except (KeyError, TypeError, ValueError) as exc:
        raise ProviderError("weekly_excel_export_observation_unverified") from exc


def _economic_rows(book, name, *, primary=False):
    rows = _sheet_rows(book, name)
    header = _ECONOMIC_HEADER if primary else _ECONOMIC_HEADER[:5]
    if (
        not rows
        or rows[0][: len(header)] != header
        or not any(
            any(isinstance(value, str) and "数据来源: Wind" in value for value in row)
            for row in rows
        )
    ):
        raise ProviderError("weekly_excel_economic_metadata_invalid")
    return rows


def _day(value):
    if isinstance(value, (datetime, date)):
        return value.date().isoformat() if isinstance(value, datetime) else value.isoformat()
    return date.fromisoformat(str(value)[:10]).isoformat()


def _raw(value):
    encoded = json.dumps(value, ensure_ascii=False, sort_keys=True, allow_nan=False).encode()
    if len(encoded) > 16 * 1024 * 1024:
        raise ProviderError("weekly_excel_response_too_large")
    return {"raw_payload": value, "raw_sha256": hashlib.sha256(encoded).hexdigest()}


def _workbook(raw):
    from openpyxl import load_workbook
    from openpyxl.utils.cell import column_index_from_string, coordinate_from_string

    with zipfile.ZipFile(BytesIO(raw)) as package:
        members = package.infolist()
        if len(members) > 1024 or sum(item.file_size for item in members) > 64 * 1024 * 1024:
            raise ProviderError("weekly_excel_export_size_limit")
        for member in members:
            if member.filename.startswith("xl/worksheets/") and member.filename.endswith(".xml"):
                for _event, element in ElementTree.iterparse(
                    BytesIO(package.read(member)), events=("start",)
                ):
                    local = element.tag.rsplit("}", 1)[-1]
                    if local == "row" and int(element.attrib.get("r", "0")) > 10010:
                        raise ProviderError("weekly_excel_export_size_limit")
                    if local == "c":
                        column, row = coordinate_from_string(element.attrib.get("r", "A1"))
                        if row > 10010 or column_index_from_string(column) > 64:
                            raise ProviderError("weekly_excel_export_size_limit")
    return load_workbook(BytesIO(raw), read_only=True, data_only=True)


def _sheet_rows(book, name):
    sheet = book[name]
    # Wind native exports can incorrectly declare dimension A1.
    sheet.reset_dimensions()
    rows: list[list[Any]] = []
    for row in sheet.iter_rows(values_only=True):
        if len(rows) >= 10010 or len(row) > 64:
            raise ProviderError("weekly_excel_export_size_limit")
        rows.append(list(row))
    return rows


def read_exports(run_root, manifest, source_contract, contract):
    """Validate numeric native exports and independently check reviewed windows."""
    exports = source_contract["excel_exports"]
    cutoff = datetime.fromisoformat(contract["cutoff"])
    result = {}
    economic_sheets = {}
    for kind, export in exports.items():
        moment = datetime.fromisoformat(export["exported_at"])
        # Fresh explicitly frozen input, never an assertion of refresh success.
        if (
            moment > datetime.now(UTC) + timedelta(minutes=5)
            or moment > cutoff
            or moment < cutoff - timedelta(days=1)
        ):
            raise ProviderError("weekly_excel_export_stale_or_future")
        raw = _resource(run_root, manifest, export["path"], immutable=True)
        if hashlib.sha256(raw).hexdigest() != export["sha256"]:
            raise ProviderError("weekly_excel_export_hash_changed")
        book = _workbook(raw)
        try:
            if kind == "economic_calendar":
                rows = _economic_rows(book, export["sheet"], primary=True)
                economic_sheets[export["sheet"]] = rows.copy()
                for sheet in export.get("additional_sheets", []):
                    additional = _economic_rows(book, sheet)
                    economic_sheets[sheet] = additional
                    rows.extend([["_rwb_source_sheet", sheet], *additional])
            else:
                rows = _sheet_rows(book, export["sheet"])
        finally:
            book.close()
        result[kind] = (rows, export["sha256"])
    economic = exports["economic_calendar"]
    observation = _export_observation(run_root, manifest, economic, contract, economic_sheets)
    start, end = contract["next_calendar_start"], contract["next_calendar_end"]
    if economic["coverage_start"] != start or economic["coverage_end"] != end:
        raise ProviderError("weekly_excel_economic_coverage_invalid")
    proof = json.loads(
        _resource(run_root, manifest, economic["coverage_evidence_path"], immutable=True)
    )
    if (
        set(proof) != {"schema_version", "source", "windows"}
        or proof["schema_version"] != 1
        or proof["source"] != "wind_excel"
    ):
        raise ProviderError("weekly_excel_economic_coverage_invalid")
    covered, exported_days = set(), set()
    if not isinstance(proof["windows"], list) or len(proof["windows"]) > 14:
        raise ProviderError("weekly_excel_economic_coverage_invalid")
    for window in proof["windows"]:
        if set(window) != {"start", "end", "resource", "sha256", "kind"} or window["kind"] not in {
            "native_export",
            "native_empty_result",
        }:
            raise ProviderError("weekly_excel_economic_coverage_invalid")
        if window["kind"] == "native_export" and (
            window["resource"] != economic["path"]
            or window["sha256"] != economic["sha256"]
            or window["start"] != observation["coverage_start"]
            or window["end"] != observation["coverage_end"]
        ):
            raise ProviderError("weekly_excel_export_observation_unverified")
        evidence = _resource(run_root, manifest, window["resource"], immutable=True)
        if hashlib.sha256(evidence).hexdigest() != window["sha256"]:
            raise ProviderError("weekly_excel_economic_coverage_invalid")
        if window["kind"] == "native_empty_result":
            empty = json.loads(evidence)
            required = {
                "schema_version",
                "source",
                "kind",
                "application",
                "source_url",
                "coverage_start",
                "coverage_end",
                "observed_at",
                "filters",
                "economic_data",
                "financial_events",
                "raw_accessibility_excerpt",
            }
            if (
                set(empty) != required
                or empty["schema_version"] != 1
                or empty["source"] != "wind_excel"
                or empty["kind"] != "native_empty_result"
                or empty["application"] != "com.wind.mac.Windotd"
                or empty["source_url"] != "http://114.80.154.45/eco.web/index.html?lan=cn#/"
                or empty["coverage_start"] != window["start"]
                or empty["coverage_end"] != window["end"]
                or window["start"] != window["end"]
                or empty["filters"]
                != {"countries": "all", "event_types": "all", "importance": ["重要", "中等", "低"]}
                or not isinstance(empty["raw_accessibility_excerpt"], str)
                or empty["coverage_start"] not in empty["raw_accessibility_excerpt"]
            ):
                raise ProviderError("weekly_excel_empty_result_unverified")
            observed = datetime.fromisoformat(empty["observed_at"])
            if (
                observed.tzinfo is None
                or observed > datetime.now(UTC) + timedelta(minutes=5)
                or observed > cutoff
                or observed < cutoff - timedelta(days=1)
            ):
                raise ProviderError("weekly_excel_empty_result_unverified")
            for key in ("economic_data", "financial_events"):
                item = empty[key]
                if (
                    not isinstance(item, dict)
                    or set(item) != {"count", "message"}
                    or type(item["count"]) is not int
                    or item["count"] != 0
                    or item["message"] != "当前日期暂无公布日程，请切换至其他日期。"
                    or item["message"] not in empty["raw_accessibility_excerpt"]
                ):
                    raise ProviderError("weekly_excel_empty_result_unverified")
        day, last = date.fromisoformat(window["start"]), date.fromisoformat(window["end"])
        if last < day or (last - day).days > 7:
            raise ProviderError("weekly_excel_economic_coverage_invalid")
        while day <= last:
            covered.add(day.isoformat())
            if window["kind"] == "native_export":
                exported_days.add(day.isoformat())
            day += _DAY
    if any(
        (date.fromisoformat(start) + timedelta(days=index)).isoformat() not in covered
        for index in range(7)
    ):
        raise ProviderError("weekly_excel_economic_coverage_incomplete")
    result["economic_export_dates"] = sorted(exported_days)
    return result


class WindExcelWeeklyProvider:
    provider_id = "wind_excel"

    def __init__(
        self,
        *,
        run_root=None,
        manifest=None,
        source_contract=None,
        client_factory=None,
        deadline_seconds=900,
    ):
        self.run_root, self.manifest = run_root, manifest
        self.source_contract = source_contract
        self.client_factory = client_factory
        self.deadline_seconds = min(1800, max(0.01, float(deadline_seconds)))

    def readiness(self):
        try:
            ready = (
                self.client_factory is not None or importlib.util.find_spec("xlwings") is not None
            )
        except (ImportError, ValueError):
            ready = False
        return {
            "provider": self.provider_id,
            "dependency_ready": ready,
            "live_verified": False,
            "code": None if ready else "wind_excel_dependency_missing",
            "adapter_revision": "wind-excel-weekly-v1",
        }

    def _calendar(self, client, market, start, contract):
        from data_layer.adapters.wind.client import _EXCEL_OPERATION_LOCK

        spec = self.source_contract["calendars"][market]
        days, requests = [], []
        cursor, last = date.fromisoformat(start), date.fromisoformat(contract["period_end"])
        while cursor <= last:
            end = min(last, cursor + timedelta(days=365))
            formula = f'=tdays("{cursor.isoformat()}","{end.isoformat()}","TradingCalendar={spec["wind_calendar"]}")'
            with _EXCEL_OPERATION_LOCK:
                client._ensure_session()
                matrix = client._execute_wsd_once(formula, 30, expected_columns=1)
            if (
                not isinstance(matrix, list)
                or len(matrix) > 1000
                or any(not isinstance(row, list) or len(row) != 1 for row in matrix)
            ):
                raise ProviderError("weekly_excel_calendar_shape_invalid")
            chunk = [_day(row[0]) for row in matrix if row and row[0] not in (None, "")]
            if chunk != sorted(set(chunk)) or any(
                not cursor.isoformat() <= day <= end.isoformat() for day in chunk
            ):
                raise ProviderError("weekly_excel_calendar_dates_invalid")
            days.extend(chunk)
            requests.append({"formula": formula, "dates": chunk})
            cursor = end + _DAY
        closes = {}
        for day in days:
            supplied = spec.get("session_closes", {}).get(day)
            if supplied is None:
                hour, minute = map(
                    int, spec.get("session_overrides", {}).get(day, spec["close_time"]).split(":")
                )
                supplied = datetime.combine(
                    date.fromisoformat(day) + timedelta(days=spec.get("close_day_offset", 0)),
                    time(hour, minute),
                    ZoneInfo(spec["timezone"]),
                ).isoformat()
            close = datetime.fromisoformat(supplied)
            if close.tzinfo is None:
                raise ProviderError("weekly_excel_session_close_unverified")
            closes[day] = supplied
        completed = [
            day
            for day in days
            if datetime.fromisoformat(closes[day]) <= datetime.fromisoformat(contract["cutoff"])
        ]
        if not completed:
            raise ProviderError("weekly_excel_calendar_empty")
        return {
            "source": "wind_excel",
            "retrieved_at": _host_now().isoformat(),
            "dates": days,
            "completed_dates": completed,
            "session_closes": closes,
            "semantics_verified": True,
            **_raw(requests),
        }

    def _series(self, spec, rows, raw, completed):
        normalized = []
        seen = set()
        for day, value in rows:
            day = _day(day)
            if day in seen:
                raise ProviderError("weekly_excel_duplicate_date")
            seen.add(day)
            if day not in completed:
                # A reviewed holiday is not a trading observation.
                if value not in (None, "", 0):
                    raise ProviderError("weekly_excel_unclosed_observation")
                continue
            if value in (None, "") or value == 0:
                continue
            if (
                isinstance(value, bool)
                or not isinstance(value, (int, float))
                or not math.isfinite(value)
                or value < 0
            ):
                raise ProviderError("weekly_excel_invalid_price")
            normalized.append({"date": day, "value": value, "published_at": None})
        return {
            **{key: spec[key] for key in ("asset", "field", "unit", "currency", "market")},
            "source": "wind_excel",
            "rows": sorted(normalized, key=lambda row: row["date"]),
            **raw,
            "retrieved_at": _host_now().isoformat(),
            "limitations": ["published_at_unknown"],
        }

    def _run(self, contract):
        from research_workbench_entrypoint.resident.admission import admission

        with admission():
            return self._run_owned(contract)

    def _run_owned(self, contract):
        snapshot = sample_snapshot(contract)
        started = _host_now()
        snapshot.update(
            source="wind_excel",
            limitations=[],
            collection_policy=COLLECTION_POLICY,
            acquisition_started_at=started.isoformat(),
            acquisition_kind="current_period",
        )
        client = None
        try:
            if (
                started.tzinfo is None
                or started.astimezone(ZoneInfo("Asia/Shanghai")).date().isoformat()
                != contract["period_end"]
                or started < datetime.fromisoformat(contract["cutoff"])
            ):
                snapshot["acquisition_kind"] = "historical_recheck"
                raise ProviderError("weekly_excel_historical_recheck_not_formal")
            exports = read_exports(self.run_root, self.manifest, self.source_contract, contract)
            if self.client_factory is None:
                from data_layer.adapters.wind.client import WindExcelClient

                client = WindExcelClient(isolated_workbook=True, existing_app_only=True)
            else:
                client = self.client_factory()
            mapping = load_weekly_mapping()
            for market in mapping["calendars"]:
                start = min(
                    [
                        contract["window_start"],
                        *[
                            spec["history_start"]
                            for spec in mapping["series"]
                            if spec["market"] == market and spec["history_start"]
                        ],
                    ]
                )
                snapshot["calendars"][market] = self._calendar(client, market, start, contract)
            edb_rows, edb_sha = exports["edb"]
            if len(edb_rows) < 9 or edb_rows[4][0] != "指标ID" or edb_rows[3][0] != "单位":
                raise ProviderError("weekly_excel_edb_metadata_missing")
            for spec in mapping["series"]:
                rows: list[tuple[Any, Any] | list[Any]]
                completed = snapshot["calendars"][spec["market"]]["completed_dates"]
                start = spec["history_start"] or contract["window_start"]
                if spec["kind"] == "edb":
                    matches = [
                        index for index, value in enumerate(edb_rows[4]) if value == spec["asset"]
                    ]
                    if len(matches) != 1:
                        raise ProviderError("weekly_excel_edb_identity_mismatch")
                    column = matches[0]
                    units = {
                        "USD/ounce": {"美元/盎司", "USD/盎司"},
                        "CNY/g": {"元/克"},
                        "USD/barrel": {"美元/桶", "USD/桶"},
                    }
                    if edb_rows[3][column] not in units.get(spec["unit"], {spec["unit"]}):
                        raise ProviderError("weekly_excel_edb_unit_mismatch")
                    rows = [
                        (row[0], row[column] if column < len(row) else None)
                        for row in edb_rows[8:]
                        if row
                        and row[0] is not None
                        and start <= _day(row[0]) <= contract["period_end"]
                    ]
                    raw = {
                        "raw_sha256": edb_sha,
                        "raw_payload": {
                            "export": self.source_contract["excel_exports"]["edb"],
                            "column": column,
                        },
                    }
                else:
                    rows, requests = [], []
                    cursor, last = date.fromisoformat(start), date.fromisoformat(completed[-1])
                    while cursor <= last:
                        end = min(last, cursor + timedelta(days=365))
                        options = f'Period=D;PriceAdj=U;Days=Trading;Fill=Blank;TradingCalendar={self.source_contract["calendars"][spec["market"]]["wind_calendar"]}'
                        from data_layer.adapters.wind.client import _EXCEL_OPERATION_LOCK

                        formula = f'=wsd("{spec["asset"]}","{spec["field"]}","{cursor.isoformat()}","{end.isoformat()}","{options}")'
                        # One accepted vendor invocation only: a timeout is unknown,
                        # and cannot silently resend the same worksheet request.
                        with _EXCEL_OPERATION_LOCK:
                            client._ensure_session()
                            matrix = client._execute_wsd_once(formula, 30, expected_columns=2)
                        if not isinstance(matrix, list) or len(matrix) > 1000:
                            raise ProviderError("weekly_excel_wsd_shape_invalid")
                        chunk = []
                        for row in matrix:
                            if (
                                not isinstance(row, list)
                                or len(row) != 2
                                or not cursor.isoformat() <= _day(row[0]) <= end.isoformat()
                            ):
                                raise ProviderError("weekly_excel_wsd_shape_invalid")
                            chunk.append([_day(row[0]), row[1]])
                        rows.extend(chunk)
                        requests.append(
                            {
                                "asset": spec["asset"],
                                "field": spec["field"],
                                "start": cursor.isoformat(),
                                "end": end.isoformat(),
                                "options": options,
                                "rows": chunk,
                            }
                        )
                        cursor = end + _DAY
                    raw = _raw(requests)
                snapshot["series"][spec["id"]] = self._series(spec, rows, raw, completed)
            economic_rows, economic_sha = exports["economic_calendar"]
            events = []
            source_sheet = self.source_contract["excel_exports"]["economic_calendar"]["sheet"]
            for row in economic_rows[1:]:
                if row and row[0] == "_rwb_source_sheet":
                    source_sheet = row[1]
                    continue
                if not row or row[0] in (None, "", "日期") or str(row[0]).startswith("数据来源"):
                    continue
                day = _day(row[0])
                if (
                    day not in exports["economic_export_dates"]
                    or len(row) < 5
                    or not all(
                        isinstance(row[index], str) and row[index].strip() for index in (2, 3, 4)
                    )
                    or row[4] not in self.source_contract["economic_calendar"]["importance_values"]
                ):
                    raise ProviderError("weekly_excel_economic_event_invalid")
                if not contract["next_calendar_start"] <= day <= contract["next_calendar_end"]:
                    continue
                events.append(
                    {
                        "date": day,
                        "source_sheet": source_sheet,
                        "country": row[2],
                        "indicator": row[3],
                        "importance": self.source_contract["economic_calendar"][
                            "importance_values"
                        ][row[4]],
                    }
                )
            snapshot["economic_calendar"] = {
                "source": "wind_excel",
                "coverage_start": contract["next_calendar_start"],
                "coverage_end": contract["next_calendar_end"],
                "coverage_verified": True,
                "events": events,
                "raw_sha256": economic_sha,
                "raw_payload": self.source_contract["excel_exports"]["economic_calendar"],
            }
            snapshot["evidence"] = [
                {"source": "wind_excel", "exports": self.source_contract["excel_exports"]}
            ]
        except Exception as exc:  # noqa: BLE001 - vendor failures become explicit partial inputs.
            log.warning("weekly_excel_collection_blocked", error_type=type(exc).__name__)
            snapshot["limitations"].append(
                str(exc) if isinstance(exc, ProviderError) else "weekly_excel_input_invalid"
            )
        finally:
            snapshot["frozen_at"] = _host_now().isoformat()
            if client is not None:
                try:
                    client.close()
                except Exception as exc:
                    with shared._WIND_LOCK:
                        shared._WIND_POISONED = True
                    log.error("weekly_excel_cleanup_unconfirmed", error_type=type(exc).__name__)
                    raise ProviderError("wind_cleanup_failed") from exc
        assessment = assess_weekly_snapshot(snapshot, contract, require_excel_acquisition=True)
        if snapshot["limitations"]:
            assessment["status"] = "blocked_data"
            assessment["issues"] = sorted({*assessment["issues"], *snapshot["limitations"]})
        return {"contract": contract, "snapshot": snapshot, "assessment": assessment}

    async def collect(self, contract, **_unused):
        with shared._WIND_LOCK:
            if shared._WIND_POISONED or (
                shared._WIND_OUTSTANDING is not None and not shared._WIND_OUTSTANDING.done()
            ):
                raise ProviderError("provider_busy")
            future = shared._WIND_EXECUTOR.submit(self._run, contract)
            shared._WIND_OUTSTANDING = future
        future.add_done_callback(shared._clear_outstanding)
        try:
            return await asyncio.wait_for(
                asyncio.shield(asyncio.wrap_future(future)), self.deadline_seconds
            )
        except TimeoutError as exc:
            log.warning("weekly_excel_deadline_reached")
            raise ProviderError("weekly_excel_timeout") from exc


async def wait_for_idle():
    """Retain vendor/close ownership until the real worker returns."""
    for future in shared.active_futures():
        try:
            await asyncio.shield(asyncio.wrap_future(future))
        except Exception as exc:  # noqa: BLE001 - drain observes the original vendor failure.
            log.warning("weekly_excel_drain_failed", error_type=type(exc).__name__)
