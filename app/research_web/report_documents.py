"""Reviewed data-only DOCX projection, shared by report sandbox and Hot exports.

The sandbox embeds this trusted source. No project imports, network or user code.
"""

import hashlib
import json
import logging
import math
import re
from copy import deepcopy
from io import BytesIO
from pathlib import Path, PurePosixPath
from zipfile import ZipFile

log = logging.getLogger("research_web.report_documents")

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_C = "http://schemas.openxmlformats.org/drawingml/2006/chart"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_S = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
_NS = {"c": _C, "s": _S}


def _office_archive(raw):
    archive = ZipFile(BytesIO(raw))
    seen, total = set(), 0
    for item in archive.infolist():
        name = PurePosixPath(item.filename)
        total += item.file_size
        if (
            item.filename in seen
            or name.is_absolute()
            or ".." in name.parts
            or "\\" in item.filename
            or total > 64 * 1024 * 1024
            or item.file_size > 32 * 1024 * 1024
            or item.external_attr >> 16 & 0o170000 == 0o120000
        ):
            archive.close()
            raise ValueError("weekly_unsafe_office_archive")
        seen.add(item.filename)
    return archive


def _set_text_node(node, text):
    from lxml import etree

    node.text = text
    node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    if "\n" in text:
        pieces = text.split("\n")
        node.text = pieces[0]
        previous = node
        for piece in pieces[1:]:
            line_break = etree.Element("{" + _W + "}br")
            previous.addnext(line_break)
            text_node = etree.Element("{" + _W + "}t")
            text_node.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
            text_node.text = piece
            line_break.addnext(text_node)
            previous = text_node


def _replace_paragraph_tokens(paragraph, values):
    nodes = [
        node
        for node in paragraph.iter("{" + _W + "}t")
        if next((p for p in node.iterancestors() if p.tag == "{" + _W + "}p"), None) is paragraph
    ]
    text = "".join(node.text or "" for node in nodes)
    spans, cursor = [], 0
    for node in nodes:
        length = len(node.text or "")
        spans.append((cursor, cursor + length, node))
        cursor += length
    replaced = set()
    pattern = re.compile(r"\{\{([^{}]+)\}\}")
    for match in reversed(list(pattern.finditer(text))):
        key = match.group(1)
        if key not in values:
            continue
        affected = [
            (start, end, node)
            for start, end, node in spans
            if start < match.end() and end > match.start()
        ]
        if not affected:
            continue
        value = str(values[key])
        for index, (start, _end, node) in enumerate(affected):
            original = node.text or ""
            begin = max(0, match.start() - start)
            end = min(len(original), match.end() - start)
            replacement = value if index == 0 else ""
            _set_text_node(node, original[:begin] + replacement + original[end:])
        replaced.add(key)
    return replaced


def build_docx_document(payload, values, template=None):
    from docx import Document

    try:
        if template:
            with _office_archive(Path(template).read_bytes()):
                pass
        document = Document(str(template)) if template else Document()
        replaced = set()
        for part in document.part.package.parts:
            if str(part.partname) == "/word/document.xml" or re.fullmatch(
                r"/word/(header|footer)\d+\.xml", str(part.partname)
            ):
                for paragraph in part.element.iter("{" + _W + "}p"):
                    replaced.update(_replace_paragraph_tokens(paragraph, values))
        if not replaced:
            document.add_heading(payload["title"], level=0)
            if payload["as_of"]:
                document.add_paragraph("数据截止时间：" + payload["as_of"])
            if payload["summary"]:
                document.add_paragraph(payload["summary"])
            for item in payload["sections"]:
                document.add_heading(item["title"], level=1)
                for line in item["content"].splitlines():
                    document.add_paragraph(line)
        if payload["missing"]:
            document.add_heading("数据缺失与限制", level=1)
            for item in payload["missing"]:
                document.add_paragraph(str(item), style="List Bullet")
        log.info("report_docx_projected")
        return document
    except Exception:
        log.warning("report_docx_projection_failed")
        raise


def _weekly_path(root, relative):
    logical = PurePosixPath(str(relative))
    if (
        logical.is_absolute()
        or ".." in logical.parts
        or "\\" in str(relative)
        or not logical.parts
        or logical.parts[0] not in {"workbooks", "snapshots"}
    ):
        raise ValueError("weekly_projection_path_invalid")
    path = root.joinpath(*logical.parts)
    if any(item.is_symlink() for item in [path, *path.parents] if item != root.parent):
        raise ValueError("weekly_projection_path_invalid")
    if not path.is_file() or not path.resolve().is_relative_to(root.resolve()):
        raise ValueError("weekly_projection_input_missing")
    if path.stat().st_size > 32 * 1024 * 1024:
        raise ValueError("weekly_projection_input_limit")
    return path


def _weekly_workbook(raw):
    import posixpath

    from lxml import etree

    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    with _office_archive(raw) as archive:
        shared = []
        if "xl/sharedStrings.xml" in archive.namelist():
            shared = [
                "".join(node.itertext())
                for node in etree.fromstring(archive.read("xl/sharedStrings.xml"), parser).findall(
                    "s:si", _NS
                )
            ]
        rels = etree.fromstring(archive.read("xl/_rels/workbook.xml.rels"), parser)
        targets = {
            node.get("Id"): node.get("Target")
            for node in rels
            if node.get("TargetMode") != "External"
        }
        workbook = etree.fromstring(archive.read("xl/workbook.xml"), parser)
        result = {}
        for sheet in workbook.findall("s:sheets/s:sheet", _NS):
            if sheet.get("name") in result:
                raise ValueError("weekly_workbook_duplicate_sheet")
            target = targets.get(sheet.get("{" + _R + "}id"), "")
            part = (
                target.lstrip("/") if target.startswith("/") else posixpath.normpath("xl/" + target)
            )
            if not part.startswith("xl/") or part not in archive.namelist():
                raise ValueError("weekly_workbook_sheet_invalid")
            cells = {}
            for cell in etree.fromstring(archive.read(part), parser).findall(
                ".//s:sheetData/s:row/s:c", _NS
            ):
                if cell.find("s:f", _NS) is not None:
                    cells[cell.get("r")] = (None, True)
                    continue
                value = cell.findtext("s:v", namespaces=_NS)
                kind = cell.get("t")
                if kind == "s":
                    value = shared[int(value)]
                elif kind == "inlineStr":
                    value = "".join(cell.find("s:is", _NS).itertext())
                elif kind not in {"str", "e", "b"} and value is not None:
                    value = float(value)
                    if not math.isfinite(value):
                        raise ValueError("weekly_workbook_nonfinite_value")
                if kind == "e":
                    raise ValueError("weekly_workbook_formula_error")
                cells[cell.get("r")] = (value, False)
            result[sheet.get("name")] = cells
    return result


def _range_cells(cells, reference):
    from openpyxl.utils.cell import get_column_letter, range_boundaries

    if not isinstance(reference, str) or not re.fullmatch(
        r"\$?[A-Z]{1,3}\$?\d+(?::\$?[A-Z]{1,3}\$?\d+)?", reference
    ):
        raise ValueError("weekly_projection_range_invalid")
    first_col, first_row, last_col, last_row = range_boundaries(reference)
    if (last_col - first_col + 1) * (last_row - first_row + 1) > 50000:
        raise ValueError("weekly_projection_range_limit")
    rows = []
    for row in range(first_row, last_row + 1):
        values = []
        for column in range(first_col, last_col + 1):
            value, formula = cells.get(get_column_letter(column) + str(row), (None, False))
            if formula:
                raise ValueError("weekly_projection_contains_formula")
            values.append(value)
        rows.append(values)
    return rows


def _chart_cache(reference, values, numeric):
    from lxml import etree

    tag = "numCache" if numeric else "strCache"
    old = reference.find("c:" + tag, _NS)
    cache = etree.Element("{" + _C + "}" + tag)
    if numeric:
        previous_format = old.find("c:formatCode", _NS) if old is not None else None
        code = etree.SubElement(cache, "{" + _C + "}formatCode")
        code.text = previous_format.text if previous_format is not None else "General"
    etree.SubElement(cache, "{" + _C + "}ptCount", val=str(len(values)))
    for index, value in enumerate(values):
        if value is None:
            if numeric:
                continue  # Different market holidays remain genuine gaps, never fabricated zeroes.
            raise ValueError("weekly_chart_category_missing")
        if numeric:
            numeric_value = float(value)
            if not math.isfinite(numeric_value):
                raise ValueError("weekly_chart_value_invalid")
            value = format(numeric_value, ".15g")
        point = etree.SubElement(cache, "{" + _C + "}pt", idx=str(index))
        etree.SubElement(point, "{" + _C + "}v").text = str(value)
    if old is not None:
        reference.replace(old, cache)
    else:
        reference.append(cache)


def _chart_projection(part, specification, workbook, sheet):
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part
    from lxml import etree
    from openpyxl import Workbook
    from openpyxl.utils.cell import range_boundaries

    cells = workbook.get(sheet)
    if cells is None:
        raise ValueError("weekly_chart_sheet_missing")
    category_range = specification["category_range"]
    value_ranges = specification["value_ranges"]
    categories = [row[0] for row in _range_cells(cells, category_range)]
    values = [[row[0] for row in _range_cells(cells, reference)] for reference in value_ranges]
    if not categories or any(len(column) != len(categories) for column in values):
        raise ValueError("weekly_chart_series_length_mismatch")
    parser = etree.XMLParser(resolve_entities=False, no_network=True)
    xml = etree.fromstring(part.blob, parser)
    series = xml.findall(".//c:ser", _NS)
    if len(series) != len(values):
        raise ValueError("weekly_chart_series_mismatch")
    embedded = Workbook()
    embedded.active.title = sheet
    for reference in [category_range, *value_ranges]:
        min_col, min_row, _, _ = range_boundaries(reference)
        for offset, row in enumerate(_range_cells(cells, reference)):
            for index, value in enumerate(row):
                embedded.active.cell(min_row + offset, min_col + index, value)
    for index, series_xml in enumerate(series):
        for key, reference, column in (
            ("cat", category_range, categories),
            ("xVal", category_range, categories),
            ("val", value_ranges[index], values[index]),
            ("yVal", value_ranges[index], values[index]),
        ):
            holder = series_xml.find("c:" + key, _NS)
            if holder is None:
                continue
            node = holder.find("c:numRef", _NS)
            numeric = node is not None
            node = node if node is not None else holder.find("c:strRef", _NS)
            if node is None:
                raise ValueError("weekly_chart_reference_missing")
            formula = node.find("c:f", _NS)
            if formula is None:
                raise ValueError("weekly_chart_reference_missing")
            formula.text = "'" + sheet.replace("'", "''") + "'!" + reference
            _chart_cache(node, column, numeric)
    if any(value is None for column in values for value in column):
        chart = xml.find("c:chart", _NS)
        blank = chart.find("c:dispBlanksAs", _NS)
        if blank is None:
            blank = etree.Element("{" + _C + "}dispBlanksAs")
            trailing = next((node for node in chart if node.tag in {
                "{" + _C + "}showDLblsOverMax", "{" + _C + "}extLst"}), None)
            chart.insert(list(chart).index(trailing), blank) if trailing is not None else chart.append(blank)
        blank.set("val", "gap")
    stream = BytesIO()
    embedded.save(stream)
    embed = Part(
        PackURI("/word/embeddings/weekly-" + PurePosixPath(str(part.partname)).stem + ".xlsx"),
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        stream.getvalue(),
        part.package,
    )
    external = xml.find("c:externalData", _NS)
    if external is None:
        external = etree.SubElement(xml, "{" + _C + "}externalData")
    for relation_id, relation in list(part.rels.items()):
        if relation.is_external or relation.reltype in {RT.OLE_OBJECT, RT.PACKAGE}:
            del part.rels[relation_id]
    relation_id = part.relate_to(embed, RT.PACKAGE)
    external.set("{" + _R + "}id", relation_id)
    update = external.find("c:autoUpdate", _NS)
    if update is None:
        update = etree.SubElement(external, "{" + _C + "}autoUpdate")
    update.set("val", "0")
    part._blob = etree.tostring(xml, encoding="UTF-8", xml_declaration=True, standalone=True)
    return {
        "sheet": sheet,
        "category_range": category_range,
        "value_ranges": value_ranges,
        "categories": categories,
        "values": values,
        "chart_xml_sha256": hashlib.sha256(part.blob).hexdigest(),
    }


def _table_rows(document, headers, rows):
    candidates = [
        table
        for table in document.tables
        if [cell.text.strip() for cell in table.rows[0].cells] == headers
    ]
    if len(candidates) != 1:
        raise ValueError("weekly_table_template_missing_or_ambiguous")
    table = candidates[0]
    prototype = deepcopy(table.rows[1]._tr if len(table.rows) > 1 else table.rows[0]._tr)
    for row in list(table.rows)[1:]:
        table._tbl.remove(row._tr)
    for values in rows:
        table._tbl.append(deepcopy(prototype))
        for cell, value in zip(table.rows[-1].cells, values):
            paragraphs = cell.paragraphs
            paragraph = paragraphs[0]
            nodes = list(paragraph._p.iter("{" + _W + "}t"))
            if nodes:
                _set_text_node(nodes[0], str(value or ""))
                for node in nodes[1:]:
                    node.text = ""
            else:
                paragraph.add_run(str(value or ""))
            for extra in paragraphs[1:]:
                cell._tc.remove(extra._p)


def apply_weekly_docx_projection(document, context, root):
    """Bind native charts and data tables to owned, verified run-copy values."""
    if context.get("weekly_production", {}).get("policy_id") != "huaan-etf-weekly-v1":
        return []
    try:
        for part in document.part.package.parts:
            element = getattr(part, "element", None)
            if element is not None:
                for paragraph in element.iter("{" + _W + "}p"):
                    text = "".join(node.text or "" for node in paragraph.iter("{" + _W + "}t"))
                    if re.search(r"\{\{[^{}]+\}\}", text):
                        raise ValueError("weekly_template_placeholder_unresolved")
        root = Path(root)
        snapshot_sha = context["dataset_snapshot_sha256"]
        snapshot = _weekly_path(root, context["dataset_snapshot"])
        if hashlib.sha256(snapshot.read_bytes()).hexdigest() != snapshot_sha:
            raise ValueError("weekly_snapshot_changed")
        projection = context["workbook_projection"]
        if projection.get("snapshot_sha256") != snapshot_sha:
            raise ValueError("weekly_workbook_snapshot_mismatch")
        workbooks = {}
        for record in projection["workbooks"]:
            path = _weekly_path(root, record["path"])
            raw = path.read_bytes()
            if hashlib.sha256(raw).hexdigest() != record["sha256"]:
                raise ValueError("weekly_workbook_changed")
            workbooks[record["path"]] = _weekly_workbook(raw)
        required = {
            block["title"]: block["id"]
            for block in context["blocks"]
            if block.get("required", True) and block.get("kind") in {"chart", "table"}
        }
        chart_specs = (
            (1, "industry_weekly_performance", "申万一级行业周涨跌幅"),
            (2, "gold_price", "黄金价格走势"),
            (3, "oil_price", "Brent vs WTI 原油期货结算价"),
        )
        parts = {str(part.partname): part for part in document.part.package.parts}
        bindings = []
        for number, key, title in chart_specs:
            spec = projection["charts"][key]
            part = parts.get(f"/word/charts/chart{number}.xml")
            if part is None or title not in required:
                raise ValueError("weekly_chart_template_missing")
            actual = _chart_projection(part, spec, workbooks[spec["workbook"]], spec["sheet"])
            bindings.append(
                {
                    "id": required[title],
                    "snapshot_sha256": snapshot_sha,
                    "content_sha256": hashlib.sha256(
                        json.dumps(
                            actual, ensure_ascii=False, sort_keys=True, separators=(",", ":")
                        ).encode()
                    ).hexdigest(),
                }
            )
        calendar = projection["economic_calendar"]
        calendar_rows = _range_cells(
            workbooks[calendar["workbook"]][calendar["sheet"]], calendar["range"]
        )
        columns = calendar_rows[0]
        indices = [columns.index(name) for name in ("日期", "国家/地区", "指标名称")]
        importance = columns.index("重要性") if "重要性" in columns else None
        calendar_rows = [
            [str(row[index] or "") for index in indices]
            for row in calendar_rows[1:]
            if importance is None or row[importance] == "重要"
        ]
        _table_rows(document, ["日期", "国家/地区", "指标名称"], calendar_rows)
        calendar_id = required.get("下周全球投资日历")
        if calendar_id is None:
            raise ValueError("weekly_calendar_block_missing")
        bindings.append(
            {
                "id": calendar_id,
                "snapshot_sha256": snapshot_sha,
                "content_sha256": hashlib.sha256(
                    json.dumps(calendar_rows, ensure_ascii=False, separators=(",", ":")).encode()
                ).hexdigest(),
            }
        )
        overseas = projection["overseas"]
        overseas_rows = _range_cells(
            workbooks[overseas["workbook"]][overseas["sheet"]], overseas["range"]
        )
        columns = overseas.get("columns")
        if columns not in (["code", "name", "return_fraction"], ["name", "code", "return_fraction"]):
            raise ValueError("weekly_overseas_columns_invalid")
        if overseas_rows and set(overseas_rows[0]) == {"指数名称", "指数代码", "周内涨跌幅"}:
            overseas_rows = overseas_rows[1:]
        _table_rows(
            document,
            ["指数名称", "指数代码", "周内涨跌幅"],
            [[row[columns.index("name")], row[columns.index("code")],
              f"{float(row[columns.index('return_fraction')]) * 100:.2f}%"] for row in overseas_rows],
        )
        if {row["id"] for row in bindings} != set(required.values()):
            raise ValueError("weekly_projection_blocks_incomplete")
        log.info("report_weekly_template_projected")
        return bindings
    except KeyError as exc:
        log.warning("report_weekly_template_projection_failed")
        raise ValueError("weekly_projection_structure_invalid") from exc
    except (IndexError, TypeError, ValueError, OSError):
        log.warning("report_weekly_template_projection_failed")
        raise
