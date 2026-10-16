"""Host-owned FinGPT execution, session-local receipts and bounded stdio events."""

from __future__ import annotations

import asyncio
import contextlib
import json
import os
import re
import shutil
import tempfile
import time
from pathlib import Path
from uuid import uuid4

from core.observability import get_logger

from ..claw_projects import ClawProjects
from ..store import StoreError

log = get_logger(__name__)
ENGINE = "fingpt-qa-v1"
ENTRY = Path(__file__).with_name("engine") / "entry.mjs"
PROJECT = Path(__file__).resolve().parents[3]


class FinGPTQA:
    def __init__(self, research, *, entry=ENTRY, node=None):
        self.research = research
        self.store = research.store
        self.entry = entry
        self.node = node or os.environ.get("RESEARCH_NODE_BINARY") or shutil.which("node")
        self.running = {}
        self.processes = {}
        self.queues = {}
        self._tiers = None
        self._tiers_at = 0
        self.locks = {}

    def owns(self, sid):
        return self.store.session(sid).get("engine") == ENGINE

    def _path(self, sid):
        return self.store.directory(sid) / "fingpt.json"

    def _read(self, sid):
        self.store.session(sid)
        path = self._path(sid)
        if path.is_symlink():
            raise StoreError("会话记录路径非法")
        try:
            return json.loads(path.read_text()) if path.exists() else {"rounds": []}
        except (OSError, ValueError):
            log.warning("fingpt_receipt_unreadable", session_id=sid)
            raise StoreError("FinGPT 会话记录无法读取", "fingpt_receipt_unavailable", 503) from None

    def _save(self, sid, value):
        parent = self.store.directory(sid)
        fd, name = tempfile.mkstemp(prefix="fingpt-", dir=parent)
        try:
            with os.fdopen(fd, "w") as stream:
                json.dump(value, stream, ensure_ascii=False)
            os.replace(name, self._path(sid))
        finally:
            if os.path.exists(name):
                os.unlink(name)

    async def _spawn(self, body):
        if not self.node:
            raise StoreError("受管 Node.js 不可用", "fingpt_node_unavailable", 503)
        private_names = {
            "MIMO_API_KEY",
            "MIMO_MODEL",
            "MINIMAX_API_KEY",
            "MINIMAX_MODEL",
            "DASHSCOPE_API_KEY",
            "QWEN_MODEL",
            "KIMI_API_KEY",
            "KIMI_MODEL",
        }
        env = {
            k: v
            for k, v in os.environ.items()
            if not k.startswith("VITE_") and k not in private_names
        }
        env["RWB_FINGPT_ENV_FILE"] = str(PROJECT / ".env.local")
        proc = await asyncio.create_subprocess_exec(
            self.node,
            str(self.entry),
            stdin=asyncio.subprocess.PIPE,
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.DEVNULL,
            env=env,
            limit=1024 * 1024,
        )
        proc.stdin.write(json.dumps(body, ensure_ascii=False).encode())
        await proc.stdin.drain()
        proc.stdin.close()
        return proc

    async def tiers(self):
        if self._tiers is None or time.monotonic() - self._tiers_at > 30:
            proc = await self._spawn({"operation": "tiers"})
            try:
                raw = await asyncio.wait_for(proc.stdout.readline(), 10)
                await asyncio.wait_for(proc.wait(), 5)
                result = json.loads(raw)["data"]
                self._tiers = {
                    "default": "lite",
                    "tiers": {
                        k: {"label": v["label"], "configured": bool(v["configured"])}
                        for k, v in result["tiers"].items()
                    },
                }
            except (ValueError, KeyError, TimeoutError):
                log.warning("fingpt_configuration_unavailable")
                raise StoreError(
                    "问答配置无法读取", "fingpt_configuration_unavailable", 503
                ) from None
            finally:
                if proc.returncode is None:
                    proc.kill()
                    await proc.wait()
        self._tiers_at = time.monotonic()
        return self._tiers

    def create(self, title="新研究", project_id="default", source_session_id=None, history=None):
        project = ClawProjects(self.research, "fingpt").get(project_id)
        if project.get("archived"):
            raise StoreError("归档项目不能新建研究", "project_archived", 409)
        row = self.store.create("fingpt", title)
        row.update(
            engine=ENGINE,
            created=True,
            model="Lite",
            project_id=project_id,
            context_source_session_id=source_session_id,
        )
        self.store.save()
        self._save(row["id"], {"rounds": [], "inherited_history": history or []})
        log.info("fingpt_session_created", session_id=row["id"])
        return self.research.summary(row)

    def detail(self, sid):
        row = self.store.session(sid)
        state = self._read(sid)
        if row["status"] == "running" and sid not in self.running:
            row["status"] = "interrupted"
            self.store.save()
        messages = []
        for item in state["rounds"]:
            messages.append({"id": item["id"] + "-u", "role": "user", "text": item["question"]})
            text = ""
            for event in item["events"]:
                if event["event"] == "update" and event["data"].get("id") == "_qa_retract":
                    text = ""
                if event["event"] == "_qa":
                    text += event["data"].get("content", "")
            if text:
                messages.append({"id": item["id"], "role": "assistant", "text": text})
        return {
            **self.research.summary(row),
            "qa_rounds": state["rounds"],
            "messages": messages,
            "files": self.store.files(sid),
            "datasets": self.research.datahub.summaries(sid),
            "activities": [],
            "questions": [],
            "approvals": [],
            "subagents": [],
            "can_cancel": sid in self.running,
            "delivery": None,
        }

    def _context(self, sid, attachments):
        files = []
        contexts = [sid]
        source = self.store.session(sid).get("context_source_session_id")
        if source:
            self.store.session(source)
            contexts.append(source)
        selected = [(sid, fid) for fid in attachments]
        if source:
            selected.extend(
                (source, fid) for fid in self.store.session(sid).get("context_source_file_ids", [])
            )
        # Explicit project file references already use the original provenance verifier.
        for context_sid in contexts:
            row = self.store.session(context_sid)
            for fid in row.get("file_references", {}):
                if fid in row.get("files", {}):
                    selected.append((context_sid, fid))
        for context_sid, fid in selected[:20]:
            path = self.store.file_path(context_sid, fid)
            stream, name = self.store.open_file(context_sid, fid)
            with stream:
                text = self._extract(stream, path.suffix.lower())
            files.append(
                {
                    "name": name,
                    "text": text,
                    "limitation": "有限摘录，未将摘录之外内容当成已读取；图片暂不解析",
                }
            )
        datasets = []
        for context_sid in contexts:
            for item in self.research.datahub.summaries(context_sid)[:12]:
                did = item.get("dataset_id") or item.get("id")
                manifest = self.research.datahub.snapshots.detail(context_sid, did)
                raw = self.research.datahub.snapshots.read(context_sid, did, "rows.json")
                datasets.append(
                    {
                        **item,
                        "manifest_sha256": manifest.get("manifest_sha256"),
                        "rows_excerpt": raw.decode()[:12000],
                        "limitation": "最多前12000字符",
                    }
                )
        return {"files": files, "datasets": datasets}

    @staticmethod
    def _extract(stream, suffix):
        if suffix in {".md", ".txt", ".csv", ".json"}:
            return stream.read(12000).decode("utf-8", errors="replace")
        if suffix == ".pdf":
            import pdfplumber

            with pdfplumber.open(stream) as pdf:
                return "\n".join((p.extract_text() or "") for p in pdf.pages[:10])[:12000]
        if suffix in {".docx", ".pptx"}:
            import zipfile
            from xml.etree import ElementTree

            with zipfile.ZipFile(stream) as archive:
                names = [
                    n
                    for n in archive.namelist()
                    if n == "word/document.xml"
                    or n.startswith("ppt/slides/slide")
                    and n.endswith(".xml")
                ][:20]
                parts = []
                for name in names:
                    if archive.getinfo(name).file_size > 2 * 1024 * 1024:
                        continue
                    parts.append(" ".join(ElementTree.fromstring(archive.read(name)).itertext()))
                return "\n".join(parts)[:12000]
        if suffix == ".xlsx":
            from openpyxl import load_workbook

            book = load_workbook(stream, read_only=True, data_only=True)
            try:
                return "\n".join(
                    str(row)
                    for sheet in book.worksheets[:5]
                    for row in sheet.iter_rows(max_row=100, values_only=True)
                )[:12000]
            finally:
                book.close()
        return "未解析此格式，请提供文字摘录。"

    async def submit(self, sid, body, key):
        async with self.locks.setdefault(sid, asyncio.Lock()):
            return await self._submit(sid, body, key)

    async def _submit(self, sid, body, key):
        if not self.owns(sid):
            raise StoreError("请新建新版 FinGPT 会话", "fingpt_legacy_session", 409)
        tiers = await self.tiers()
        if not tiers["tiers"][body.tier]["configured"]:
            raise StoreError("所选模型未配置", "fingpt_tier_unconfigured", 409)
        state = self._read(sid)
        signature = {
            "question": body.msg,
            "mode": body.mode,
            "tier": body.tier,
            "attachment_ids": body.attachment_ids,
        }
        prior = next((r for r in state["rounds"] if r["key"] == key), None)
        if prior:
            if any(prior.get(k) != v for k, v in signature.items()):
                raise StoreError("重复请求内容冲突", "idempotency_conflict", 409)
            return prior["id"]
        if len(self.running) >= 4:
            raise StoreError("研究任务已满，请稍后发送", "fingpt_busy", 429)
        if sid in self.running:
            raise StoreError("本会话仍在回答", "session_busy", 409)
        scope = self.store.session(sid).get("source_scope", ["datahub", "web", "uploads"])
        if body.attachment_ids and "uploads" not in scope:
            raise StoreError("当前研究未选择上传资料")
        context = self._context(sid, body.attachment_ids)
        if "uploads" not in scope:
            context["files"] = []
        if "datahub" not in scope:
            context["datasets"] = []
        history = [
            *state.get("inherited_history", []),
            *[{"role": m["role"], "content": m["text"]} for m in self.detail(sid)["messages"]],
        ][-6:]
        rid = uuid4().hex
        item = {"id": rid, "key": key, **signature, "events": [], "status": "running"}
        state["rounds"].append(item)
        self._save(sid, state)
        row = self.store.session(sid)
        row.update(status="running", model=body.tier.title(), updated_at=time.time())
        self.store.save()
        request = {
            "operation": "ask",
            "msg": body.msg,
            "mode": body.mode,
            "tier": body.tier,
            "context": context,
            "history": history,
            "source_scope": scope,
        }
        self.research.running[sid] = True
        self.running[sid] = asyncio.create_task(self._run(sid, rid, request))
        log.info("fingpt_research_started", session_id=sid, mode=body.mode, tier=body.tier)
        return rid

    async def _run(self, sid, rid, request):
        final = False
        proc = None
        try:
            proc = await self._spawn(request)
            self.processes[sid] = proc
            async with asyncio.timeout(900):
                while line := await proc.stdout.readline():
                    event = json.loads(line)
                    if event.get("event") not in {
                        "setup",
                        "start_of_agent",
                        "step_start",
                        "tool_call",
                        "tool_call_result",
                        "_cot",
                        "_qa",
                        "update",
                        "answer_source",
                        "error",
                    }:
                        raise ValueError("invalid_event")
                    self._append(sid, rid, event)
                    final |= event["event"] == "update" and event["data"].get("id") == "_final"
                await proc.wait()
            if proc.returncode or not final:
                raise ValueError("incomplete_stream")
            self.store.session(sid)["status"] = "completed"
        except asyncio.CancelledError:
            self.store.session(sid)["status"] = "interrupted"
            self._append(
                sid, rid, {"event": "error", "data": {"content": "已停止回答", "code": "STOPPED"}}
            )
        except (OSError, ValueError, TimeoutError):
            log.warning("fingpt_research_failed", session_id=sid)
            self.store.session(sid)["status"] = "failed"
            self._append(
                sid,
                rid,
                {
                    "event": "error",
                    "data": {"content": "研究未完成，请检查配置或重试", "code": "RESEARCH_FAILED"},
                },
            )
        finally:
            if proc and proc.returncode is None:
                with contextlib.suppress(ProcessLookupError):
                    proc.terminate()
                try:
                    await asyncio.wait_for(proc.wait(), 3)
                except TimeoutError:
                    proc.kill()
                    await proc.wait()
            self.research.running[sid] = False
            self.running.pop(sid, None)
            self.processes.pop(sid, None)
            state = self._read(sid)
            next(r for r in state["rounds"] if r["id"] == rid)["status"] = self.store.session(sid)[
                "status"
            ]
            self._save(sid, state)
            self.store.session(sid)["updated_at"] = time.time()
            self.store.save()
            self._notify(sid)
            log.info(
                "fingpt_research_finished", session_id=sid, status=self.store.session(sid)["status"]
            )

    def _artifact(self, sid, name, text):
        if not re.fullmatch(r"[a-z0-9-]+\.md", name):
            raise StoreError("产物名非法")
        path = self.store.directory(sid) / "outputs" / name
        if path.is_symlink():
            raise StoreError("产物路径非法")
        fd = os.open(
            path, os.O_WRONLY | os.O_CREAT | os.O_TRUNC | getattr(os, "O_NOFOLLOW", 0), 0o600
        )
        with os.fdopen(fd, "w") as stream:
            stream.write(text)

    def _append(self, sid, rid, event):
        state = self._read(sid)
        item = next(r for r in state["rounds"] if r["id"] == rid)
        event["seq"] = len(item["events"]) + 1
        item["events"].append(event)
        data = event["data"]
        if event["event"] == "update" and data.get("id") == "work_note":
            self._artifact(
                sid,
                f"note-{rid}-{data.get('sub_id', 0)}.md",
                str(data.get("content", {}).get("body", "")),
            )
        if event["event"] == "update" and data.get("id") == "_final":
            text = ""
            for current in item["events"]:
                if current["event"] == "update" and current["data"].get("id") == "_qa_retract":
                    text = ""
                elif current["event"] == "_qa":
                    text += current["data"].get("content", "")
            self._artifact(sid, f"report-{rid}.md", text)
        self._save(sid, state)
        self._notify(sid)

    def _notify(self, sid):
        for changed in self.queues.get(sid, set()):
            changed.set()
        for changed in self.research.listeners:
            changed.set()

    async def cancel(self, sid):
        task = self.running.get(sid)
        if task:
            task.cancel()
            await task
        return {"status": self.store.session(sid)["status"]}

    async def close(self):
        for sid in list(self.running):
            await self.cancel(sid)
