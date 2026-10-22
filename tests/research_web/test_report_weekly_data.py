"""Closed weekly data checks using synthetic vendor-shaped inputs, never live calls."""

import copy
import hashlib
import json
import tempfile
import unittest
from datetime import date, datetime, timedelta
from pathlib import Path
from zoneinfo import ZoneInfo

from openpyxl import Workbook, load_workbook
from openpyxl.chart import BarChart, LineChart, Reference

from app.research_web.report_workflows.weekly_data import (
    assess_weekly_snapshot,
    build_equivalence_proof,
    build_weekly_data_contract,
    load_weekly_mapping,
    materialize_weekly_workbooks,
    merge_weekly_sources,
    weekly_code_identity,
    weekly_snapshot_sha256,
)


def synthetic_snapshot(contract, *, short_week=False):
    """Explicit synthetic calendars/values; never reused as vendor proof."""
    mapping = load_weekly_mapping()
    all_days = []
    current, end = date(2010, 1, 1), date.fromisoformat(contract["period_end"])
    while current <= end:
        if current.weekday() < 5 and not (
            short_week and current.isoformat() in {"2026-10-21", "2026-10-22"}
        ):
            all_days.append(current.isoformat())
        current += timedelta(days=1)
    calendars = {}
    for market in mapping["calendars"]:
        earliest = min(
            [
                contract["window_start"],
                *[
                    spec["history_start"]
                    for spec in mapping["series"]
                    if spec["market"] == market and spec["history_start"]
                ],
            ]
        )
        days = [day for day in all_days if day >= earliest]
        calendars[market] = {
            "source": "wind_excel",
            "raw_sha256": "a" * 64,
            "dates": days,
            "completed_dates": days,
            "session_closes": {day: day + "T10:00:00+08:00" for day in days},
            "semantics_verified": True,
        }
    snapshot = {
        "schema_version": 1,
        "workflow_id": "huaan-etf-weekly",
        "cutoff": contract["cutoff"],
        "mapping_sha256": mapping_hash(),
        "source": "wind_excel",
        "calendars": calendars,
        "series": {},
        "economic_calendar": {
            "source": "wind_excel",
            "coverage_start": contract["next_calendar_start"],
            "coverage_end": contract["next_calendar_end"],
            "raw_sha256": "b" * 64,
            "coverage_verified": True,
            "events": [
                {
                    "date": contract["next_calendar_start"],
                    "country": "合成市场",
                    "indicator": "合成测试事件",
                    "importance": "重要",
                }
            ],
        },
        "evidence": [{"synthetic": True}],
    }
    for spec in mapping["series"]:
        days = [
            day for day in all_days if day >= (spec["history_start"] or contract["window_start"])
        ]
        snapshot["series"][spec["id"]] = {
            **{key: spec[key] for key in ("asset", "field", "unit", "currency", "market")},
            "source": "wind_excel",
            "raw_sha256": "c" * 64,
            "rows": [
                {
                    "date": day,
                    "value": "100" if spec["group"] != "turnover" else "300",
                    "published_at": day + "T10:00:00+08:00",
                }
                for day in days
            ],
        }
    return snapshot


def mapping_hash():
    from app.research_web.report_workflows.weekly_data import MAPPING_SHA256

    return MAPPING_SHA256


class WeeklyDataTests(unittest.TestCase):
    def contract(self):
        return build_weekly_data_contract("2026-10-23T15:30:00+08:00")

    def test_complete_synthetic_snapshot_passes_without_model_claims(self):
        snapshot = synthetic_snapshot(self.contract())
        result = assess_weekly_snapshot(snapshot, self.contract())
        self.assertEqual(result["issues"], [])
        self.assertEqual(result["status"], "passed")
        self.assertEqual(len(result["observations"]), 57)

    def test_short_week_turnover_uses_actual_three_sessions_and_raw_operands(self):
        snapshot = synthetic_snapshot(self.contract(), short_week=True)
        result = assess_weekly_snapshot(snapshot, self.contract())
        self.assertEqual(result["status"], "passed")
        turnover = next(row for row in result["observations"] if row["kind"] == "turnover")
        self.assertEqual(turnover["trading_days"], 3)
        self.assertEqual(turnover["value"], "300")
        self.assertEqual(len(turnover["current_rows"]), 3)

    def test_one_session_lag_allowed_but_two_blocks(self):
        snapshot = synthetic_snapshot(self.contract())
        snapshot["series"]["domestic_000300"]["rows"].pop()
        self.assertEqual(assess_weekly_snapshot(snapshot, self.contract())["status"], "passed")
        snapshot["series"]["domestic_000300"]["rows"].pop()
        result = assess_weekly_snapshot(snapshot, self.contract())
        self.assertIn("domestic_000300:weekly_series_stale", result["issues"])

    def test_fake_quality_flag_and_stock_identity_do_not_clear_missing_data(self):
        snapshot = synthetic_snapshot(self.contract())
        snapshot["series"]["domestic_000300"]["asset"] = "600519.SH"
        snapshot["status"] = "passed"
        result = assess_weekly_snapshot(snapshot, self.contract())
        self.assertEqual(result["status"], "blocked_data")
        self.assertIn("domestic_000300:weekly_series_not_equivalent", result["issues"])

    def test_calendar_close_after_cutoff_does_not_count_as_completed(self):
        snapshot = synthetic_snapshot(self.contract())
        snapshot["calendars"]["US"]["session_closes"]["2026-10-23"] = "2026-10-23T16:00:00-04:00"
        result = assess_weekly_snapshot(snapshot, self.contract())
        self.assertTrue(any("weekly_calendar_future_session" in item for item in result["issues"]))

    def test_economic_calendar_must_cover_next_week(self):
        snapshot = synthetic_snapshot(self.contract())
        snapshot["economic_calendar"]["coverage_start"] = "2026-10-19"
        self.assertIn(
            "weekly_economic_calendar_incomplete",
            assess_weekly_snapshot(snapshot, self.contract())["issues"],
        )

    def test_only_equal_source_groups_migrate(self):
        baseline = synthetic_snapshot(self.contract())
        api = copy.deepcopy(baseline)
        api["source"] = "wind_client_api"
        for value in api["series"].values():
            value["source"] = "wind_client_api"
        api["series"]["gold_london"]["rows"][-1]["value"] = "999"
        merged = merge_weekly_sources(api, baseline)
        self.assertNotIn("gold", merged["migration_groups"])
        self.assertIn("domestic", merged["migration_groups"])
        self.assertEqual(merged["series"]["gold_london"]["source"], "wind_excel")
        self.assertEqual(assess_weekly_snapshot(merged, self.contract())["status"], "passed")

    def test_prior_verified_inputs_can_admit_future_api_without_new_excel_refresh(self):
        baseline = synthetic_snapshot(self.contract())
        api = copy.deepcopy(baseline)
        api["source"] = "wind_client_api"
        for value in api["series"].values():
            value["source"] = "wind_client_api"
        identities = {
            "configuration_sha256": "d" * 64,
            "source_contract_sha256": "e" * 64,
            "sdk_identity_sha256": "f" * 64,
        }
        proof = build_equivalence_proof(api, baseline, **identities)
        future_contract = build_weekly_data_contract("2026-10-30T15:30:00+08:00")
        future = synthetic_snapshot(future_contract)
        future["source"] = "wind_client_api"
        for value in future["series"].values():
            value["source"] = "wind_client_api"
        future.update(identities)
        future.update(weekly_code_identity())
        future["equivalence_proof"] = proof
        self.assertEqual(assess_weekly_snapshot(future, future_contract)["status"], "passed")
        from unittest.mock import patch

        with patch(
            "app.research_web.report_workflows.weekly_data.ADAPTER_REVISION", "changed-adapter"
        ):
            self.assertIn(
                "weekly_equivalence_adapter_changed",
                assess_weekly_snapshot(future, future_contract)["issues"],
            )
        future["configuration_sha256"] = "0" * 64
        self.assertIn(
            "weekly_equivalence_configuration_changed",
            assess_weekly_snapshot(future, future_contract)["issues"],
        )

    def test_cutoff_is_last_friday_1530_not_wall_clock_today(self):
        contract = build_weekly_data_contract(
            datetime(2026, 10, 23, 15, 42, tzinfo=ZoneInfo("Asia/Shanghai"))
        )
        self.assertEqual(contract["cutoff"], "2026-10-23T15:30:00+08:00")
        self.assertEqual(contract["period_start"], "2026-10-19")
        self.assertEqual(contract["next_calendar_start"], "2026-10-26")
        earlier = build_weekly_data_contract(
            datetime(2026, 10, 23, 15, 29, tzinfo=ZoneInfo("Asia/Shanghai"))
        )
        self.assertEqual(earlier["cutoff"], "2026-10-16T15:30:00+08:00")

    def test_naive_cutoff_rejected(self):
        with self.assertRaisesRegex(ValueError, "timezone"):
            build_weekly_data_contract(datetime(2026, 10, 23, 15, 30))

    def test_mapping_preserves_exact_edb_and_all_industry_codes(self):
        mapping = load_weekly_mapping()
        self.assertEqual(
            {row["asset"] for row in mapping["series"] if row["kind"] == "edb"},
            {"S0270980", "S0035819", "S0031525", "M0000005"},
        )
        self.assertEqual(sum(row["group"] == "industry" for row in mapping["series"]), 31)
        self.assertIn("000985.CSI", {row["asset"] for row in mapping["series"]})

    def test_old_workbook_cells_do_not_certify_trading_days(self):
        contract = build_weekly_data_contract(
            datetime(2026, 10, 23, 15, 30, tzinfo=ZoneInfo("Asia/Shanghai"))
        )
        assessment = assess_weekly_snapshot({"workbooks": [{"cells": {}}]}, contract)
        self.assertEqual(assessment["status"], "blocked_data")
        self.assertIn("weekly_snapshot_schema_invalid", assessment["issues"])

    def test_hash_is_stable_but_changes_when_real_input_changes(self):
        source = {"series": {"x": {"rows": [{"date": "2026-10-23", "value": 1}]}}}
        other = copy.deepcopy(source)
        self.assertEqual(weekly_snapshot_sha256(source), weekly_snapshot_sha256(other))
        other["series"]["x"]["rows"][0]["value"] = 2
        self.assertNotEqual(weekly_snapshot_sha256(source), weekly_snapshot_sha256(other))

    def test_workbook_projection_preserves_package_and_updates_chart_caches(self):
        snapshot = synthetic_snapshot(self.contract())
        with tempfile.TemporaryDirectory() as temporary:
            root = Path(temporary)
            (root / "workbooks").mkdir()
            (root / "snapshots").mkdir()
            raw = json.dumps(snapshot, ensure_ascii=False, indent=2).encode()
            (root / "snapshots/report-data.json").write_bytes(raw)
            self.make_workbooks(root)
            resources = [
                {"path": "workbooks/" + name}
                for name in ("周报数据.xlsx", "周报图表.xlsx", "全球经济日历.xlsx")
            ]
            result = materialize_weekly_workbooks(
                snapshot, root, {"workflow_id": "huaan-etf-weekly", "resources": resources}
            )
            self.assertEqual(result["snapshot_sha256"], hashlib.sha256(raw).hexdigest())
            self.assertEqual(result["overseas"]["range"], "A2:C13")
            actual = load_workbook(root / "workbooks/周报数据.xlsx", data_only=True)
            self.assertEqual(actual["市场成交"]["B2"].value, 300 / 1e12)
            self.assertEqual(actual["国内"]["C2"].value, 0)
            self.assertEqual(actual["国内"]["C2"].number_format, "0.0000")
            from zipfile import ZipFile

            from lxml import etree

            with ZipFile(root / "workbooks/周报图表.xlsx") as package:
                chart = etree.fromstring(package.read("xl/charts/chart2.xml"))
                references = chart.xpath("//*[local-name()='f']/text()")
                self.assertIn(str(result["charts"]["gold_price"]["last_row"]), " ".join(references))

    @staticmethod
    def make_workbooks(root):
        primary = Workbook()
        primary.remove(primary.active)
        for name in ("日期", "国内", "市场成交", "海外", "黄金", "石油"):
            primary.create_sheet(name)
        for row in range(2, 7):
            primary["国内"][f"C{row}"] = -999
            primary["国内"][f"C{row}"].number_format = "0.0000"
        primary.save(root / "workbooks/周报数据.xlsx")
        # Actual Wind s_info_name formulas use cached t=str values. openpyxl
        # cannot author formula caches, so insert the vendor-shaped cell part.
        from zipfile import ZipFile

        from lxml import etree

        target = root / "workbooks/周报数据.xlsx"
        with ZipFile(target) as package:
            parts = {
                item.filename: (item, package.read(item.filename)) for item in package.infolist()
            }
        info, raw = parts["xl/worksheets/sheet2.xml"]
        sheet = etree.fromstring(raw)
        namespace = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
        row = sheet.find(f".//{{{namespace}}}row[@r='2']")
        cell = etree.SubElement(row, f"{{{namespace}}}c", r="B2", t="str")
        etree.SubElement(cell, f"{{{namespace}}}f").text = "s_info_name(A2)"
        etree.SubElement(cell, f"{{{namespace}}}v").text = "沪深300"
        parts[info.filename] = (info, etree.tostring(sheet))
        with ZipFile(target, "w") as package:
            for entry, raw in parts.values():
                package.writestr(entry, raw)
        graph = Workbook()
        graph.remove(graph.active)
        industry = graph.create_sheet("申万一级行业指数")
        industry.append(["日期", "代码", "全名", "名称", "周收益"])
        for row in range(2, 33):
            industry.append([0, "code", "industry", "合成行业" + str(row), -999])
        bar = BarChart()
        bar.add_data(Reference(industry, min_col=5, min_row=1, max_row=32), titles_from_data=True)
        bar.set_categories(Reference(industry, min_col=4, min_row=2, max_row=32))
        industry.add_chart(bar, "G2")
        for name in ("伦敦金_Au9999", "期货结算价(连续)_布伦特原油"):
            sheet = graph.create_sheet(name)
            sheet.append(["header"])
            sheet.append(["date", "price1", "price2"])
            for day in range(3, 7):
                sheet.append([day, -999, -999])
            chart = LineChart()
            chart.add_data(
                Reference(sheet, min_col=2, max_col=3, min_row=2, max_row=6), titles_from_data=True
            )
            chart.set_categories(Reference(sheet, min_col=1, min_row=3, max_row=6))
            sheet.add_chart(chart, "E2")
        graph.save(root / "workbooks/周报图表.xlsx")
        calendar = Workbook()
        calendar.active.title = "经济数据"
        calendar.active.append(
            ["日期", "时间", "国家/地区", "指标名称", "重要性", "前值", "预测值", "今值"]
        )
        calendar.active.append(["stale", None, "stale", "stale", "重要"])
        calendar.save(root / "workbooks/全球经济日历.xlsx")


if __name__ == "__main__":
    unittest.main()
