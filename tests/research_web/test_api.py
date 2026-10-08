"""BFF acceptance tests: native transport is replaced only for deterministic regressions."""

import asyncio
import json
import time

import pytest
from fastapi.testclient import TestClient

from app.research_web.client import RuntimeFailure
from app.research_web.main import create_app
from app.research_web.model_credentials import MODEL_REF, docker_backend, execute
from app.research_web.model_file_store import ModelStoreError
from app.research_web.service import ResearchService
from app.research_web.store import Store


@pytest.fixture(autouse=True)
def isolated_keyring(monkeypatch):
    """Every test owns its credentials; never consult the host Keychain."""
    import keyring

    values = {}
    monkeypatch.setattr(
        keyring, "get_password", lambda service, account: values.get((service, account))
    )
    monkeypatch.setattr(
        keyring,
        "set_password",
        lambda service, account, value: values.__setitem__((service, account), value),
    )
    monkeypatch.setattr(
        keyring, "delete_password", lambda service, account: values.pop((service, account), None)
    )


class NativeFixture:
    def __init__(self):
        self.calls = []
        self.fail_prompt = False
        self.confirm_delete = True
        self.running = set()

    async def rpc(self, method, payload):
        self.calls.append((method, payload))
        if method == "host.describe":
            return {"version": "fixture", "model": "fixture", "provider": "fixture"}
        if method == "credentials.describe":
            return {"credentials": {"RESEARCH_DSH_API_KEY": {"configured": True, "writable": True}}}
        if method == "llm.models":
            return {
                "groups": [
                    {
                        "id": "deepseek-official",
                        "models": [{"id": "deepseek-flash"}, {"id": "new-model"}],
                    }
                ],
                "failures": [],
            }
        if method == "session.prompt" and self.fail_prompt:
            raise RuntimeFailure("connection lost")
        if method == "subagent.list":
            return {"entries": []}
        if method == "session.list":
            return {
                "items": [
                    {"sessionId": session_id, "running": True}
                    for session_id in sorted(self.running)
                ]
            }
        if method == "session.delete":
            return {"deletedSessionIds": [payload["sessionId"]] if self.confirm_delete else []}
        if method == "skill.list":
            return {"skills": []}
        return {"accepted": True}

    async def plugin_json(self, method, path, *, params=None, payload=None):
        self.calls.append(
            (f"plugin:{method}", {"path": path, "params": params, "payload": payload})
        )
        if path == "/research/tabbit/status":
            return {
                "status": "ready",
                "pluginVersion": "0.3.4",
                "browserVersion": "1.13.23",
                "launcherPresent": True,
                "onlineInstances": 1,
                "selectedInstance": "ABCDEF0123456789",
            }
        if path == "/research/tabbit/access":
            return {"accepted": True}
        if path == "/research/tabbit/tabs":
            return {
                "instanceId": "ABCDEF0123456789",
                "tabs": [
                    {
                        "tabId": 7,
                        "title": "Live research page",
                        "url": "https://example.com/live",
                        "active": True,
                        "state": "available",
                    }
                ],
            }
        if path == "/research/tabbit/live-extract":
            return {
                "markers": [
                    {
                        "tabId": tab_id,
                        "title": "Live research page",
                        "marker": f"@[Live research page](rwb-tabbit:{tab_id}-token)",
                    }
                    for tab_id in payload["tabIds"]
                ]
            }
        raise AssertionError(path)

    async def history(self, sid):
        return []

    async def close(self):
        pass

    async def frames(self, channel):
        yield {"type": "connected", "channel": channel}
        await asyncio.Event().wait()


class DockerModelFixture(NativeFixture):
    """Real temporary model bridge/store; only native transport/generation is fake."""

    def __init__(self, root, installation_id="a" * 32):
        super().__init__()
        self.root = root
        self.root.mkdir(parents=True, mode=0o700, exist_ok=True)
        self.backend = docker_backend(root / "models" / installation_id, installation_id)
        self.backend_failure = False
        self.write_uncertain = False
        self.outcome = "completed"

    async def rpc(self, method, payload):
        if method.startswith("credentials."):
            if self.backend_failure:
                raise RuntimeFailure("private fixture detail", "credential/rejected")
            op = method.removeprefix("credentials.")
            request = {"op": op, "ref": MODEL_REF}
            if op == "set":
                request["value"] = payload["value"]
            try:
                response = execute(self.root, request, self.backend)
            except ModelStoreError:
                raise RuntimeFailure("凭据操作失败", "credential/rejected") from None
            if op in {"set", "unset"} and self.write_uncertain:
                raise RuntimeFailure("model_credential_commit_uncertain", "credential/rejected")
            if op == "describe":
                return {"credentials": {MODEL_REF: response}}
            return response
        response = await super().rpc(method, payload)
        if method == "session.prompt":
            sid = payload["sessionId"]
            events = [
                {"event": {"seq": 1, "type": "user/message", "data": {"content": "临时研究问题"}}},
            ]
            if self.outcome == "completed":
                events.append(
                    {
                        "event": {
                            "seq": 2,
                            "type": "assistant/message",
                            "data": {
                                "turn": 1,
                                "step": 1,
                                "message": {
                                    "content": [{"type": "text", "text": "确定性研究结果"}]
                                },
                            },
                        }
                    }
                )
            reason = {"kind": self.outcome}
            if self.outcome == "failed":
                reason = {"kind": "error", "error": {"message": self.failure_message}}
            if self.outcome == "running":
                self.running.add(sid)
                events.append({"event": {"seq": 2, "type": "turn/start", "data": {}}})
            else:
                events.append({"event": {"seq": 3, "type": "turn/end", "data": {"reason": reason}}})
            (self.root / f"history-{sid}.json").write_text(json.dumps(events), encoding="utf-8")
        elif method == "session.cancel":
            sid = payload["sessionId"]
            events = await self.history(sid)
            events.append(
                {"event": {"seq": 3, "type": "turn/end", "data": {"reason": {"kind": "aborted"}}}}
            )
            (self.root / f"history-{sid}.json").write_text(json.dumps(events), encoding="utf-8")
            self.running.discard(sid)
        return response

    async def history(self, sid):
        path = self.root / f"history-{sid}.json"
        return json.loads(path.read_text(encoding="utf-8")) if path.exists() else []


def test_docker_model_public_settings_keep_blank_replace_and_clear_separate(tmp_path):
    native = DockerModelFixture(tmp_path / "private")
    service = ResearchService(native, Store(tmp_path / "data"))
    with TestClient(create_app(service)) as client:
        service.connected = {"mux", "host"}
        endpoint = "/api/research/runtime/model"
        # The existing browser omits a blank input, rather than submitting an empty new Key.
        for key in ("synthetic-first", None, "synthetic-replacement"):
            response = client.put(
                endpoint,
                json={"model": "new-model", **({"api_key": key} if key is not None else {})},
            )
            assert response.status_code == 200
            assert key is None or key not in response.text
            expected = key or "synthetic-first"
            assert (
                execute(native.root, {"op": "resolve", "ref": MODEL_REF}, native.backend)["value"]
                == expected
            )
        native.backend_failure = True
        response = client.get("/api/research/runtime")
        assert response.status_code == 200
        assert response.json()["health_check_passed"] is True
        assert response.json()["credential_configured"] is None
        assert client.get("/").status_code == 200
        native.backend_failure = False
        response = client.put(endpoint, json={"model": "new-model", "clear_api_key": True})
        assert response.status_code == 200
        assert client.get("/api/research/runtime").json()["credential_configured"] is False


@pytest.mark.asyncio
async def test_docker_model_bridge_runtime_and_cold_research_recovery(tmp_path):
    native = DockerModelFixture(tmp_path / "private")
    service = ResearchService(native, Store(tmp_path / "data"))
    service.connected = {"mux", "host"}
    status = await service.runtime()
    assert status["health_check_passed"] is True
    assert status["credential_storage"] == "docker_private_file"
    assert status["credential_configured"] is False
    old = await service.create()
    with pytest.raises(RuntimeFailure) as missing:
        await service.send(old["id"], "问题", "docker-missing-1")
    assert missing.value.code == "model_credentials_missing"
    await service.configure_model("deepseek-official", "new-model", "temporary-synthetic-secret")
    await service.configure_model("deepseek-official", "new-model")
    assert (
        execute(native.root, {"op": "resolve", "ref": MODEL_REF}, native.backend)["value"]
        == "temporary-synthetic-secret"
    )
    assert service.store.session(old["id"])["model"] == "deepseek-flash"
    row = await service.create()
    assert row["model"] == "new-model"
    for _ in range(2):
        await service.send(row["id"], "问题", "docker-research-1")
    assert sum(method == "session.prompt" for method, _ in native.calls) == 1
    cold_native = DockerModelFixture(native.root)
    cold = ResearchService(cold_native, Store(tmp_path / "data"))
    cold.connected = {"mux", "host"}
    restored = await cold.detail(row["id"])
    assert restored["messages"][-1]["text"] == "确定性研究结果"
    assert restored["status"] == "completed"
    status = await cold.runtime()
    assert status["credential_configured"] is True
    assert status["last_model_test"] is None
    assert "temporary-synthetic-secret" not in json.dumps(status)
    assert str(native.root) not in json.dumps(status)
    other = DockerModelFixture(native.root, "b" * 32)
    assert (await ResearchService(other, Store(tmp_path / "other-data")).runtime())[
        "credential_configured"
    ] is False
    await cold.configure_model("deepseek-official", "new-model", clear_api_key=True)
    assert (
        execute(native.root, {"op": "describe", "ref": MODEL_REF}, native.backend)["configured"]
        is False
    )


@pytest.mark.asyncio
async def test_docker_model_backend_and_postcommit_uncertainty_remain_fail_closed(tmp_path):
    native = DockerModelFixture(tmp_path / "private")
    service = ResearchService(native, Store(tmp_path / "data"))
    service.connected = {"mux", "host"}
    await service.configure_model("deepseek-official", "new-model", "synthetic-old")
    native.backend_failure = True
    status = await service.runtime()
    assert status["connected"] is True and status["health_check_passed"] is True
    assert status["credential_configured"] is None
    assert status["credential_storage"] == "unknown"
    assert "private fixture detail" not in json.dumps(status)
    native.backend_failure = False
    native.write_uncertain = True
    with pytest.raises(RuntimeFailure) as uncertain:
        await service.configure_model("deepseek-official", "new-model", "synthetic-new")
    assert uncertain.value.code == "model_configuration_uncertain"
    assert (
        execute(native.root, {"op": "resolve", "ref": MODEL_REF}, native.backend)["value"]
        == "synthetic-new"
    )
    cold = ResearchService(DockerModelFixture(native.root), Store(tmp_path / "data"))
    cold.connected = {"mux", "host"}
    assert (await cold.runtime())["configuration_uncertain"] is True
    row = await cold.create()
    with pytest.raises(RuntimeFailure) as rejected:
        await cold.send(row["id"], "问题", "docker-uncertain-1")
    assert rejected.value.code == "model_configuration_uncertain"
    assert not any(method == "session.prompt" for method, _ in cold.client.calls)


@pytest.mark.asyncio
@pytest.mark.parametrize(
    "message,code",
    [("401 authentication", "model_auth_failed"), ("network timed out", "model_network_timeout")],
)
async def test_docker_model_generation_failures_are_not_health_failures(tmp_path, message, code):
    native = DockerModelFixture(tmp_path / "private")
    native.outcome = "failed"
    native.failure_message = message
    service = ResearchService(native, Store(tmp_path / "data"))
    service.connected = {"mux", "host"}
    await service.configure_model("deepseek-official", "new-model", "synthetic-key")
    result = await service.test_model()
    assert result["status"] == "failed" and result["code"] == code
    assert (await service.runtime())["health_check_passed"] is True


@pytest.mark.asyncio
async def test_docker_model_cancelled_commit_is_durable_and_clear_recovers(tmp_path):
    native = DockerModelFixture(tmp_path / "private")
    service = ResearchService(native, Store(tmp_path / "data"))
    original = native.rpc

    async def cancelled(method, payload):
        result = await original(method, payload)
        if method == "credentials.set":
            raise asyncio.CancelledError
        return result

    native.rpc = cancelled
    with pytest.raises(asyncio.CancelledError):
        await service.configure_model("deepseek-official", "new-model", "synthetic-committed")
    cold = ResearchService(DockerModelFixture(native.root), Store(tmp_path / "data"))
    assert cold.store.data["model_configuration_uncertain"] is True
    assert (
        execute(native.root, {"op": "resolve", "ref": MODEL_REF}, native.backend)["value"]
        == "synthetic-committed"
    )
    await cold.configure_model("deepseek-official", "new-model", clear_api_key=True)
    assert not cold.store.data.get("model_configuration_uncertain")
    assert (await cold.runtime())["credential_configured"] is False


@pytest.mark.asyncio
@pytest.mark.parametrize("active_child", [False, True])
async def test_docker_model_active_parent_or_child_refuses_key_changes(tmp_path, active_child):
    native = DockerModelFixture(tmp_path / "private")
    service = ResearchService(native, Store(tmp_path / "data"))
    row = await service.create()
    original = native.rpc
    if active_child:

        async def child(method, payload):
            if method == "subagent.list":
                return {"entries": [{"activity": "running"}]}
            return await original(method, payload)

        native.rpc = child
    else:
        native.running.add(row["id"])
    with pytest.raises(RuntimeFailure) as busy:
        await service.configure_model("deepseek-official", "new-model", "synthetic-key")
    assert busy.value.code == "model_change_busy"
    assert (
        execute(native.root, {"op": "describe", "ref": MODEL_REF}, native.backend)["configured"]
        is False
    )


@pytest.mark.asyncio
async def test_docker_model_cancelled_research_history_survives_fresh_service(tmp_path):
    native = DockerModelFixture(tmp_path / "private")
    native.outcome = "running"
    service = ResearchService(native, Store(tmp_path / "data"))
    service.connected = {"mux", "host"}
    await service.configure_model("deepseek-official", "new-model", "synthetic-key")
    row = await service.create()
    await service.send(row["id"], "问题", "docker-cancel-1")
    assert (await service.cancel(row["id"]))["accepted"] is True
    cold = ResearchService(DockerModelFixture(native.root), Store(tmp_path / "data"))
    cold.connected = {"mux", "host"}
    assert (await cold.detail(row["id"]))["status"] == "cancelled"
    assert not any(method == "session.prompt" for method, _ in cold.client.calls)


@pytest.fixture
def api(tmp_path):
    native = NativeFixture()
    service = ResearchService(native, Store(tmp_path))
    with TestClient(create_app(service)) as client:
        service.connected = {"mux", "host"}
        yield client, native, service


def test_create_submit_duplicate_never_replays(api):
    client, native, _ = api
    created = client.post("/api/research/sessions", json={"mode": "fingpt"})
    assert created.status_code == 201
    sid = created.json()["id"]
    endpoint = f"/api/research/sessions/{sid}/messages"
    for _ in range(2):
        response = client.post(
            endpoint, json={"text": "你好"}, headers={"Idempotency-Key": "abcdefgh"}
        )
        assert response.status_code == 202, response.text
    assert len([call for call in native.calls if call[0] == "session.prompt"]) == 1
    assert (
        client.post(
            endpoint, json={"text": "别的问题"}, headers={"Idempotency-Key": "abcdefgh"}
        ).status_code
        == 400
    )


def test_tabbit_status_access_inventory_and_live_message(api):
    client, native, _ = api
    sid = client.post("/api/research/sessions", json={}).json()["id"]

    status = client.get("/api/research/runtime/tabbit")
    assert status.status_code == 200
    assert status.json()["status"] == "ready"
    assert status.json()["web_fetch_enabled"] is False

    denied = client.get(f"/api/research/sessions/{sid}/tabbit-tabs")
    assert denied.status_code == 403
    assert denied.json()["error"]["code"] == "tabbit_page_access_required"

    grant = client.post(f"/api/research/sessions/{sid}/tabbit-access", json={"decision": "approve"})
    assert grant.status_code == 200
    tabs = client.get(f"/api/research/sessions/{sid}/tabbit-tabs?q=research")
    assert tabs.status_code == 200
    assert tabs.json()["items"][0]["tab_id"] == 7

    response = client.post(
        f"/api/research/sessions/{sid}/messages",
        headers={"Idempotency-Key": "tabbit-message-1"},
        json={
            "text": "总结这个页面",
            "tabbit_tabs": [{"tab_id": 7, "instance_id": "ABCDEF0123456789"}],
            "tabbit_live_confirmed": True,
        },
    )
    assert response.status_code == 202, response.text
    prompt = [call for call in native.calls if call[0] == "session.prompt"][-1][1]
    sent_text = prompt["content"][0]["text"]
    assert "@[Live research page](rwb-tabbit:7-token)" in sent_text


def test_tabbit_live_message_requires_confirmation_and_limits_selection(api):
    client, _, _ = api
    sid = client.post("/api/research/sessions", json={}).json()["id"]
    client.post(f"/api/research/sessions/{sid}/tabbit-access", json={"decision": "approve"})

    unconfirmed = client.post(
        f"/api/research/sessions/{sid}/messages",
        headers={"Idempotency-Key": "tabbit-message-2"},
        json={
            "text": "总结",
            "tabbit_tabs": [{"tab_id": 7, "instance_id": "ABCDEF0123456789"}],
        },
    )
    assert unconfirmed.status_code == 409
    assert unconfirmed.json()["error"]["code"] == "tabbit_claim_confirmation_required"

    too_many = client.post(
        f"/api/research/sessions/{sid}/messages",
        headers={"Idempotency-Key": "tabbit-message-3"},
        json={
            "text": "总结",
            "tabbit_tabs": [
                {"tab_id": index, "instance_id": "ABCDEF0123456789"} for index in range(9)
            ],
            "tabbit_live_confirmed": True,
        },
    )
    assert too_many.status_code == 422


def test_tabbit_settings_validate_dependency_and_report_restart(api):
    client, _, _ = api
    invalid = client.put(
        "/api/research/runtime/tabbit",
        json={"browser_enabled": False, "web_fetch_enabled": True},
    )
    assert invalid.status_code == 409
    assert invalid.json()["error"]["code"] == "tabbit_web_fetch_requires_browser"

    saved = client.put(
        "/api/research/runtime/tabbit",
        json={
            "browser_enabled": True,
            "web_fetch_enabled": True,
            "instance_id": "ABCDEF0123456789",
        },
    )
    assert saved.status_code == 200
    assert saved.json()["restart_required"] is True


def test_session_soft_delete_restore_is_recoverable_and_never_archives_native_history(api):
    client, native, service = api
    created = client.post("/api/research/sessions", json={}).json()
    sid = created["id"]
    artifact = service.store.directory(sid) / "outputs" / "retained.md"
    artifact.write_text("保留的研究产物", encoding="utf-8")

    deleted = client.delete(f"/api/research/sessions/{sid}")
    assert deleted.status_code == 200
    assert deleted.json()["deleted_at"]
    assert sid not in {item["id"] for item in client.get("/api/research/sessions").json()["items"]}
    deleted_items = client.get("/api/research/sessions?view=deleted").json()["items"]
    assert [item["id"] for item in deleted_items] == [sid]
    assert deleted_items[0]["mode"] == "fingpt"
    assert client.get(f"/api/research/sessions/{sid}").status_code == 410
    assert client.get(f"/api/research/sessions/{sid}").json()["error"]["code"] == "session_deleted"
    assert artifact.read_text(encoding="utf-8") == "保留的研究产物"
    assert not any(call[0] == "workspace.archiveSession" for call in native.calls)

    repeated = client.delete(f"/api/research/sessions/{sid}")
    assert repeated.status_code == 200
    assert repeated.json()["deleted_at"] == deleted.json()["deleted_at"]

    restored = client.post(f"/api/research/sessions/{sid}/restore")
    assert restored.status_code == 200
    assert "deleted_at" not in restored.json()
    assert sid in {item["id"] for item in client.get("/api/research/sessions").json()["items"]}
    assert client.get(f"/api/research/sessions/{sid}").status_code == 200
    assert artifact.read_text(encoding="utf-8") == "保留的研究产物"
    assert client.post(f"/api/research/sessions/{sid}/restore").status_code == 200


def test_session_soft_delete_rejects_busy_and_unknown_sessions(api):
    client, native, service = api
    sid = client.post("/api/research/sessions", json={}).json()["id"]
    service.store.session(sid)["status"] = "running"
    service.store.save()
    native.running.add(sid)

    busy = client.delete(f"/api/research/sessions/{sid}")
    assert busy.status_code == 409
    assert busy.json()["error"]["code"] == "session_busy"
    assert client.get(f"/api/research/sessions/{sid}").status_code == 200

    unknown = client.delete("/api/research/sessions/00000000-0000-0000-0000-000000000000")
    assert unknown.status_code == 400
    assert unknown.json()["error"]["code"] == "invalid_resource"
    assert client.get("/api/research/sessions?view=unknown").status_code == 422


def test_session_soft_delete_reconciles_stale_cached_running_state(api):
    client, _, service = api
    sid = client.post("/api/research/sessions", json={}).json()["id"]
    service.store.data["sessions"][sid]["status"] = "running"
    service.store.save()

    response = client.delete(f"/api/research/sessions/{sid}")

    assert response.status_code == 200
    assert response.json()["deleted_at"]
    assert response.json()["status"] == "idle"


def test_deleted_session_can_be_permanently_deleted_with_native_confirmation(api):
    client, native, service = api
    sid = client.post("/api/research/sessions", json={}).json()["id"]
    artifact = service.store.directory(sid) / "outputs" / "removed.md"
    artifact.write_text("永久删除", encoding="utf-8")
    client.delete(f"/api/research/sessions/{sid}")

    response = client.delete(f"/api/research/sessions/{sid}/permanent")

    assert response.status_code == 200
    assert response.json() == {"id": sid, "purged": True}
    assert ("session.delete", {"sessionId": sid}) in native.calls
    assert not artifact.exists()
    assert sid not in service.store.data["sessions"]
    assert client.get(f"/api/research/sessions/{sid}").status_code == 400


def test_permanent_delete_removes_sealed_capabilities_and_private_datahub_state(api):
    client, _, service = api
    sid = client.post("/api/research/sessions", json={}).json()["id"]
    session_root = service.store.directory(sid)
    sealed = session_root / "resources" / "capabilities" / "fund-research-workflow" / "2"
    sealed.mkdir(parents=True)
    capability = sealed / "SKILL.md"
    capability.write_text("# sealed capability", encoding="utf-8")
    capability.chmod(0o400)
    sealed.chmod(0o500)

    private_paths = [
        service.store.root / ".control" / "snapshots" / sid,
        service.store.root / ".control" / "calls" / sid,
    ]
    for private in private_paths:
        private.mkdir(parents=True)
        (private / "record.json").write_text("{}", encoding="utf-8")

    deleted = client.delete(f"/api/research/sessions/{sid}")
    assert deleted.status_code == 200, deleted.text

    response = client.delete(f"/api/research/sessions/{sid}/permanent")

    assert response.status_code == 200, response.text
    assert not session_root.exists()
    assert all(not private.exists() for private in private_paths)


def test_expired_deleted_session_is_purged_when_deleted_view_is_loaded(api):
    client, _, service = api
    sid = client.post("/api/research/sessions", json={}).json()["id"]
    client.delete(f"/api/research/sessions/{sid}")
    service.store.session(sid, include_deleted=True)["deleted_at"] = 0
    service.store.save()

    response = client.get("/api/research/sessions?view=deleted")

    assert response.status_code == 200
    assert response.json()["items"] == []
    assert sid not in service.store.data["sessions"]


def test_expired_deleted_session_cannot_be_restored_before_the_next_purge_cycle(api):
    client, _, service = api
    sid = client.post("/api/research/sessions", json={}).json()["id"]
    client.delete(f"/api/research/sessions/{sid}")
    service.store.session(sid, include_deleted=True)["deleted_at"] = 0
    service.store.save()

    response = client.post(f"/api/research/sessions/{sid}/restore")

    assert response.status_code == 410
    assert response.json()["error"]["code"] == "session_restore_expired"
    assert service.store.session(sid, include_deleted=True)["deleted_at"] == 0


def test_permanent_delete_requires_soft_delete(api):
    client, _, _ = api
    sid = client.post("/api/research/sessions", json={}).json()["id"]

    response = client.delete(f"/api/research/sessions/{sid}/permanent")

    assert response.status_code == 409
    assert response.json()["error"]["code"] == "session_not_deleted"


def test_permanent_delete_keeps_tombstone_when_dsh_does_not_confirm(api):
    client, native, service = api
    sid = client.post("/api/research/sessions", json={}).json()["id"]
    artifact = service.store.directory(sid) / "outputs" / "retained.md"
    artifact.write_text("仍可重试", encoding="utf-8")
    client.delete(f"/api/research/sessions/{sid}")
    native.confirm_delete = False

    response = client.delete(f"/api/research/sessions/{sid}/permanent")

    assert response.status_code == 503
    assert response.json()["error"]["code"] == "native_session_delete_unconfirmed"
    assert service.store.session(sid, include_deleted=True)["deleted_at"]
    assert artifact.read_text(encoding="utf-8") == "仍可重试"


@pytest.mark.asyncio
async def test_online_retention_loop_purges_without_opening_deleted_view(tmp_path, monkeypatch):
    import app.research_web.service as service_module

    monkeypatch.setattr(service_module, "SESSION_PURGE_INTERVAL_SECONDS", 0.01)
    service = ResearchService(NativeFixture(), Store(tmp_path))
    row = service.store.create("fingpt", "到期研究")
    service.store.soft_delete(row["id"])
    service.store.session(row["id"], include_deleted=True)["deleted_at"] = 0
    service.store.save()

    task = asyncio.create_task(service._retention_loop())
    try:
        for _ in range(100):
            if row["id"] not in service.store.data["sessions"]:
                break
            await asyncio.sleep(0.01)
        assert row["id"] not in service.store.data["sessions"]
    finally:
        task.cancel()
        with pytest.raises(asyncio.CancelledError):
            await task


def test_datahub_read_only_catalog_authenticated_queries_and_upgrade(api):
    import httpx
    from test_datahub import nav_page, nav_row

    from app.research_web.datahub import BusinessQuery

    client, native, service = api
    catalog = client.get("/api/research/data/catalog")
    assert catalog.status_code == 200
    assert len(catalog.json()["capabilities"]) == 15
    sid = client.post("/api/research/sessions", json={}).json()["id"]
    assert client.get(f"/api/research/sessions/{sid}").json()["datasets"] == []
    called = []

    def provider(request):
        called.append(request)
        return httpx.Response(200, json=nav_page([nav_row("2025-12-31")], 1))

    service.datahub.transport = httpx.MockTransport(provider)
    payload = {
        "session_id": sid,
        "call_id": "api-test",
        "query": BusinessQuery(
            capability="fund_data", parameters={"dataset": "nav", "code": "000001"}
        ).model_dump(),
    }
    for path in ("business-query", "cancel"):
        assert client.post(f"/api/research/internal/data/{path}", json=payload).status_code == 403
    assert called == []
    headers = {"X-Research-Data-Key": service.datahub.control["token"]}
    response = client.post(
        "/api/research/internal/data/business-query", json=payload, headers=headers
    )
    assert response.status_code == 200, response.text
    dataset = response.json()
    did = dataset["dataset_id"]
    assert len(called) == 1
    prefix = f"/api/research/sessions/{sid}/datasets/{did}"
    assert client.get(prefix + "/rows?offset=0&limit=1").json()["total"] == 1
    assert client.get(prefix + "/rows?limit=501").status_code in {400, 422}
    detail = client.get(prefix).json()
    assert client.get(detail["files"][0]["url"]).status_code == 200
    assert client.get(f"/api/research/sessions/{sid}/files").json()["items"] == []
    assert len(client.get(f"/api/research/sessions/{sid}").json()["datasets"]) == 1
    source_text = f"请保留这段原文，继续分析 inputs/datasets/{did}/rows.json。"

    async def history_with_original_dataset(requested_sid):
        if requested_sid != sid:
            return []
        return [
            {
                "event": {
                    "seq": 1,
                    "type": "user/message",
                    "data": {
                        "id": "original-user-message",
                        "content": [{"type": "text", "text": source_text}],
                    },
                }
            }
        ]

    native.history = history_with_original_dataset
    service.loaded.discard(sid)
    upgraded = client.post(f"/api/research/sessions/{sid}/upgrade").json()
    new_sid = upgraded["id"]
    copied = client.get(f"/api/research/sessions/{new_sid}/datasets").json()["items"][0]
    assert copied["dataset_id"] != did
    copied_detail = client.get(
        f"/api/research/sessions/{new_sid}/datasets/{copied['dataset_id']}"
    ).json()
    assert copied_detail["origin_dataset_id"] == did
    assert copied_detail["retrieved_at"] == detail["retrieved_at"]
    assert copied_detail["origin_manifest_sha256"] == detail["manifest_sha256"]
    original_draft = "继续研究以下会话：\nuser: " + source_text
    assert upgraded["draft"].startswith(original_draft)
    handoff = upgraded["draft"][len(original_draft) :]
    assert f'"origin_dataset_id": "{did}"' in handoff
    assert f'"dataset_id": "{copied["dataset_id"]}"' in handoff
    assert copied_detail["origin_manifest_sha256"] in handoff
    assert copied_detail["manifest_sha256"] in handoff
    assert copied_detail["retrieved_at"] in handoff
    for item in copied_detail["files"]:
        assert item["path"] in handoff and item["sha256"] in handoff
        assert (service.store.directory(new_sid) / item["path"]).is_file()
    assert f"inputs/datasets/{did}/" not in handoff
    assert str(service.store.directory(sid)) not in handoff
    assert client.get(f"/api/research/sessions/{sid}").json()["messages"][0]["text"] == source_text
    assert client.get(copied_detail["files"][0]["url"]).status_code == 200
    assert client.get(prefix.replace(sid, new_sid)).status_code == 400
    assert len(called) == 1


def test_unknown_admission_not_retried(api):
    client, native, _ = api
    sid = client.post("/api/research/sessions", json={}).json()["id"]
    native.fail_prompt = True
    for _ in range(2):
        response = client.post(
            f"/api/research/sessions/{sid}/messages",
            json={"text": "你好"},
            headers={"Idempotency-Key": "abcdefgh"},
        )
        assert response.status_code == 503
    assert len([call for call in native.calls if call[0] == "session.prompt"]) == 1


@pytest.mark.parametrize("extension", ["docx", "pptx"])
def test_office_document_upload_preserves_bytes_and_session_ownership(api, extension):
    from io import BytesIO

    stream = BytesIO()
    if extension == "docx":
        from docx import Document

        document = Document()
        document.add_heading("Word 验收报告", level=0)
        document.add_paragraph("这段内容保持不变。")
        document.save(stream)
    else:
        import zipfile

        # Upload is byte transport, not an Office renderer. Reuse the minimal
        # slide-package fixture shape used by test_report_rendering; this does
        # not certify that PowerPoint opened or edited the fixture.
        with zipfile.ZipFile(stream, "w") as package:
            package.writestr(
                "ppt/slides/slide1.xml",
                '<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main" '
                'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">'
                "<p:cSld><p:spTree><p:sp><p:txBody><a:p><a:r><a:t>研究概览</a:t>"
                "</a:r></a:p></p:txBody></p:sp></p:spTree></p:cSld></p:sld>",
            )
    body = stream.getvalue()
    client, _, _ = api
    sid = client.post("/api/research/sessions", json={}).json()["id"]
    response = client.post(
        f"/api/research/sessions/{sid}/uploads",
        files=[("files", (f"中文 验收.{extension}", body, "application/octet-stream"))],
    )
    assert response.status_code == 200, response.text
    item = response.json()["items"][0]
    assert client.get(item["url"]).content == body
    other = client.post("/api/research/sessions", json={}).json()["id"]
    assert client.get(item["url"].replace(sid, other)).status_code == 400


def test_upload_ownership_and_preview_sandbox(api):
    client, _, _ = api
    sid = client.post("/api/research/sessions", json={}).json()["id"]
    uploaded = client.post(
        f"/api/research/sessions/{sid}/uploads",
        files=[("files", ("../note.md", b"# Notes", "text/markdown"))],
    )
    assert uploaded.status_code == 200, uploaded.text
    file = uploaded.json()["items"][0]
    preview = client.get(file["preview_url"])
    assert preview.status_code == 200
    assert "sandbox" in preview.headers["content-security-policy"]
    assert "allow-same-origin" not in preview.headers["content-security-policy"]
    other = client.post("/api/research/sessions", json={}).json()["id"]
    assert client.get(file["url"].replace(sid, other)).status_code == 400


def test_document_business_api_generates_download_and_reads_owned_word(api, monkeypatch):
    from io import BytesIO

    from docx import Document

    client, _, service = api

    async def unavailable_native(_session_id, _document):
        return {"outcome": "failed", "code": "native_document_executor_not_ready"}

    monkeypatch.setattr(service.local_integrations, "run_document", unavailable_native)
    sid = client.post("/api/research/sessions", json={}).json()["id"]
    endpoint = f"/api/research/sessions/{sid}/document-operations"
    body = {
        "format": "docx",
        "mode": "file",
        "operation": "generate",
        "content": {
            "title": "Word 验收报告",
            "paragraphs": ["这段内容保持不变。", "报告版本 A。"],
            "tables": [[["项目", "数值"], ["样本", "2"]]],
        },
    }
    response = client.post(endpoint, json=body)
    assert response.status_code == 200, response.text
    result = response.json()
    assert result["status"] == "completed", result
    downloaded = client.get(result["output"]["url"])
    assert Document(BytesIO(downloaded.content)).paragraphs[2].text == "报告版本 A。"
    read = client.post(
        endpoint,
        json={
            "format": "docx",
            "mode": "file",
            "operation": "read",
            "file_id": result["output"]["id"],
        },
    ).json()
    assert read["document"]["tables"][0][1][1] == "2"
    unavailable = client.post(endpoint, json={**body, "mode": "native"}).json()
    assert unavailable["status"] == "failed"
    assert unavailable["code"] == "native_document_executor_not_ready"
    assert unavailable["mode"] == "native"
    assert client.post(endpoint, json={**body, "path": "/etc/passwd"}).status_code == 422


def test_cross_origin_and_invalid_credentials_do_not_leak_input(api):
    client, _, _ = api
    assert (
        client.post(
            "/api/research/sessions", json={}, headers={"Origin": "https://evil.test"}
        ).status_code
        == 403
    )
    response = client.put(
        "/api/research/runtime/model", json={"provider": "bad", "api_key": "secret-canary"}
    )
    assert response.status_code == 422
    assert "secret-canary" not in response.text


def test_no_dsh_connection_no_fake_success(api):
    client, _, service = api
    sid = client.post("/api/research/sessions", json={}).json()["id"]
    service.connected.clear()
    response = client.post(
        f"/api/research/sessions/{sid}/messages",
        json={"text": "问题"},
        headers={"Idempotency-Key": "abcdefgh"},
    )
    assert response.status_code == 503


def test_model_selection_is_restored(tmp_path):
    store = Store(tmp_path)
    store.data["model"] = {"provider": "deepseek-official", "model": "custom-model"}
    store.save()
    assert (
        ResearchService(NativeFixture(), Store(tmp_path)).default_model["model"] == "custom-model"
    )


@pytest.mark.asyncio
async def test_unowned_instance_is_read_only(tmp_path):
    service = ResearchService(NativeFixture(), Store(tmp_path), expected_cwd=tmp_path / "owned")
    with pytest.raises(RuntimeFailure, match="专属"):
        await service.create()
    assert not any(method == "session.create" for method, _ in service.client.calls)


def test_skill_catalog_resumes_native_session_after_runtime_restart(api):
    client, native, _ = api
    sid = client.post("/api/research/sessions", json={}).json()["id"]
    original = native.rpc
    resumed = False

    async def cold_rpc(method, payload):
        nonlocal resumed
        if method == "session.models":
            assert payload == {"sessionId": sid}
            resumed = True
        if method == "skill.list" and not resumed:
            raise RuntimeFailure("not attached", "session-not-found")
        return await original(method, payload)

    native.rpc = cold_rpc
    response = client.get("/api/research/skills")
    assert response.status_code == 200, response.text
    assert resumed


def test_child_activity_comes_from_native_parent_scoped_history(api):
    client, native, _ = api
    sid = client.post("/api/research/sessions", json={"mode": "claw"}).json()["id"]
    original = native.rpc

    async def child_rpc(method, payload):
        if method == "subagent.list":
            return {
                "entries": [
                    {
                        "kind": "child",
                        "id": "child-1",
                        "mode": "continuable",
                        "label": "行业分析",
                        "activity": "inactive",
                    }
                ]
            }
        if method == "subagent.history":
            assert payload["parentSessionId"] == sid
            assert payload["childSessionId"] == "child-1"
            return {
                "events": [
                    {
                        "event": {
                            "seq": 1,
                            "type": "tool/call",
                            "data": {"callId": "tool-1", "name": "web_search", "arguments": "{}"},
                        }
                    },
                    {
                        "event": {
                            "seq": 2,
                            "type": "turn/end",
                            "data": {
                                "reason": {
                                    "kind": "error",
                                    "error": {"message": "search unavailable"},
                                }
                            },
                        }
                    },
                ],
                "hasMore": False,
            }
        return await original(method, payload)

    native.rpc = child_rpc
    detail = client.get(f"/api/research/sessions/{sid}").json()
    assert detail["subagents"][0]["status"] == "failed"
    assert detail["subagents"][0]["error"] == "search unavailable"
    assert detail["activities"][0]["agent_id"] == "child-1"


def test_active_child_keeps_parent_stoppable_in_detail_and_history(api):
    client, native, service = api
    sid = client.post("/api/research/sessions", json={"mode": "claw"}).json()["id"]
    service.store.session(sid)["status"] = "completed"
    original = native.rpc

    async def child_rpc(method, payload):
        if method == "subagent.list":
            return {
                "entries": [
                    {
                        "kind": "child",
                        "id": "active-child",
                        "mode": "continuable",
                        "label": "分析",
                        "activity": "running",
                    }
                ]
            }
        if method == "subagent.history":
            return {"events": [], "hasMore": False}
        return await original(method, payload)

    native.rpc = child_rpc
    assert client.get(f"/api/research/sessions/{sid}").json()["status"] == "running"
    assert client.get("/api/research/sessions").json()["items"][0]["status"] == "running"


def test_session_list_bounds_child_queries_and_preserves_store_order(api):
    client, native, service = api
    session_ids = []
    for index in range(50):
        row = service.store.create("fingpt", f"研究 {index}")
        row["created"] = True
        row["status"] = "idle"
        row["updated_at"] = index
        session_ids.append(row["id"])
    not_created = service.store.create("fingpt", "未创建研究")
    not_created["updated_at"] = 50
    deleted = service.store.create("fingpt", "已删除研究")
    deleted["created"] = True
    deleted["updated_at"] = 51
    deleted["deleted_at"] = time.time()
    target = session_ids[24]
    service.store.session(target)["status"] = "completed"
    service.store.save()

    original = native.rpc
    active = 0
    max_active = 0
    queried_parents = []

    async def bounded_rpc(method, payload):
        nonlocal active, max_active
        if method == "session.list":
            return {"items": []}
        if method == "subagent.list":
            queried_parents.append(payload["parentSessionId"])
            active += 1
            max_active = max(max_active, active)
            try:
                await asyncio.sleep(0.01)
                entries = (
                    [{"id": "child", "activity": "running"}]
                    if payload["parentSessionId"] == target
                    else []
                )
                return {"entries": entries}
            finally:
                active -= 1
        return await original(method, payload)

    native.rpc = bounded_rpc
    response = client.get("/api/research/sessions?view=all")

    assert response.status_code == 200, response.text
    items = response.json()["items"]
    assert [item["id"] for item in items] == [
        deleted["id"],
        not_created["id"],
        *reversed(session_ids),
    ]
    assert next(item for item in items if item["id"] == target)["status"] == "running"
    assert all(item["status"] == "idle" for item in items if item["id"] != target)
    assert sorted(queried_parents) == sorted(session_ids)
    assert not_created["id"] not in queried_parents
    assert deleted["id"] not in queried_parents
    assert 1 < max_active <= 8


@pytest.mark.asyncio
async def test_session_list_reconciles_only_created_idle_stale_running_rows(tmp_path, monkeypatch):
    native = NativeFixture()
    service = ResearchService(native, Store(tmp_path))
    uncreated = service.store.create("fingpt", "未创建但缓存运行")
    native_running = service.store.create("fingpt", "原生仍在运行")
    child_running = service.store.create("fingpt", "子任务仍在运行")
    eligible = service.store.create("fingpt", "应恢复的陈旧运行会话")
    rows = [uncreated, native_running, child_running, eligible]
    for index, row in enumerate(rows):
        row["status"] = "running"
        row["updated_at"] = index
    for row in (native_running, child_running, eligible):
        row["created"] = True
    service.store.save()

    original = native.rpc

    async def eligibility_rpc(method, payload):
        if method == "session.list":
            return {"items": [{"sessionId": native_running["id"], "running": True}]}
        if method == "subagent.list":
            entries = (
                [{"id": "active-child", "activity": "running"}]
                if payload["parentSessionId"] == child_running["id"]
                else []
            )
            return {"entries": entries}
        return await original(method, payload)

    detail_calls = []

    async def tracked_detail(sid):
        detail_calls.append(sid)
        if sid == uncreated["id"]:
            raise RuntimeFailure("该会话未成功创建，请新建研究", "session_create_failed")
        return {"status": "completed"}

    native.rpc = eligibility_rpc
    monkeypatch.setattr(service, "detail", tracked_detail)
    items = await service.list_sessions()

    assert [item["id"] for item in items] == [row["id"] for row in reversed(rows)]
    statuses = {item["id"]: item["status"] for item in items}
    assert statuses[uncreated["id"]] == "running"
    assert statuses[native_running["id"]] == "running"
    assert statuses[child_running["id"]] == "running"
    assert statuses[eligible["id"]] == "completed"
    assert uncreated["id"] not in detail_calls
    assert native_running["id"] not in detail_calls
    assert child_running["id"] not in detail_calls
    assert detail_calls == [eligible["id"]]


@pytest.mark.asyncio
async def test_session_list_reconciles_stale_running_details_concurrently_in_order(
    tmp_path, monkeypatch
):
    native = NativeFixture()
    service = ResearchService(native, Store(tmp_path))
    session_ids = []
    terminal_statuses = {}
    for index in range(10):
        row = service.store.create("fingpt", f"陈旧运行会话 {index}")
        row["created"] = True
        row["status"] = "running"
        row["updated_at"] = index
        session_ids.append(row["id"])
        terminal_statuses[row["id"]] = "completed" if index % 2 == 0 else "failed"
    service.store.save()

    active = 0
    max_active = 0
    finished = set()
    delay_seconds = 0.04

    async def delayed_detail(sid):
        nonlocal active, max_active
        active += 1
        max_active = max(max_active, active)
        try:
            await asyncio.sleep(delay_seconds)
            return {"status": terminal_statuses[sid]}
        finally:
            active -= 1
            finished.add(sid)

    monkeypatch.setattr(service, "detail", delayed_detail)
    started = time.monotonic()
    items = await service.list_sessions()
    elapsed = time.monotonic() - started

    concurrency_ok = 1 < max_active <= 8
    duration_ok = elapsed < delay_seconds * 6
    assert concurrency_ok and duration_ok, (
        f"max_active={max_active}, elapsed={elapsed:.3f}s, " f"limit={delay_seconds * 6:.3f}s"
    )
    assert 1 < max_active <= 8
    assert [item["id"] for item in items] == list(reversed(session_ids))
    assert [item["status"] for item in items] == [
        terminal_statuses[sid] for sid in reversed(session_ids)
    ]
    assert active == 0
    assert finished == set(session_ids)


@pytest.mark.asyncio
async def test_session_list_cancels_stale_detail_tasks_before_reconciliation_failure(
    tmp_path, monkeypatch
):
    native = NativeFixture()
    service = ResearchService(native, Store(tmp_path))
    rows = [service.store.create("fingpt", f"陈旧运行会话 {index}") for index in range(10)]
    for index, row in enumerate(rows):
        row["created"] = True
        row["status"] = "running"
        row["updated_at"] = index
    service.store.save()

    failed_sid = rows[-1]["id"]
    active = 0
    started = set()
    completed = set()
    cancelled = set()

    async def failing_detail(sid):
        nonlocal active
        active += 1
        started.add(sid)
        try:
            if sid == failed_sid:
                await asyncio.sleep(0.01)
                raise RuntimeFailure("detail reconciliation failed", "detail_reconciliation_failed")
            await asyncio.sleep(0.2)
            completed.add(sid)
            return {"status": "completed"}
        except asyncio.CancelledError:
            cancelled.add(sid)
            raise
        finally:
            active -= 1

    monkeypatch.setattr(service, "detail", failing_detail)
    with pytest.raises(RuntimeFailure, match="detail reconciliation failed"):
        await service.list_sessions()
    active_after_failure = active
    completed_after_failure = set(completed)
    await asyncio.sleep(0.3)

    assert len(started) > 1
    assert active_after_failure == 0
    assert completed_after_failure == set()
    assert active == 0
    assert completed == set()
    assert cancelled == started - {failed_sid}


def test_session_list_cancels_slow_child_lookups_before_runtime_failure_response(api):
    client, native, service = api
    rows = [service.store.create("fingpt", f"研究 {index}") for index in range(20)]
    for index, row in enumerate(rows):
        row["created"] = True
        row["updated_at"] = index
    service.store.save()
    failed_parent = rows[-1]["id"]
    original = native.rpc
    active = 0
    completed = []

    async def failing_rpc(method, payload):
        nonlocal active
        if method == "session.list":
            return {"items": []}
        if method == "subagent.list":
            parent_id = payload["parentSessionId"]
            active += 1
            try:
                if parent_id == failed_parent:
                    await asyncio.sleep(0.01)
                    raise RuntimeFailure("child lookup failed", "child_lookup_failed")
                await asyncio.sleep(0.1)
                completed.append(parent_id)
                return {"entries": []}
            finally:
                active -= 1
        return await original(method, payload)

    native.rpc = failing_rpc
    response = client.get("/api/research/sessions")
    active_after_response = active
    completed_after_response = list(completed)
    time.sleep(0.4)

    assert response.status_code == 503
    assert response.json() == {
        "error": {"code": "child_lookup_failed", "message": "child lookup failed"}
    }
    assert active_after_response == 0
    assert completed_after_response == []
    assert active == 0
    assert completed == []


@pytest.mark.asyncio
async def test_model_configuration_serializes_with_session_creation(tmp_path):
    native = NativeFixture()
    service = ResearchService(native, Store(tmp_path))
    await service.create()
    entered, release = asyncio.Event(), asyncio.Event()
    original = native.rpc

    async def blocked_rpc(method, payload):
        if method == "llm.models" and not release.is_set():
            entered.set()
            await release.wait()
        return await original(method, payload)

    native.rpc = blocked_rpc
    configuring = asyncio.create_task(service.configure_model("deepseek-official", "new-model"))
    await asyncio.wait_for(entered.wait(), 1)
    creating = asyncio.create_task(service.create())
    await asyncio.sleep(0)
    assert len(service.store.data["sessions"]) == 1
    assert service.default_model["model"] != "new-model"
    release.set()
    await asyncio.gather(configuring, creating)
    assert len(service.store.data["sessions"]) == 2
    assert service.store.data["model"]["model"] == "new-model"


@pytest.mark.asyncio
async def test_model_save_does_not_reselect_existing_sessions(tmp_path):
    native = NativeFixture()
    service = ResearchService(native, Store(tmp_path))
    old = await service.create()
    native.calls.clear()
    result = await service.configure_model("deepseek-official", "new-model")
    assert result["configured"] is True
    assert not any(method == "session.selectModel" for method, _ in native.calls)
    assert service.store.session(old["id"])["model"] == "deepseek-flash"
    assert (await service.create())["model"] == "new-model"


@pytest.mark.asyncio
async def test_runtime_reachability_does_not_claim_model_application(tmp_path):
    native = NativeFixture()
    service = ResearchService(native, Store(tmp_path))
    service.connected = {"mux", "host"}
    await service.configure_model("deepseek-official", "new-model")
    assert (await service.runtime())["configuration_saved"] is True
    assert (await service.runtime())["runtime_applied"] is False
    await service.create()
    assert (await service.runtime())["runtime_applied"] is True


@pytest.mark.asyncio
async def test_model_save_refuses_busy_and_unknown_model_before_credentials(tmp_path):
    native = NativeFixture()
    service = ResearchService(native, Store(tmp_path))
    native.running.add("active-child")
    with pytest.raises(RuntimeFailure, match="运行"):
        await service.configure_model("deepseek-official", "new-model", "fixture-key")
    assert not any(method == "credentials.set" for method, _ in native.calls)
    native.running.clear()
    native.calls.clear()
    with pytest.raises(RuntimeFailure) as error:
        await service.configure_model("deepseek-official", "unknown", "fixture-key")
    assert error.value.code == "model_unavailable"
    assert not any(method == "credentials.set" for method, _ in native.calls)


@pytest.mark.asyncio
async def test_model_failed_credential_write_restores_saved_default(tmp_path):
    native = NativeFixture()
    service = ResearchService(native, Store(tmp_path))
    original = native.rpc

    async def rejected(method, payload):
        if method == "credentials.set":
            raise RuntimeFailure("rejected", "credential/rejected")
        return await original(method, payload)

    native.rpc = rejected
    with pytest.raises(RuntimeFailure):
        await service.configure_model("deepseek-official", "new-model", "fixture-key")
    assert service.default_model["model"] == "deepseek-flash"
    assert Store(tmp_path).data.get("model", {}).get("model", "deepseek-flash") == "deepseek-flash"
    assert Store(tmp_path).data.get("model_configuration_uncertain") is True


@pytest.mark.asyncio
async def test_read_only_credential_refusal_preserves_usable_configuration(tmp_path):
    native = NativeFixture()
    service = ResearchService(native, Store(tmp_path))
    original = native.rpc

    async def read_only(method, payload):
        if method == "credentials.describe":
            return {
                "credentials": {"RESEARCH_DSH_API_KEY": {"configured": True, "writable": False}}
            }
        return await original(method, payload)

    native.rpc = read_only
    with pytest.raises(RuntimeFailure) as error:
        await service.configure_model("deepseek-official", "new-model", "fixture-key")
    assert error.value.code == "model_credentials_read_only"
    assert not service.store.data.get("model_configuration_uncertain")
    assert service.default_model["model"] == "deepseek-flash"
    assert not any(method == "credentials.set" for method, _ in native.calls)


def test_model_unknown_id_and_clear_contract(api):
    client, native, _ = api
    response = client.put(
        "/api/research/runtime/model", json={"model": "unknown", "api_key": "fixture-key"}
    )
    assert response.status_code == 503
    assert response.json()["error"]["code"] == "model_unavailable"
    assert not any(method == "credentials.set" for method, _ in native.calls)
    response = client.put(
        "/api/research/runtime/model", json={"model": "deepseek-flash", "clear_api_key": True}
    )
    assert response.status_code == 200
    assert ("credentials.unset", {"ref": "RESEARCH_DSH_API_KEY"}) in native.calls
    response = client.put(
        "/api/research/runtime/model",
        json={"model": "deepseek-flash", "clear_api_key": True, "api_key": "fixture-key"},
    )
    assert response.status_code == 422


def test_model_masks_are_rejected_before_credential_write(api):
    client, native, _ = api
    for placeholder in ("********", "••••••••", "[REDACTED]"):
        response = client.put(
            "/api/research/runtime/model",
            json={"model": "deepseek-flash", "api_key": placeholder},
        )
        assert response.status_code == 422
        assert response.json()["error"]["code"] == "invalid_request"
        assert not any(method == "credentials.set" for method, _ in native.calls)


def test_empty_credential_blocks_regular_research_before_native_prompt(api):
    client, native, _ = api
    sid = client.post("/api/research/sessions", json={}).json()["id"]
    original = native.rpc

    async def no_credential(method, payload):
        if method == "credentials.describe":
            return {
                "credentials": {"RESEARCH_DSH_API_KEY": {"configured": False, "writable": True}}
            }
        return await original(method, payload)

    native.rpc = no_credential
    response = client.post(
        f"/api/research/sessions/{sid}/messages",
        json={"text": "非业务测试", "formats": []},
        headers={"Idempotency-Key": "closeout-missing-key"},
    )
    assert response.status_code == 503
    assert response.json()["error"]["code"] == "model_credentials_missing"
    assert not any(method == "session.prompt" for method, _ in native.calls)


@pytest.mark.asyncio
async def test_credential_value_selection_clear_and_cold_recovery_are_instance_local(tmp_path):
    class MemoryNative(NativeFixture):
        def __init__(self):
            super().__init__()
            self.value = None
            self.writes = 0
            self.lose_next_response = False

        async def rpc(self, method, payload):
            if method == "credentials.describe":
                return {
                    "credentials": {
                        "RESEARCH_DSH_API_KEY": {
                            "configured": self.value is not None,
                            "writable": True,
                        }
                    }
                }
            if method == "credentials.set":
                self.value = payload["value"]
                self.writes += 1
                if self.lose_next_response:
                    self.lose_next_response = False
                    raise RuntimeFailure("fixture response lost")
            if method == "credentials.unset":
                self.value = None
            return await super().rpc(method, payload)

    def settings_service(root, native):
        # This regression uses only the existing settings/admission state machine,
        # avoiding unrelated capability catalog construction and vendor access.
        service = object.__new__(ResearchService)
        service.store = Store(root)
        service.client = native
        service.owned = True
        service.expected_cwd = None
        service.lock = asyncio.Lock()
        service.model_test_lock = asyncio.Lock()
        service.running = {}
        service.default_model = service.store.data.get(
            "model", {"provider": "deepseek-official", "model": "deepseek-flash"}
        )
        return service

    native_a, native_b = MemoryNative(), MemoryNative()
    a = settings_service(tmp_path / "a", native_a)
    b = settings_service(tmp_path / "b", native_b)
    old_sid = a.store.create("fingpt", "fixture old session")["id"]
    await a.configure_model("deepseek-official", "deepseek-flash", "synthetic-a")
    await b.configure_model("deepseek-official", "deepseek-flash", "synthetic-b")
    await a.configure_model("deepseek-official", "deepseek-flash")
    assert native_a.value == "synthetic-a" and native_a.writes == 1
    await a.configure_model("deepseek-official", "deepseek-flash", "synthetic-replacement")
    assert native_a.value == "synthetic-replacement"
    assert native_b.value == "synthetic-b"
    native_a.lose_next_response = True
    with pytest.raises(RuntimeFailure) as error:
        await a.configure_model("deepseek-official", "deepseek-flash", "synthetic-unknown")
    assert error.value.code == "model_configuration_uncertain"
    cold = settings_service(tmp_path / "a", native_a)
    with pytest.raises(RuntimeFailure):
        await cold.configure_model("deepseek-official", "deepseek-flash")
    await cold.configure_model("deepseek-official", "deepseek-flash", clear_api_key=True)
    assert native_a.value is None and native_b.value == "synthetic-b"
    cold = settings_service(tmp_path / "a", native_a)
    new_sid = cold.store.create("fingpt", "fixture new session")["id"]
    for sid in (old_sid, new_sid):
        with pytest.raises(RuntimeFailure) as error:
            await cold.send(sid, "fixture", "new-submission")
        assert error.value.code == "model_credentials_missing"
    await cold.configure_model("deepseek-official", "deepseek-flash", "synthetic-restored")
    assert native_a.value == "synthetic-restored"
    assert not cold.store.data.get("model_configuration_uncertain")
    assert cold.store.data["model_credential_cleared"] is False
    assert not any(method == "session.prompt" for method, _ in native_a.calls)


@pytest.mark.asyncio
async def test_runtime_reports_actual_credential_source_without_file_assumption(tmp_path):
    native = NativeFixture()
    service = object.__new__(ResearchService)
    service.client = native
    service.store = Store(tmp_path)
    service.default_model = {"provider": "deepseek-official", "model": "deepseek-flash"}
    service.expected_cwd = None
    service.owned = True
    service.connected = {"mux", "host"}
    service.last_runtime_success_at = None
    original = native.rpc

    async def environment_source(method, payload):
        if method == "credentials.describe":
            return {
                "credentials": {
                    "RESEARCH_DSH_API_KEY": {"configured": True, "source": "env", "writable": False}
                }
            }
        return await original(method, payload)

    native.rpc = environment_source
    assert (await service.runtime())["credential_storage"] == "environment"


@pytest.mark.asyncio
async def test_model_backend_failure_preserves_runtime_health_and_unknown_credential(tmp_path):
    native = NativeFixture()
    service = object.__new__(ResearchService)
    service.client = native
    service.store = Store(tmp_path)
    service.default_model = {"provider": "deepseek-official", "model": "deepseek-flash"}
    service.expected_cwd = None
    service.owned = True
    service.connected = {"mux", "host"}
    service.last_runtime_success_at = None
    original = native.rpc

    async def unavailable(method, payload):
        if method == "credentials.describe":
            raise RuntimeFailure("凭据服务不可用", "credential/rejected")
        return await original(method, payload)

    native.rpc = unavailable
    status = await service.runtime()
    assert status["connected"] is True
    assert status["health_check_passed"] is True
    assert status["credential_configured"] is None
    assert status["credential_storage"] == "unknown"
    assert status["credential_code"] == "model_credential_backend_unavailable"


@pytest.mark.asyncio
async def test_model_ambiguous_write_blocks_new_requests(tmp_path):
    native = NativeFixture()
    service = ResearchService(native, Store(tmp_path))
    original = native.rpc

    async def disconnected(method, payload):
        if method == "credentials.set":
            raise RuntimeFailure("response lost")
        return await original(method, payload)

    native.rpc = disconnected
    with pytest.raises(RuntimeFailure) as error:
        await service.configure_model("deepseek-official", "new-model", "fixture-key")
    assert error.value.code == "model_configuration_uncertain"
    assert (await service.runtime())["runtime_applied"] is False
    with pytest.raises(RuntimeFailure) as error:
        await service.test_model()
    assert error.value.code == "model_configuration_uncertain"


@pytest.mark.asyncio
async def test_model_smoke_requires_native_final_text(tmp_path):
    native = NativeFixture()
    service = ResearchService(native, Store(tmp_path))
    service.connected = {"mux", "host"}

    async def history(_sid):
        return [
            {
                "event": {
                    "seq": 1,
                    "type": "assistant/message",
                    "data": {
                        "turn": 1,
                        "step": 1,
                        "message": {"content": [{"type": "text", "text": "模型生成测试完成"}]},
                    },
                }
            },
            {"event": {"seq": 2, "type": "turn/end", "data": {"reason": {"kind": "completed"}}}},
        ]

    native.history = history
    assert (await service.test_model())["status"] == "passed"
    assert service.store.data["model_test"]["status"] == "passed"

    async def empty_history(_sid):
        return [
            {"event": {"seq": 2, "type": "turn/end", "data": {"reason": {"kind": "completed"}}}}
        ]

    native.history = empty_history
    assert (await service.test_model())["status"] == "failed"


@pytest.mark.asyncio
async def test_model_cancelled_write_is_uncertain_after_cold_reload(tmp_path):
    native = NativeFixture()
    service = ResearchService(native, Store(tmp_path))
    entered = asyncio.Event()
    original = native.rpc

    async def suspended(method, payload):
        if method == "credentials.set":
            entered.set()
            await asyncio.Event().wait()
        return await original(method, payload)

    native.rpc = suspended
    task = asyncio.create_task(
        service.configure_model("deepseek-official", "new-model", "fixture-key")
    )
    await entered.wait()
    task.cancel()
    with pytest.raises(asyncio.CancelledError):
        await task
    assert Store(tmp_path).data.get("model_configuration_uncertain") is True
    assert Store(tmp_path).data.get("model", service.default_model)["model"] == "deepseek-flash"


@pytest.mark.asyncio
async def test_model_test_cannot_admit_during_configuration(tmp_path):
    native = NativeFixture()
    service = ResearchService(native, Store(tmp_path))
    entered, release = asyncio.Event(), asyncio.Event()
    original = native.rpc

    async def suspended(method, payload):
        if method == "llm.models":
            entered.set()
            await release.wait()
        return await original(method, payload)

    native.rpc = suspended
    task = asyncio.create_task(service.configure_model("deepseek-official", "new-model"))
    await entered.wait()
    try:
        with pytest.raises(RuntimeFailure) as error:
            await asyncio.wait_for(service.test_model(), 0.2)
        assert error.value.code == "model_test_busy"
    finally:
        release.set()
        await task


@pytest.mark.asyncio
async def test_model_smoke_rejects_partial_text_with_empty_final_message(tmp_path):
    native = NativeFixture()
    service = ResearchService(native, Store(tmp_path))
    service.connected = {"mux", "host"}

    async def history(_sid):
        return [
            {
                "event": {
                    "seq": 1,
                    "type": "assistant/chunk",
                    "data": {
                        "turn": 1,
                        "step": 1,
                        "chunk": {"type": "text-delta", "index": 0, "text": "partial"},
                    },
                }
            },
            {
                "event": {
                    "seq": 2,
                    "type": "assistant/message",
                    "data": {"turn": 1, "step": 2, "message": {"content": []}},
                }
            },
            {"event": {"seq": 3, "type": "turn/end", "data": {"reason": {"kind": "completed"}}}},
        ]

    native.history = history
    assert (await service.test_model())["status"] == "failed"


def test_native_question_reply_requires_ownership_and_complete_answers(api):
    client, native, service = api
    sid = client.post("/api/research/sessions", json={}).json()["id"]
    other = client.post("/api/research/sessions", json={}).json()["id"]
    service.questions["question-rpc"] = {
        "sessionId": sid,
        "questions": [{"id": "period", "question": "研究期间？"}],
    }
    replies = []

    async def respond(rpc_id, value):
        replies.append((rpc_id, value))
        return {"accepted": True}

    native.respond = respond
    endpoint = f"/api/research/sessions/{sid}/questions/question-rpc"
    body = {"answers": [{"id": "period", "selected": [], "custom": "2025"}]}
    assert client.post(endpoint.replace(sid, other), json=body).status_code == 400
    assert (
        client.post(endpoint, json={"answers": [{"id": "wrong", "custom": "2025"}]}).status_code
        == 400
    )
    assert not replies
    assert client.post(endpoint, json=body).status_code == 200
    assert replies == [("question-rpc", {"sessionId": sid, "answer": body})]


def test_image_attachment_is_forwarded_as_native_image_not_just_a_path(api):
    import base64

    client, native, _ = api
    sid = client.post("/api/research/sessions", json={}).json()["id"]
    image_bytes = b"\x89PNG\r\n\x1a\nfixture"
    uploaded = client.post(
        f"/api/research/sessions/{sid}/uploads",
        files=[("files", ("chart.png", image_bytes, "image/png"))],
    ).json()["items"][0]
    response = client.post(
        f"/api/research/sessions/{sid}/messages",
        json={"text": "解读图片", "attachment_ids": [uploaded["id"]]},
        headers={"Idempotency-Key": "image-request"},
    )
    assert response.status_code == 202
    prompt = next(payload for method, payload in native.calls if method == "session.prompt")
    assert prompt["content"][1]["type"] == "image"
    assert base64.b64decode(prompt["content"][1]["data"]) == image_bytes


def test_each_claw_turn_appends_the_visible_language_contract_last(api):
    client, native, _ = api
    sid = client.post("/api/research/sessions", json={"mode": "claw"}).json()["id"]

    response = client.post(
        f"/api/research/sessions/{sid}/messages",
        json={"text": "分析市场", "tool_ids": ["research_run_script"]},
        headers={"Idempotency-Key": "language-contract"},
    )

    assert response.status_code == 202
    prompt = next(payload for method, payload in native.calls if method == "session.prompt")
    text = prompt["content"][0]["text"]
    assert text.index("用户选择的研究工具意图") < text.index("使用 DSH 原生子 Agent")
    assert text.rstrip().endswith("不得直接以英文回答透传。")
    assert "所有可见过程说明、工具调用前后说明、提问、错误解释、总结和最终答复" in text


def test_upgraded_empty_instance_uses_catalog_default_without_rewriting_saved_model(tmp_path):
    empty = ResearchService(NativeFixture(), Store(tmp_path / "empty"))
    assert empty.default_model == {"provider": "deepseek-official", "model": "deepseek-flash"}
    old = Store(tmp_path / "old")
    old.data["model"] = {"provider": "deepseek-official", "model": "deepseek-v4-flash"}
    old.save()
    preserved = ResearchService(NativeFixture(), old)
    assert preserved.default_model["model"] == "deepseek-v4-flash"
