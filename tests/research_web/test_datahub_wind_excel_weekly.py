"""Offline Excel inputs; synthetic workdays never certify real market semantics."""

import asyncio
import hashlib
import json
import re
import threading
from datetime import UTC, date, datetime, timedelta
from io import BytesIO
from types import SimpleNamespace
from unittest.mock import MagicMock, patch
from zipfile import ZipFile

import pytest
from openpyxl import Workbook

from app.research_web.datahub import providers_wind as shared
from app.research_web.datahub import providers_wind_excel_weekly as excel_provider
from app.research_web.datahub.providers import ProviderError
from app.research_web.datahub.providers_wind_excel_weekly import (
    WindExcelWeeklyProvider,
    read_exports,
)
from app.research_web.report_workflows.models import WeeklyProductionPolicy, WorkflowError
from app.research_web.report_workflows.weekly_data import (
    MAPPING_SHA256,
    build_weekly_data_contract,
    load_weekly_mapping,
)
from app.research_web.report_workflows.weekly_sources import (
    SOURCE_CONTRACT,
    load_run_source_contract,
)
from app.research_web.store import StoreError


@pytest.fixture(autouse=True)
def synthetic_live_acquisition_clock(monkeypatch):
    # Transport fixtures represent an actual collection on the report Friday,
    # independently of when this offline regression happens to execute.
    monkeypatch.setattr(
        excel_provider,
        "_host_now",
        lambda: datetime.fromisoformat("2026-10-16T16:00:00+08:00"),
        raising=False,
    )


def days(start, end):
    cursor = date.fromisoformat(start)
    while cursor.isoformat() <= end:
        if cursor.weekday() < 5:
            yield cursor.isoformat()
        cursor += timedelta(days=1)


def workbook_bytes(rows):
    book = Workbook()
    for row in rows:
        book.active.append(row)
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def native_export_observation(export_sha, contract, *, start=None, end=None):
    start = start or contract["next_calendar_start"]
    end = end or contract["next_calendar_end"]
    pages = {}
    for key, label in (("economic_data", "经济数据"), ("financial_events", "财经大事")):
        empty = key == "financial_events"
        message = "当前日期暂无公布日程，请切换至其他日期。" if empty else None
        pages[key] = {
            "sheet": None if empty else "Sheet",
            "result": "empty" if empty else "exported",
            "filters": {
                "countries": "all",
                "event_types": "all",
                "importance": ["重要", "中等", "低"],
            },
            "message": message,
            "raw_accessibility_excerpt": f"合成夹具 {label} {start} 至 {end} 国家/地区:全部 事件类型:全部 重要 中等 低 {message or ''}",
        }
    return {
        "schema_version": 1,
        "source": "wind_excel",
        "kind": "native_export_observation",
        "application": "com.wind.mac.Windotd",
        "source_url": "http://114.80.154.45/eco.web/index.html?lan=cn#/",
        "export_path": "inputs/calendar.xlsx",
        "export_sha256": export_sha,
        "coverage_start": start,
        "coverage_end": end,
        "observed_at": contract["cutoff"],
        "pages": pages,
    }


def package(root, cutoff="2026-10-16T16:00:00+08:00"):
    contract = build_weekly_data_contract(cutoff)
    resources = []

    def save(path, raw):
        target = root / path
        target.parent.mkdir(parents=True, exist_ok=True)
        target.write_bytes(raw)
        sha = hashlib.sha256(raw).hexdigest()
        resources.append({"path": path, "sha256": sha})
        return sha

    edb = [spec for spec in load_weekly_mapping()["series"] if spec["kind"] == "edb"]
    rows = [
        ["国家", "英国", "中国", "英国", "美国"],
        ["指标名称", "金", "金", "油", "油"],
        ["频率", "日", "日", "日", "日"],
        ["单位", "USD/盎司", "元/克", "USD/桶", "USD/桶"],
        ["指标ID", *[spec["asset"] for spec in edb]],
        ["时间区间"],
        ["来源", "Wind"],
        ["更新时间", "2026-10-16"],
    ]
    rows += [
        [date.fromisoformat(day), 100, 100, 100, 100]
        for day in days("2010-01-01", contract["period_end"])
    ]
    sha = save("inputs/edb.xlsx", workbook_bytes(rows))
    economic_sha = save(
        "inputs/calendar.xlsx",
        workbook_bytes(
            [
                ["日期", "时间", "国家/地区", "指标名称", "重要性", "前值", "预测值", "今值"],
                [
                    date.fromisoformat(contract["next_calendar_start"]),
                    "09:00",
                    "中国",
                    "合成事件",
                    "重要",
                    None,
                    None,
                    None,
                ],
                ["数据来源: Wind"],
            ]
        ),
    )
    proof = {
        "schema_version": 1,
        "source": "wind_excel",
        "windows": [
            {
                "start": contract["next_calendar_start"],
                "end": contract["next_calendar_end"],
                "resource": "inputs/calendar.xlsx",
                "sha256": economic_sha,
                "kind": "native_export",
            }
        ],
    }
    proof_sha = save("inputs/coverage.json", json.dumps(proof).encode())
    observation_sha = save(
        "inputs/calendar-observation.json",
        json.dumps(native_export_observation(economic_sha, contract)).encode(),
    )
    source = {
        "schema_version": 2,
        "data_source": "wind_excel",
        "mapping_sha256": MAPPING_SHA256,
        "verification": {
            "status": "reviewed",
            "evidence_path": "inputs/review.txt",
            "evidence_sha256": save("inputs/review.txt", b"synthetic review only"),
        },
        "calendars": {
            market: {
                "semantics_verified": True,
                "wind_calendar": "SSE",
                "timezone": "Asia/Shanghai",
                "close_time": "10:00",
            }
            for market in load_weekly_mapping()["calendars"]
        },
        "economic_calendar": {
            "method": "wind_excel_export",
            "importance_values": {"重要": "重要", "中等": "较重要", "低": "一般"},
        },
        "excel_exports": {
            "edb": {
                "path": "inputs/edb.xlsx",
                "sha256": sha,
                "exported_at": contract["cutoff"],
                "sheet": "Sheet",
            },
            "economic_calendar": {
                "path": "inputs/calendar.xlsx",
                "sha256": economic_sha,
                "exported_at": contract["cutoff"],
                "sheet": "Sheet",
                "coverage_start": contract["next_calendar_start"],
                "coverage_end": contract["next_calendar_end"],
                "coverage_evidence_path": "inputs/coverage.json",
                "coverage_evidence_sha256": proof_sha,
                "observation_path": "inputs/calendar-observation.json",
                "observation_sha256": observation_sha,
            },
        },
    }
    save(SOURCE_CONTRACT, json.dumps(source).encode())
    return contract, source, {"resources": resources}, save


class Client:
    def __init__(self):
        self.calls = []
        self.closed = False

    def _ensure_session(self):
        pass

    def _execute_wsd_once(self, formula, timeout, expected_columns):
        if expected_columns == 2:
            return self.execute_wsd(*re.findall(r'"([^"]+)"', formula))
        assert expected_columns == 1
        start, end = re.findall(r'"([0-9-]+)"', formula)
        self.calls.append(("calendar", start, end))
        return [[date.fromisoformat(day)] for day in days(start, end)]

    def execute_wsd(self, asset, field, start, end, options):
        assert "TradingCalendar=SSE" in options
        assert (date.fromisoformat(end) - date.fromisoformat(start)).days <= 365
        self.calls.append((asset, start, end))
        return [[date.fromisoformat(day), 100] for day in days(start, end)]

    def close(self):
        self.closed = True


def test_policy_is_opt_in_and_old_api_default_remains():
    assert WeeklyProductionPolicy().data_source == "wind_client_api"
    assert WeeklyProductionPolicy(data_source="wind_excel").data_source == "wind_excel"


def test_excel_contract_has_no_wset_or_api_gate(tmp_path):
    contract, source, manifest, _ = package(tmp_path)
    loaded, sha = load_run_source_contract(tmp_path, manifest)
    assert loaded == source and len(sha) == 64
    client = Client()
    provider = WindExcelWeeklyProvider(
        run_root=tmp_path, manifest=manifest, source_contract=source, client_factory=lambda: client
    )
    with patch(
        "app.research_web.datahub.providers_wind_api._sdk_default",
        side_effect=AssertionError("SDK must not be imported"),
    ):
        result = asyncio.run(provider.collect(contract))
    assert result["assessment"]["status"] == "passed", result["assessment"]["issues"]
    assert len(result["snapshot"]["series"]) == 57
    assert result["snapshot"]["source"] == "wind_excel"
    assert client.closed
    assert all(
        (date.fromisoformat(end) - date.fromisoformat(start)).days <= 365
        for _, start, end in client.calls
    )


def test_missing_sunday_proof_blocks_even_if_max_row_is_monday(tmp_path):
    contract, source, manifest, save = package(tmp_path)
    proof = {
        "schema_version": 1,
        "source": "wind_excel",
        "windows": [
            {
                "start": contract["next_calendar_start"],
                "end": (
                    date.fromisoformat(contract["next_calendar_end"]) - timedelta(days=1)
                ).isoformat(),
                "resource": "inputs/calendar.xlsx",
                "sha256": source["excel_exports"]["economic_calendar"]["sha256"],
                "kind": "native_export",
            }
        ],
    }
    # Replace the owned declaration too, representing a distinct unpublished fixture.
    manifest["resources"][:] = [
        r for r in manifest["resources"] if r["path"] != "inputs/coverage.json"
    ]
    source["excel_exports"]["economic_calendar"]["coverage_evidence_sha256"] = save(
        "inputs/coverage.json", json.dumps(proof).encode()
    )
    save_observation(
        tmp_path,
        source,
        manifest,
        save,
        native_export_observation(
            source["excel_exports"]["economic_calendar"]["sha256"],
            contract,
            end=proof["windows"][0]["end"],
        ),
    )
    with pytest.raises(ProviderError, match="coverage_incomplete"):
        read_exports(tmp_path, manifest, source, contract)


@pytest.mark.parametrize("failure", ["stale", "hash", "units", "identity"])
def test_invalid_native_inputs_block_without_finalization(tmp_path, failure):
    contract, source, manifest, save = package(tmp_path)
    if failure == "stale":
        source["excel_exports"]["edb"]["exported_at"] = "2026-10-01T16:00:00+08:00"
    elif failure == "hash":
        (tmp_path / "inputs/edb.xlsx").write_bytes(b"bad")
    else:
        rows = [
            ["国家"],
            ["指标名称"],
            ["频率"],
            ["单位", "USD/克" if failure == "units" else "USD/盎司"],
            ["指标ID", "s1" if failure == "identity" else "S0270980"],
            ["时间区间"],
            ["来源"],
            ["更新时间"],
            [date(2026, 10, 16), 100],
        ]
        manifest["resources"][:] = [
            r for r in manifest["resources"] if r["path"] != "inputs/edb.xlsx"
        ]
        source["excel_exports"]["edb"]["sha256"] = save("inputs/edb.xlsx", workbook_bytes(rows))
    result = asyncio.run(
        WindExcelWeeklyProvider(
            run_root=tmp_path, manifest=manifest, source_contract=source, client_factory=Client
        ).collect(contract)
    )
    assert result["assessment"]["status"] == "blocked_data"
    assert result["assessment"]["issues"]


def test_excel_dependency_readiness_does_not_probe_sdk():
    with patch(
        "importlib.util.find_spec", side_effect=lambda name: object() if name == "xlwings" else None
    ):
        assert WindExcelWeeklyProvider().readiness()["dependency_ready"]


def test_default_factory_requires_existing_app_and_owns_only_temporary_workbook(tmp_path):
    contract, source, manifest, _ = package(tmp_path)
    client = Client()
    with patch("data_layer.adapters.wind.client.WindExcelClient", return_value=client) as factory:
        result = asyncio.run(
            WindExcelWeeklyProvider(
                run_root=tmp_path, manifest=manifest, source_contract=source
            ).collect(contract)
        )
    factory.assert_called_once_with(isolated_workbook=True, existing_app_only=True)
    assert result["assessment"]["status"] == "passed", result["assessment"]["issues"]
    assert client.closed


def test_timeout_retains_shared_slot_until_close_returns(tmp_path):
    contract, source, manifest, _ = package(tmp_path)
    entered, release = threading.Event(), threading.Event()

    class Slow(Client):
        def close(self):
            entered.set()
            assert release.wait(10)
            super().close()

    async def scenario():
        provider = WindExcelWeeklyProvider(
            run_root=tmp_path,
            manifest=manifest,
            source_contract=source,
            client_factory=Slow,
            deadline_seconds=0.01,
        )
        with pytest.raises(ProviderError, match="weekly_excel_timeout"):
            await provider.collect(contract)
        assert await asyncio.to_thread(entered.wait, 10)
        assert shared.active_futures()
        with pytest.raises(ProviderError, match="provider_busy"):
            await provider.collect(contract)
        release.set()
        from app.research_web.datahub.providers_wind_excel_weekly import wait_for_idle

        await wait_for_idle()
        assert not shared.active_futures()

    try:
        asyncio.run(scenario())
    finally:
        release.set()


def test_host_collection_freezes_source_identity_and_never_calibrates(tmp_path):
    import yaml

    from app.research_web.datahub import DataHub
    from app.research_web.store import Store

    async def scenario():
        store = Store(tmp_path)
        sid = store.create("claw", "synthetic excel weekly")["id"]
        rid = "a" * 32
        root = tmp_path / "report-workflows" / "runs" / rid
        contract, source, manifest, _ = package(root)
        manifest.update(
            workflow_id="huaan-etf-weekly",
            version=1,
            weekly_production={"policy_id": "huaan-etf-weekly-v1", "data_source": "wind_excel"},
        )
        marker = {
            "workflow_id": "huaan-etf-weekly",
            "version": 1,
            "run_id": rid,
            "weekly_data_source": "wind_excel",
            "owner_id": "synthetic-owner",
        }
        store.session(sid)["report_workflow"] = marker
        parent = tmp_path / "report-workflows"
        catalog = {
            "runs": {
                rid: {
                    "weekly_data_source": "wind_excel",
                    "owner_id": "synthetic-owner",
                    "session_id": sid,
                    "version": 1,
                }
            },
            "workflows": {
                "huaan-etf-weekly": {
                    "current_version": 1,
                    "status": "enabled",
                    "versions": {"1": {"published": True}},
                }
            },
        }

        def save_version():
            (parent / "catalog.json").write_text(json.dumps(catalog))
            path = parent / "huaan-etf-weekly" / "versions" / str(manifest["version"])
            path.mkdir(parents=True, exist_ok=True)
            (path / "manifest.yaml").write_text(yaml.safe_dump(manifest))

        save_version()
        hub = DataHub(store)
        hub.connections.statuses = lambda: {
            "wind": {
                "configured": True,
                "preferred_adapter": "excel",
                "credential_store_available": True,
            }
        }
        configuration = ["d" * 64]
        hub.source_configuration_digest = lambda _source: configuration[0]
        client = Client()
        provider = WindExcelWeeklyProvider(
            run_root=root, manifest=manifest, source_contract=source, client_factory=lambda: client
        )
        kwargs = {
            "run_root": root,
            "manifest": manifest,
            "provider": provider,
            "prepare_excel": lambda: pytest.fail("API calibration must not run"),
        }
        first = await hub.collect_report_weekly(sid, "excel-call", contract["cutoff"], **kwargs)
        assert first["assessment"]["status"] == "passed"
        count = len(client.calls)
        second = await hub.collect_report_weekly(sid, "excel-call", contract["cutoff"], **kwargs)
        assert second["snapshot"] == first["snapshot"]
        assert len(client.calls) == count
        assert not (tmp_path / ".control" / "report-weekly-equivalence").exists()
        receipt = hub.snapshots.receipt(sid, "excel-call")
        hub.snapshots.receipt(sid, "excel-call", {**receipt, "status": "pending"})
        with pytest.raises(StoreError) as caught:
            await hub.collect_report_weekly(sid, "excel-call", contract["cutoff"], **kwargs)
        assert caught.value.code == "weekly_call_unknown"
        hub.snapshots.receipt(sid, "excel-call", receipt)
        configuration[0] = "e" * 64
        with pytest.raises(StoreError) as caught:
            await hub.collect_report_weekly(sid, "excel-call", contract["cutoff"], **kwargs)
        assert caught.value.code == "data_cache_authorization_changed"
        configuration[0] = "d" * 64
        marker["owner_id"] = "another-owner"
        with pytest.raises(StoreError) as caught:
            await hub.collect_report_weekly(sid, "another-call", contract["cutoff"], **kwargs)
        assert caught.value.code == "weekly_resource_invalid"
        marker["owner_id"] = "synthetic-owner"
        manifest["version"] = marker["version"] = 2
        catalog["runs"][rid]["version"] = 2
        workflow = catalog["workflows"]["huaan-etf-weekly"]
        workflow["current_version"] = 2
        workflow["versions"]["2"] = {"published": True}
        save_version()
        with pytest.raises(StoreError) as caught:
            await hub.collect_report_weekly(sid, "excel-call", contract["cutoff"], **kwargs)
        assert caught.value.code == "weekly_call_conflict"
        await hub.close()

    asyncio.run(scenario())


@pytest.mark.parametrize("invalid", [False, "scope", "after_cutoff"])
def test_native_empty_sunday_requires_all_scope_and_both_original_pages(tmp_path, invalid):
    contract, source, manifest, save = package(tmp_path)
    sunday = contract["next_calendar_end"]
    native_start = (
        date.fromisoformat(contract["next_calendar_start"]) - timedelta(days=1)
    ).isoformat()
    message = "当前日期暂无公布日程，请切换至其他日期。"
    empty = {
        "schema_version": 1,
        "source": "wind_excel",
        "kind": "native_empty_result",
        "application": "com.wind.mac.Windotd",
        "source_url": "http://114.80.154.45/eco.web/index.html?lan=cn#/",
        "coverage_start": sunday,
        "coverage_end": sunday,
        "observed_at": contract["cutoff"],
        "filters": {"countries": "all", "event_types": "all", "importance": ["重要", "中等", "低"]},
        "economic_data": {"count": 0, "message": message},
        "financial_events": {"count": 0, "message": message},
        "raw_accessibility_excerpt": sunday + " 经济数据 财经大事 " + message,
    }
    if invalid == "scope":
        empty["filters"]["countries"] = "中国"
    elif invalid == "after_cutoff":
        observed = datetime.fromisoformat(contract["cutoff"]) + timedelta(minutes=1)
        assert observed < datetime.now(UTC)
        empty["observed_at"] = observed.isoformat()
    empty_sha = save("inputs/sunday.json", json.dumps(empty).encode())
    proof = {
        "schema_version": 1,
        "source": "wind_excel",
        "windows": [
            {
                "start": native_start,
                "end": (date.fromisoformat(sunday) - timedelta(days=1)).isoformat(),
                "resource": "inputs/calendar.xlsx",
                "sha256": source["excel_exports"]["economic_calendar"]["sha256"],
                "kind": "native_export",
            },
            {
                "start": sunday,
                "end": sunday,
                "resource": "inputs/sunday.json",
                "sha256": empty_sha,
                "kind": "native_empty_result",
            },
        ],
    }
    manifest["resources"][:] = [
        r for r in manifest["resources"] if r["path"] != "inputs/coverage.json"
    ]
    source["excel_exports"]["economic_calendar"]["coverage_evidence_sha256"] = save(
        "inputs/coverage.json", json.dumps(proof).encode()
    )
    save_observation(
        tmp_path,
        source,
        manifest,
        save,
        native_export_observation(
            source["excel_exports"]["economic_calendar"]["sha256"],
            contract,
            start=native_start,
            end=proof["windows"][0]["end"],
        ),
    )
    if invalid:
        with pytest.raises(ProviderError, match="empty_result_unverified"):
            read_exports(tmp_path, manifest, source, contract)
    else:
        assert read_exports(tmp_path, manifest, source, contract)["economic_calendar"]


def replace_resource(manifest, save, path, raw):
    manifest["resources"][:] = [r for r in manifest["resources"] if r["path"] != path]
    return save(path, raw)


def save_observation(root, source, manifest, save, observation):
    export = source["excel_exports"]["economic_calendar"]
    export["observation_sha256"] = replace_resource(
        manifest, save, export["observation_path"], json.dumps(observation).encode()
    )


def test_native_export_observation_accepts_actual_vendor_filter_labels(tmp_path):
    contract, source, manifest, save = package(tmp_path)
    observation = native_export_observation(
        source["excel_exports"]["economic_calendar"]["sha256"], contract
    )
    for page in observation["pages"].values():
        page["raw_accessibility_excerpt"] = page["raw_accessibility_excerpt"].replace(
            "国家/地区:全部 事件类型:全部 重要 中等 低",
            "选择国家: 全部 数据/事件: 全部 重要性: 重要 中等 低",
        )
    save_observation(tmp_path, source, manifest, save, observation)
    assert read_exports(tmp_path, manifest, source, contract)["economic_export_dates"]


@pytest.mark.parametrize(
    "invalid", ["single_day", "filtered", "forged_ax_date", "missing_page", "wrong_export"]
)
def test_native_export_scope_requires_independent_ui_receipt(tmp_path, invalid):
    contract, source, manifest, save = package(tmp_path)
    export = source["excel_exports"]["economic_calendar"]
    observation = json.loads((tmp_path / export["observation_path"]).read_bytes())
    if invalid == "single_day":
        observation["coverage_end"] = observation["coverage_start"]
    elif invalid == "filtered":
        observation["pages"]["economic_data"]["filters"]["countries"] = "中国"
    elif invalid == "forged_ax_date":
        observation["pages"]["economic_data"]["raw_accessibility_excerpt"] = observation["pages"][
            "economic_data"
        ]["raw_accessibility_excerpt"].replace(observation["coverage_end"], "2026-10-01")
    elif invalid == "missing_page":
        del observation["pages"]["financial_events"]
    else:
        observation["export_sha256"] = "f" * 64
    save_observation(tmp_path, source, manifest, save, observation)
    with pytest.raises(ProviderError, match="observation_unverified"):
        read_exports(tmp_path, manifest, source, contract)


@pytest.mark.parametrize("kind", ["edb", "economic_calendar"])
def test_static_export_after_historical_cutoff_is_rejected_even_before_current_time(tmp_path, kind):
    contract, source, manifest, _ = package(tmp_path)
    exported = datetime.fromisoformat(contract["cutoff"]) + timedelta(minutes=1)
    assert exported < datetime.now(UTC)
    source["excel_exports"][kind]["exported_at"] = exported.isoformat()
    with pytest.raises(ProviderError, match="weekly_excel_export_stale_or_future"):
        read_exports(tmp_path, manifest, source, contract)


@pytest.mark.parametrize(
    "when,accepted", [("2026-10-16T16:00:00+08:00", True), ("2026-10-23T16:00:00+08:00", False)]
)
def test_host_records_current_period_freeze_and_blocks_historical_online_recheck(
    tmp_path, monkeypatch, when, accepted
):
    contract, source, manifest, _ = package(tmp_path)
    monkeypatch.setattr(excel_provider, "_host_now", lambda: datetime.fromisoformat(when))
    client = Client()
    result = asyncio.run(
        WindExcelWeeklyProvider(
            run_root=tmp_path,
            manifest=manifest,
            source_contract=source,
            client_factory=lambda: client,
        ).collect(contract)
    )
    snapshot = result["snapshot"]
    assert snapshot["acquisition_started_at"] == when
    assert snapshot["frozen_at"] == when
    if accepted:
        assert result["assessment"]["status"] == "passed", result["assessment"]["issues"]
        assert snapshot["acquisition_kind"] == "current_period"
        assert snapshot["cutoff"] != snapshot["frozen_at"]
    else:
        assert result["assessment"]["status"] == "blocked_data"
        assert snapshot["acquisition_kind"] == "historical_recheck"
        assert "weekly_excel_historical_recheck_not_formal" in result["assessment"]["issues"]
        assert client.calls == []


@pytest.mark.parametrize(
    "changed", ["missing", "future_freeze", "before_start", "retrieved_after_freeze"]
)
def test_independent_host_assessment_rejects_changed_acquisition_times(tmp_path, changed):
    from app.research_web.report_workflows.weekly_data import assess_weekly_snapshot

    contract, source, manifest, _ = package(tmp_path)
    snapshot = asyncio.run(
        WindExcelWeeklyProvider(
            run_root=tmp_path, manifest=manifest, source_contract=source, client_factory=Client
        ).collect(contract)
    )["snapshot"]
    if changed == "missing":
        snapshot.pop("acquisition_started_at", None)
    elif changed == "future_freeze":
        snapshot["frozen_at"] = (datetime.now(UTC) + timedelta(days=1)).isoformat()
    elif changed == "before_start":
        snapshot["frozen_at"] = "2026-10-16T15:59:59+08:00"
    else:
        next(iter(snapshot["series"].values()))["retrieved_at"] = "2026-10-16T16:00:01+08:00"
    assert (
        assess_weekly_snapshot(snapshot, contract, require_excel_acquisition=True)["status"]
        == "blocked_data"
    )


def test_excel_importance_mapping_must_cover_native_three_levels(tmp_path):
    _, source, manifest, save = package(tmp_path)
    source["economic_calendar"]["importance_values"] = {"重要": "重要"}
    replace_resource(manifest, save, SOURCE_CONTRACT, json.dumps(source).encode())
    with pytest.raises(WorkflowError, match="周报数据源合同未核实"):
        load_run_source_contract(tmp_path, manifest)


def test_coverage_cannot_substitute_unrelated_immutable_resource(tmp_path):
    contract, source, manifest, save = package(tmp_path)
    proof = json.loads((tmp_path / "inputs/coverage.json").read_bytes())
    proof["windows"][0].update(
        resource="inputs/unrelated.txt", sha256=save("inputs/unrelated.txt", b"unrelated")
    )
    source["excel_exports"]["economic_calendar"]["coverage_evidence_sha256"] = replace_resource(
        manifest, save, "inputs/coverage.json", json.dumps(proof).encode()
    )
    with pytest.raises(ProviderError, match="coverage_invalid|observation_unverified"):
        read_exports(tmp_path, manifest, source, contract)


def test_native_wrong_dimension_still_reads_all_edb_observations(tmp_path):
    contract, source, manifest, save = package(tmp_path)
    stream = BytesIO()
    with ZipFile(tmp_path / "inputs/edb.xlsx") as original, ZipFile(stream, "w") as changed:
        for entry in original.infolist():
            raw = original.read(entry)
            if entry.filename == "xl/worksheets/sheet1.xml":
                raw = re.sub(rb'<dimension ref="[^"]+"', b'<dimension ref="A1"', raw)
            changed.writestr(entry, raw)
    source["excel_exports"]["edb"]["sha256"] = replace_resource(
        manifest, save, "inputs/edb.xlsx", stream.getvalue()
    )
    result = asyncio.run(
        WindExcelWeeklyProvider(
            run_root=tmp_path, manifest=manifest, source_contract=source, client_factory=Client
        ).collect(contract)
    )
    assert result["assessment"]["status"] == "passed", result["assessment"]["issues"]
    assert max(len(item["rows"]) for item in result["snapshot"]["series"].values()) > 4000


@pytest.mark.parametrize("financial_columns", [5, 8])
def test_both_native_economic_sheets_preserve_events_and_original_level(
    tmp_path, financial_columns
):
    contract, source, manifest, save = package(tmp_path)
    book = Workbook()
    book.active.title = "经济数据"
    second = book.create_sheet("财经大事")
    header = ["日期", "时间", "国家/地区", "指标名称", "重要性", "前值", "预测值", "今值"]
    for sheet, importance in ((book.active, "中等"), (second, "低")):
        sheet.append(header if sheet is book.active else header[:financial_columns])
        sheet.append(
            [
                date.fromisoformat(contract["next_calendar_start"]),
                "09:00",
                "中国",
                sheet.title,
                importance,
            ]
        )
        sheet.append(["数据来源: Wind"])
    stream = BytesIO()
    book.save(stream)
    book.close()
    sha = replace_resource(manifest, save, "inputs/calendar.xlsx", stream.getvalue())
    source["excel_exports"]["economic_calendar"].update(
        sha256=sha, sheet="经济数据", additional_sheets=["财经大事"]
    )
    observation = native_export_observation(sha, contract)
    observation["pages"]["economic_data"]["sheet"] = "经济数据"
    observation["pages"]["financial_events"].update(
        sheet="财经大事", result="exported", message=None
    )
    save_observation(tmp_path, source, manifest, save, observation)
    proof = json.loads((tmp_path / "inputs/coverage.json").read_bytes())
    proof["windows"][0]["sha256"] = sha
    source["excel_exports"]["economic_calendar"]["coverage_evidence_sha256"] = replace_resource(
        manifest, save, "inputs/coverage.json", json.dumps(proof).encode()
    )
    result = asyncio.run(
        WindExcelWeeklyProvider(
            run_root=tmp_path, manifest=manifest, source_contract=source, client_factory=Client
        ).collect(contract)
    )
    assert result["assessment"]["status"] == "passed", result["assessment"]["issues"]
    assert [
        (item["source_sheet"], item["importance"])
        for item in result["snapshot"]["economic_calendar"]["events"]
    ] == [("经济数据", "较重要"), ("财经大事", "一般")]


def test_cleanup_failure_poison_blocks_shared_worker_without_replay(tmp_path, monkeypatch):
    contract, source, manifest, _ = package(tmp_path)
    monkeypatch.setattr(shared, "_WIND_POISONED", False)
    monkeypatch.setattr(shared, "_WIND_OUTSTANDING", None)

    class UnconfirmedClose(Client):
        def close(self):
            raise RuntimeError("synthetic close failure")

    client = UnconfirmedClose()
    provider = WindExcelWeeklyProvider(
        run_root=tmp_path, manifest=manifest, source_contract=source, client_factory=lambda: client
    )

    async def scenario():
        with pytest.raises(ProviderError, match="wind_cleanup_failed"):
            await provider.collect(contract)
        assert shared._WIND_POISONED is True
        count = len(client.calls)
        with pytest.raises(ProviderError, match="provider_busy"):
            await provider.collect(contract)
        assert len(client.calls) == count

    asyncio.run(scenario())


def test_real_client_unknown_book_creation_poison_survives_finally_and_new_call_id(
    tmp_path, monkeypatch
):
    from data_layer.adapters.wind import client as wind_client

    contract, source, manifest, _ = package(tmp_path)
    monkeypatch.setattr(shared, "_WIND_POISONED", False)
    monkeypatch.setattr(shared, "_WIND_OUTSTANDING", None)
    user_book = SimpleNamespace(name="existing-user.xlsx", close=MagicMock())
    books = MagicMock()
    books.__iter__.side_effect = lambda: iter([user_book])
    books.add.side_effect = TimeoutError("synthetic unknown books.add outcome")
    application = SimpleNamespace(books=books, quit=MagicMock())
    xlwings = SimpleNamespace(App=MagicMock())
    monkeypatch.setattr(wind_client, "running_excel_apps", lambda _xw: [application])
    monkeypatch.setattr(wind_client, "prepare_macos_excel_compat", lambda: None)
    client = wind_client.WindExcelClient(isolated_workbook=True, existing_app_only=True)
    factory = MagicMock(return_value=client)
    provider = WindExcelWeeklyProvider(
        run_root=tmp_path, manifest=manifest, source_contract=source, client_factory=factory
    )

    async def scenario():
        with pytest.raises(ProviderError, match="wind_cleanup_failed"):
            await provider.collect(contract, call_id="unknown-create-first")
        assert client._creation_unknown is True
        assert shared._WIND_POISONED is True
        assert books.add.call_count == 1
        assert [book.name for book in books] == ["existing-user.xlsx"]
        with pytest.raises(ProviderError, match="provider_busy"):
            await provider.collect(contract, call_id="unknown-create-second")
        factory.assert_called_once_with()
        books.add.assert_called_once_with()
        user_book.close.assert_not_called()
        application.quit.assert_not_called()
        xlwings.App.assert_not_called()

    with patch.dict("sys.modules", {"xlwings": xlwings}):
        asyncio.run(scenario())


def test_native_calendar_window_can_overlap_previous_sunday_without_including_it(tmp_path):
    contract, source, manifest, save = package(tmp_path)
    previous_sunday = (
        date.fromisoformat(contract["next_calendar_start"]) - timedelta(days=1)
    ).isoformat()
    rows = [
        ["日期", "时间", "国家/地区", "指标名称", "重要性", "前值", "预测值", "今值"],
        [date.fromisoformat(previous_sunday), "09:00", "中国", "前周日", "低"],
        [date.fromisoformat(contract["next_calendar_start"]), "09:00", "中国", "下周一", "重要"],
        ["数据来源: Wind"],
    ]
    sha = replace_resource(manifest, save, "inputs/calendar.xlsx", workbook_bytes(rows))
    source["excel_exports"]["economic_calendar"]["sha256"] = sha
    save_observation(
        tmp_path,
        source,
        manifest,
        save,
        native_export_observation(sha, contract, start=previous_sunday),
    )
    proof = json.loads((tmp_path / "inputs/coverage.json").read_bytes())
    proof["windows"][0].update(start=previous_sunday, sha256=sha)
    source["excel_exports"]["economic_calendar"]["coverage_evidence_sha256"] = replace_resource(
        manifest, save, "inputs/coverage.json", json.dumps(proof).encode()
    )
    result = asyncio.run(
        WindExcelWeeklyProvider(
            run_root=tmp_path, manifest=manifest, source_contract=source, client_factory=Client
        ).collect(contract)
    )
    assert result["assessment"]["status"] == "passed", result["assessment"]["issues"]
    assert [item["indicator"] for item in result["snapshot"]["economic_calendar"]["events"]] == [
        "下周一"
    ]


@pytest.mark.parametrize("kind", ["calendar", "series"])
def test_vendor_matrix_over_one_thousand_rows_is_rejected(tmp_path, kind):
    contract, source, manifest, _ = package(tmp_path)

    class Oversized(Client):
        def _execute_wsd_once(self, formula, timeout, expected_columns):
            if (kind == "calendar") == (expected_columns == 1):
                return [[date(2026, 10, 16), *([100] if expected_columns == 2 else [])]] * 1001
            return super()._execute_wsd_once(formula, timeout, expected_columns)

    client = Oversized()
    result = asyncio.run(
        WindExcelWeeklyProvider(
            run_root=tmp_path,
            manifest=manifest,
            source_contract=source,
            client_factory=lambda: client,
        ).collect(contract)
    )
    assert result["assessment"]["status"] == "blocked_data"
    assert (
        f"weekly_excel_{'calendar' if kind == 'calendar' else 'wsd'}_shape_invalid"
        in result["assessment"]["issues"]
    )
    assert client.closed


def test_unreviewed_market_calendar_never_passes_source_contract(tmp_path):
    _, source, manifest, save = package(tmp_path)
    next(iter(source["calendars"].values()))["semantics_verified"] = False
    replace_resource(manifest, save, SOURCE_CONTRACT, json.dumps(source).encode())
    with pytest.raises(WorkflowError) as caught:
        load_run_source_contract(tmp_path, manifest)
    assert caught.value.code == "weekly_source_contract_unverified"


def test_holiday_zeros_do_not_become_prices_and_unclosed_values_are_blocked():
    spec = load_weekly_mapping()["series"][0]
    provider = WindExcelWeeklyProvider()
    value = provider._series(spec, [["2026-10-17", 0], ["2026-10-16", 100]], {}, ["2026-10-16"])
    assert value["rows"] == [{"date": "2026-10-16", "value": 100, "published_at": None}]
    with pytest.raises(ProviderError, match="unclosed_observation"):
        provider._series(spec, [["2026-10-17", 100]], {}, ["2026-10-16"])
