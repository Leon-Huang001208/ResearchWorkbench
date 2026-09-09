"""Reviewed report projection regressions."""

import asyncio
import json
import sys
import zipfile
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook, load_workbook

from app.research_web.report_rendering import render_report_payload
from app.research_web.store import Store


@pytest.mark.skipif(
    sys.platform != "darwin", reason="renderer requires native macOS sandbox"
)
def test_report_payload_projects_to_real_office_and_html_files(tmp_path: Path):
    store = Store(tmp_path)
    session = store.create("claw", "report projection")
    root = store.directory(session["id"])
    (root / "outputs" / "report_payload.json").write_text(
        json.dumps(
            {
                "title": "真实周报",
                "as_of": "2026-09-04",
                "summary": "基于锁定资料形成。",
                "sections": [
                    {
                        "id": "market",
                        "title": "市场回顾",
                        "content": "市场数据存在部分缺失。",
                    },
                    {
                        "id": "risk",
                        "title": "风险提示",
                        "content": "历史数据不代表未来。",
                    },
                ],
                "sources": [{"title": "DataHub 快照", "as_of": "2026-09-04"}],
                "missing": ["行业数据尚未更新"],
            },
            ensure_ascii=False,
        )
    )

    result = asyncio.run(
        render_report_payload(store, session["id"], ["docx", "html", "xlsx"])
    )

    assert result["status"] == "completed"
    for name in ("report.docx", "report.html", "report.xlsx"):
        assert (root / "outputs" / name).stat().st_size > 0


@pytest.mark.skipif(
    sys.platform != "darwin", reason="renderer requires native macOS sandbox"
)
def test_report_projection_requires_structured_payload(tmp_path: Path):
    store = Store(tmp_path)
    session = store.create("claw", "missing payload")
    result = asyncio.run(render_report_payload(store, session["id"], ["html"]))
    assert result["status"] == "missing_payload"


@pytest.mark.skipif(
    sys.platform != "darwin", reason="renderer requires native macOS sandbox"
)
def test_report_workflow_projection_uses_locked_template_and_refreshed_workbook(
    tmp_path: Path,
):
    store = Store(tmp_path)
    session = store.create("claw", "locked workflow projection")
    root = store.directory(session["id"])
    templates = root / "inputs" / "report-workflow" / "templates"
    workbooks = root / "inputs" / "report-workflow" / "workbooks"
    templates.mkdir(parents=True)
    workbooks.mkdir(parents=True)

    document = Document()
    document.add_paragraph("{{market}}")
    document.save(templates / "report.docx")
    workbook = Workbook()
    workbook.active["A1"] = "locked-refreshed-data"
    workbook.save(workbooks / "source.xlsx")
    (root / "inputs" / "report-workflow" / "run-context.json").write_text(
        json.dumps({"primary_workbook": "workbooks/source.xlsx"})
    )
    (root / "outputs" / "report_payload.json").write_text(
        json.dumps(
            {
                "title": "锁定周报",
                "as_of": "2026-09-07",
                "sections": [{"id": "market", "title": "市场", "content": "真实正文"}],
            },
            ensure_ascii=False,
        )
    )

    result = asyncio.run(
        render_report_payload(store, session["id"], ["docx", "html", "xlsx"])
    )

    assert result["status"] == "completed"
    rendered = Document(root / "outputs" / "report.docx")
    assert "真实正文" in "\n".join(item.text for item in rendered.paragraphs)
    delivered = load_workbook(root / "outputs" / "report.xlsx", data_only=False)
    assert delivered.active["A1"].value == "locked-refreshed-data"


@pytest.mark.skipif(
    sys.platform != "darwin", reason="renderer requires native macOS sandbox"
)
def test_report_projection_normalizes_block_families_and_reports_missing_blocks(
    tmp_path: Path,
):
    store = Store(tmp_path)
    session = store.create("claw", "block-family report projection")
    root = store.directory(session["id"])
    workflow = root / "inputs" / "report-workflow"
    workflow.mkdir(parents=True)
    (workflow / "run-context.json").write_text(
        json.dumps(
            {
                "name": "华安ETF周报",
                "primary_workbook": None,
                "blocks": [
                    {"id": "block_001", "title": "市场回顾", "required": True},
                    {"id": "block_002", "title": "行业观察", "required": True},
                    {"id": "block_003", "title": "数据表", "required": True},
                ],
            },
            ensure_ascii=False,
        )
    )
    (root / "outputs" / "report_payload.json").write_text(
        json.dumps(
            {
                "workflow_id": "huaan-etf-weekly",
                "report_period": {"end_date_trading": "2026-09-04"},
                "blocks": {
                    "block_001": {"title": "市场回顾", "text": "市场真实正文。"}
                },
                "date_blocks": {
                    "block_002": {
                        "title": "行业观察",
                        "status": "missing_unresolved",
                        "text": "模型提供了缺失解释，但没有可交付正文。",
                    }
                },
                "table_blocks": {
                    "block_003": {"title": "数据表", "content": "表格口径说明。"}
                },
                "missing": [
                    {
                        "block": "block_002",
                        "title": "行业观察",
                        "reason": "缺少新闻证据",
                    }
                ],
                "sources": ["锁定的数据快照"],
            },
            ensure_ascii=False,
        )
    )

    result = asyncio.run(render_report_payload(store, session["id"], ["docx", "html"]))

    assert result["status"] == "completed"
    assert result["missing_blocks"] == ["block_002"]
    rendered = Document(root / "outputs" / "report.docx")
    text = "\n".join(item.text for item in rendered.paragraphs)
    assert "华安ETF周报" in text
    assert "市场真实正文" in text
    assert "缺少新闻证据" in text


@pytest.mark.skipif(
    sys.platform != "darwin", reason="renderer requires native macOS sandbox"
)
def test_report_projection_replaces_pptx_placeholder_split_across_text_runs(
    tmp_path: Path,
):
    store = Store(tmp_path)
    session = store.create("claw", "split PPTX placeholder")
    root = store.directory(session["id"])
    templates = root / "inputs" / "report-workflow" / "templates"
    templates.mkdir(parents=True)
    with zipfile.ZipFile(templates / "deck.pptx", "w") as package:
        package.writestr("[Content_Types].xml", "<Types/>")
        package.writestr(
            "ppt/slides/slide1.xml",
            (
                '<p:sld xmlns:p="urn:p" xmlns:a="urn:a"><p:cSld><a:p>'
                "<a:r><a:t>前缀 {{市</a:t></a:r>"
                "<a:r><a:t>场概览}}</a:t></a:r>"
                "<a:r><a:t> 后缀</a:t></a:r>"
                "</a:p></p:cSld></p:sld>"
            ),
        )
    (root / "outputs" / "report_payload.json").write_text(
        json.dumps(
            {
                "title": "投资风向标",
                "sections": [
                    {"id": "market", "title": "市场概览", "content": "真实市场正文"}
                ],
            },
            ensure_ascii=False,
        )
    )

    result = asyncio.run(render_report_payload(store, session["id"], ["pptx"]))

    assert result["status"] == "completed"
    with zipfile.ZipFile(root / "outputs" / "report.pptx") as package:
        slide = package.read("ppt/slides/slide1.xml").decode("utf-8")
    assert "{{市场概览}}" not in slide
    assert "真实市场正文" in slide
