"""Real-record replay: analysis quantities cannot bypass system declarations."""

import copy
import hashlib
import json
from fractions import Fraction
from pathlib import Path

import pytest

from app.research_web.resources.research_report import archive_report_version, list_report_versions
from app.research_web.resources.research_report_calculations import derive_business_observations
from app.research_web.resources.research_report_plan import (
    build_business_report_plan,
    write_business_report_plan,
)

FIXTURE = Path(__file__).parent / "fixtures/fingpt-20261020"


def recorded_case():
    manifest = json.loads((FIXTURE / "manifest.json").read_text())
    for name, sha in manifest["files"].items():
        assert hashlib.sha256((FIXTURE / name).read_bytes()).hexdigest() == sha
    bundle = json.loads((FIXTURE / "research_evidence.json").read_text())
    context = {
        "task_id": bundle["task_id"],
        "contract_sha256": bundle["contract_sha256"],
        "as_of": "2026-10-16",
        "requirement_ids": ["r1"],
        "formats": ["html", "docx"],
    }
    plan = {
        "title": "中国年度GDP研究",
        "missing": ["最终核实资料尚未取得，跨版本可比性未核验"],
        "sections": [
            {
                "id": "analysis",
                "heading": "结论与计算",
                "requirement_ids": ["r1"],
                "evidence_ids": [
                    "gdp_current_price_preliminary@2024",
                    "gdp_current_price_preliminary@2025",
                    "gdp_real_growth_preliminary@2025",
                ],
                "analysis": "仅比较已取得的初步核算资料，跨版本可比性仍待核验。",
            }
        ],
    }
    return context, plan, bundle


@pytest.mark.parametrize(
    "text",
    [
        "名义增速3.91%，实际增速5.0%，隐含平减指数约−1.05%。",
        "现价总量增加52795亿元。",
        "现价总量增加五万二千七百九十五亿元。",
        "隐含指数负一点零五百分比。",
        "增长３．９１％。",
        "增长3&#46;91%。",
        "增长²%。",
        "名义增速 ≈ 实际增速 + 价格变化。",
        "[[unbound:52795]]",
    ],
)
def test_unbound_analysis_quantities_and_formulas_reject_before_publication(text):
    context, plan, bundle = recorded_case()
    plan["sections"][0]["analysis"] = text
    with pytest.raises(ValueError, match="research_report_analysis_unbound"):
        build_business_report_plan(context, plan, bundle)


def test_exact_recorded_wrong_derivation_is_not_a_system_calculation():
    context, _plan, bundle = recorded_case()
    failed = json.loads((FIXTURE / "failed-report.json").read_text())
    assert any("−1.05%" in s["text"] for s in failed["sections"])
    assert all(r["operation"] != "gdp_implicit_deflator" for r in failed["business_calculations"])
    current, prior, real = (
        "gdp_current_price_preliminary@2025",
        "gdp_current_price_preliminary@2024",
        "gdp_real_growth_preliminary@2025",
    )
    recipes = [
        {
            "id": "increment",
            "operation": "annual_change",
            "operands": [current, prior],
            "decimal_places": 0,
        },
        {
            "id": "deflator",
            "operation": "gdp_implicit_deflator",
            "operands": [current, prior, real],
            "decimal_places": 2,
        },
    ]
    rows = derive_business_observations(bundle["observations"], recipes, task_id=context["task_id"])
    assert rows[0]["value"] == "52795"
    assert rows[1]["value"] == "-1.03"
    exact = (Fraction(1401879, 1349084) / (1 + Fraction(5, 100)) - 1) * 100
    assert Fraction(rows[1]["exact_value"]) == exact
    assert len(rows[1]["input_ids"]) == 3


def bound_plan():
    context, plan, bundle = recorded_case()
    plan["calculations"] = [
        {
            "id": "nominal",
            "operation": "percent_change",
            "operands": [
                "gdp_current_price_preliminary@2025",
                "gdp_current_price_preliminary@2024",
            ],
            "decimal_places": 2,
        },
    ]
    plan["sections"][0]["evidence_ids"].append("$nominal")
    plan["sections"][0][
        "analysis"
    ] = "现价计算结果：{{fact:$nominal}}。计算过程：{{formula:$nominal}}。可比性尚未核验，计算正确不代表全部验收通过。"
    return context, plan, bundle


@pytest.mark.parametrize(
    "replacement",
    [
        "{{fact:unknown}}",
        "{{fact:gdp_real_growth_preliminary@2024}}",
        "{{value:$nominal}}",
        "{{formula:gdp_current_price_preliminary@2024}}",
        "{{fact:$nominal}}3.91%",
    ],
)
def test_unknown_unselected_or_value_only_placeholders_rejected(replacement):
    context, plan, bundle = bound_plan()
    plan["sections"][0]["analysis"] = replacement
    with pytest.raises(ValueError):
        build_business_report_plan(context, plan, bundle)


def test_bound_analysis_chinese_body_audit_archive_and_recovery(tmp_path, monkeypatch):
    context, plan, bundle = bound_plan()
    monkeypatch.chdir(tmp_path)
    Path("resources").mkdir()
    Path("resources/research_evidence.json").write_text(json.dumps(bundle))
    result = write_business_report_plan(context, plan)
    assert "audit" in result["written"]
    canonical = json.loads(Path(result["written"]["json"]).read_text())
    html = Path(result["written"]["html"]).read_text()
    audit = Path(result["written"]["audit"]).read_text()
    assert "Extracted from" not in html and "Calculated only from" not in html
    assert "Extracted from" in audit and "Calculated only from" in audit
    assert "3.91" in html and "亿元" in html and "现价" in html
    assert "{{fact:" not in html
    assert len(canonical["numeric_claims"]) == 4
    assert all(c["text"] in canonical["sections"][0]["text"] for c in canonical["numeric_claims"])
    from docx import Document

    doc = Document(result["written"]["docx"])
    assert "Extracted from" not in "\n".join(p.text for p in doc.paragraphs)
    assert len(doc.tables[0].rows) == 5
    archived = archive_report_version(Path("outputs"))
    assert len(archived["files"]) == 4
    assert any(r["name"] == Path(result["written"]["audit"]).name for r in archived["files"])
    assert list_report_versions(Path("outputs"))[0]["version_id"] == archived["version_id"]
    before = {k: Path(p).read_bytes() for k, p in result["written"].items()}
    plan["sections"][0]["analysis"] = "未经计算增加52795亿元。"
    with pytest.raises(ValueError):
        write_business_report_plan(context, plan)
    assert all(Path(result["written"][k]).read_bytes() == b for k, b in before.items())
    # A changed appendix invalidates the exact archived report, never a normal evidence gap.
    audit_row = next(r for r in archived["files"] if "audit" in r["name"])
    target = Path("outputs") / archived["directory"] / audit_row["stored_name"]
    target.chmod(0o600)
    target.write_text("corrupt")
    assert list_report_versions(Path("outputs"))[0]["quality"] == "invalid"


@pytest.mark.parametrize(
    "text",
    [
        "占比为四分之一。",
        "增量为五万。",
        "The increase is five trillion.",
        "增长五。",
        "增长一半。",
        "同比下降三成。",
        "Ⅴ个百分点",
        "两倍。",
        "收入翻番。",
        "占比为俩成。",
        "增量为仨。",
        "价格对折。",
        "It doubled.",
    ],
)
def test_non_arabic_quantity_bypasses_rejected(text):
    context, plan, bundle = recorded_case()
    plan["sections"][0]["analysis"] = text
    with pytest.raises(ValueError, match="research_report_analysis_unbound"):
        build_business_report_plan(context, plan, bundle)


@pytest.mark.parametrize("field", ["title", "heading", "missing"])
def test_other_author_fields_cannot_smuggle_numerical_conclusions(field):
    context, plan, bundle = recorded_case()
    if field == "heading":
        plan["sections"][0][field] = "增量52795亿元"
    else:
        plan[field] = ["增量52795亿元"] if field == "missing" else "增量52795亿元"
    with pytest.raises(ValueError):
        build_business_report_plan(context, plan, bundle)


@pytest.mark.parametrize("question", ["约1200字中文", "约1200字的中文报告", "中文正文约1200字"])
def test_frozen_chinese_target_and_length_failure(question):
    from app.research_web.resources.research_report_plan import required_body_target

    assert required_body_target({"question": question}) == 1200
    context, plan, bundle = bound_plan()
    context["body_target"] = 1200
    with pytest.raises(ValueError, match="research_report_chinese_body_length"):
        build_business_report_plan(context, plan, bundle)


@pytest.mark.parametrize("kind", ["html", "docx", "pdf"])
def test_actual_render_cannot_be_replaced_with_self_consistent_hashes(kind):
    from app.research_web.resources.research_report import (
        docx_bytes,
        html_bytes,
        pdf_bytes,
        verify_rendered_artifact,
    )

    context, plan, bundle = bound_plan()
    payload = build_business_report_plan(context, plan, bundle)["payload"]
    raw = {"html": html_bytes, "docx": docx_bytes, "pdf": pdf_bytes}[kind](payload)
    verify_rendered_artifact(payload, kind, raw)
    altered = copy.deepcopy(payload)
    altered["presentation"][0]["text"] = "伪造正文增长52795亿元。"
    forged = {"html": html_bytes, "docx": docx_bytes, "pdf": pdf_bytes}[kind](altered)
    with pytest.raises(ValueError, match="research_report_render_binding_changed"):
        verify_rendered_artifact(payload, kind, forged)
    with pytest.raises(ValueError, match="research_report_render_binding_changed"):
        verify_rendered_artifact(payload, kind, b"broken file")


def host_recorded_case(tmp_path, monkeypatch):
    from test_report_system_flow import case

    from app.research_web.research_quality_service import write_business_snapshot

    service, sid, receipt, context, _ = case(tmp_path, monkeypatch, acquire=False)
    receipt["report_binding_policy"] = "bound-analysis-v1"
    _, plan, _recorded = bound_plan()
    root = service.store.directory(sid)
    for source in json.loads((FIXTURE / "source-records.json").read_text()):
        path = "inputs/web-" + source["id"].split(":")[1] + ".txt"
        raw = (FIXTURE / Path(path).name).read_bytes()
        (root / path).write_bytes(raw)
        fid = hashlib.sha256(path.encode()).hexdigest()[:24]
        service.store.session(sid)["files"][fid] = path
        service.store.session(sid).setdefault("web_evidence", {})[source["id"]] = {
            **source,
            "source_id": source["id"],
            "path": path,
            "file_id": fid,
            "sha256": hashlib.sha256(raw).hexdigest(),
        }
    service.store.save()
    write_business_snapshot(service, sid, receipt["research_contract"])
    plan["sections"][0]["requirement_ids"] = context["requirement_ids"]
    plan["sections"][0]["analysis"] = "{{fact:$nominal}}"
    plan["missing"] = []
    return service, sid, receipt, context, plan


@pytest.mark.parametrize(
    "tamper",
    [
        None,
        "analysis",
        "claim",
        "presentation",
        "audit",
        "audit_missing",
        "rendered_html",
        "rendered_docx",
        "downgrade",
    ],
)
def test_real_owned_bodies_to_host_binding_and_recovery(tmp_path, monkeypatch, tamper):
    import asyncio

    from app.research_web.research_quality_service import check_quality, quality_versions
    from app.research_web.resources.research_report_plan import audit_content

    async def run():
        service, sid, _receipt, context, plan = host_recorded_case(tmp_path, monkeypatch)
        result = write_business_report_plan(context, plan)
        first = await check_quality(service, sid)
        assert first["status"] == "passed", first
        path = Path(result["written"]["json"])
        saved = json.loads(path.read_text())
        original_files = {Path(p): Path(p).read_bytes() for p in result["written"].values()}
        if tamper is None:
            return
        current = copy.deepcopy(saved)
        if tamper == "analysis":
            current["system_plan"]["sections"][0]["analysis"] = "增加52795亿元。"
        elif tamper == "claim":
            current["numeric_claims"][0]["value"] = "9"
        elif tamper == "presentation":
            current["presentation"][0]["text"] = "伪造52795亿元。"
        elif tamper == "audit":
            Path(result["written"]["audit"]).write_text("{}")
        elif tamper == "audit_missing":
            Path(result["written"]["audit"]).unlink()
        elif tamper.startswith("rendered_"):
            kind = tamper.split("_")[1]
            raw = b"<html>unbound 52795</html>" if kind == "html" else b"broken docx"
            Path(result["written"][kind]).write_bytes(raw)
            next(r for r in current["artifacts"] if r["name"].endswith("." + kind))["sha256"] = (
                hashlib.sha256(raw).hexdigest()
            )
        elif tamper == "downgrade":
            current.pop("system_plan")
        if tamper not in {"audit", "audit_missing"}:
            # Self-consistent appendix hashes cannot certify forged report content.
            raw = audit_content(current)
            Path(result["written"]["audit"]).write_bytes(raw)
            current["audit_attachment"]["sha256"] = hashlib.sha256(raw).hexdigest()
            path.write_text(json.dumps(current))
        bad = await service.detail(sid)
        assert bad["research_report"]["quality"] == "invalid", bad
        assert bad["research_report"]["file_state"] == "invalid"
        # Restore only the controlled fixture; the original archive and quality survive.
        for target, raw in original_files.items():
            target.write_bytes(raw)
        repaired = write_business_report_plan(context, plan)
        good = await check_quality(service, sid)
        assert good["status"] == "passed", good
        versions = quality_versions(service, sid)
        assert versions and any(v["quality"]["status"] == "passed" for v in versions)
        assert Path(repaired["written"]["audit"]).exists()

    asyncio.run(run())


def test_continuation_preserves_frozen_body_target_and_conflicts_fail_closed():
    from app.research_web.resources.research_report_plan import required_body_target

    contract = {
        "question": "补充最终核实资料，请继续研究。",
        "requirements": [
            {"id": "original", "text": "使用实际证据，生成约1200字中文HTML和Word报告。"}
        ],
    }
    assert required_body_target(contract) == 1200
    with pytest.raises(ValueError, match="research_report_body_target_conflict"):
        required_body_target({**contract, "question": "改写为约2000字中文报告。"})
