"""Offline success and recovery of the GDP path that exhausted the live budget."""

import asyncio
import copy
import hashlib
import json
from pathlib import Path
from types import SimpleNamespace

import pytest

from app.research_web.research_quality import make_contract
from app.research_web.research_quality_service import (
    check_quality,
    quality_result,
    quality_versions,
    write_business_snapshot,
)
from app.research_web.research_web_sources import stored_sources
from app.research_web.resources.research_report_plan import write_business_report_plan
from app.research_web.store import Store, StoreError

FINAL = """发布时间：2025-12-26
经最终核实，2024年，GDP现价总量为1348066亿元，比初步核算数减少1018亿元；按不变价格计算，比上年增长5.0%，与初步核算数持平。
"""
PRELIMINARY = """中华人民共和国2025年国民经济和社会发展统计公报
2026 02/28 19:44:03
来源：新华网
一、综合
初步核算，全年国内生产总值[2]1401879亿元，比上年增长5.0%。
[2]国内生产总值、三次产业及相关行业增加值、地区生产总值、人均国内生产总值和国民总收入绝对数按现价计算，增长速度按不变价格计算。
"""


def case(tmp_path, monkeypatch, *, acquire=True):
    store = Store(tmp_path / "data")
    sid = store.create("fingpt", "offline GDP flow")["id"]
    store.session(sid)["delivery_key"] = "one-turn"
    contract = make_contract(
        "a" * 24,
        "GDP，只生成HTML和DOCX，不发送到外部。",
        ["来源与数据", "计算"],
        "2026-10-16",
        "deep",
    )
    receipt = {"research_contract": contract}
    store.data["receipts"][sid + ":one-turn"] = receipt
    service = SimpleNamespace(store=store, lock=asyncio.Lock())
    root = store.directory(sid)
    monkeypatch.chdir(root)
    # The run begins without observations. Acquiring bodies refreshes this same snapshot.
    write_business_snapshot(service, sid, contract)
    assert json.loads((root / "resources/research_evidence.json").read_text())["observations"] == []
    for ordinal, (text, url) in enumerate(
        [(FINAL, "https://www.gov.cn/owned-gdp"), (PRELIMINARY, "https://www.news.cn/owned-gdp")]
        if acquire
        else []
    ):
        identity = ("a" if ordinal == 0 else "b") * 24
        path = "inputs/web-" + identity + ".txt"
        (root / path).write_text(text)
        fid = hashlib.sha256(path.encode()).hexdigest()[:24]
        store.session(sid)["files"][fid] = path
        store.session(sid).setdefault("web_evidence", {})["web:" + identity] = {
            "id": "web:" + identity,
            "source_id": "web:" + identity,
            "path": path,
            "file_id": fid,
            "title": "GDP final" if ordinal == 0 else "GDP preliminary",
            "url": url,
            "sha256": hashlib.sha256(text.encode()).hexdigest(),
            "as_of": "unknown",
        }
    store.save()
    write_business_snapshot(service, sid, contract)

    async def detail(_sid):
        current = json.loads((root / "outputs/research-report.json").read_text())
        delivery = {"files": [{**a, "valid": True} for a in current["artifacts"]]}
        files = store.files(sid)
        result = quality_result(service, sid, receipt, files, stored_sources(store, sid), delivery)
        return {"research_quality": result, "research_report": result["report"]}

    service.detail = detail
    context = {
        "task_id": contract["task_id"],
        "contract_sha256": contract["sha256"],
        "as_of": contract["as_of"],
        "requirement_ids": [r["id"] for r in contract["requirements"]],
        "formats": ["html", "docx"],
    }
    plan = {
        "title": "GDP observation and independent calculation",
        "sections": [
            {
                "id": "facts",
                "heading": "数据",
                "evidence_ids": [
                    "gdp_current_price_final@2024",
                    "gdp_current_price_preliminary@2025",
                ],
                "requirement_ids": ["r1"],
            },
            {
                "id": "calc",
                "heading": "计算",
                "evidence_ids": ["$nominal"],
                "requirement_ids": ["r2"],
            },
        ],
        "calculations": [
            {
                "id": "nominal",
                "operation": "percent_change",
                "operands": ["gdp_current_price_preliminary@2025", "gdp_current_price_final@2024"],
                "decimal_places": 2,
            }
        ],
        "missing": [],
    }
    return service, sid, receipt, context, plan


def test_acquire_to_system_report_check_archive_and_restart(tmp_path, monkeypatch):
    async def run():
        service, sid, receipt, context, plan = case(tmp_path, monkeypatch)
        written = write_business_report_plan(context, plan)
        assert set(written["written"]) == {"html", "docx", "json", "md", "audit"}
        payload = json.loads(Path(written["written"]["json"]).read_text())
        assert len(payload["numeric_claims"]) == 3
        assert payload["numeric_claims"][-1]["value"] == "3.99"
        html = Path(written["written"]["html"]).read_text()
        assert "<table>" in html
        assert "(1401879 − 1348066) ÷ 1348066 × 100 ≈ 3.99 %" in html
        assert "来源记录发布日期：2025-12-26" in html
        assert "来源记录发布日期：2026-02-28" in html
        from docx import Document

        doc = Document(written["written"]["docx"])
        assert doc.tables
        assert any("≈ 3.99 %" in p.text for p in doc.paragraphs)
        # This controlled contract certifies acquired quantities and arithmetic only.
        # The original user's full question has extra coverage/analysis requirements.
        first = await check_quality(service, sid)
        assert first["status"] == "passed", first
        assert first["report_delivery"]["status"] == "bound"
        assert len(first["report_delivery"]["files"]) == 4
        assert await check_quality(service, sid) == first
        for f in first["report_delivery"]["files"]:
            assert (
                hashlib.sha256((service.store.directory(sid) / f["path"]).read_bytes()).hexdigest()
                == f["sha256"]
            )
        service.store.save()
        reopened = Store(service.store.root)
        assert (
            reopened.data["receipts"][sid + ":one-turn"]["quality_last_result"]
            == receipt["quality_last_result"]
        )
        assert (
            quality_versions(SimpleNamespace(store=reopened), sid)[0]["quality"]["status"]
            == "passed"
        )

    asyncio.run(run())


def test_actual_source_callback_refreshes_empty_inventory_without_a_second_route(
    tmp_path, monkeypatch
):
    from app.research_web import research_web_sources as reader
    from app.research_web.resources.research_report import source_id

    async def run():
        service, sid, _, context, plan = case(tmp_path, monkeypatch, acquire=False)
        original_detail = service.detail
        urls = ["https://www.gov.cn/owned-gdp", "https://www.news.cn/owned-gdp"]
        sources = [
            {
                "id": source_id(url),
                "url": url,
                "title": "GDP source " + str(i),
                "state": "search_summary",
            }
            for i, url in enumerate(urls)
        ]
        calls = []

        async def detail(identity):
            current = service.store.directory(sid) / "outputs/research-report.json"
            if not current.exists():
                return {"research_sources": sources}
            return await original_detail(identity)

        async def fetch(url):
            calls.append(url)
            i = urls.index(url)
            return {
                "content": [FINAL, PRELIMINARY][i],
                "as_of": ["2025-12-26", "2026-02-28"][i],
                "final_url": url,
                "response_sha256": "c" * 64,
            }

        service.detail = detail
        monkeypatch.setattr(reader, "fetch_owned_document", fetch)
        for source in sources:
            value = await reader.read_source(service, sid, "read", source["id"])
            assert len(value["numeric_observations"]) == 2
            assert "report_plan" in value["numeric_claim_guidance"]
            assert "build_numeric_claim" not in value["numeric_claim_guidance"]
        assert calls == urls
        write_business_report_plan(context, plan)
        result = await check_quality(service, sid)
        assert result["status"] == "passed", result
        assert len(result["report_delivery"]["files"]) == 4

    asyncio.run(run())


def test_plan_validation_and_same_allowance_repair_preserve_prior_version(tmp_path, monkeypatch):
    async def run():
        service, sid, _, context, plan = case(tmp_path, monkeypatch)
        plan["missing"] = ["需要补充核对"]
        written = write_business_report_plan(context, plan)
        canonical = service.store.directory(sid) / written["written"]["json"]
        original = canonical.read_bytes()
        first = await check_quality(service, sid)
        assert first["status"] == "incomplete" and first["repair_count"] == 1
        broken = copy.deepcopy(plan)
        broken["calculations"][0]["operands"][0] = "invented"
        with pytest.raises(ValueError):
            write_business_report_plan(context, broken)
        assert canonical.read_bytes() == original
        plan["missing"] = []
        write_business_report_plan(context, plan)
        second = await check_quality(service, sid)
        assert second["status"] == "passed" and second["repair_count"] == 1
        versions = quality_versions(service, sid)
        assert len(versions) == 2
        assert next(v for v in versions if not v["is_current"])["quality"]["status"] == "incomplete"

    asyncio.run(run())


@pytest.mark.parametrize("tamper", ["claim", "source", "analysis", "formats", "formula", "date"])
def test_independent_host_never_trusts_written_values_or_skipped_requirements(
    tmp_path, monkeypatch, tamper
):
    async def run():
        service, sid, _, context, plan = case(tmp_path, monkeypatch)
        if tamper == "analysis":
            plan["sections"][0]["analysis"] = "该经济已经进入持续繁荣。"
        written = write_business_report_plan(context, plan)
        path = service.store.directory(sid) / written["written"]["json"]
        payload = json.loads(path.read_text())
        if tamper == "claim":
            payload["numeric_claims"][-1]["value"] = "4.99"
            path.write_text(json.dumps(payload))
        elif tamper == "source":
            p = service.store.directory(sid) / ("inputs/web-" + "a" * 24 + ".txt")
            p.write_text(p.read_text().replace("1348066", "1340000"))
        elif tamper == "formats":
            payload["report_formats"] = ["html"]
            payload["artifacts"] = [a for a in payload["artifacts"] if a["name"].endswith(".html")]
            path.write_text(json.dumps(payload))
        elif tamper == "formula":
            payload["business_calculation_trace"][0]["formula"] = "1 + 1 = 3"
            path.write_text(json.dumps(payload))
        elif tamper == "date":
            payload["sources"][0]["as_of"] = "2020-01-01"
            path.write_text(json.dumps(payload))
        if tamper == "source":
            with pytest.raises(StoreError) as failure:
                await check_quality(service, sid)
            assert failure.value.code == "research_source_integrity_failed"
            return
        result = await check_quality(service, sid)
        assert result["status"] != "passed"

    asyncio.run(run())


def test_zero_supported_evidence_still_delivers_explicit_unpassed_partial(tmp_path, monkeypatch):
    async def run():
        service, sid, _, context, plan = case(tmp_path, monkeypatch)
        plan.pop("calculations")
        plan["sections"] = [
            {
                "id": "gap",
                "heading": "缺口",
                "evidence_ids": [],
                "requirement_ids": ["r1", "r2"],
                "analysis": "所需证据尚未取得，暂不作事实或数值结论。",
            }
        ]
        plan["missing"] = ["尚无支持所需结论的观测"]
        write_business_report_plan(context, plan)
        result = await check_quality(service, sid)
        assert result["status"] == "incomplete"
        assert result["report_delivery"]["status"] == "bound"
        assert any(i.startswith("section_uncited") for i in result["issues"])

    asyncio.run(run())


@pytest.mark.parametrize("failed_name", ["report.md", "research-report.json"])
def test_two_format_publication_failure_retains_current_and_can_retry(
    tmp_path, monkeypatch, failed_name
):
    from app.research_web.resources import research_report_plan as module

    _service, _sid, _, context, plan = case(tmp_path, monkeypatch)
    written = write_business_report_plan(context, plan)
    before = {key: Path(path).read_bytes() for key, path in written["written"].items()}
    replace = module.os.replace

    def fail(src, dst):
        if (
            Path(src).name == failed_name
            and Path(src).parent.name.startswith(".report-plan-")
            and Path(dst).parent == Path("outputs")
        ):
            raise OSError("injected publish failure")
        return replace(src, dst)

    changed = {**plan, "title": "revised report"}
    with monkeypatch.context() as scoped:
        scoped.setattr(module.os, "replace", fail)
        with pytest.raises(OSError, match="injected publish failure"):
            module.write_business_report_plan(context, changed)
    for key, raw in before.items():
        assert Path(written["written"][key]).read_bytes() == raw
    retried = write_business_report_plan(context, changed)
    assert json.loads(Path(retried["written"]["json"]).read_text())["title"] == "revised report"


def test_system_plan_executes_in_original_macos_sandbox(tmp_path, monkeypatch):
    import sys

    from app.research_web.sandbox import SandboxConfig, run_script

    if sys.platform != "darwin":
        pytest.skip("original macOS sandbox belongs to native Mac acceptance")
    service, sid, _, context, plan = case(tmp_path, monkeypatch)
    code = (
        "import json\nfrom research_report_plan import write_business_report_plan\n"
        + f"result=write_business_report_plan({context!r},{plan!r})\n"
        + "assert set(result['written']) == {'json','md','html','docx','audit'}\n"
        + "print('SYSTEM_REPORT_NATIVE_OK')\n"
    )
    result = run_script(
        SandboxConfig(service.store.root, Path(sys.executable)), service.store.directory(sid), code
    )
    assert result.status == "completed", result.stderr
    assert "SYSTEM_REPORT_NATIVE_OK" in result.stdout
