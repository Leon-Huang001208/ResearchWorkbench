"""Delivery admission and real sandbox parsing; native model execution is not simulated."""

import asyncio
import sys
from pathlib import Path

import pytest
from fastapi.testclient import TestClient
from test_api import NativeFixture

from app.research_web.main import create_app
from app.research_web.service import ResearchService
from app.research_web.store import Store


@pytest.fixture
def delivery_api(tmp_path):
    native = NativeFixture()
    original = native.rpc

    async def rpc(method, payload):
        if method == "skill.list":
            return {"skills": [{"name": "fund-evaluation"}]}
        return await original(method, payload)

    native.rpc = rpc
    service = ResearchService(native, Store(tmp_path))
    with TestClient(create_app(service)) as client:
        service.connected = {"mux", "host"}
        sid = client.post("/api/research/sessions", json={}).json()["id"]
        yield client, native, service, sid


def submit(api, formats=None, key="delivery-one", skill=None):
    client, _, _, sid = api
    body = {"text": "交付本次研究"}
    if formats is not None:
        body["expected_formats"] = formats
    if skill:
        body["skill_id"] = skill
    return client.post(
        f"/api/research/sessions/{sid}/messages",
        json=body,
        headers={"Idempotency-Key": key},
    )


def finish(api, reason="completed"):
    _, native, service, sid = api
    prompt = [payload for method, payload in native.calls if method == "session.prompt"][-1]
    first = max(service.events.get(sid, {}), default=0) + 1
    service.events.setdefault(sid, {}).update(
        {
            first: {
                "event": {
                    "seq": first,
                    "type": "user/message",
                    "data": {"content": prompt["content"]},
                }
            },
            first
            + 1: {
                "event": {
                    "seq": first + 1,
                    "type": "turn/end",
                    "data": {"reason": {"kind": reason}},
                }
            },
        }
    )
    service.running[sid] = False
    service.loaded.add(sid)


def detail(api):
    client, _, _, sid = api
    response = client.get(f"/api/research/sessions/{sid}")
    assert response.status_code == 200, response.text
    return response.json()


def test_format_contract_rejects_unknown_format(delivery_api):
    assert submit(delivery_api, ["exe"]).status_code == 422


def test_explicit_no_files_overrides_skill_in_native_prompt(delivery_api):
    _, native, _, _ = delivery_api
    assert submit(delivery_api, [], skill="fund-evaluation").status_code == 202
    prompt = [payload for method, payload in native.calls if method == "session.prompt"][-1]
    assert "无需文件" in prompt["content"][0]["text"]
    assert detail(delivery_api)["delivery"]["required_formats"] == []


def test_fund_default_and_explicit_override_are_bound_to_receipt(delivery_api):
    assert submit(delivery_api, skill="fund-evaluation").status_code == 202
    row = detail(delivery_api)
    assert row["delivery"]["required_formats"] == ["docx", "html", "xlsx"]
    assert row["delivery"]["status"] == "pending"
    finish(delivery_api)
    assert detail(delivery_api)["delivery"]["status"] == "incomplete"
    assert (
        submit(delivery_api, ["md"], key="delivery-two", skill="fund-evaluation").status_code == 202
    )
    assert detail(delivery_api)["delivery"]["required_formats"] == ["md"]


def test_completed_turn_with_missing_files_is_not_delivered(delivery_api):
    assert submit(delivery_api, ["docx", "xlsx"]).status_code == 202
    finish(delivery_api)
    row = detail(delivery_api)
    assert row["status"] == "completed"
    assert row["delivery"]["status"] == "incomplete"
    assert row["delivery"]["missing_formats"] == ["docx", "xlsx"]
    assert row["delivery"]["files"] == []


@pytest.mark.skipif(sys.platform != "darwin", reason="strict parser requires native macOS sandbox")
def test_old_files_inputs_and_same_name_hash_update(delivery_api):
    _, _, service, sid = delivery_api
    root = service.store.directory(sid)
    (root / "outputs/report.md").write_text("old report")
    (root / "inputs/input.html").write_text("<p>Not an output</p>")
    assert submit(delivery_api, ["md", "html"]).status_code == 202
    finish(delivery_api)
    row = detail(delivery_api)["delivery"]
    assert row["status"] == "incomplete"
    assert row["missing_formats"] == ["md", "html"]
    assert not row["files"]
    assert submit(delivery_api, ["md"], key="delivery-two").status_code == 202
    (root / "outputs/report.md").write_text("new report")
    finish(delivery_api)
    row = detail(delivery_api)["delivery"]
    assert row["status"] == "completed"
    assert row["files"][0]["valid"] is True
    saved = Store(service.store.root).receipt(sid, "delivery-two")["delivery"]
    assert saved["files"] == row["files"]


@pytest.mark.skipif(sys.platform != "darwin", reason="strict parser requires native macOS sandbox")
def test_real_sandbox_reopens_documents_and_rejects_corrupt_empty_files(delivery_api):
    from app.research_web.sandbox import SandboxConfig, run_script

    _, _, service, sid = delivery_api
    assert submit(delivery_api, ["docx", "xlsx", "html", "md", "png"]).status_code == 202
    root = service.store.directory(sid)
    generated = run_script(
        SandboxConfig(service.store.root, Path(sys.executable)),
        root,
        """
from docx import Document
from openpyxl import Workbook
from PIL import Image
d = Document(); d.add_paragraph('Actual research'); d.save('outputs/good.docx')
Document().save('outputs/empty.docx')
w = Workbook(); w.active.append(['metric', 'value']); w.active.append(['return', 0]); w.save('outputs/good.xlsx')
w = Workbook(); w.active.append(['metric', 'value']); w.save('outputs/empty.xlsx')
open('outputs/bad.docx', 'wb').write(b'not an office file')
open('outputs/report.html', 'w').write('<html><body><p>Actual report</p></body></html>')
open('outputs/empty.html', 'w').write('<html><script>hidden text</script></html>')
open('outputs/empty.md', 'w').write('   ')
Image.new('RGB', (2, 2)).save('outputs/chart.png')
""",
    )
    assert generated.status == "completed", generated
    finish(delivery_api)
    row = detail(delivery_api)["delivery"]
    assert row["status"] == "incomplete"
    assert row["missing_formats"] == ["md"]
    files = {item["name"]: item for item in row["files"]}
    for name in ["good.docx", "good.xlsx", "report.html", "chart.png"]:
        assert files[name]["valid"] is True, files[name]
    for name in ["bad.docx", "empty.docx", "empty.xlsx", "empty.html", "empty.md"]:
        assert files[name]["valid"] is False, files[name]
        assert files[name]["reason"]


@pytest.mark.skipif(sys.platform != "darwin", reason="strict parser requires native macOS sandbox")
@pytest.mark.parametrize(
    "metadata_sheet", ["sources", "source", " Sources ", "来源", "说明", "参考资料"]
)
def test_metadata_only_workbook_does_not_satisfy_analysis_delivery(delivery_api, metadata_sheet):
    from app.research_web.sandbox import SandboxConfig, run_script

    _, _, service, sid = delivery_api
    assert submit(delivery_api, ["xlsx"]).status_code == 202
    generated = run_script(
        SandboxConfig(service.store.root, Path(sys.executable)),
        service.store.directory(sid),
        "from openpyxl import Workbook\n"
        "book = Workbook(); book.active.title = 'analysis'\n"
        f"sources = book.create_sheet({metadata_sheet!r})\n"
        "sources.append(['https://example.com/one']); sources.append(['https://example.com/two'])\n"
        "book.save('outputs/report.xlsx')\n",
    )
    assert generated.status == "completed", generated
    finish(delivery_api)
    delivery = detail(delivery_api)["delivery"]
    assert delivery["status"] == "incomplete"
    assert delivery["missing_formats"] == ["xlsx"]
    assert "非元数据" in delivery["files"][0]["reason"]


@pytest.mark.parametrize("multi", [None, False, True])
def test_native_question_single_selection_is_validated(delivery_api, multi):
    client, native, service, sid = delivery_api
    question = {
        "id": "period",
        "question": "期间",
        "options": [{"label": "2024"}, {"label": "2025"}],
    }
    if multi is not None:
        question["multiSelect"] = multi
    service.questions["rpc"] = {"sessionId": sid, "questions": [question]}
    responses = []

    async def respond(rpc_id, body):
        responses.append((rpc_id, body))
        return {"accepted": True}

    native.respond = respond
    response = client.post(
        f"/api/research/sessions/{sid}/questions/rpc",
        json={"answers": [{"id": "period", "selected": ["2024", "2025"], "custom": ""}]},
    )
    assert response.status_code == (200 if multi else 400)
    assert len(responses) == (1 if multi else 0)


@pytest.mark.skipif(sys.platform != "darwin", reason="strict parser requires native macOS sandbox")
def test_analysis_sheet_with_sources_and_zero_value_is_valid(delivery_api):
    from app.research_web.sandbox import SandboxConfig, run_script

    _, _, service, sid = delivery_api
    assert submit(delivery_api, ["xlsx"]).status_code == 202
    generated = run_script(
        SandboxConfig(service.store.root, Path(sys.executable)),
        service.store.directory(sid),
        "from openpyxl import Workbook\n"
        "book = Workbook(); book.active.title = 'analysis'\n"
        "book.active.append(['metric', 'value']); book.active.append(['return', 0])\n"
        "sources = book.create_sheet('sources'); sources.append(['url']); sources.append(['https://example.com'])\n"
        "book.save('outputs/report.xlsx')\n",
    )
    assert generated.status == "completed", generated
    finish(delivery_api)
    assert detail(delivery_api)["delivery"]["status"] == "completed"


def test_idempotency_unknown_and_serial_admission(delivery_api):
    _, native, _, _ = delivery_api
    assert submit(delivery_api, ["md"]).status_code == 202
    assert submit(delivery_api, ["md"]).status_code == 202
    assert submit(delivery_api, ["html"]).status_code == 400
    assert submit(delivery_api, ["md"], key="delivery-two").status_code == 503
    assert len([call for call in native.calls if call[0] == "session.prompt"]) == 1
    finish(delivery_api)
    detail(delivery_api)
    native.fail_prompt = True
    assert submit(delivery_api, ["md"], key="delivery-two").status_code == 503
    assert submit(delivery_api, ["md"], key="delivery-two").status_code == 503
    assert len([call for call in native.calls if call[0] == "session.prompt"]) == 2
    assert detail(delivery_api)["delivery"]["status"] == "admission_unknown"


def test_plain_chat_does_not_invoke_parser(delivery_api, monkeypatch):
    from app.research_web import sandbox

    def forbidden(*args):
        raise AssertionError("plain chat must not parse files")

    monkeypatch.setattr(sandbox, "run_script", forbidden)
    assert submit(delivery_api).status_code == 202
    finish(delivery_api)
    assert detail(delivery_api)["delivery"]["status"] == "not_required"


def test_restart_restores_exact_receipt_and_unknown_admission_from_native_history(delivery_api):
    _, native, service, sid = delivery_api
    native.fail_prompt = True
    assert submit(delivery_api, ["md"]).status_code == 503
    finish(delivery_api)
    entries = list(service.events[sid].values())

    async def history(_sid):
        return entries

    native.history = history
    restored = ResearchService(native, Store(service.store.root))
    with TestClient(create_app(restored)) as after_restart:
        restored.connected = {"mux", "host"}
        state = after_restart.get(f"/api/research/sessions/{sid}").json()
        assert state["delivery"]["status"] == "incomplete"
        assert restored.store.receipt(sid, "delivery-one")["status"] == "accepted"
        response = after_restart.post(
            f"/api/research/sessions/{sid}/messages",
            json={"text": "交付本次研究", "expected_formats": ["md"]},
            headers={"Idempotency-Key": "delivery-one"},
        )
        assert response.status_code == 202
    assert len([call for call in native.calls if call[0] == "session.prompt"]) == 1


def test_historical_completion_without_current_marker_cannot_finalize(delivery_api):
    _, _, service, sid = delivery_api
    service.events[sid] = {
        1: {"event": {"seq": 1, "type": "turn/end", "data": {"reason": {"kind": "completed"}}}}
    }
    assert submit(delivery_api, ["md"]).status_code == 202
    service.running[sid] = False
    assert detail(delivery_api)["delivery"]["status"] == "pending"
    assert submit(delivery_api, [], key="delivery-two").status_code == 503


def test_running_child_prevents_finalization_after_parent_ends(delivery_api):
    _, native, _, _ = delivery_api
    assert submit(delivery_api, ["md"]).status_code == 202
    finish(delivery_api)
    original = native.rpc

    async def rpc(method, payload):
        if method == "subagent.list":
            return {
                "entries": [
                    {"id": "child", "kind": "child", "activity": "running", "mode": "continuable"}
                ]
            }
        if method == "subagent.history":
            return {"events": [], "hasMore": False}
        return await original(method, payload)

    native.rpc = rpc
    assert detail(delivery_api)["delivery"]["status"] == "pending"
    assert submit(delivery_api, [], key="delivery-two").status_code == 503


def test_verifier_failure_is_not_success_and_is_not_retried_implicitly(delivery_api, monkeypatch):
    from app.research_web import sandbox

    _, _, service, sid = delivery_api
    calls = []

    def unavailable(*args):
        calls.append(args)
        return sandbox.ScriptResult("failed", error="sandbox unavailable")

    monkeypatch.setattr(sandbox, "run_script", unavailable)
    assert submit(delivery_api, ["md"]).status_code == 202
    (service.store.directory(sid) / "outputs/report.md").write_text("real output")
    finish(delivery_api)
    assert detail(delivery_api)["delivery"]["status"] == "verification_failed"
    assert detail(delivery_api)["delivery"]["status"] == "verification_failed"
    assert len(calls) == 1


@pytest.mark.skipif(sys.platform != "darwin", reason="strict parser requires native macOS sandbox")
def test_parser_detects_same_name_swap_and_rejects_linked_files(delivery_api, monkeypatch):
    import os

    from app.research_web import sandbox

    _, _, service, sid = delivery_api
    assert submit(delivery_api, ["md"]).status_code == 202
    root = service.store.directory(sid)
    report = root / "outputs/report.md"
    report.write_text("before validation")
    (root / "inputs/source.md").write_text("input contents")
    os.link(root / "inputs/source.md", root / "outputs/linked.md")
    (root / "outputs/symlink.md").symlink_to(root / "inputs/source.md")
    original = sandbox.run_script

    def swap(config, directory, code):
        report.write_text("changed during validation")
        return original(config, directory, code)

    monkeypatch.setattr(sandbox, "run_script", swap)
    finish(delivery_api)
    row = detail(delivery_api)["delivery"]
    assert row["status"] == "incomplete"
    files = {item["name"]: item for item in row["files"]}
    assert files["report.md"]["valid"] is False
    assert "变化" in files["report.md"]["reason"]
    assert files["linked.md"]["valid"] is False
    assert "symlink.md" not in files


@pytest.mark.skipif(sys.platform != "darwin", reason="strict parser requires native macOS sandbox")
@pytest.mark.parametrize("mutation", ["empty", "same_size", "delete"])
def test_files_changed_after_parser_success_cannot_be_published_as_valid(
    delivery_api, monkeypatch, mutation
):
    import json

    from app.research_web import sandbox

    _, _, service, sid = delivery_api
    assert submit(delivery_api, ["md"]).status_code == 202
    report = service.store.directory(sid) / "outputs/report.md"
    report.write_text("valid report")
    original = sandbox.run_script

    def parse_then_mutate(config, directory, code):
        result = original(config, directory, code)
        assert result.status == "completed"
        assert json.loads(result.stdout)[0]["valid"] is True
        if mutation == "delete":
            report.unlink()
        else:
            report.write_text("" if mutation == "empty" else "other report")
        return result

    monkeypatch.setattr(sandbox, "run_script", parse_then_mutate)
    finish(delivery_api)
    row = detail(delivery_api)["delivery"]
    assert row["status"] == "incomplete"
    assert row["missing_formats"] == ["md"]
    assert row["files"][0]["valid"] is False
    assert "解析后" in row["files"][0]["reason"]
    stored = Store(service.store.root).receipt(sid, "delivery-one")["delivery"]
    assert stored["status"] == "incomplete"
    assert stored["files"][0]["valid"] is False


@pytest.mark.asyncio
async def test_concurrent_detail_requests_parse_once(tmp_path, monkeypatch):
    from app.research_web import sandbox
    from app.research_web.delivery import Delivery

    store = Store(tmp_path)
    sid = store.create("fingpt", "test")["id"]
    delivery = Delivery(store)
    metadata = delivery.begin(sid, "key", ["md"])
    store.reserve(sid, "key", "digest", metadata)
    store.receipt(sid, "key", "accepted")
    (store.directory(sid) / "outputs/result.md").write_text("content")
    calls = []

    def unavailable(*args):
        calls.append(args)
        return sandbox.ScriptResult("failed")

    monkeypatch.setattr(sandbox, "run_script", unavailable)
    entries = [
        {"event": {"seq": 1, "type": "user/message", "data": {"content": metadata["marker"]}}},
        {"event": {"seq": 2, "type": "turn/end", "data": {"reason": {"kind": "completed"}}}},
    ]
    results = await asyncio.gather(
        *[delivery.refresh(sid, entries, busy=False, connected=True) for _ in range(3)]
    )
    assert len(calls) == 1
    assert all(result["status"] == "verification_failed" for result in results)
