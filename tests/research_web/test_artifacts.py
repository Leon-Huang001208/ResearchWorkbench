"""Actual document files, not chat-text substitutes."""

import builtins
from pathlib import Path

import pytest

from app.research_web.resources.research_helpers import read_pdf, write_deliverables


def test_markdown_only_delivery_avoids_office_imports_and_extra_files(tmp_path, monkeypatch):
    original_import = builtins.__import__

    def deny_office(name, *args, **kwargs):
        if name in {"docx", "openpyxl", "pdfplumber"}:
            raise AssertionError(f"Unrequested format imported: {name}")
        return original_import(name, *args, **kwargs)

    monkeypatch.setattr(builtins, "__import__", deny_office)
    outputs = write_deliverables(
        tmp_path, "研究", [{"heading": "结论", "text": "正文"}], [], [], formats=["md"]
    )
    assert set(outputs) == {"markdown"}
    assert isinstance(outputs["markdown"], Path)
    assert list(tmp_path.iterdir()) == [outputs["markdown"]]
    assert "正文" in outputs["markdown"].read_text()


@pytest.mark.parametrize("formats", [[], ["csv"], "md", [None], {"md": True}])
def test_invalid_format_selection_fails_before_creating_output(tmp_path, formats):
    output = tmp_path / "outputs"
    with pytest.raises(ValueError, match="formats"):
        write_deliverables(output, "研究", [], [], [], formats=formats)
    assert not output.exists()


def test_pdf_flag_cannot_add_unselected_format(tmp_path):
    with pytest.raises(ValueError, match="pdf=True"):
        write_deliverables(tmp_path / "outputs", "研究", [], [], [], formats=["md"], pdf=True)
    assert not (tmp_path / "outputs").exists()


def test_explicit_pdf_only_is_real_readable_file(tmp_path):
    outputs = write_deliverables(
        tmp_path,
        "Research",
        [{"heading": "Result", "text": "Revenue 123"}],
        [],
        [],
        formats=["pdf"],
    )
    assert set(outputs) == {"pdf"}
    assert list(tmp_path.iterdir()) == [outputs["pdf"]]
    assert "123" in read_pdf(outputs["pdf"])[0]["text"]


@pytest.mark.parametrize("formats", [("html",), ["docx"], ["xlsx"], ["md", "md", "html"]])
def test_explicit_formats_generate_only_selected_readable_outputs(tmp_path, formats):
    outputs = write_deliverables(
        tmp_path,
        "Research",
        [{"heading": "Result", "text": "123"}],
        [{"value": 123}],
        [],
        formats=formats,
    )
    expected = {"markdown" if item == "md" else item for item in formats}
    assert set(outputs) == expected
    assert set(tmp_path.iterdir()) == set(outputs.values())
    if "docx" in outputs:
        from docx import Document

        assert Document(outputs["docx"]).paragraphs[0].text == "Research"
    if "xlsx" in outputs:
        from openpyxl import load_workbook

        book = load_workbook(outputs["xlsx"])
        assert book["analysis"]["A2"].value == 123
        book.close()


def test_legacy_pdf_flag_keeps_all_four_default_formats(tmp_path):
    outputs = write_deliverables(
        tmp_path, "Research", [{"heading": "Result", "text": "123"}], [], [], pdf=True
    )
    assert set(outputs) == {"markdown", "html", "docx", "xlsx", "pdf"}
    assert set(tmp_path.iterdir()) == set(outputs.values())


def test_requested_office_files_are_real_openable_documents(tmp_path):
    from docx import Document
    from openpyxl import load_workbook

    outputs = write_deliverables(
        tmp_path,
        "研究",
        [{"heading": "结论", "text": "缺少行情数据，不做估值结论"}],
        [{"metric": "PE", "value": None}],
        ["https://example.com/source"],
    )
    assert set(outputs) == {"markdown", "html", "docx", "xlsx"}
    assert Document(outputs["docx"]).paragraphs[0].text == "研究"
    workbook = load_workbook(outputs["xlsx"])
    assert workbook["analysis"]["B2"].value is None
    assert "https://example.com/source" in outputs["html"].read_text()


def test_pdf_extraction_includes_real_page_numbers(tmp_path):
    from matplotlib.figure import Figure

    path = tmp_path / "input.pdf"
    figure = Figure()
    figure.text(0.1, 0.8, "Revenue 123")
    figure.savefig(path)
    pages = read_pdf(path)
    assert pages[0]["page"] == 1
    assert "123" in pages[0]["text"]
