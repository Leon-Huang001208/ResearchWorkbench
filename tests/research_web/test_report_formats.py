"""Native report formats and numeric tables without additional display facts."""

import json
import re

import pytest
from docx import Document

from app.research_web.resources.research_report import (
    archive_report_version,
    list_report_versions,
    required_report_formats,
    validate_payload,
    write_research_report,
)


def payload():
    text = "2025年全年GDP现价总量为1401879亿元。"
    return {
        "task_id": "a" * 24,
        "title": "GDP研究",
        "as_of": "2026-10-16",
        "sections": [{"id": "gdp", "heading": "数据", "text": text, "citations": ["web:one"]}],
        "sources": [
            {"id": "web:one", "title": "原始<script>资料", "url": "https://example.com/a?x=1&y=2"}
        ],
        "notes": [{"stage_id": "gdp", "text": text, "citations": ["web:one"]}],
        "missing": ["最终核实资料未取得"],
        "numeric_claims": [
            {
                "id": "gdp_2025",
                "section_id": "gdp",
                "text": text,
                "value": "1401879",
                "period": "2025年",
                "unit": "亿元",
                "operation": "identity",
                "operands": [
                    {
                        "observation_id": "gdp2025",
                        "value": "1401879",
                        "period": "2025年",
                        "unit": "亿元",
                        "quote": text,
                    }
                ],
            }
        ],
    }


@pytest.mark.parametrize(
    "question,expected",
    [
        ("只生成HTML和DOCX，不发送到外部。", ("html", "docx")),
        ({"question": "仅生成 docx、html。"}, ("html", "docx")),
        ("仅生成PDF。", ("pdf",)),
        ("仅生成HTML、DOCX和PDF。", ("html", "docx", "pdf")),
        ("生成HTML和DOCX。", ("html", "docx", "pdf")),
        ("不要只生成HTML和DOCX。", ("html", "docx", "pdf")),
        ("不仅生成HTML和DOCX。", ("html", "docx", "pdf")),
        ("仅生成HTML，DOCX，XLSX。", ("html", "docx", "pdf")),
        ("仅生成HTML和XLSX。", ("html", "docx", "pdf")),
        ("仅生成HTML和DOCX兼容格式。", ("html", "docx", "pdf")),
        ("仅生成HTML。仅生成PDF。", ("html", "docx", "pdf")),
        ({"report_formats": ["html"]}, ("html", "docx", "pdf")),
    ],
)
def test_only_explicit_bounded_format_declaration_changes_legacy_default(question, expected):
    assert required_report_formats(question) == expected


def test_explicit_native_formats_only_from_the_same_frozen_delivery():
    contract = {"task_id": "a" * 24, "question": "只生成HTML和DOCX。"}
    delivery = {"task_id": contract["task_id"], "required_formats": ["html"]}
    assert required_report_formats(contract, delivery) == ("html",)
    assert required_report_formats(contract, {**delivery, "required_formats": []}) == (
        "html",
        "docx",
    )
    assert required_report_formats(contract) == ("html", "docx")
    for changed in [
        {**delivery, "task_id": "b" * 24},
        {**delivery, "required_formats": ["html", "exe"]},
    ]:
        with pytest.raises(ValueError, match="research_report_delivery_binding"):
            required_report_formats(contract, changed)


def test_two_formats_native_table_preserves_canonical_text_and_evidence(tmp_path):
    value = payload()
    paths = write_research_report(tmp_path, value, formats=("html", "docx"))
    assert set(paths) == {"json", "html", "docx"}
    assert not list(tmp_path.glob("*.pdf"))
    saved = json.loads(paths["json"].read_text())
    assert saved["report_formats"] == ["html", "docx"]
    assert saved["sections"] == value["sections"]
    assert saved["notes"] == value["notes"]
    assert saved["sources"] == value["sources"]
    assert saved["missing"] == value["missing"]
    html = paths["html"].read_text()
    assert value["sections"][0]["text"] in html
    assert "<script>" not in html and "&lt;script&gt;" in html
    assert "最终核实资料未取得" in html
    assert "https://example.com/a?x=1&amp;y=2" in html
    assert "<table>" in html
    assert re.findall(r"<td>(.*?)</td>", html) == [
        value["numeric_claims"][0]["text"],
        "2025年",
        "1401879",
        "亿元",
    ]
    document = Document(paths["docx"])
    assert value["sections"][0]["text"] in [p.text for p in document.paragraphs]
    assert [[cell.text for cell in row.cells] for row in document.tables[0].rows][1] == [
        value["numeric_claims"][0]["text"],
        "2025年",
        "1401879",
        "亿元",
    ]
    version = archive_report_version(tmp_path)
    assert version["missing_formats"] == []
    assert list_report_versions(tmp_path)[0]["version_id"] == version["version_id"]


@pytest.mark.parametrize(
    "field,wrong",
    [("text", "2025年GDP为42亿元。"), ("value", "42"), ("unit", "美元"), ("period", "2024年")],
)
def test_inconsistent_claim_never_adds_numeric_table(tmp_path, field, wrong):
    value = payload()
    value["numeric_claims"][0][field] = wrong
    paths = write_research_report(tmp_path, value, formats=("html", "docx"))
    assert "<table>" not in paths["html"].read_text()
    assert not Document(paths["docx"]).tables
    assert value["sections"][0]["text"] in paths["html"].read_text()


def test_archive_retains_missing_selected_format_without_requiring_pdf(tmp_path):
    paths = write_research_report(tmp_path, payload(), formats=("html", "docx"))
    paths["docx"].unlink()
    version = archive_report_version(tmp_path)
    assert version["missing_formats"] == ["docx"]
    assert version["missing_artifacts"] == [paths["docx"].name]
    assert list_report_versions(tmp_path)[0]["quality"] == "incomplete"


def test_table_escapes_claim_markup_and_preserves_full_decimal(tmp_path):
    value = payload()
    text = "<b>2025年名义增速为3.912345678901234567890123456%。</b>"
    value["sections"][0]["text"] = text
    value["numeric_claims"][0].update(text=text, value="3.912345678901234567890123456", unit="%")
    paths = write_research_report(tmp_path, value, formats=("html", "docx"))
    html = paths["html"].read_text()
    assert "<b>" not in html and "&lt;b&gt;" in html
    assert "<td>3.912345678901234567890123456</td>" in html
    assert Document(paths["docx"]).tables[0].cell(1, 2).text == "3.912345678901234567890123456"


def test_default_remains_three_formats_and_payload_metadata_cannot_select(tmp_path):
    value = payload()
    value["report_formats"] = ["html"]
    paths = write_research_report(tmp_path, value)
    assert set(paths) == {"json", "html", "docx", "pdf"}
    assert "report_formats" not in json.loads(paths["json"].read_text())
    assert archive_report_version(tmp_path)["missing_formats"] == []


@pytest.mark.parametrize(
    "formats",
    [[], ["html", "html"], ["xlsx"], ["HTML"], "html", {"html"}, ["html", "docx", "pdf", "html"]],
)
def test_invalid_format_selection_is_rejected_before_writing(tmp_path, formats):
    with pytest.raises(ValueError, match="research_report_formats_invalid"):
        write_research_report(tmp_path, payload(), formats=formats)
    assert list(tmp_path.iterdir()) == []
    value = payload()
    value["report_formats"] = formats
    with pytest.raises(ValueError, match="research_report_formats_invalid"):
        validate_payload(value)
