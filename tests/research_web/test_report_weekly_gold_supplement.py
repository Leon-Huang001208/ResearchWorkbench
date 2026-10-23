"""Synthetic raw replay of the single reviewed London gold history repair."""

import copy
import hashlib
import json
from datetime import date, datetime, timedelta
from io import BytesIO
from zipfile import ZIP_DEFLATED, ZipFile

import pytest
from openpyxl import load_workbook
from test_datahub_wind_excel_weekly import days, exchange_package, replace_resource
from test_report_weekly_data import synthetic_snapshot

from app.research_web.datahub.providers import ProviderError
from app.research_web.report_workflows import weekly_sources as sources
from app.research_web.report_workflows.models import WorkflowError
from app.research_web.report_workflows.weekly_data import assess_weekly_snapshot


def gold_package(root):
    contract, source, manifest, save, _ = exchange_package(root, "2026-10-23T15:30:00+08:00")
    path = source["excel_exports"]["edb"]["path"]
    book = load_workbook(root / path)
    for row in book.active:
        if row[0].value == datetime(2017, 1, 30) or row[0].value == datetime(2026, 10, 23):
            row[1].value = None
    book.active.append(["数据来源：Wind"])
    stream = BytesIO()
    book.save(stream)
    book.close()
    edb_sha = replace_resource(manifest, save, path, stream.getvalue())
    source["excel_exports"]["edb"]["sha256"] = edb_sha
    source["calendars"]["LONDON_GOLD"].update(
        wind_calendar="IDC", timezone="Europe/London", close_time="17:00"
    )
    proof = {
        "schema_version": 1,
        "source": "wind_excel",
        "series_id": "history_gold_london",
        "original_edb_sha256": edb_sha,
        "unit": "USD/ounce",
        "currency": "USD",
        "period_end": "2026-10-23",
        "coverage_start": "2017-01-01",
        "coverage_end": "2026-10-22",
        "observed_at": "2026-10-23T15:10:00+08:00",
        "history_requests": [],
        "calendar_requests": [],
        "recheck": {
            "formula": '=wsd("SPTAUUSDOZ.IDC","close","2017-01-30","2017-01-30","Period=D","PriceAdj=U","Days=Trading","Fill=Blank","TradingCalendar=IDC")',
            "rows": [["2017-01-30 00:00:00", 100]],
        },
    }
    for kind, end in (("history", "2026-10-22"), ("calendar", "2026-10-22")):
        cursor, last = date(2017, 1, 1), date.fromisoformat(end)
        while cursor <= last:
            finish = min(cursor + timedelta(days=365), last)
            start, stop = cursor.isoformat(), finish.isoformat()
            if kind == "history":
                formula = f'=wsd("SPTAUUSDOZ.IDC","close","{start}","{stop}","Period=D","PriceAdj=U","Days=Trading","Fill=Blank","TradingCalendar=IDC")'
                request = {
                    "formula": formula,
                    "raw_rows": [[d + " 00:00:00", 100] for d in days(start, stop)],
                }
            else:
                formula = f'=tdays("{start}","{stop}","TradingCalendar=IDC")'
                request = {
                    "formula": formula,
                    "raw_rows": [[d + " 00:00:00"] for d in days(start, stop)],
                }
            proof[kind + "_requests"].append(request)
            cursor = finish + timedelta(days=1)
    source["schema_version"] = 4
    source["history_supplement"] = {
        "path": "inputs/gold-supplement.json",
        "sha256": save("inputs/gold-supplement.json", json.dumps(proof).encode()),
    }
    replace_resource(manifest, save, sources.SOURCE_CONTRACT, json.dumps(source).encode())
    snapshot = synthetic_snapshot(contract)
    calendar = snapshot["calendars"]["LONDON_GOLD"]
    calendar["session_closes"] = {d: d + "T17:00:00+01:00" for d in calendar["dates"]}
    calendar["completed_dates"] = [d for d in calendar["dates"] if d <= "2026-10-22"]
    for spec in __import__(
        "app.research_web.report_workflows.weekly_data", fromlist=["load_weekly_mapping"]
    ).load_weekly_mapping()["series"]:
        if spec["market"] == "LONDON_GOLD":
            snapshot["series"][spec["id"]]["rows"] = [
                row for row in snapshot["series"][spec["id"]]["rows"] if row["date"] <= "2026-10-22"
            ]
    return contract, source, manifest, save, proof, snapshot


def install(root, source, manifest, save, proof):
    source["history_supplement"]["sha256"] = replace_resource(
        manifest, save, "inputs/gold-supplement.json", json.dumps(proof).encode()
    )
    replace_resource(manifest, save, sources.SOURCE_CONTRACT, json.dumps(source).encode())


def replay(root, source, manifest, contract, snapshot):
    return sources.read_gold_history_supplement(
        root, manifest, source, contract, snapshot["calendars"]["LONDON_GOLD"]["completed_dates"]
    )


def test_schema4_and_single_missing_point_replay_preserve_edb_bytes(tmp_path):
    contract, source, manifest, _, _, snapshot = gold_package(tmp_path)
    before = (tmp_path / "inputs/edb.xlsx").read_bytes()
    loaded, _ = sources.load_run_source_contract(tmp_path, manifest)
    assert loaded == source
    rows, lineage = replay(tmp_path, source, manifest, contract, snapshot)
    assert next(row for row in rows if row[0] == "2017-01-30")[1] == 100
    assert lineage["original_edb_sha256"] == hashlib.sha256(before).hexdigest()
    assert lineage["supplement_sha256"] == source["history_supplement"]["sha256"]
    assert (tmp_path / "inputs/edb.xlsx").read_bytes() == before


def test_schema4_requires_original_schema3_india_calendar(tmp_path):
    _, source, manifest, save, _, _ = gold_package(tmp_path)
    source["calendars"]["IN"] = {
        "semantics_verified": True,
        "wind_calendar": "SSE",
        "timezone": "Asia/Shanghai",
        "close_time": "10:00",
    }
    replace_resource(manifest, save, sources.SOURCE_CONTRACT, json.dumps(source).encode())
    with pytest.raises(WorkflowError, match="数据源合同未核实"):
        sources.load_run_source_contract(tmp_path, manifest)


def test_replay_preserves_original_edb_non_session_observation(tmp_path):
    contract, source, manifest, save, proof, snapshot = gold_package(tmp_path)
    path = tmp_path / "inputs/edb.xlsx"
    book = load_workbook(path)
    footer_row = book.active.max_row
    book.active.insert_rows(footer_row)
    book.active.cell(footer_row, 1, datetime(2017, 1, 29))
    book.active.cell(footer_row, 2, 1190.5)
    stream = BytesIO()
    book.save(stream)
    book.close()
    sha = replace_resource(manifest, save, "inputs/edb.xlsx", stream.getvalue())
    source["excel_exports"]["edb"]["sha256"] = proof["original_edb_sha256"] = sha
    install(tmp_path, source, manifest, save, proof)
    rows, lineage = replay(tmp_path, source, manifest, contract, snapshot)
    assert ("2017-01-29", 1190.5) in rows
    assert (
        lineage["overlap_count"] == len(snapshot["calendars"]["LONDON_GOLD"]["completed_dates"]) - 1
    )


def test_gold_replay_checks_workbook_budget_before_formula_scan(tmp_path, monkeypatch):
    contract, source, manifest, save, proof, snapshot = gold_package(tmp_path)
    stream = BytesIO()
    with (
        ZipFile(stream, "w", compression=ZIP_DEFLATED) as archive,
        archive.open("xl/worksheets/sheet1.xml", "w") as worksheet,
    ):
        chunk = b"0" * (1024 * 1024)
        for _ in range(65):
            worksheet.write(chunk)
    sha = replace_resource(manifest, save, "inputs/edb.xlsx", stream.getvalue())
    source["excel_exports"]["edb"]["sha256"] = proof["original_edb_sha256"] = sha
    install(tmp_path, source, manifest, save, proof)
    scans = []

    def forbidden_formula_scan(*_args, **_kwargs):
        scans.append("attempted")
        raise AssertionError("formula_scan_before_budget")

    monkeypatch.setattr(sources, "ZipFile", forbidden_formula_scan)
    with pytest.raises(WorkflowError) as failure:
        replay(tmp_path, source, manifest, contract, snapshot)
    assert isinstance(failure.value.__cause__, ProviderError)
    assert str(failure.value.__cause__) == "weekly_excel_export_size_limit"
    assert scans == []


def test_gold_replay_still_rejects_formula_after_budget_check(tmp_path):
    contract, source, manifest, save, proof, snapshot = gold_package(tmp_path)
    book = load_workbook(tmp_path / "inputs/edb.xlsx")
    book.active["F10"] = "=1+1"
    stream = BytesIO()
    book.save(stream)
    book.close()
    sha = replace_resource(manifest, save, "inputs/edb.xlsx", stream.getvalue())
    source["excel_exports"]["edb"]["sha256"] = proof["original_edb_sha256"] = sha
    install(tmp_path, source, manifest, save, proof)
    with pytest.raises(WorkflowError) as failure:
        replay(tmp_path, source, manifest, contract, snapshot)
    assert str(failure.value.__cause__) == "gold_edb_formula"


@pytest.mark.parametrize(
    "value,reason",
    [
        (999999, "gold_edb_unclosed_observation"),
        ("unverified", "gold_edb_unclosed_observation"),
        (True, "gold_edb_bool"),
        (False, "gold_edb_bool"),
    ],
)
def test_gold_replay_rejects_unclosed_nonempty_original_edb_value(tmp_path, value, reason):
    contract, source, manifest, save, proof, snapshot = gold_package(tmp_path)
    book = load_workbook(tmp_path / "inputs/edb.xlsx")
    for row in book.active:
        if row[0].value == datetime(2026, 10, 23):
            row[1].value = value
    stream = BytesIO()
    book.save(stream)
    book.close()
    sha = replace_resource(manifest, save, "inputs/edb.xlsx", stream.getvalue())
    source["excel_exports"]["edb"]["sha256"] = proof["original_edb_sha256"] = sha
    install(tmp_path, source, manifest, save, proof)
    with pytest.raises(WorkflowError) as failure:
        replay(tmp_path, source, manifest, contract, snapshot)
    assert str(failure.value.__cause__) == reason


@pytest.mark.parametrize("value", [None, "", 0])
def test_gold_replay_skips_only_legacy_empty_unclosed_values(tmp_path, value):
    contract, source, manifest, save, proof, snapshot = gold_package(tmp_path)
    book = load_workbook(tmp_path / "inputs/edb.xlsx")
    for row in book.active:
        if row[0].value == datetime(2026, 10, 23):
            row[1].value = value
    stream = BytesIO()
    book.save(stream)
    book.close()
    sha = replace_resource(manifest, save, "inputs/edb.xlsx", stream.getvalue())
    source["excel_exports"]["edb"]["sha256"] = proof["original_edb_sha256"] = sha
    install(tmp_path, source, manifest, save, proof)
    rows, _ = replay(tmp_path, source, manifest, contract, snapshot)
    assert all(day != "2026-10-23" for day, _ in rows)
    assert ("2017-01-30", 100) in rows


def test_gold_replay_keeps_legacy_period_end_filter(tmp_path):
    contract, source, manifest, save, proof, snapshot = gold_package(tmp_path)
    book = load_workbook(tmp_path / "inputs/edb.xlsx")
    footer_row = book.active.max_row
    book.active.insert_rows(footer_row)
    book.active.cell(footer_row, 1, datetime(2026, 10, 24))
    book.active.cell(footer_row, 2, 999999)
    stream = BytesIO()
    book.save(stream)
    book.close()
    sha = replace_resource(manifest, save, "inputs/edb.xlsx", stream.getvalue())
    source["excel_exports"]["edb"]["sha256"] = proof["original_edb_sha256"] = sha
    install(tmp_path, source, manifest, save, proof)
    rows, _ = replay(tmp_path, source, manifest, contract, snapshot)
    assert all(day != "2026-10-24" for day, _ in rows)


def test_gold_replay_rejects_duplicate_unclosed_empty_day(tmp_path):
    contract, source, manifest, save, proof, snapshot = gold_package(tmp_path)
    book = load_workbook(tmp_path / "inputs/edb.xlsx")
    footer_row = book.active.max_row
    book.active.insert_rows(footer_row)
    book.active.cell(footer_row, 1, datetime(2026, 10, 23))
    book.active.cell(footer_row, 2, 0)
    stream = BytesIO()
    book.save(stream)
    book.close()
    sha = replace_resource(manifest, save, "inputs/edb.xlsx", stream.getvalue())
    source["excel_exports"]["edb"]["sha256"] = proof["original_edb_sha256"] = sha
    install(tmp_path, source, manifest, save, proof)
    with pytest.raises(WorkflowError) as failure:
        replay(tmp_path, source, manifest, contract, snapshot)
    assert str(failure.value.__cause__) == "gold_edb_duplicate"


@pytest.mark.parametrize("original,duplicate", [(100, 0), (0, 100), (100, None), (None, 100)])
def test_gold_replay_rejects_duplicate_completed_day_even_with_empty_cell(
    tmp_path, original, duplicate
):
    contract, source, manifest, save, proof, snapshot = gold_package(tmp_path)
    book = load_workbook(tmp_path / "inputs/edb.xlsx")
    for row in book.active:
        if row[0].value == datetime(2017, 1, 31):
            row[1].value = original
    footer_row = book.active.max_row
    book.active.insert_rows(footer_row)
    book.active.cell(footer_row, 1, datetime(2017, 1, 31))
    book.active.cell(footer_row, 2, duplicate)
    stream = BytesIO()
    book.save(stream)
    book.close()
    sha = replace_resource(manifest, save, "inputs/edb.xlsx", stream.getvalue())
    source["excel_exports"]["edb"]["sha256"] = proof["original_edb_sha256"] = sha
    install(tmp_path, source, manifest, save, proof)
    with pytest.raises(WorkflowError) as failure:
        replay(tmp_path, source, manifest, contract, snapshot)
    assert str(failure.value.__cause__) == "gold_edb_duplicate"


@pytest.mark.parametrize(
    "attack",
    [
        "edb_hash",
        "unit",
        "asset",
        "formula",
        "bool",
        "infinite",
        "trim",
        "overlap",
        "second_missing",
        "zero",
        "replace",
        "future_period",
        "calendar_trim",
        "pass_flag",
    ],
)
def test_raw_replay_rejects_unbound_or_incomplete_inputs(tmp_path, attack):
    contract, source, manifest, save, proof, snapshot = gold_package(tmp_path)
    if attack == "edb_hash":
        proof["original_edb_sha256"] = "0" * 64
    elif attack == "unit":
        proof["unit"] = "CNY/g"
    elif attack in {"asset", "formula"}:
        proof["recheck"]["formula"] = proof["recheck"]["formula"].replace(
            "SPTAUUSDOZ.IDC" if attack == "asset" else "close", "RT_PRE_CLOSE"
        )
    elif attack in {"bool", "infinite"}:
        proof["recheck"]["rows"][0][1] = True if attack == "bool" else float("inf")
    elif attack == "trim":
        proof["history_requests"][0]["raw_rows"].pop()
    elif attack == "overlap":
        proof["history_requests"][0]["raw_rows"][0][1] = 101
    elif attack in {"second_missing", "zero", "replace"}:
        path = tmp_path / "inputs/edb.xlsx"
        book = load_workbook(path)
        for row in book.active:
            if row[0].value == datetime(
                2017, 1, 31 if attack in {"second_missing", "zero"} else 30
            ):
                row[1].value = (
                    None if attack == "second_missing" else 0 if attack == "zero" else 100
                )
        stream = BytesIO()
        book.save(stream)
        book.close()
        sha = replace_resource(manifest, save, "inputs/edb.xlsx", stream.getvalue())
        source["excel_exports"]["edb"]["sha256"] = proof["original_edb_sha256"] = sha
    elif attack == "future_period":
        contract = dict(contract, period_end="2026-10-30")
    elif attack == "calendar_trim":
        proof["calendar_requests"][0]["raw_rows"].pop()
    else:
        proof["passed"] = True
    install(tmp_path, source, manifest, save, proof)
    with pytest.raises(WorkflowError, match="黄金历史补查"):
        replay(tmp_path, source, manifest, contract, snapshot)


def test_host_assessment_replays_resource_even_after_lineage_removed(tmp_path):
    contract, source, manifest, _, _, snapshot = gold_package(tmp_path)
    rows, lineage = replay(tmp_path, source, manifest, contract, snapshot)
    series = snapshot["series"]["history_gold_london"]
    series.update(
        rows=[{"date": day, "value": value, "published_at": None} for day, value in rows],
        raw_sha256=lineage["original_edb_sha256"],
        history_supplement=lineage,
    )
    series["raw_payload"] = {"export": source["excel_exports"]["edb"], "column": 1}
    assert (
        assess_weekly_snapshot(snapshot, contract, run_root=tmp_path, manifest=manifest)["status"]
        == "passed"
    )
    for attack in ("value", "target_value", "lineage_delete", "lineage_edit", "resource"):
        changed = copy.deepcopy(snapshot)
        if attack == "value":
            changed["series"]["history_gold_london"]["rows"][0]["value"] = 101
        elif attack == "target_value":
            next(
                row
                for row in changed["series"]["history_gold_london"]["rows"]
                if row["date"] == "2017-01-30"
            )["value"] = 101
        elif attack == "lineage_delete":
            del changed["series"]["history_gold_london"]["history_supplement"]
        elif attack == "lineage_edit":
            changed["series"]["history_gold_london"]["history_supplement"]["value"] = "101"
        else:
            (tmp_path / source["history_supplement"]["path"]).write_text("{}")
        result = assess_weekly_snapshot(changed, contract, run_root=tmp_path, manifest=manifest)
        assert result["status"] == "blocked_data"
        assert "weekly_gold_supplement_unverified" in result["issues"]
    assert assess_weekly_snapshot(snapshot, contract)["status"] == "blocked_data"
