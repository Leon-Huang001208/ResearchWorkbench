"""Trusted template projection preserves formatting and binds actual chart data."""

import hashlib
from io import BytesIO
from zipfile import ZipFile

import pytest
from docx import Document
from docx.opc.constants import RELATIONSHIP_TYPE as RT
from docx.opc.packuri import PackURI
from docx.opc.part import Part
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from openpyxl import Workbook, load_workbook

from app.research_web.report_documents import (
    apply_weekly_docx_projection,
    build_docx_document,
)


def _payload():
    return {"title": "周报", "as_of": "2026-10-23", "summary": "", "sections": [], "missing": []}


def test_split_tokens_keep_runs_table_header_footer_and_line_breaks(tmp_path):
    template = tmp_path / "template.docx"
    document = Document()
    paragraph = document.add_paragraph()
    paragraph.add_run("前缀 ").italic = True
    paragraph.add_run("{{市场").bold = True
    paragraph.add_run("回顾}} 后缀").italic = True
    document.sections[0].header.paragraphs[0].add_run("{{页眉}}").bold = True
    document.sections[0].footer.paragraphs[0].add_run("{{页脚}}").italic = True
    document.add_table(rows=1, cols=1).cell(0, 0).text = "{{表格}}"
    document.save(template)
    projected = build_docx_document(
        _payload(),
        {"市场回顾": "本期\n更新", "页眉": "品牌", "页脚": "日期", "表格": "数值"},
        template,
    )
    result = projected.paragraphs[0]
    assert result.text == "前缀 本期\n更新 后缀"
    assert result.runs[0].italic is True
    assert result.runs[1].bold is True
    assert result.runs[2].italic is True
    assert projected.sections[0].header.paragraphs[0].text == "品牌"
    assert projected.sections[0].header.paragraphs[0].runs[0].bold is True
    assert projected.sections[0].footer.paragraphs[0].text == "日期"
    assert projected.tables[0].cell(0, 0).text == "数值"


def _chart(document, number, sheet):
    xml = f"""<c:chartSpace xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships"><c:chart><c:plotArea><c:lineChart><c:ser><c:cat><c:strRef><c:f>'{sheet}'!$A$2:$A$3</c:f><c:strCache><c:ptCount val="2"/><c:pt idx="0"><c:v>OLD</c:v></c:pt></c:strCache></c:strRef></c:cat><c:val><c:numRef><c:f>'{sheet}'!$B$2:$B$3</c:f><c:numCache><c:formatCode>General</c:formatCode><c:ptCount val="2"/><c:pt idx="0"><c:v>999</c:v></c:pt></c:numCache></c:numRef></c:val></c:ser></c:lineChart></c:plotArea></c:chart><c:externalData r:id="rIdOld"><c:autoUpdate val="1"/></c:externalData></c:chartSpace>""".encode()
    part = Part(
        PackURI(f"/word/charts/chart{number}.xml"),
        "application/vnd.openxmlformats-officedocument.drawingml.chart+xml",
        xml,
        document.part.package,
    )
    part.load_rel(RT.OLE_OBJECT, "file:///old/source.xlsx", "rIdOld", is_external=True)
    relation = document.part.relate_to(part, RT.CHART)
    node = OxmlElement("c:chart")
    node.set(qn("r:id"), relation)
    document.add_paragraph()._p.append(node)


@pytest.fixture
def projection(tmp_path):
    root = tmp_path / "inputs" / "report-workflow"
    (root / "workbooks").mkdir(parents=True)
    (root / "snapshots").mkdir()
    snapshot = b'{"schema_version":1,"cutoff":"2026-10-23T15:30:00+08:00"}'
    (root / "snapshots" / "report-data.json").write_bytes(snapshot)
    snapshot_sha = hashlib.sha256(snapshot).hexdigest()
    book = Workbook()
    book.active.title = "行业"
    for name in ("黄金", "原油"):
        book.create_sheet(name)
    for sheet in book:
        sheet.append(["日期", "值"])
        sheet.append(["2026-10-22", 10])
        sheet.append(["2026-10-23", 20])
    calendar = book.create_sheet("经济数据")
    calendar.append(["日期", "国家/地区", "指标名称", "重要性"])
    calendar.append(["2026-10-26", "中国", "经济指标", "重要"])
    overseas = book.create_sheet("海外")
    overseas.append(["指数名称", "指数代码", "周内涨跌幅"])
    overseas.append(["恒生指数", "HSI", 0.025])
    path = root / "workbooks" / "fresh.xlsx"
    book.save(path)
    metadata = {
        "snapshot_sha256": snapshot_sha,
        "workbooks": [
            {
                "path": "workbooks/fresh.xlsx",
                "sha256": hashlib.sha256(path.read_bytes()).hexdigest(),
            }
        ],
        "charts": {
            key: {
                "workbook": "workbooks/fresh.xlsx",
                "sheet": sheet,
                "category_range": "A2:A3",
                "value_ranges": ["B2:B3"],
            }
            for key, sheet in (
                ("industry_weekly_performance", "行业"),
                ("gold_price", "黄金"),
                ("oil_price", "原油"),
            )
        },
        "economic_calendar": {
            "workbook": "workbooks/fresh.xlsx",
            "sheet": "经济数据",
            "range": "A1:D2",
        },
        "overseas": {
            "workbook": "workbooks/fresh.xlsx",
            "sheet": "海外",
            "range": "A1:C2",
            "columns": ["name", "code", "return_fraction"],
        },
    }
    document = Document()
    document.add_paragraph("本期正文")
    for i, sheet in enumerate(("行业", "黄金", "原油"), 1):
        _chart(document, i, sheet)
    table = document.add_table(rows=2, cols=3)
    for cell, value in zip(table.rows[0].cells, ["日期", "国家/地区", "指标名称"]):
        cell.text = value
    for cell in table.rows[1].cells:
        cell.text = "OLD"
    table = document.add_table(rows=2, cols=3)
    for cell, value in zip(table.rows[0].cells, ["指数名称", "指数代码", "周内涨跌幅"]):
        cell.text = value
    context = {
        "dataset_snapshot_sha256": snapshot_sha,
        "dataset_snapshot": "snapshots/report-data.json",
        "weekly_production": {"policy_id": "huaan-etf-weekly-v1"},
        "workbook_projection": metadata,
        "blocks": [
            {"id": key, "title": title, "kind": kind, "required": True}
            for key, title, kind in (
                ("industry", "申万一级行业周涨跌幅", "chart"),
                ("gold", "黄金价格走势", "chart"),
                ("oil", "Brent vs WTI 原油期货结算价", "chart"),
                ("calendar", "下周全球投资日历", "table"),
            )
        ],
    }
    return document, context, root


def test_native_chart_caches_and_embedded_books_use_fresh_values(projection, tmp_path):
    document, context, root = projection
    proof = apply_weekly_docx_projection(document, context, root)
    assert {row["id"] for row in proof} == {"industry", "gold", "oil", "calendar"}
    assert all(row["snapshot_sha256"] == context["dataset_snapshot_sha256"] for row in proof)
    output = tmp_path / "report.docx"
    document.save(output)
    with ZipFile(output) as archive:
        for i in (1, 2, 3):
            chart = archive.read(f"word/charts/chart{i}.xml").decode()
            assert ">999<" not in chart and ">20<" in chart
            relationships = archive.read(f"word/charts/_rels/chart{i}.xml.rels").decode()
            assert "file:" not in relationships and 'TargetMode="External"' not in relationships
        assert (
            len([name for name in archive.namelist() if name.startswith("word/embeddings/")]) == 3
        )
        for name in archive.namelist():
            if name.startswith("word/embeddings/"):
                embedded = load_workbook(BytesIO(archive.read(name)), data_only=True)
                assert embedded.active["B3"].value == 20
                assert embedded.active["B3"].data_type != "f"
    assert document.tables[0].cell(1, 0).text == "2026-10-26"
    assert document.tables[1].cell(1, 2).text == "2.50%"


def test_changed_workbook_or_snapshot_is_never_bound(projection):
    document, context, root = projection
    (root / "workbooks" / "fresh.xlsx").write_bytes(b"changed")
    with pytest.raises(ValueError, match="weekly_"):
        apply_weekly_docx_projection(document, context, root)


def test_cross_market_holiday_gap_is_blank_in_cache_and_embedded_workbook(projection):
    from lxml import etree

    document, context, root = projection
    path = root / "workbooks/fresh.xlsx"
    book = load_workbook(path)
    book["黄金"]["B3"] = None
    book.save(path)
    context["workbook_projection"]["workbooks"][0]["sha256"] = hashlib.sha256(
        path.read_bytes()
    ).hexdigest()
    apply_weekly_docx_projection(document, context, root)
    stream = BytesIO()
    document.save(stream)
    with ZipFile(stream) as archive:
        xml = etree.fromstring(archive.read("word/charts/chart2.xml"))
        ns = {"c": "http://schemas.openxmlformats.org/drawingml/2006/chart"}
        assert xml.find(".//c:numCache/c:ptCount", ns).get("val") == "2"
        assert len(xml.findall(".//c:numCache/c:pt", ns)) == 1
        assert xml.find(".//c:dispBlanksAs", ns).get("val") == "gap"
        embedded = load_workbook(
            BytesIO(archive.read("word/embeddings/weekly-chart2.xlsx")), data_only=True
        )
        assert embedded.active["B3"].value is None


def test_overseas_real_workbook_code_name_order_is_mapped_explicitly(projection):
    document, context, root = projection
    path = root / "workbooks/fresh.xlsx"
    book = load_workbook(path)
    book["海外"]["A2"] = "HSI"
    book["海外"]["B2"] = "恒生指数"
    book.save(path)
    context["workbook_projection"]["workbooks"][0]["sha256"] = hashlib.sha256(
        path.read_bytes()
    ).hexdigest()
    context["workbook_projection"]["overseas"].update(
        range="A2:C2", columns=["code", "name", "return_fraction"]
    )
    apply_weekly_docx_projection(document, context, root)
    assert document.tables[1].cell(1, 0).text == "恒生指数"
    assert document.tables[1].cell(1, 1).text == "HSI"


def test_missing_chart_mapping_is_never_bound(projection):
    document, context, root = projection
    context["workbook_projection"]["charts"].pop("gold_price")
    with pytest.raises(ValueError, match="weekly_"):
        apply_weekly_docx_projection(document, context, root)


def test_unmapped_required_block_is_never_marked_projected(projection):
    document, context, root = projection
    context["blocks"].append(
        {"id": "unmapped", "title": "未映射图表", "kind": "chart", "required": True}
    )
    with pytest.raises(ValueError, match="weekly_projection_blocks_incomplete"):
        apply_weekly_docx_projection(document, context, root)


def test_unsafe_template_archive_is_rejected_before_loading(tmp_path):
    template = tmp_path / "unsafe.docx"
    with ZipFile(template, "w") as archive:
        archive.writestr("../outside.xml", "<xml/>")
    with pytest.raises(ValueError, match="weekly_unsafe_office_archive"):
        build_docx_document(_payload(), {}, template)


def test_run_split_unresolved_placeholder_is_not_certified(projection):
    document, context, root = projection
    paragraph = document.add_paragraph()
    paragraph.add_run("{{遗漏")
    paragraph.add_run("字段}}")
    with pytest.raises(ValueError, match="weekly_template_placeholder_unresolved"):
        apply_weekly_docx_projection(document, context, root)
