"""Huaan production binds canonical content, trusted inputs and final files."""

import asyncio
import copy
import hashlib
import json
import sys
from datetime import datetime
from types import SimpleNamespace

import pytest
from pydantic import ValidationError
from test_report_weekly_admission import freeze_scope, weekly
from test_report_workflow_integration import isolated_keyring

from app.research_web.report_workflows.models import ReportWorkflowManifest, WorkflowError
from app.research_web.report_workflows.weekly_quality import (
    POLICY_ID,
    bind_weekly_projection,
    canonical_projection,
    freeze_weekly_production,
    stage_weekly_projection,
    trusted_weekly_context,
    verify_weekly_rendered_artifact,
)

__all__ = ["isolated_keyring", "weekly"]


def _service(tmp_path):
    root = tmp_path / "session"
    (root / "inputs" / "report-workflow" / "snapshots").mkdir(parents=True)
    (root / "inputs" / "report-workflow" / "templates").mkdir()
    (root / "outputs").mkdir()
    template = root / "inputs" / "report-workflow" / "templates" / "original.docx"
    template.write_bytes(b"original template")
    snapshot = json.dumps({"source": "test", "rows": []}).encode()
    (root / "inputs" / "report-workflow" / "snapshots" / "report-data.json").write_bytes(snapshot)
    frozen = {
        "policy_id": POLICY_ID,
        "workflow_id": "huaan-etf-weekly",
        "version": 1,
        "as_of": "2026-10-16",
        "requirements": [],
        "data_contract": {},
    }
    context = {
        "blocks": [{"id": "market", "title": "市场", "kind": "narrative", "required": True}],
        "weekly_production": frozen,
        "template_resources": [
            {"path": "templates/original.docx", "sha256": _digest(template.read_bytes())}
        ],
    }
    context_raw = json.dumps(context).encode()
    (root / "inputs" / "report-workflow" / "run-context.json").write_bytes(context_raw)
    run = {
        "id": "r",
        "workflow_id": "huaan-etf-weekly",
        "version": 1,
        "session_id": "s",
        "weekly_production": frozen,
        "weekly_context_sha256": _digest(context_raw),
        "dataset_snapshot_sha256": _digest(snapshot),
    }
    marker = {
        "quality_policy": POLICY_ID,
        "workflow_id": "huaan-etf-weekly",
        "version": 1,
        "run_id": "r",
    }
    session = {"delivery_key": "key", "report_workflow": marker}
    receipt = {"research_contract": {"task_id": "a" * 24, "sha256": "c" * 64}}
    store = SimpleNamespace(
        directory=lambda _sid: root,
        session=lambda _sid: session,
        data={"receipts": {"s:key": receipt}},
        save=lambda: None,
    )
    runtime = SimpleNamespace(_run=lambda _rid: copy.deepcopy(run))
    return (
        SimpleNamespace(store=store, report_workflows=SimpleNamespace(runtime=runtime)),
        run,
        context,
    )


def _digest(raw):
    return hashlib.sha256(raw).hexdigest()


def _manifest(workflow_id="huaan-etf-weekly"):
    return {
        "workflow_id": workflow_id,
        "name": "华安ETF周报",
        "version": 1,
        "blocks": [{"id": "market", "title": "市场", "kind": "narrative"}],
        "delivery": {"formats": ["docx", "html"]},
        "weekly_production": {"policy_id": POLICY_ID, "news_block_ids": ["market"]},
    }


def _canonical():
    return {
        "task_id": "a" * 24,
        "title": "华安ETF周报",
        "as_of": "2026-10-16",
        "sections": [
            {"id": "market", "heading": "市场", "text": "市场资料。", "citations": ["file:abc"]}
        ],
        "sources": [{"id": "file:abc", "title": "行情"}],
    }


def test_weekly_policy_is_explicit_and_scoped_to_huaan():
    assert (
        ReportWorkflowManifest.model_validate(_manifest()).weekly_production.policy_id == POLICY_ID
    )
    with pytest.raises(ValidationError):
        ReportWorkflowManifest.model_validate(_manifest("other-weekly"))
    value = _manifest()
    value["weekly_production"]["news_block_ids"] = ["unknown"]
    with pytest.raises(ValidationError):
        ReportWorkflowManifest.model_validate(value)
    legacy = _manifest()
    legacy.pop("weekly_production")
    assert ReportWorkflowManifest.model_validate(legacy).weekly_production is None


def test_canonical_projection_preserves_body_and_rejects_wrong_task_or_missing_block():
    value = _canonical()
    context = {"blocks": [{"id": "market", "title": "市场", "required": True}]}
    binding = {"task_id": "a" * 24, "canonical_sha256": "b" * 64}
    projected = canonical_projection(value, context, binding)
    assert projected["sections"] == [{"id": "market", "title": "市场", "content": "市场资料。"}]
    assert projected["binding"] == binding
    assert value == _canonical()
    with pytest.raises(WorkflowError, match="规范报告"):
        canonical_projection({**value, "task_id": "c" * 24}, context, binding)
    with pytest.raises(WorkflowError, match="区块"):
        canonical_projection(value, {"blocks": [{"id": "absent", "required": True}]}, binding)


def test_projection_uses_canonical_chinese_presentation_and_preserves_audit_body():
    canonical = _canonical()
    canonical["presentation"] = [
        {
            "id": "market",
            "heading": "市场",
            "text": "用于交付的中文内容。",
            "citations": ["file:abc"],
        }
    ]
    projected = canonical_projection(
        canonical, {"blocks": [{"id": "market", "required": True}]}, {"task_id": "a" * 24}
    )
    assert projected["sections"][0]["content"] == "用于交付的中文内容。"
    assert canonical["sections"][0]["text"] == "市场资料。"
    canonical["presentation"][0]["id"] = "wrong"
    with pytest.raises(WorkflowError):
        canonical_projection(
            canonical, {"blocks": [{"id": "market", "required": True}]}, {"task_id": "a" * 24}
        )


def test_projection_restores_only_the_host_frozen_section_title():
    canonical = _canonical()
    context = {"blocks": [{"id": "market", "title": "申万一级行业周涨跌幅", "required": True}]}
    projected = canonical_projection(canonical, context, {"task_id": "a" * 24})
    assert projected["sections"][0]["title"] == "申万一级行业周涨跌幅"
    assert canonical["sections"][0]["heading"] == "市场"


def test_only_frozen_date_metadata_is_inserted_without_model_prose():
    value = _manifest()
    value["blocks"].extend(
        [
            {"id": "start", "title": "开始日期", "kind": "workbook"},
            {"id": "end", "title": "结束日期", "kind": "workbook"},
        ]
    )
    manifest = ReportWorkflowManifest.model_validate(value)
    frozen = freeze_weekly_production(
        manifest,
        {
            "cutoff": "2026-10-23T15:30:00+08:00",
            "period_start": "2026-10-19",
            "period_end": "2026-10-23",
        },
    )
    assert frozen["requirements"] == [{"id": "market", "text": "市场"}]
    context = {"blocks": value["blocks"], "weekly_production": frozen}
    projected = canonical_projection(_canonical(), context, {"task_id": "a" * 24})
    assert projected["sections"][-2:] == [
        {"id": "start", "title": "开始日期", "content": "2026-10-19"},
        {"id": "end", "title": "结束日期", "content": "2026-10-23"},
    ]
    value["blocks"].append({"id": "other", "title": "其他底稿", "kind": "workbook"})
    with pytest.raises(WorkflowError):
        canonical_projection(
            _canonical(),
            {"blocks": value["blocks"], "weekly_production": frozen},
            {"task_id": "a" * 24},
        )


@pytest.mark.parametrize("formal", [False, True])
def test_refresh_uses_bound_frozen_run_only_for_explicit_weekly_policy(tmp_path, formal):
    from io import BytesIO

    from openpyxl import Workbook, load_workbook

    from app.research_web.report_workflows.catalog import ReportWorkflowService

    value = _manifest()
    if not formal:
        value.pop("weekly_production")
    value["workbook_policies"] = [
        {"workbook": "workbooks/source.xlsx", "required_cells": ["Sheet1!A1"]}
    ]
    catalog = ReportWorkflowService(tmp_path)
    catalog.create_draft(ReportWorkflowManifest.model_validate(value))
    book = Workbook()
    book.active.title = "Sheet1"
    book.active["A1"] = "mother"
    stream = BytesIO()
    book.save(stream)
    catalog.upload_resource("huaan-etf-weekly", "workbooks/source.xlsx", stream.getvalue())
    catalog.create_version("huaan-etf-weekly")
    catalog.publish_version("huaan-etf-weekly", 1)
    workspace = catalog.create_run_workspace("huaan-etf-weekly", 1)
    rid = workspace["run_id"]
    candidate = catalog.run_path(rid) / "workbooks" / "source.xlsx"
    book.active["A1"] = "frozen"
    book.save(candidate)
    with catalog._exclusive():
        catalog.data["runs"][rid].update(
            weekly_excel_sources={"workbooks/source.xlsx": _digest(candidate.read_bytes())},
            weekly_excel_cutoff="cutoff",
            weekly_data_contract={"cutoff": "cutoff"},
        )
        catalog._save()
    captured = {}

    class CaptureRefresh:
        def refresh(self, source, _root, _policy, **_kwargs):
            captured["source"] = source
            captured["value"] = load_workbook(source, data_only=True).active["A1"].value
            return SimpleNamespace(status=SimpleNamespace(value="ready"))

    refreshed = catalog.refresh_workbook(
        rid, "workbooks/source.xlsx", refresh_service=CaptureRefresh()
    )
    assert refreshed.status.value == "ready"
    assert captured["value"] == ("frozen" if formal else "mother")
    mother = catalog.resource_path("huaan-etf-weekly", 1, "workbooks/source.xlsx")
    assert captured["source"] == (candidate if formal else mother)
    assert load_workbook(mother, data_only=True).active["A1"].value == "mother"


def test_refresh_rejects_changed_run_source_before_provider_call(tmp_path):
    from io import BytesIO

    from openpyxl import Workbook

    from app.research_web.report_workflows.catalog import ReportWorkflowService

    value = _manifest()
    value["workbook_policies"] = [{"workbook": "workbooks/source.xlsx"}]
    catalog = ReportWorkflowService(tmp_path)
    catalog.create_draft(ReportWorkflowManifest.model_validate(value))
    stream = BytesIO()
    Workbook().save(stream)
    catalog.upload_resource("huaan-etf-weekly", "workbooks/source.xlsx", stream.getvalue())
    catalog.create_version("huaan-etf-weekly")
    catalog.publish_version("huaan-etf-weekly", 1)
    rid = catalog.create_run_workspace("huaan-etf-weekly", 1)["run_id"]
    with catalog._exclusive():
        catalog.data["runs"][rid].update(
            weekly_excel_sources={"workbooks/source.xlsx": "f" * 64},
            weekly_excel_cutoff="cutoff",
            weekly_data_contract={"cutoff": "cutoff"},
        )
        catalog._save()
    with pytest.raises(WorkflowError) as error:
        catalog.refresh_workbook(rid, "workbooks/source.xlsx")
    assert error.value.code == "weekly_excel_source_changed"


def test_trusted_context_does_not_accept_browser_or_model_metadata(tmp_path):
    row = {
        "report_workflow": {
            "quality_policy": POLICY_ID,
            "workflow_id": "huaan-etf-weekly",
            "version": 1,
            "run_id": "r",
        }
    }
    run = {
        "id": "r",
        "workflow_id": "huaan-etf-weekly",
        "version": 1,
        "session_id": "s",
        "weekly_production": {"policy_id": POLICY_ID},
    }
    store = SimpleNamespace(session=lambda _sid: row, directory=lambda _sid: tmp_path)
    runtime = SimpleNamespace(_run=lambda _rid: copy.deepcopy(run))
    service = SimpleNamespace(store=store, report_workflows=SimpleNamespace(runtime=runtime))
    with pytest.raises(WorkflowError):
        trusted_weekly_context(service, "s")
    row["report_workflow"]["workflow_id"] = "other-weekly"
    with pytest.raises(WorkflowError):
        trusted_weekly_context(service, "s")
    row["report_workflow"].pop("quality_policy")
    assert trusted_weekly_context(service, "s") is None


def test_model_authored_projection_receipt_never_authenticates_files(tmp_path):
    root = tmp_path / "s"
    (root / "outputs").mkdir(parents=True)
    raw = json.dumps(_canonical()).encode()
    (root / "outputs" / "research-report.json").write_bytes(raw)
    receipt = {"research_contract": {"task_id": "a" * 24, "sha256": "c" * 64}}
    store = SimpleNamespace(
        directory=lambda _sid: root,
        session=lambda _sid: {
            "delivery_key": "key",
            "report_workflow": {"quality_policy": POLICY_ID},
        },
        data={"receipts": {"s:key": receipt}},
    )
    with pytest.raises(ValueError, match="weekly"):
        verify_weekly_rendered_artifact(store, "s", _canonical(), "html", b"fake")
    # Adding the receipt to model output cannot supply the private Host record.
    candidate = {**_canonical(), "weekly_projection": {"canonical_sha256": _digest(raw)}}
    with pytest.raises(ValueError, match="weekly"):
        verify_weekly_rendered_artifact(store, "s", candidate, "html", b"fake")


@pytest.mark.parametrize(
    "changed", ["canonical", "snapshot", "context", "template", "artifact", "task"]
)
def test_host_projection_is_bound_to_every_current_input_and_output(tmp_path, changed):
    service, _run, _context = _service(tmp_path)
    root = service.store.directory("s")
    (root / "outputs" / "research-report.json").write_text(json.dumps(_canonical()))
    body = "<html>市场资料。</html>".encode()
    (root / "outputs" / "report.html").write_bytes(body)
    staged = stage_weekly_projection(service, "s")
    bind_weekly_projection(service, "s", staged, {"files": ["outputs/report.html"]})
    canonical = json.loads((root / "outputs" / "research-report.json").read_text())
    verify_weekly_rendered_artifact(service.store, "s", canonical, "html", body)
    paths = {
        "canonical": "outputs/research-report.json",
        "snapshot": "inputs/report-workflow/snapshots/report-data.json",
        "context": "inputs/report-workflow/run-context.json",
        "template": "inputs/report-workflow/templates/original.docx",
    }
    if changed in paths:
        path = root / paths[changed]
        path.write_bytes(path.read_bytes() + b" ")
    elif changed == "artifact":
        body += b"changed"
        (root / "outputs" / "report.html").write_bytes(body)
    else:
        service.store.data["receipts"]["s:key"]["research_contract"]["task_id"] = "d" * 24
    with pytest.raises(ValueError, match="weekly"):
        verify_weekly_rendered_artifact(service.store, "s", canonical, "html", body)


def test_required_chart_cannot_be_accepted_by_reusing_old_template(tmp_path):
    service, run, context = _service(tmp_path)
    root = service.store.directory("s")
    context["blocks"].append({"id": "chart1", "title": "行业图", "kind": "chart", "required": True})
    raw = json.dumps(context).encode()
    (root / "inputs" / "report-workflow" / "run-context.json").write_bytes(raw)
    run["weekly_context_sha256"] = _digest(raw)
    canonical = _canonical()
    canonical["sections"].append(
        {"id": "chart1", "heading": "行业图", "text": "行业图数据。", "citations": ["file:abc"]}
    )
    (root / "outputs" / "research-report.json").write_text(json.dumps(canonical))
    (root / "outputs" / "report.html").write_text("<html>未更新的图</html>")
    staged = stage_weekly_projection(service, "s")
    with pytest.raises(WorkflowError, match="图表"):
        bind_weekly_projection(service, "s", staged, {"files": ["outputs/report.html"]})
    assert "weekly_projection" not in service.store.data["receipts"]["s:key"]
    assert json.loads((root / "outputs" / "research-report.json").read_text()) == canonical


def test_template_cannot_drop_a_required_narrative_even_with_valid_file(tmp_path):
    service, _run, _context = _service(tmp_path)
    root = service.store.directory("s")
    (root / "outputs" / "research-report.json").write_text(json.dumps(_canonical()))
    (root / "outputs" / "report.html").write_text("<html>遗漏规范正文</html>")
    staged = stage_weekly_projection(service, "s")
    with pytest.raises(WorkflowError) as error:
        bind_weekly_projection(service, "s", staged, {"files": ["outputs/report.html"]})
    assert error.value.code == "weekly_narrative_projection_incomplete"


def test_failed_final_save_preserves_previous_final_pointer(tmp_path):
    from contextlib import nullcontext

    from app.research_web.report_workflows.runtime import ReportWorkflowRuntime

    workflow = {"last_finalized_run": "previous"}
    run = {
        "id": "new",
        "workflow_id": "huaan-etf-weekly",
        "status": "validating",
        "owner_id": "owner",
    }

    def failed_save():
        raise OSError("disk full")

    runtime = ReportWorkflowRuntime.__new__(ReportWorkflowRuntime)
    runtime.catalog = SimpleNamespace(
        _exclusive=nullcontext,
        _row=lambda _wid: workflow,
        data={"runs": {"new": run}},
        _save=failed_save,
    )
    runtime._require_weekly_owner = lambda _owner: {"id": "owner"}
    with pytest.raises(OSError, match="disk full"):
        runtime._finish_run(
            "new",
            {
                "status": "completed",
                "finalization_status": "finalized",
                "finalization_binding": {"canonical_sha256": "a" * 64},
            },
        )
    assert workflow["last_finalized_run"] == "previous"
    assert run["status"] == "validating" and "finalization_status" not in run


def _runtime_fixture(
    tmp_path,
    monkeypatch,
    *,
    revoke_during_collection=False,
    unknown_send=False,
    native_accept=False,
    data_source="wind_client_api",
):
    from app.research_web.report_workflows.catalog import ReportWorkflowService
    from app.research_web.report_workflows.runtime import ReportWorkflowRuntime
    from app.research_web.store import Store

    contract = {"cutoff": "2026-10-16T15:30:00+08:00"}

    def materialize(_snapshot, root, _manifest, **_policy):
        return {"snapshot_sha256": _digest((root / "snapshots" / "report-data.json").read_bytes())}

    monkeypatch.setitem(
        sys.modules,
        "app.research_web.report_workflows.weekly_data",
        SimpleNamespace(
            build_weekly_data_contract=lambda _now: copy.deepcopy(contract),
            materialize_weekly_workbooks=materialize,
        ),
    )
    catalog = ReportWorkflowService(tmp_path / "catalog")
    declaration = _manifest()
    declaration["weekly_production"]["data_source"] = data_source
    manifest = ReportWorkflowManifest.model_validate(declaration)
    catalog.create_draft(manifest)
    catalog.create_version(manifest.workflow_id)
    catalog.publish_version(manifest.workflow_id, 1)
    with catalog._exclusive():
        catalog._row(manifest.workflow_id)["status"] = "enabled"
        catalog._save()
    owner = {"enabled": True}
    sent = []
    store = Store(tmp_path / "state")

    def authorize(expected_owner_id=None):
        if not owner["enabled"] or expected_owner_id not in {None, "owner"}:
            raise WorkflowError("管理员已停用", "weekly_owner_denied", 403)
        return {"id": "owner", "role": "admin"}

    async def create(mode, title):
        return store.create(mode, title)

    async def collect(_sid, _call_id, *, cutoff, run_root, manifest, prepare_excel):
        assert cutoff == contract["cutoff"]
        assert run_root.is_dir() and manifest.workflow_id == "huaan-etf-weekly"
        assert callable(prepare_excel)
        marker = store.session(_sid)["report_workflow"]
        assert (
            marker["workflow_id"] == manifest.workflow_id and marker["version"] == manifest.version
        )
        assert marker["run_id"] == run_root.name and "quality_policy" not in marker
        if revoke_during_collection:
            owner["enabled"] = False
        return {
            "contract": copy.deepcopy(contract),
            "snapshot": {"source": "test"},
            "assessment": {
                "status": (
                    "passed"
                    if revoke_during_collection or unknown_send or native_accept
                    else "blocked_data"
                ),
                "issues": ["calendar_unverified"],
            },
        }

    async def send(*args, **kwargs):
        sid, _text, key = args
        sent.append((args, kwargs))
        task = hashlib.sha256(f"{sid}:{key}".encode()).hexdigest()[:24]
        store.reserve(
            sid,
            key,
            "synthetic",
            {"task_id": task, "marker": f"[RESEARCH_TASK:{task}]", "required_formats": []},
        )
        if native_accept:
            store.receipt(sid, key, "accepted")
        if unknown_send:
            raise ConnectionError("native response lost")

    async def detail(_sid):
        return {"status": "idle", "delivery": {"status": "admission_unknown"}}

    service = SimpleNamespace(
        store=store,
        create=create,
        send=send,
        detail=detail,
        datahub=SimpleNamespace(collect_report_weekly=collect),
        events={},
    )
    service.report_workflows = SimpleNamespace(require_weekly_owner=authorize)
    runtime = ReportWorkflowRuntime(service, catalog, None)
    service.report_workflows.runtime = runtime

    return runtime, owner, sent


def test_unverified_data_never_sends_model_and_period_is_not_started_twice(tmp_path, monkeypatch):
    runtime, _owner, sent = _runtime_fixture(tmp_path, monkeypatch)

    async def run():
        first = await runtime.start_run("huaan-etf-weekly")
        await runtime.tasks[first["id"]]
        second = await runtime.start_run("huaan-etf-weekly")
        assert first["id"] == second["id"]
        assert runtime._run(first["id"])["status"] == "blocked_data"
        assert len(runtime.catalog.data["runs"]) == 1

    asyncio.run(run())
    assert sent == []


def test_owner_revocation_during_collection_blocks_model_admission(tmp_path, monkeypatch):
    runtime, _owner, sent = _runtime_fixture(tmp_path, monkeypatch, revoke_during_collection=True)

    async def run():
        created = await runtime.start_run("huaan-etf-weekly")
        await runtime.tasks[created["id"]]
        row = runtime._run(created["id"])
        assert row["failure_code"] == "weekly_owner_denied"
        assert row["finalization_status"] == "blocked"

    asyncio.run(run())
    assert sent == []


def test_unknown_model_submission_is_preserved_for_attention_without_retry(weekly, monkeypatch):
    from app.research_web.client import RuntimeFailure

    service, native, run, _manifest = weekly
    runtime = service.report_workflows.runtime
    freeze_scope(weekly)
    original = native.rpc

    async def lost(method, payload):
        result = await original(method, payload)
        if method == "session.prompt":
            raise RuntimeFailure("lost ACK", "transport_lost")
        return result

    monkeypatch.setattr(native, "rpc", lost)

    async def execute():
        with pytest.raises(RuntimeFailure):
            await runtime._submit_weekly(run["id"])
        checked = await runtime._inspect_weekly_run(run["id"])
        assert checked["status"] == "needs_attention"
        assert checked["research_submission_state"] == "unknown"

    asyncio.run(execute())
    assert sum(name == "session.prompt" for name, _ in native.calls) == 1


def test_legacy_success_cannot_enable_formal_automation(tmp_path, monkeypatch):
    runtime, _owner, sent = _runtime_fixture(tmp_path, monkeypatch)
    runtime._new_run("legacy", "huaan-etf-weekly", 1, "manual", "completed")
    runtime._update_run("legacy", delivery_status="complete")
    with pytest.raises(WorkflowError) as error:
        runtime.put_schedule(
            "huaan-etf-weekly",
            {"kind": "weekly", "enabled": True, "weekday": 4, "hour": 15, "minute": 30},
        )
    assert error.value.code == "manual_run_required"
    with pytest.raises(WorkflowError) as error:
        asyncio.run(runtime.start_run("huaan-etf-weekly", trigger="automation"))
    assert error.value.code == "manual_run_required"
    assert sent == []


def test_excel_production_cannot_enable_schedule_even_after_manual_finalization(
    tmp_path, monkeypatch
):
    runtime, _owner, sent = _runtime_fixture(tmp_path, monkeypatch)
    original = runtime.catalog.manifest
    runtime.catalog.manifest = lambda *args: original(*args).model_copy(
        update={
            "weekly_production": original(*args).weekly_production.model_copy(
                update={"data_source": "wind_excel"}
            )
        }
    )
    runtime._manual_weekly_passed = lambda *_args: True
    with pytest.raises(WorkflowError) as caught:
        runtime.put_schedule(
            "huaan-etf-weekly",
            {"kind": "weekly", "enabled": True, "weekday": 4, "hour": 15, "minute": 30},
        )
    assert caught.value.code == "weekly_excel_schedule_disabled"
    with pytest.raises(WorkflowError) as caught:
        asyncio.run(runtime.start_run("huaan-etf-weekly", trigger="schedule"))
    assert caught.value.code == "weekly_excel_schedule_disabled"
    assert sent == []
    runtime.put_schedule(
        "huaan-etf-weekly",
        {"kind": "weekly", "enabled": False, "weekday": 4, "hour": 15, "minute": 30},
    )
    with runtime.catalog._exclusive():
        runtime.catalog._row("huaan-etf-weekly")["schedule"].update(
            enabled=True, next_run_at="2026-10-16T07:30:00+00:00"
        )
        runtime.catalog._save()
    asyncio.run(runtime.tick(datetime.fromisoformat("2026-10-16T16:00:00+08:00")))
    schedule = runtime.schedule("huaan-etf-weekly")
    assert schedule["enabled"] is False and schedule["next_run_at"] is None
    assert schedule["last_error_code"] == "weekly_excel_schedule_disabled"
    assert sent == [] and runtime.catalog.data["runs"] == {}


def test_failed_admission_journal_never_calls_native_rpc(weekly, monkeypatch):
    from app.research_web.report_workflows import weekly_execution

    service, native, run, _manifest = weekly
    runtime = service.report_workflows.runtime
    freeze_scope(weekly)

    def fail(*_args):
        raise WorkflowError("journal fsync failed", "weekly_attempt_failed", 409)

    monkeypatch.setattr(weekly_execution, "create_attempt", fail)
    with pytest.raises(WorkflowError):
        asyncio.run(runtime._submit_weekly(run["id"]))
    current = runtime._run(run["id"])
    assert current["research_submission_state"] == "preflight_rejected"
    assert "research_admission_at" not in current
    assert not any(name == "session.prompt" for name, _ in native.calls)


def test_startup_recovers_accepted_rpc_without_sending_again(weekly, monkeypatch):
    from app.research_web.report_workflows.runtime import ReportWorkflowRuntime

    service, native, run, _manifest = weekly
    runtime = service.report_workflows.runtime
    freeze_scope(weekly)
    runtime._update_run(run["id"], status="preparing_data")
    original_save = runtime.catalog._save
    failed = []

    def save():
        if not failed and any(
            row.get("research_started_at") for row in runtime.catalog.data["runs"].values()
        ):
            failed.append(True)
            raise OSError("crash after accepted RPC")
        return original_save()

    monkeypatch.setattr(runtime.catalog, "_save", save)
    monitored = []

    async def execute():
        with pytest.raises(OSError):
            await runtime._submit_weekly(run["id"])
        recovered = ReportWorkflowRuntime(service, runtime.catalog, None)
        service.report_workflows.runtime = recovered

        async def monitor(run_id):
            monitored.append(run_id)

        monkeypatch.setattr(recovered, "_resume_delivery", monitor)
        await recovered.start()
        await asyncio.gather(*list(recovered.tasks.values()))
        await recovered.close()

    asyncio.run(execute())
    assert monitored == [run["id"]]
    assert failed == [True]
    assert sum(name == "session.prompt" for name, _ in native.calls) == 1


def test_startup_retains_unknown_rpc_without_model_or_vendor_replay(weekly, monkeypatch):
    from app.research_web.client import RuntimeFailure
    from app.research_web.report_workflows.runtime import ReportWorkflowRuntime

    service, native, run, _manifest = weekly
    runtime = service.report_workflows.runtime
    freeze_scope(weekly)
    original = native.rpc

    async def lost(method, payload):
        result = await original(method, payload)
        if method == "session.prompt":
            raise RuntimeFailure("lost ACK", "transport_lost")
        return result

    monkeypatch.setattr(native, "rpc", lost)

    async def execute():
        with pytest.raises(RuntimeFailure):
            await runtime._submit_weekly(run["id"])
        recovered = ReportWorkflowRuntime(service, runtime.catalog, None)
        service.report_workflows.runtime = recovered
        await recovered.start()
        assert recovered.tasks == {}
        assert recovered._run(run["id"])["status"] == "needs_attention"
        await recovered.close()

    asyncio.run(execute())
    assert sum(name == "session.prompt" for name, _ in native.calls) == 1


def test_pending_data_receipt_is_only_inspected_not_replayed(tmp_path, monkeypatch):
    runtime, _owner, sent = _runtime_fixture(tmp_path, monkeypatch)

    async def run():
        created = await runtime.start_run("huaan-etf-weekly")
        await runtime.tasks[created["id"]]
        runtime._update_run(created["id"], status="needs_attention")
        runtime.service.datahub.snapshots = SimpleNamespace(
            receipt=lambda *_args: {"status": "pending"}
        )
        checked = await runtime.retry(created["id"])
        assert checked["status"] == "needs_attention"
        assert checked["failure_code"] == "weekly_data_admission_unknown"
        assert runtime.tasks == {}

    asyncio.run(run())
    assert sent == []
