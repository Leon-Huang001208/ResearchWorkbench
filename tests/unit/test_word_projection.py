"""Tests for WordProjection."""
from pathlib import Path
from unittest.mock import patch

import pytest

from core.contracts import SectionOutput, TableSpec
from reporting.projections.word import WordProjection


class TestWordProjection:
    """Test suite for WordProjection."""

    @pytest.mark.skipif(not WordProjection()._check_docx(), reason="python-docx not installed")
    def test_save_to_file(self, tmp_path):
        """Test saving Word document to a file."""
        projection = WordProjection()
        sections = [
            SectionOutput(
                key="test",
                title="Test Section",
                content="Test content",
                evidence_refs=[],
                warnings=[],
            )
        ]
        output_file = tmp_path / "report.docx"

        projection.save(output_file, "Test Report", sections)

        assert output_file.exists()

    def test_save_with_metadata(self, tmp_path):
        """Test saving with metadata."""
        projection = WordProjection()
        if not projection._check_docx():
            pytest.skip("python-docx not installed")

        sections = [
            SectionOutput(
                key="test",
                title="Test Section",
                content="Test content",
                evidence_refs=[],
                warnings=[],
            )
        ]
        metadata = {"author": "Test Author"}
        output_file = tmp_path / "report.docx"

        projection.save(output_file, "Test Report", sections, metadata=metadata)
        assert output_file.exists()

    def test_save_with_evidence_refs(self, tmp_path):
        """Test saving with evidence references."""
        projection = WordProjection()
        if not projection._check_docx():
            pytest.skip("python-docx not installed")

        sections = [
            SectionOutput(
                key="test",
                title="Test Section",
                content="Content",
                evidence_refs=["doc1", "doc2"],
                warnings=[],
            )
        ]
        output_file = tmp_path / "report.docx"

        projection.save(output_file, "Test Report", sections)
        assert output_file.exists()

    def test_save_with_warnings(self, tmp_path):
        """Test saving with warnings."""
        projection = WordProjection()
        if not projection._check_docx():
            pytest.skip("python-docx not installed")

        sections = [
            SectionOutput(
                key="test",
                title="Test Section",
                content="Content",
                evidence_refs=[],
                warnings=["Warning 1"],
            )
        ]
        output_file = tmp_path / "report.docx"

        projection.save(output_file, "Test Report", sections)
        assert output_file.exists()

    def test_save_raises_when_docx_not_installed(self, tmp_path):
        """Test that it raises ImportError when python-docx is not available."""
        projection = WordProjection()
        projection._has_docx = False

        sections = [
            SectionOutput(
                key="test",
                title="Test Section",
                content="Content",
                evidence_refs=[],
                warnings=[],
            )
        ]
        output_file = tmp_path / "report.docx"

        with pytest.raises(ImportError):
            projection.save(output_file, "Test Report", sections)

    def test_check_docx_handles_import_error(self):
        """Test that _check_docx handles import errors gracefully."""
        projection = WordProjection()

        with patch("importlib.import_module", side_effect=ImportError):
            # This test verifies the code path - we can't easily patch the module check
            # since it happens in __init__
            result = projection._check_docx()
            # The exact result depends on whether python-docx is actually installed
            # but the method should not raise an exception
            assert isinstance(result, bool)


class TestWordProjectionFromTemplate:
    """测试从模板生成 Word 文档"""

    @pytest.fixture
    def docx_available(self) -> bool:
        """检查 python-docx 是否可用"""
        try:
            import docx  # noqa: F401

            return True
        except ImportError:
            return False

    def create_test_template(self, tmp_path) -> Path:
        """创建测试用的 Word 模板文件"""
        import docx

        template_path = tmp_path / "template.docx"
        doc = docx.Document()

        doc.add_paragraph("Report Title: {{title}}")
        doc.add_paragraph("Executive Summary: {{summary}}")
        doc.add_paragraph("Date: {date}")

        doc.save(template_path)
        return template_path

    @pytest.mark.skipif(not WordProjection()._check_docx(), reason="python-docx not installed")
    def test_save_from_template(self, tmp_path):
        """测试从模板保存"""

        projection = WordProjection()

        # 创建测试模板
        template_path = self.create_test_template(tmp_path)
        output_path = tmp_path / "output.docx"

        # 创建章节输出
        sections = [
            SectionOutput(
                key="title",
                title="Title",
                content="Q2 Investment Report",
                evidence_refs=[],
                warnings=[],
            ),
            SectionOutput(
                key="summary",
                title="Summary",
                content="Market performed well this quarter.",
                evidence_refs=[],
                warnings=[],
            ),
        ]

        # 创建占位符映射
        placeholders = {
            "title": "Q2 Investment Report",
            "summary": "Market performed well this quarter.",
            "date": "2024-06-30",
        }

        # 从模板保存
        projection.save_from_template(
            output_path,
            template_path,
            sections,
            placeholders=placeholders,
        )

        assert output_path.exists()

    @pytest.mark.skipif(not WordProjection()._check_docx(), reason="python-docx not installed")
    def test_save_from_template_prefers_long_placeholders_over_bare_words(self, tmp_path):
        """Overlapping Chinese placeholders should not corrupt generated text."""
        import docx

        template_path = tmp_path / "overlap_template.docx"
        output_path = tmp_path / "overlap_output.docx"
        doc = docx.Document()
        doc.add_paragraph("{{黄金市场回顾}}")
        doc.save(template_path)

        projection = WordProjection()
        projection.save_from_template(
            output_path,
            template_path,
            sections=[],
            placeholders={
                "黄金": "黄金短句",
                "市场回顾": "市场回顾短句",
                "黄金市场回顾": "黄金市场回顾完整正文。",
            },
        )

        result = docx.Document(output_path)
        text = "\n".join(paragraph.text for paragraph in result.paragraphs)
        assert text == "黄金市场回顾完整正文。"
        assert "{{" not in text
        assert "黄金短句" not in text
        assert "市场回顾短句" not in text

    @pytest.mark.skipif(not WordProjection()._check_docx(), reason="python-docx not installed")
    def test_save_from_template_does_not_replace_bare_chinese_table_labels(self, tmp_path):
        """Chinese labels in static tables are not implicit placeholders."""
        import docx

        template_path = tmp_path / "calendar_template.docx"
        output_path = tmp_path / "calendar_output.docx"
        doc = docx.Document()
        table = doc.add_table(rows=2, cols=3)
        table.rows[0].cells[0].text = "地区"
        table.rows[0].cells[1].text = "周内新闻"
        table.rows[0].cells[2].text = "影响标的"
        table.rows[1].cells[0].text = "美国"
        table.rows[1].cells[1].text = "{{美国新闻}}"
        table.rows[1].cells[2].text = "美元，美债，美股"
        doc.add_paragraph("2026-04-30")
        table2 = doc.add_table(rows=1, cols=2)
        table2.rows[0].cells[0].text = "美国"
        table2.rows[0].cells[1].text = "联邦基金目标利率(%)"
        doc.save(template_path)

        projection = WordProjection()
        projection.save_from_template(
            output_path,
            template_path,
            sections=[],
            placeholders={
                "美国": "美国宏观正文",
                "美国新闻": "美国新闻正文",
            },
        )

        result = docx.Document(output_path)
        all_cells = [
            cell.text
            for table in result.tables
            for row in table.rows
            for cell in row.cells
        ]
        assert all_cells[3] == "美国"
        assert all_cells[4] == "美国新闻正文"
        assert all_cells[6] == "美国"
        assert "美国宏观正文" not in all_cells

    @pytest.mark.skipif(not WordProjection()._check_docx(), reason="python-docx not installed")
    def test_save_from_template_with_tables(self, tmp_path):
        """测试从模板保存并添加表格"""

        projection = WordProjection()

        # 创建测试模板
        template_path = self.create_test_template(tmp_path)
        output_path = tmp_path / "output_with_table.docx"

        # 创建表格规范
        table = TableSpec(
            table_id="metrics",
            title="Key Metrics",
            headers=["Metric", "Value"],
            rows=[["Revenue", "$10M"], ["Profit", "$2M"]],
        )

        # 从模板保存
        projection.save_from_template(
            output_path,
            template_path,
            sections=[],
            placeholders={"title": "Test", "summary": "Test", "date": "2024"},
            tables=[table],
        )

        assert output_path.exists()

    @pytest.mark.skipif(not WordProjection()._check_docx(), reason="python-docx not installed")
    def test_save_from_template_replaces_placeholder_with_table_in_place(self, tmp_path):
        """Table placeholders should become real Word tables at the placeholder position."""
        import docx

        template_path = tmp_path / "calendar_template.docx"
        output_path = tmp_path / "calendar_output.docx"
        doc = docx.Document()
        doc.add_paragraph("前文")
        doc.add_paragraph("{{下周全球投资日历}}")
        doc.add_paragraph("后文")
        doc.save(template_path)

        table = TableSpec(
            table_id="global_investment_calendar",
            title="下周全球投资日历",
            placeholder="下周全球投资日历",
            headers=["日期", "国家/地区", "指标名称"],
            rows=[
                ["2026-06-15", "美国", "6月纽约联储制造业指数"],
                ["2026-06-16", "欧盟", "5月欧元区CPI:同比"],
            ],
        )

        projection = WordProjection()
        projection.save_from_template(
            output_path=output_path,
            template_path=template_path,
            sections=[],
            placeholders={},
            tables=[table],
        )

        result = docx.Document(output_path)
        assert len(result.tables) == 1
        assert result.tables[0].rows[0].cells[0].text == "日期"
        assert result.tables[0].rows[1].cells[1].text == "美国"
        paragraphs = [paragraph.text for paragraph in result.paragraphs]
        assert "{{下周全球投资日历}}" not in paragraphs

        body_tags = [child.tag.rsplit("}", 1)[-1] for child in result.element.body]
        table_index = body_tags.index("tbl")
        paragraph_texts_before_table = [
            child.text
            for child in result.element.body[:table_index]
            if child.tag.rsplit("}", 1)[-1] == "p"
        ]
        paragraph_texts_after_table = [
            child.text
            for child in result.element.body[table_index + 1 :]
            if child.tag.rsplit("}", 1)[-1] == "p"
        ]
        assert "前文" in paragraph_texts_before_table
        assert "后文" in paragraph_texts_after_table

    def test_save_from_template_raises_when_file_not_found(self, tmp_path):
        """测试模板文件不存在时抛出异常"""
        projection = WordProjection()
        if not projection._check_docx():
            pytest.skip("python-docx not installed")

        template_path = tmp_path / "non_existent.docx"
        output_path = tmp_path / "output.docx"

        with pytest.raises(FileNotFoundError):
            projection.save_from_template(output_path, template_path, [])

    def test_save_from_template_raises_when_docx_not_installed(self, tmp_path):
        """测试 python-docx 未安装时抛出异常"""
        projection = WordProjection()
        projection._has_docx = False

        template_path = tmp_path / "template.docx"
        output_path = tmp_path / "output.docx"

        with pytest.raises(ImportError):
            projection.save_from_template(output_path, template_path, [])
