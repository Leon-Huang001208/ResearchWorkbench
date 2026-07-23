"""
UnifiedPipeline 测试套件.

测试 Build → Validate → Render 三阶段流水线.
"""

import io
from pathlib import Path
from unittest.mock import MagicMock

import pytest

from core.contracts.content_element import (
    ContentBlock,
    HeadingElement,
    ParagraphElement,
)
from core.contracts.document import Document, Section
from core.contracts.reporting import ReportTask, SectionSpec, TemplateConfig
from reporting.builder.content_builder import ContentBuilder
from reporting.builder.pipeline import (
    PipelineResult,
    UnifiedPipeline,
    ValidationResult,
)

# ============================================================================
# Fixtures
# ============================================================================


@pytest.fixture
def pipeline():
    return UnifiedPipeline(output_dir=Path("output"))


@pytest.fixture
def sample_document():
    return Document(
        document_id="test_doc",
        title="Test Document",
        subtitle="A test document",
        sections=[
            Section(
                section_id="sec1",
                title="Section One",
                blocks=[
                    ContentBlock(
                        block_id="b1",
                        elements=[
                            HeadingElement(text="Heading", level=2),
                            ParagraphElement.plain("Hello world."),
                        ],
                    ),
                ],
            ),
        ],
    )


@pytest.fixture
def sample_task():
    return ReportTask(
        task_id="task_001",
        template_name="weekly_report",
        context={"market_summary": "Markets were up this week."},
    )


# ============================================================================
# ValidationResult 测试
# ============================================================================


class TestValidationResult:
    def test_default_is_valid(self):
        vr = ValidationResult()
        assert vr.is_valid is True
        assert vr.errors == []
        assert vr.warnings == []

    def test_with_errors(self):
        vr = ValidationResult(
            is_valid=False,
            errors=["Missing title", "No sections"],
        )
        assert not vr.is_valid
        assert len(vr.errors) == 2

    def test_with_stats(self):
        vr = ValidationResult(
            stats={"sections": 3, "elements": 15},
        )
        assert vr.stats["sections"] == 3


# ============================================================================
# PipelineResult 测试
# ============================================================================


class TestPipelineResult:
    def test_default_success(self):
        pr = PipelineResult()
        assert pr.success is True
        assert pr.outputs == {}
        assert pr.error is None

    def test_with_document(self, sample_document):
        pr = PipelineResult(document=sample_document)
        assert pr.document is not None
        assert pr.document.title == "Test Document"

    def test_with_outputs(self, sample_document):
        pr = PipelineResult(
            document=sample_document,
            outputs={"word": Path("output/test.docx")},
        )
        assert "word" in pr.outputs

    def test_failure_with_error(self):
        pr = PipelineResult(
            success=False,
            error="Build phase failed",
            warnings=["Missing metadata"],
        )
        assert not pr.success
        assert pr.error == "Build phase failed"
        assert len(pr.warnings) == 1


# ============================================================================
# UnifiedPipeline 初始化 测试
# ============================================================================


class TestPipelineInit:
    def test_default_init(self):
        p = UnifiedPipeline()
        assert p.output_dir == Path("output")
        assert isinstance(p.content_builder, ContentBuilder)

    def test_custom_output_dir(self):
        p = UnifiedPipeline(output_dir=Path("/tmp/reports"))
        assert p.output_dir == Path("/tmp/reports")

    def test_register_renderer(self, pipeline):
        mock_renderer = MagicMock()
        mock_renderer.format_name = "Mock"
        pipeline.register_renderer("mock", mock_renderer)
        assert "mock" in pipeline._renderers

    def test_register_template(self, pipeline):
        tpl = MagicMock()
        pipeline.register_template("test_tpl", tpl)
        assert "test_tpl" in pipeline._templates


# ============================================================================
# Validate 阶段测试
# ============================================================================


class TestValidate:
    def test_valid_document(self, pipeline, sample_document):
        result = pipeline._validate(sample_document)
        assert result.is_valid
        assert result.stats["sections"] == 1
        assert result.stats["elements"] == 2

    def test_missing_title(self, pipeline):
        doc = Document(title="")
        result = pipeline._validate(doc)
        assert not result.is_valid
        assert any("title" in e.lower() for e in result.errors)

    def test_no_sections(self, pipeline):
        doc = Document(
            document_id="d1",
            title="Empty Doc",
            sections=[],
        )
        result = pipeline._validate(doc)
        assert not result.is_valid
        assert any("at least one section" in e.lower() for e in result.errors)

    def test_duplicate_section_ids(self, pipeline):
        doc = Document(
            document_id="d1",
            title="Dup IDs",
            sections=[
                Section(section_id="s1", title="First"),
                Section(section_id="s1", title="Second"),
            ],
        )
        result = pipeline._validate(doc)
        assert not result.is_valid
        assert any("duplicate" in e.lower() for e in result.errors)

    def test_missing_block_id_warning(self, pipeline):
        doc = Document(
            document_id="d1",
            title="No Block ID",
            sections=[
                Section(
                    section_id="s1",
                    title="S1",
                    blocks=[
                        ContentBlock(
                            block_id="",
                            elements=[ParagraphElement.plain("ok")],
                        ),
                    ],
                ),
            ],
        )
        result = pipeline._validate(doc)
        assert result.is_valid  # 缺少 block_id 仅是警告，不是错误
        assert any("missing block_id" in w.lower() for w in result.warnings)

    def test_empty_block_warning(self, pipeline):
        doc = Document(
            document_id="d1",
            title="Empty Block",
            sections=[
                Section(
                    section_id="s1",
                    title="S1",
                    blocks=[
                        ContentBlock(block_id="b1", elements=[]),
                    ],
                ),
            ],
        )
        result = pipeline._validate(doc)
        assert result.is_valid  # 空 block 仅是警告
        assert any("no elements" in w.lower() for w in result.warnings)

    def test_missing_section_title_warning(self, pipeline):
        doc = Document(
            document_id="d1",
            title="Doc",
            sections=[
                Section(
                    section_id="s1",
                    title="",
                    blocks=[
                        ContentBlock(
                            block_id="b1",
                            elements=[ParagraphElement.plain("ok")],
                        ),
                    ],
                ),
            ],
        )
        result = pipeline._validate(doc)
        assert any("missing title" in w.lower() for w in result.warnings)

    def test_stats_count_elements(self, pipeline):
        doc = Document(
            document_id="d1",
            title="Stats Doc",
            sections=[
                Section(
                    section_id="s1",
                    title="S1",
                    blocks=[
                        ContentBlock(
                            block_id="b1",
                            elements=[
                                ParagraphElement.plain("a"),
                                ParagraphElement.plain("b"),
                                ParagraphElement.plain("c"),
                            ],
                        ),
                        ContentBlock(
                            block_id="b2",
                            elements=[HeadingElement(text="H", level=3)],
                        ),
                    ],
                ),
                Section(
                    section_id="s2",
                    title="S2",
                    blocks=[
                        ContentBlock(
                            block_id="b3",
                            elements=[ParagraphElement.plain("x")],
                        ),
                    ],
                ),
            ],
        )
        result = pipeline._validate(doc)
        assert result.stats["sections"] == 2
        assert result.stats["blocks"] == 3
        assert result.stats["elements"] == 5


# ============================================================================
# Render 阶段测试
# ============================================================================

# 标记需要 python-pptx 的测试
try:
    import pptx  # noqa: F401

    PPT_AVAILABLE = True
except ImportError:
    PPT_AVAILABLE = False


class TestRender:
    def test_resolve_renderer_word(self, pipeline):
        renderer = pipeline._resolve_renderer("word")
        assert renderer.format_name == "Word"

    @pytest.mark.skipif(not PPT_AVAILABLE, reason="python-pptx not installed")
    def test_resolve_renderer_ppt(self, pipeline):
        renderer = pipeline._resolve_renderer("ppt")
        assert renderer.format_name == "PPT"

    def test_resolve_renderer_markdown(self, pipeline):
        renderer = pipeline._resolve_renderer("markdown")
        assert renderer.format_name == "Markdown"

    def test_resolve_renderer_aliases(self, pipeline):
        assert pipeline._resolve_renderer("docx").format_name == "Word"
        if PPT_AVAILABLE:
            assert pipeline._resolve_renderer("pptx").format_name == "PPT"
        assert pipeline._resolve_renderer("md").format_name == "Markdown"

    def test_resolve_renderer_registered(self, pipeline):
        mock = MagicMock()
        mock.format_name = "Custom"
        pipeline.register_renderer("custom", mock)
        assert pipeline._resolve_renderer("custom") is mock

    def test_resolve_renderer_unknown(self, pipeline):
        with pytest.raises(ValueError, match="Unsupported format"):
            pipeline._resolve_renderer("pdf")

    def test_format_extension(self, pipeline):
        assert pipeline._format_extension("word") == "docx"
        assert pipeline._format_extension("ppt") == "pptx"
        assert pipeline._format_extension("markdown") == "md"
        assert pipeline._format_extension("docx") == "docx"
        assert pipeline._format_extension("unknown") == "unknown"


# ============================================================================
# Execute 端到端测试
# ============================================================================


class TestExecute:
    def test_execute_markdown(self, pipeline, sample_document, tmp_path):
        pipeline.register_renderer("markdown", pipeline._resolve_renderer("markdown"))
        output_dir = tmp_path / "reports"
        output_dir.mkdir()

        result = pipeline.execute(
            task=sample_document,
            output_formats=["markdown"],
            output_dir=output_dir,
            strategy="document",
        )
        assert result.success
        assert "markdown" in result.outputs
        assert result.outputs["markdown"].exists()

    def test_execute_validation_failure(self, pipeline, tmp_path):
        bad_doc = Document(title="")  # 缺少标题
        output_dir = tmp_path / "reports"
        output_dir.mkdir()

        result = pipeline.execute(
            task=bad_doc,
            output_formats=["markdown"],
            output_dir=output_dir,
            strategy="document",
        )
        assert not result.success
        assert "Validation failed" in result.error

    def test_execute_with_warnings(self, pipeline, sample_document, tmp_path):
        pipeline.register_renderer("markdown", pipeline._resolve_renderer("markdown"))
        output_dir = tmp_path / "reports"
        output_dir.mkdir()

        # 删除 document_id 触发 warning 但不是 error
        doc = sample_document.model_copy(update={"document_id": ""})
        result = pipeline.execute(
            task=doc,
            output_formats=["markdown"],
            output_dir=output_dir,
            strategy="document",
        )
        # 应该成功但有警告
        assert result.success
        assert len(result.warnings) > 0

    def test_execute_to_buffer(self, pipeline, sample_document):
        pipeline.register_renderer("markdown", pipeline._resolve_renderer("markdown"))
        result = pipeline.execute_to_buffer(
            task=sample_document,
            output_formats=["markdown"],
            strategy="document",
        )
        assert result.success
        assert "markdown" in result.buffers
        assert isinstance(result.buffers["markdown"], io.BytesIO)
        content = result.buffers["markdown"].getvalue().decode("utf-8")
        assert "Test Document" in content

    def test_execute_unknown_strategy(self, pipeline, sample_document):
        result = pipeline.execute(
            task=sample_document,
            output_formats=["markdown"],
            strategy="unsupported_strategy",
        )
        assert not result.success
        assert result.error is not None
        assert "Unknown build strategy" in result.error

    def test_execute_template_strategy(self, pipeline, sample_task, tmp_path):
        pipeline.register_renderer("markdown", pipeline._resolve_renderer("markdown"))
        output_dir = tmp_path / "reports"
        output_dir.mkdir()

        # 注册模板
        tpl = TemplateConfig(
            name="weekly_report",
            description="Weekly Report",
            version="1.0",
            sections=[
                SectionSpec(
                    key="market_summary",
                    title="Market Summary",
                    target_words=200,
                ),
            ],
        )
        pipeline.register_template("weekly_report", tpl)

        result = pipeline.execute(
            task=sample_task,
            output_formats=["markdown"],
            output_dir=output_dir,
            strategy="template",
        )
        assert result.success
        assert "markdown" in result.outputs


# ============================================================================
# quick_render 便捷函数测试
# ============================================================================


class TestQuickRender:
    def test_quick_render_markdown(self, sample_document, tmp_path):
        output_dir = tmp_path / "quick"
        output_dir.mkdir()

        # quick_render pattern: pass document directly with strategy="document"
        # 需要直接构造一个能通过 _build 的流程
        pipeline = UnifiedPipeline(output_dir=output_dir)
        pipeline.register_renderer("markdown", pipeline._resolve_renderer("markdown"))
        result = pipeline.execute(
            task=sample_document,
            output_formats=["markdown"],
            output_dir=output_dir,
            strategy="document",
        )
        assert result.success
        assert "markdown" in result.outputs
