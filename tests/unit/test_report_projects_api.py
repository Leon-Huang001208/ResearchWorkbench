"""Tests for report project API routes."""

import hashlib
import json
import re
import threading
import time
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path
from unittest.mock import patch

import yaml
from fastapi.testclient import TestClient

from app.api.main import app
from core.interfaces.model_gateway import ModelResponse
from reporting.projects.chart_generation import GeneratedChartInfo
from reporting.projects.generation import (
    EvidenceSnippet,
    GeneratedSectionInfo,
    PromptTemplateBlock,
    ReportGenerationResult,
    ReportProjectGenerationService,
    RetrievalConfig,
    _dedupe_data_sentence_fuzzy,
    _strip_instruction_leaks,
    apply_report_defaults_to_placeholder,
    build_market_hotspot_messages,
    build_retrieval_config,
    compute_report_period,
    filter_and_rank_evidence,
    render_generation_constraints,
    render_writing_parameters,
    resolve_report_generation_scope,
)
from reporting.projects.keyword_profiles import (
    apply_keyword_profile_to_config,
    keyword_profiles_for_api,
)
from reporting.projects.project_manager import ReportProject, ReportProjectManager
from reporting.projects.run import ReportProjectRunRequest, ReportProjectRunService

client = TestClient(app)


def test_huaan_prompt_placeholders_use_report_level_retrieval_defaults():
    """华安周报公共检索和重排配置应放在 defaults，placeholder 只保留关键词差异。"""
    config_path = (
        Path(__file__).resolve().parents[2]
        / "report_projects"
        / "华安ETF周报"
        / "config"
        / "report_config.yaml"
    )
    source = config_path.read_text(encoding="utf-8")
    config = yaml.safe_load(source)

    assert "&id" not in source
    assert "*id" not in source

    defaults = config["defaults"]
    retrieval_defaults = defaults["retrieval"]
    assert retrieval_defaults["mode"] == "hybrid"
    assert retrieval_defaults["top_k"] == 10
    assert retrieval_defaults["candidate_k"] == 40
    assert retrieval_defaults["semantic_candidate_k"] == 80
    assert retrieval_defaults["keyword_weight"] == 0.6
    assert retrieval_defaults["semantic_weight"] == 0.4
    assert retrieval_defaults["embedding_model"] == "data/models/embeddings/bge-large-zh-v1.5"
    rerank_defaults = defaults["rerank"]
    assert rerank_defaults["enabled"] is True
    assert rerank_defaults["provider"] == "bge-reranker"
    assert rerank_defaults["model"] == "data/models/rerankers/bge-reranker-large"
    assert rerank_defaults["top_n"] == 30
    assert rerank_defaults["min_score"] == 0.35
    assert defaults["validators"]["forbid_external_facts"] is True
    assert defaults["validators"]["forbid_direct_investment_advice"] is True

    prompt_placeholders = {
        name: item
        for name, item in config["placeholders"].items()
        if item.get("prompt_template") and item.get("type") in {"paragraph", "prompt"}
    }

    assert prompt_placeholders
    for name, item in prompt_placeholders.items():
        retrieval = item.get("retrieval")
        if not retrieval:
            retrieval = next(
                (
                    component.get("retrieval")
                    for component in item.get("components", [])
                    if component.get("type") == "llm_writing"
                ),
                None,
            )
        assert retrieval, f"{name} missing retrieval"
        assert retrieval.get("keywords") or retrieval.get(
            "keyword_profile"
        ), f"{name} missing retrieval keywords/profile"
        assert "query_terms" not in retrieval
        assert "mode" not in retrieval
        assert "top_k" not in retrieval
        assert "fusion" not in retrieval
        assert "rerank" not in retrieval


def test_report_defaults_are_merged_before_building_retrieval_config():
    """生成侧应将 defaults.retrieval / defaults.validators 合并进单个占位符。"""
    report_config = {
        "defaults": {
            "query_mode": "retrieval_query_embedded",
            "validators": {
                "forbid_wind_data": True,
                "no_newline": True,
                "forbidden_phrases": ["根据文件"],
            },
            "retrieval": {
                "mode": "hybrid",
                "top_k": 10,
                "candidate_k": 40,
                "semantic_candidate_k": 80,
                "keyword_weight": 0.6,
                "semantic_weight": 0.4,
                "embedding_model": "BAAI/bge-large-zh-v1.5",
            },
            "rerank": {
                "enabled": True,
                "provider": "bge-reranker",
                "model": "BAAI/bge-reranker-large",
                "top_n": 30,
                "min_score": 0.35,
            },
        }
    }
    placeholder_config = {
        "type": "prompt",
        "title": "航天",
        "retrieval": {"keywords": ["航天", "卫星"]},
    }

    merged = apply_report_defaults_to_placeholder(report_config, placeholder_config)
    retrieval_config = build_retrieval_config(merged, default_top_k=4)

    assert merged["query_mode"] == "retrieval_query_embedded"
    assert merged["validators"]["forbid_wind_data"] is True
    assert merged["validators"]["forbidden_phrases"] == ["根据文件"]
    assert retrieval_config.mode == "hybrid"
    assert retrieval_config.top_k == 10
    assert retrieval_config.candidate_k == 40
    assert retrieval_config.must_any == ["航天", "卫星"]
    assert retrieval_config.embedding_model == "BAAI/bge-large-zh-v1.5"
    assert retrieval_config.rerank_enabled is True
    assert retrieval_config.rerank_provider == "bge-reranker"
    assert retrieval_config.rerank_model == "BAAI/bge-reranker-large"
    assert retrieval_config.rerank_top_n == 30
    assert retrieval_config.min_rerank_score == 0.35


def test_generation_constraints_and_writing_parameters_are_rendered_separately():
    """共用生成约束和单段写作参数应分层进入 prompt。"""
    config = {
        "generation_constraints": [
            "严格依据上传材料和 evidence，不添加外部知识或虚构数据",
            "生成一段正文，不输出换行符",
        ],
        "target_words": 250,
        "max_words": 320,
        "min_news_count": 5,
    }

    constraints = render_generation_constraints(config)
    writing_parameters = render_writing_parameters(config)

    assert "严格依据上传材料和 evidence" in constraints
    assert "生成一段正文，不输出换行符" in constraints
    assert "目标字数" not in constraints
    assert "至少使用" not in constraints
    assert "目标字数：约 250 字" in writing_parameters
    assert "最大字数：不超过 320 字" in writing_parameters
    assert "至少使用 5 条 evidence/news 信息" in writing_parameters


def test_market_hotspot_prompt_uses_component_structure_without_metadata(tmp_path: Path):
    """A股市场回顾续写 prompt 应使用固定开头和后续结构，不暴露项目元信息。"""
    project = ReportProject(
        name="测试周报",
        slug="test-weekly-report",
        project_dir=tmp_path,
        word_template_path=tmp_path / "report_template.docx",
        excel_workbook_path=tmp_path / "report_data.xlsx",
        report_config_path=tmp_path / "report_config.yaml",
        output_dir=tmp_path / "generated",
        run_log_dir=tmp_path / "runs",
    )
    template = build_market_template_for_test()
    config = {
        "generation_constraints": [
            "严格依据上传材料和 evidence，不添加外部知识或虚构数据",
            "生成一段正文，不输出换行符",
        ],
        "target_words": 250,
        "max_words": 320,
        "min_news_count": 5,
        "components": [
            {
                "name": "市场热点与趋势判断",
                "type": "llm_writing",
                "writing_structure": [
                    "接在固定开头之后，概括本周市场热点板块或概念",
                    "描述板块轮动特征",
                ],
            }
        ],
    }

    messages = build_market_hotspot_messages(
        project=project,
        placeholder="A股市场回顾",
        title="A股市场回顾",
        template=template,
        data_sentence=("本周A股市场整体呈现分化趋势，主要指数表现不一：沪深300涨0.19%。" "交易面，A股市场本周日均成交额在2.40万亿左右，市场投资热情回落。"),
        params={},
        max_words=320,
        config=config,
        evidence=[
            EvidenceSnippet(
                source="ingestion:news",
                title="热点",
                content="算力硬件、新能源和商业航天反复活跃。",
                published_at="2026-06-05",
            )
        ],
    )

    assert len(messages) == 2
    assert messages[0]["role"] == "system"
    assert messages[1]["role"] == "user"
    system_prompt = messages[0]["content"]
    user_prompt = messages[1]["content"]
    assert "不要重复或改写已经给出的市场数据开头" in system_prompt
    assert user_prompt.index("生成约束：") < user_prompt.index("写作参数：")
    assert "已完成的市场数据开头" in user_prompt
    assert "续写要求：" in user_prompt
    assert "接在固定开头之后，概括本周市场热点板块或概念" in user_prompt
    assert "请只输出续写正文" in user_prompt
    assert "项目：" not in user_prompt
    assert "Word 占位符" not in user_prompt
    assert "段落标题" not in user_prompt
    assert "检索 Query" not in user_prompt
    # 旧版 prompt 中的易泄漏短语不应出现在新 prompt 中
    assert "固定开头：" not in user_prompt
    assert "以下内容已由系统" not in user_prompt
    assert "不得改写、删减或重复" not in user_prompt
    assert "配置参数" not in user_prompt


def test_strip_instruction_leaks_removes_actual_prompt_wording():
    """应移除 LLM 输出的 '以下内容已由系统根据Excel数据生成...' 指令文本。"""
    text = "以下内容已由系统根据Excel数据生成，必须作为正文开头，不得改写、删减或重复" "：本周市场热点集中于避险与政策驱动方向，半导体板块遭遇集中抛售。"
    result = _strip_instruction_leaks(text)
    assert "以下内容已由系统" not in result
    assert "不得改写" not in result
    assert "删减或重复" not in result
    assert result.startswith("本周市场热点")


def test_strip_instruction_leaks_removes_parenthesized_variant():
    """应移除括号包裹的旧版指令文本变体。"""
    text = "（以上内容由系统从 Excel 底稿自动生成，请以本段为固定开头）" "本周市场热点集中于医药板块。"
    result = _strip_instruction_leaks(text)
    assert "以上内容由系统从" not in result
    assert "Excel 底稿" not in result
    assert result.startswith("本周市场热点")


def test_strip_instruction_leaks_removes_bare_parenthetical():
    """应移除裸括号内的 '（不得改写、删减或重复）'。"""
    text = "（不得改写、删减或重复）本周市场整体呈调整趋势。"
    result = _strip_instruction_leaks(text)
    assert "不得改写" not in result
    assert result.startswith("本周市场")


def test_strip_instruction_leaks_handles_no_leaks():
    """无泄漏的正文应原样返回。"""
    text = "本周A股市场整体呈现调整趋势，市场热点集中于避险方向。"
    result = _strip_instruction_leaks(text)
    assert result == text


def test_dedupe_data_sentence_fuzzy_exact_match():
    """精确重复的数据句应被去重（调用方已有精确匹配，但验证模糊函数不破坏正常流程）。"""
    ds = "沪深300涨1.72%，市场投资热情回落。"
    hotspot = ds + "市场热点集中于避险方向。"
    result = _dedupe_data_sentence_fuzzy(ds, hotspot)
    assert result == hotspot


def test_dedupe_data_sentence_fuzzy_with_transition_word():
    """过渡词后重复数据句应被去重。"""
    ds = "本周A股市场整体呈现调整趋势，主要指数表现不一：沪深300跌1.72%。"
    hotspot = "具体来看，" + ds + "市场热点集中于避险与政策驱动方向。"
    result = _dedupe_data_sentence_fuzzy(ds, hotspot)
    assert "沪深300跌1.72%" not in result or result.startswith("市场热点")


def test_dedupe_data_sentence_fuzzy_no_duplicate():
    """无重复的正文应原样返回。"""
    ds = "沪深300涨0.86%，中证500涨0.41%。"
    hotspot = "市场热点集中于医药和消费板块，资金风险偏好有所回升。"
    result = _dedupe_data_sentence_fuzzy(ds, hotspot)
    assert result == hotspot


def test_dedupe_data_sentence_fuzzy_short_text():
    """过短文本不做模糊去重，原样返回。"""
    ds = "沪深300涨0.86%。"
    hotspot = "市场反弹。"
    result = _dedupe_data_sentence_fuzzy(ds, hotspot)
    assert result == hotspot


def build_market_template_for_test():
    return PromptTemplateBlock(
        title="A股市场回顾",
        retrieval_query="请检索市场热点",
        writing_requirements="旧写作要求",
        raw_text="旧模板",
    )


def test_keyword_profiles_are_available_for_report_project_workbench():
    """项目详情应返回可复用关键词 profile，供新模板占位符初始化检索词。"""
    profiles = keyword_profiles_for_api()

    assert "人工智能" in profiles
    assert "AI" in profiles["人工智能"]["keywords"]
    assert profiles["人工智能"]["query"]
    assert profiles["人工智能"]["threshold"] > 0

    response = client.get("/api/report-projects/")
    assert response.status_code == 200
    projects = response.json()["projects"]
    assert projects
    assert "人工智能" in projects[0]["keyword_profiles"]


def test_missing_retrieval_keywords_are_filled_from_keyword_profile():
    """占位符未显式写关键词时，生成侧应从 keyword_profiles 继承检索关键词。"""
    config = {
        "title": "人工智能",
        "type": "prompt",
        "prompt_template": "人工智能",
        "retrieval": {"mode": "hybrid"},
    }

    enriched = apply_keyword_profile_to_config("人工智能", config)
    retrieval_config = build_retrieval_config(enriched, default_top_k=8)

    assert "人工智能" in retrieval_config.must_any
    assert "AI" in retrieval_config.must_any
    assert enriched["retrieval"]["keyword_profile"] == "人工智能"
    assert enriched["retrieval"]["keywords"]
    assert "query_terms" not in enriched["retrieval"]


def test_explicit_keyword_profile_fills_keywords_when_keywords_are_missing():
    """配置了 keyword_profile 但未写 keywords 时，应从指定 profile 展开关键词。"""
    config = {
        "title": "自定义标题",
        "type": "prompt",
        "retrieval": {"keyword_profile": "航天"},
    }

    enriched = apply_keyword_profile_to_config("自定义标题", config)
    retrieval_config = build_retrieval_config(enriched, default_top_k=8)

    assert "航天" in retrieval_config.must_any
    assert enriched["retrieval"]["keyword_profile"] == "航天"
    assert enriched["retrieval"]["keywords"]


def test_compute_report_period_uses_report_date_and_week_monday():
    """报告周期结束日取报告日，开始日取同周周一。"""
    period = compute_report_period("2026-06-05")

    assert period.start_date == "2026-06-01"
    assert period.end_date == "2026-06-05"


def test_resolve_report_generation_scope_prefilters_before_retrieval():
    """生成范围先统一解析，后续检索只接收已经确定的周期。"""
    report_config = {
        "defaults": {
            "report_period": {
                "report_date": "2026-06-05",
                "start_date": "2026-05-30",
                "end_date": "2026-06-05",
            }
        }
    }

    yaml_scope = resolve_report_generation_scope(report_config)

    assert yaml_scope.report_period.start_date == "2026-05-30"
    assert yaml_scope.report_period.end_date == "2026-06-05"
    assert yaml_scope.data_scope == "custom"

    override_scope = resolve_report_generation_scope(
        report_config,
        report_date="2026-06-10",
        start_date="2026-06-03",
        end_date="2026-06-10",
    )

    assert override_scope.report_period.start_date == "2026-06-03"
    assert override_scope.report_period.end_date == "2026-06-10"
    assert override_scope.data_scope == "custom"


def write_minimal_docx(path: Path, text: str) -> None:
    """Write a tiny docx package with one document.xml body."""
    document_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body><w:p><w:r><w:t>{text}</w:t></w:r></w:p></w:body>
</w:document>"""
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("word/document.xml", document_xml)
        archive.writestr(
            "word/_rels/document.xml.rels",
            """<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>""",
        )


def write_minimal_xlsx(path: Path) -> None:
    """Write a tiny xlsx package with workbook metadata and one sheet."""
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr(
            "xl/workbook.xml",
            """<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main"
                xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
                <sheets><sheet name="基本信息" sheetId="1" r:id="rId1"/></sheets>
            </workbook>""",
        )
        archive.writestr(
            "xl/_rels/workbook.xml.rels",
            """<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
                <Relationship Id="rId1" Type="worksheet" Target="worksheets/sheet1.xml"/>
            </Relationships>""",
        )
        archive.writestr(
            "xl/worksheets/sheet1.xml",
            """<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
                <dimension ref="A1:C3"/>
                <sheetData>
                  <row r="1"><c r="A1" t="inlineStr"><is><t>日期</t></is></c></row>
                  <row r="2"><c r="B2" t="inlineStr"><is><t>创业板50</t></is></c></row>
                </sheetData>
            </worksheet>""",
        )


def write_minimal_pptx(path: Path, text: str) -> None:
    """Write a tiny pptx package with one slide text box."""
    slide_xml = f"""<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<p:sld xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"
       xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main">
  <p:cSld>
    <p:spTree>
      <p:nvGrpSpPr><p:cNvPr id="1" name=""/><p:cNvGrpSpPr/><p:nvPr/></p:nvGrpSpPr>
      <p:grpSpPr/>
      <p:sp>
        <p:nvSpPr><p:cNvPr id="2" name="Title"/><p:cNvSpPr/><p:nvPr/></p:nvSpPr>
        <p:txBody>
          <a:bodyPr/><a:lstStyle/>
          <a:p><a:r><a:t>{text}</a:t></a:r></a:p>
        </p:txBody>
      </p:sp>
    </p:spTree>
  </p:cSld>
</p:sld>"""
    with zipfile.ZipFile(path, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("ppt/slides/slide1.xml", slide_xml)


def write_market_review_xlsx(path: Path) -> None:
    """Write cached market data used by composite market review generation."""
    from openpyxl import Workbook

    workbook = Workbook()
    date_sheet = workbook.active
    date_sheet.title = "日期"
    date_sheet.append(["本周五", "上周五"])
    date_sheet.append(["2026-06-05", "2026-05-29"])

    domestic = workbook.create_sheet("国内")
    domestic.append(["指数代码", "指数名称", "周内涨跌幅"])
    domestic.append(["000300.SH", "沪深300", 0.86])
    domestic.append(["000905.SH", "中证500", 0.41])
    domestic.append(["000852.SH", "中证1000", -0.04])
    domestic.append(["399673.SZ", "创业板50", -0.42])
    domestic.append(["000688.SH", "科创50", 2.13])

    turnover = workbook.create_sheet("市场成交")
    turnover.append(["指数代码", "本周日均成交额", "上周日均成交额"])
    turnover.append(["000985.CSI", 2.55, 2.30])

    gold = workbook.create_sheet("黄金")
    gold.append(["代码", "简称", "周收盘价", "周涨跌幅"])
    gold.append(["SPTAUUSDOZ.IDC", "伦敦金现", 4704.743, -2.744129703627285])
    gold.append(["AU9999.SGE", "SGE黄金9999", 1033.25, -1.7486972728310162])

    oil = workbook.create_sheet("石油")
    oil.append(["代码", "简称", "周涨跌幅", "周收盘价", "上周收盘价", "涨跌"])
    oil.append(["B.IPE", "ICE布油", 17.293649037397675, 106.01, 92.42, 13.59])
    oil.append(["CL.NYM", "NYMEX WTI原油", 17.447632885337217, 97, 84, 13])

    workbook.save(path)


def write_global_calendar_xlsx(path: Path) -> None:
    """Write a cached global economic calendar workbook."""
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "经济数据"
    sheet.append(["日期", "时间", "国家/地区", "指标名称", "重要性", "前值", "预测值", "今值"])
    sheet.append(["2026-06-15", "21:30", "美国", "6月纽约联储制造业指数", "重要", -9.2, None, None])
    sheet.append(["2026-06-16", "17:00", "欧盟", "5月欧元区CPI:同比", "重要", 2.2, None, None])
    sheet.append(["2026-06-17", "08:00", "日本", "低优先级指标", "一般", 1.0, None, None])
    workbook.save(path)


def test_list_report_projects_returns_project_assets(tmp_path: Path, monkeypatch):
    """报告项目接口应返回项目包资产、配置和历史报告。"""
    project_dir = tmp_path / "创业板50周报"
    (project_dir / "templates").mkdir(parents=True)
    (project_dir / "data").mkdir()
    (project_dir / "config").mkdir()
    (project_dir / "generated").mkdir()
    (project_dir / "runs").mkdir()
    (project_dir / "templates" / "report_template.docx").write_bytes(b"docx")
    (project_dir / "data" / "创业板50周报（iFind版）.xlsx").write_bytes(b"xlsx")
    (project_dir / "data" / "domestic.json").write_text("{}", encoding="utf-8")
    (project_dir / "data" / "周报图表.xlsx").write_bytes(b"chart")
    (project_dir / "data" / "周报页眉.png").write_bytes(b"png")
    (project_dir / "config" / "report_config.yaml").write_text("sections: []\n", encoding="utf-8")
    (project_dir / "config" / "prompt_templates.md").write_text("前缀提示词", encoding="utf-8")
    (project_dir / "generated" / "20260605_创业板50周报.docx").write_bytes(b"report")
    (project_dir / "project.yaml").write_text(
        "\n".join(
            [
                "name: 创业板50周报",
                "active_word_template: templates/report_template.docx",
                "active_excel_workbook: data/创业板50周报（iFind版）.xlsx",
                "report_config: config/report_config.yaml",
                "prompt_templates: config/prompt_templates.md",
                "data_sources:",
                "  - data/domestic.json",
                "output_dir: generated",
                "run_log_dir: runs",
            ]
        ),
        encoding="utf-8",
    )

    import app.api.routes.report_projects as report_projects_route

    monkeypatch.setattr(
        report_projects_route,
        "report_project_manager",
        ReportProjectManager(projects_root=tmp_path),
    )

    response = client.get("/api/report-projects/")

    assert response.status_code == 200
    data = response.json()
    assert data["total"] == 1
    project = data["projects"][0]
    assert project["name"] == "创业板50周报"
    assert project["word_template_filename"] == "report_template.docx"
    assert project["excel_workbook_filename"] == "创业板50周报（iFind版）.xlsx"
    assert project["report_config_filename"] == "report_config.yaml"
    assert project["prompt_templates_filename"] == "prompt_templates.md"
    assert project["data_source_files"] == ["domestic.json"]
    assert [asset["file_name"] for asset in project["data_assets"]] == [
        "创业板50周报（iFind版）.xlsx",
        "周报图表.xlsx",
        "domestic.json",
        "周报页眉.png",
    ]
    assert project["data_assets"][0]["kind"] == "primary_excel"
    assert project["data_assets"][1]["kind"] == "workbook"
    assert project["data_assets"][2]["kind"] == "query_json"
    assert project["data_assets"][3]["kind"] == "image"
    assert project["generated_reports"][0]["file_name"] == "20260605_创业板50周报.docx"


def test_get_report_project_returns_real_template_asset_summary(tmp_path: Path, monkeypatch):
    """项目详情应返回真实 Word 占位符、原始 section YAML 和 Excel sheet 摘要。"""
    project_dir = tmp_path / "创业板50周报"
    (project_dir / "templates").mkdir(parents=True)
    (project_dir / "data").mkdir()
    (project_dir / "config").mkdir()
    (project_dir / "generated").mkdir()
    (project_dir / "runs").mkdir()
    write_minimal_docx(
        project_dir / "templates" / "report_template.docx",
        "标题 {{ title }} 日期 {{ start_date }} 内容 {{ content1 }}",
    )
    write_minimal_xlsx(project_dir / "data" / "cyb50.xlsx")
    section_yaml = "\n".join(
        [
            "name: 创业板50周报模板",
            "sections:",
            "- key: content1",
            "  title: 正文观点",
            "  placeholder: content1",
            "  type: paragraph",
            "  prompt_template: 正文观点",
            "  retrieval:",
            "    keywords:",
            "    - 创业板50",
        ]
    )
    (project_dir / "config" / "report_config.yaml").write_text(section_yaml, encoding="utf-8")
    prompt_source = "## 正文观点\n检索 Query：创业板50\n\n写作要求：严格依据上传文件"
    (project_dir / "config" / "prompt_templates.md").write_text(prompt_source, encoding="utf-8")
    (project_dir / "project.yaml").write_text(
        "\n".join(
            [
                "name: 创业板50周报",
                "active_word_template: templates/report_template.docx",
                "active_excel_workbook: data/cyb50.xlsx",
                "report_config: config/report_config.yaml",
                "prompt_templates: config/prompt_templates.md",
                "output_dir: generated",
                "run_log_dir: runs",
            ]
        ),
        encoding="utf-8",
    )

    import app.api.routes.report_projects as report_projects_route

    monkeypatch.setattr(
        report_projects_route,
        "report_project_manager",
        ReportProjectManager(projects_root=tmp_path),
    )

    response = client.get("/api/report-projects/创业板50周报")

    assert response.status_code == 200
    project = response.json()
    assert project["word_placeholders"] == ["title", "start_date", "content1"]
    assert project["report_config_source"] == section_yaml
    assert project["prompt_templates_source"] == prompt_source
    assert project["report_config"]["sections"][0]["placeholder"] == "content1"
    assert project["excel_sheets"][0]["name"] == "基本信息"
    assert project["excel_sheets"][0]["dimension"] == "A1:C3"
    assert project["excel_sheets"][0]["nonempty_count"] == 2
    assert "A1=日期" in project["excel_sheets"][0]["sample_cells"]
    assert project["compiled_plan"]["ready"] is True
    assert project["compiled_plan"]["placeholders"][0]["placeholder"] == "content1"
    assert project["compiled_plan"]["placeholders"][0]["prompt_found"] is True
    assert project["compiled_plan"]["placeholders"][0]["retrieval_ready"] is True


def test_get_report_project_returns_ppt_template_placeholders(tmp_path: Path, monkeypatch):
    """PPT 项目详情应返回 PPT 模板资产和占位符。"""
    project_dir = tmp_path / "月度PPT"
    (project_dir / "templates").mkdir(parents=True)
    (project_dir / "config").mkdir()
    (project_dir / "generated").mkdir()
    (project_dir / "runs").mkdir()
    write_minimal_pptx(
        project_dir / "templates" / "report_template.pptx",
        "标题 {{title}} 期间 {{period}}",
    )
    (project_dir / "config" / "report_config.yaml").write_text(
        "placeholders:\n  title:\n    type: static\n    value: 月度PPT\n",
        encoding="utf-8",
    )
    (project_dir / "project.yaml").write_text(
        "\n".join(
            [
                "name: 月度PPT",
                "project_type: ppt",
                "active_ppt_template: templates/report_template.pptx",
                "report_config: config/report_config.yaml",
                "output_dir: generated",
                "run_log_dir: runs",
            ]
        ),
        encoding="utf-8",
    )

    import app.api.routes.report_projects as report_projects_route

    monkeypatch.setattr(
        report_projects_route,
        "report_project_manager",
        ReportProjectManager(projects_root=tmp_path),
    )

    response = client.get("/api/report-projects/月度PPT")

    assert response.status_code == 200
    project = response.json()
    assert project["project_type"] == "ppt"
    assert project["template_filename"] == "report_template.pptx"
    assert project["ppt_template_filename"] == "report_template.pptx"
    assert project["word_template_filename"] == ""
    assert project["ppt_placeholders"] == ["title", "period"]
    assert project["word_placeholders"] == []


def test_docx_placeholders_keep_word_first_seen_order(tmp_path: Path, monkeypatch):
    """占位符地图应按 Word 正文首次出现顺序展示，而不是按名称排序。"""
    project_dir = tmp_path / "排序周报"
    (project_dir / "templates").mkdir(parents=True)
    (project_dir / "data").mkdir()
    (project_dir / "config").mkdir()
    (project_dir / "generated").mkdir()
    (project_dir / "runs").mkdir()
    write_minimal_docx(
        project_dir / "templates" / "report_template.docx",
        "{{z_last}} {{a_first}} {{middle}} {{a_first}}",
    )
    write_minimal_xlsx(project_dir / "data" / "data.xlsx")
    (project_dir / "config" / "report_config.yaml").write_text("sections: []\n", encoding="utf-8")
    (project_dir / "project.yaml").write_text(
        "\n".join(
            [
                "name: 排序周报",
                "active_word_template: templates/report_template.docx",
                "active_excel_workbook: data/data.xlsx",
                "report_config: config/report_config.yaml",
                "output_dir: generated",
                "run_log_dir: runs",
            ]
        ),
        encoding="utf-8",
    )

    import app.api.routes.report_projects as report_projects_route

    monkeypatch.setattr(
        report_projects_route,
        "report_project_manager",
        ReportProjectManager(projects_root=tmp_path),
    )

    response = client.get("/api/report-projects/排序周报")

    assert response.status_code == 200
    assert response.json()["word_placeholders"] == ["z_last", "a_first", "middle"]


def test_update_report_project_source_persists_prompt_templates(tmp_path: Path, monkeypatch):
    """源码保存应写回项目文件，而不是只保存浏览器草稿。"""
    project_dir = tmp_path / "华安ETF周报"
    (project_dir / "templates").mkdir(parents=True)
    (project_dir / "data").mkdir()
    (project_dir / "config").mkdir()
    (project_dir / "generated").mkdir()
    (project_dir / "runs").mkdir()
    write_minimal_docx(project_dir / "templates" / "report_template.docx", "{{ content1 }}")
    write_minimal_xlsx(project_dir / "data" / "data.xlsx")
    (project_dir / "config" / "report_config.yaml").write_text("sections: []\n", encoding="utf-8")
    (project_dir / "config" / "prompt_templates.md").write_text("旧 prompt", encoding="utf-8")
    (project_dir / "project.yaml").write_text(
        "\n".join(
            [
                "name: 华安ETF周报",
                "active_word_template: templates/report_template.docx",
                "active_excel_workbook: data/data.xlsx",
                "report_config: config/report_config.yaml",
                "prompt_templates: config/prompt_templates.md",
                "output_dir: generated",
                "run_log_dir: runs",
            ]
        ),
        encoding="utf-8",
    )

    import app.api.routes.report_projects as report_projects_route

    monkeypatch.setattr(
        report_projects_route,
        "report_project_manager",
        ReportProjectManager(projects_root=tmp_path),
    )

    response = client.put(
        "/api/report-projects/华安ETF周报/source",
        json={"source_kind": "prompt_templates", "content": "新 prompt\n{{query}}"},
    )

    assert response.status_code == 200
    data = response.json()
    assert data["prompt_templates_source"] == "新 prompt\n{{query}}"
    assert (project_dir / "config" / "prompt_templates.md").read_text(
        encoding="utf-8"
    ) == "新 prompt\n{{query}}"


def test_render_report_project_writes_to_project_generated_dir(tmp_path: Path, monkeypatch):
    """项目级生成应写入该项目自己的 generated 目录。"""
    project_dir = tmp_path / "创业板50周报"
    (project_dir / "templates").mkdir(parents=True)
    (project_dir / "data").mkdir()
    (project_dir / "config").mkdir()
    (project_dir / "generated").mkdir()
    (project_dir / "runs").mkdir()
    (project_dir / "templates" / "report_template.docx").write_bytes(b"docx")
    (project_dir / "data" / "创业板50周报（iFind版）.xlsx").write_bytes(b"xlsx")
    (project_dir / "config" / "report_config.yaml").write_text("sections: []\n", encoding="utf-8")
    (project_dir / "project.yaml").write_text(
        "\n".join(
            [
                "name: 创业板50周报",
                "active_word_template: templates/report_template.docx",
                "active_excel_workbook: data/创业板50周报（iFind版）.xlsx",
                "report_config: config/report_config.yaml",
                "output_dir: generated",
                "run_log_dir: runs",
            ]
        ),
        encoding="utf-8",
    )

    import app.api.routes.report_projects as report_projects_route

    monkeypatch.setattr(
        report_projects_route,
        "report_project_manager",
        ReportProjectManager(projects_root=tmp_path),
    )

    def fake_save_from_template(output_path, template_path, sections, placeholders, **kwargs):
        output_path.write_bytes(b"rendered")

    with patch(
        "reporting.projections.word.WordProjection.save_from_template",
        side_effect=fake_save_from_template,
    ):
        response = client.post(
            "/api/report-projects/创业板50周报/render",
            json={"placeholders": {"market_summary_placeholder": "测试正文"}},
        )

    assert response.status_code == 200
    data = response.json()
    output_path = Path(data["file_path"])
    assert output_path.parent == project_dir / "generated"
    assert output_path.name.endswith("_创业板50周报.docx")
    assert output_path.read_bytes() == b"rendered"
    assert data["download_url"].startswith("/api/report-projects/创业板50周报/download/")
    assert data["preview_url"].startswith("/api/report-projects/创业板50周报/preview/")


def test_report_project_run_service_renders_word_project(tmp_path: Path):
    """项目运行 module 应独立完成 Word 投影、run log 和 warning 聚合。"""
    project_dir = tmp_path / "华安ETF周报"
    (project_dir / "templates").mkdir(parents=True)
    (project_dir / "data").mkdir()
    (project_dir / "config").mkdir()
    (project_dir / "generated").mkdir()
    (project_dir / "runs").mkdir()
    (project_dir / "templates" / "report_template.docx").write_bytes(b"docx")
    (project_dir / "data" / "data.xlsx").write_bytes(b"xlsx")
    (project_dir / "config" / "report_config.yaml").write_text(
        "placeholders:\n  人工智能:\n    title: 人工智能\n",
        encoding="utf-8",
    )
    (project_dir / "config" / "prompt_templates.md").write_text(
        "## 人工智能\n检索 Query：AI\n\n写作要求：短句",
        encoding="utf-8",
    )
    (project_dir / "project.yaml").write_text(
        "\n".join(
            [
                "name: 华安ETF周报",
                "active_word_template: templates/report_template.docx",
                "active_excel_workbook: data/data.xlsx",
                "report_config: config/report_config.yaml",
                "prompt_templates: config/prompt_templates.md",
                "output_dir: generated",
                "run_log_dir: runs",
            ]
        ),
        encoding="utf-8",
    )
    project = ReportProjectManager(projects_root=tmp_path).get_project("华安ETF周报")

    class FakeGenerationService:
        def generate_placeholders(self, **kwargs):
            assert kwargs["project"] == project
            assert kwargs["manual_placeholders"] == {"manual": "手工值"}
            assert kwargs["report_period"].start_date == "2026-06-01"
            return ReportGenerationResult(
                placeholders={"人工智能": "AI 生成段落", "manual": "手工值"},
                sections=[
                    GeneratedSectionInfo(
                        placeholder="人工智能",
                        title="人工智能",
                        prompt_template="人工智能",
                        retrieval_query="AI",
                        evidence_count=2,
                        model_name="deepseek-chat",
                        provider="deepseek",
                        tokens_used=88,
                        warnings=["段落 warning"],
                        evidence=[
                            EvidenceSnippet(
                                source="ingestion:cls",
                                title="AI 新闻",
                                content="AI 产业链本周有新增证据。",
                                published_at="2026-06-05",
                            )
                        ],
                    )
                ],
                warnings=["生成 warning"],
            )

    class FakeChartService:
        def generate_and_embed(self, **kwargs):
            assert kwargs["project"] == project
            assert kwargs["docx_path"].exists()
            return [
                GeneratedChartInfo(
                    chart_id="chart1",
                    title="图表",
                    workbook="data.xlsx",
                    source_chart="xl/charts/chart1.xml",
                    replace_kind="chart_to_image",
                    point_count=10,
                    warnings=["图表 warning"],
                )
            ]

    def fake_table_builder(**kwargs):
        assert kwargs["project"] == project
        return [], [{"table_id": "calendar", "title": "日历", "warnings": ["表格 warning"]}]

    captured = {}

    def fake_save_from_template(output_path, template_path, sections, placeholders, **kwargs):
        captured["output_path"] = output_path
        captured["placeholders"] = placeholders
        output_path.write_bytes(b"rendered")

    service = ReportProjectRunService(
        generation_service=FakeGenerationService(),
        chart_service=FakeChartService(),
        table_builder=fake_table_builder,
        word_projection_factory=lambda: type(
            "FakeWordProjection",
            (),
            {"save_from_template": staticmethod(fake_save_from_template)},
        )(),
    )

    result = service.execute(
        project=project,
        report_config={"placeholders": {"人工智能": {"title": "人工智能"}}},
        prompt_templates_source="## 人工智能\n检索 Query：AI",
        request=ReportProjectRunRequest(
            placeholders={"manual": "手工值"},
            generate_from_config=True,
            report_date="2026-06-05",
        ),
    )

    assert captured["placeholders"] == {"人工智能": "AI 生成段落", "manual": "手工值"}
    assert result.output_path == captured["output_path"]
    assert result.file_name.endswith("_华安ETF周报.docx")
    assert result.generated_placeholder_count == 2
    assert result.evidence_count == 2
    assert result.warnings == ["生成 warning", "段落 warning", "图表 warning", "表格 warning"]
    assert result.run_log_path.exists()
    run_record = json.loads(result.run_log_path.read_text(encoding="utf-8"))
    assert run_record["generation"]["sections"][0]["evidence"][0]["title"] == "AI 新闻"
    assert run_record["charts"][0]["chart_id"] == "chart1"
    assert run_record["tables"][0]["table_id"] == "calendar"


def test_render_ppt_report_project_writes_pptx_to_generated_dir(tmp_path: Path, monkeypatch):
    """PPT 项目级生成应替换占位符并输出 pptx。"""
    project_dir = tmp_path / "月度PPT"
    (project_dir / "templates").mkdir(parents=True)
    (project_dir / "config").mkdir()
    (project_dir / "generated").mkdir()
    (project_dir / "runs").mkdir()
    write_minimal_pptx(project_dir / "templates" / "report_template.pptx", "{{title}} / {{period}}")
    (project_dir / "config" / "report_config.yaml").write_text(
        "placeholders: {}\n", encoding="utf-8"
    )
    (project_dir / "project.yaml").write_text(
        "\n".join(
            [
                "name: 月度PPT",
                "project_type: ppt",
                "active_ppt_template: templates/report_template.pptx",
                "report_config: config/report_config.yaml",
                "output_dir: generated",
                "run_log_dir: runs",
            ]
        ),
        encoding="utf-8",
    )

    import app.api.routes.report_projects as report_projects_route

    monkeypatch.setattr(
        report_projects_route,
        "report_project_manager",
        ReportProjectManager(projects_root=tmp_path),
    )

    response = client.post(
        "/api/report-projects/月度PPT/render",
        json={
            "generate_from_config": False,
            "placeholders": {"title": "华安ETF月报", "period": "2026年4月"},
        },
    )

    assert response.status_code == 200
    data = response.json()
    output_path = Path(data["file_path"])
    assert output_path.parent == project_dir / "generated"
    assert output_path.name.endswith("_月度PPT.pptx")
    assert data["download_url"].endswith(f"/download/{output_path.name}")
    assert data["preview_url"].endswith(f"/preview/{output_path.name}")
    assert data["generated_placeholder_count"] == 2
    with zipfile.ZipFile(output_path) as archive:
        slide_xml = archive.read("ppt/slides/slide1.xml").decode("utf-8")
    assert "华安ETF月报" in slide_xml
    assert "2026年4月" in slide_xml
    assert "{{title}}" not in slide_xml


def test_generation_service_uses_prompt_query_evidence_and_reporting_model(tmp_path: Path):
    """生成服务应把 Prompt 内置 Query 检索结果交给 reporting 模型路由生成。"""
    project_dir = tmp_path / "华安ETF周报"
    (project_dir / "templates").mkdir(parents=True)
    (project_dir / "data").mkdir()
    (project_dir / "config").mkdir()
    (project_dir / "generated").mkdir()
    (project_dir / "runs").mkdir()
    write_minimal_docx(project_dir / "templates" / "report_template.docx", "{{人工智能}}")
    write_minimal_xlsx(project_dir / "data" / "data.xlsx")
    (project_dir / "config" / "report_config.yaml").write_text(
        "placeholders: {}\n", encoding="utf-8"
    )
    (project_dir / "project.yaml").write_text(
        "\n".join(
            [
                "name: 华安ETF周报",
                "active_word_template: templates/report_template.docx",
                "active_excel_workbook: data/data.xlsx",
                "report_config: config/report_config.yaml",
                "output_dir: generated",
                "run_log_dir: runs",
            ]
        ),
        encoding="utf-8",
    )
    project = ReportProjectManager(projects_root=tmp_path).get_project("华安ETF周报")

    class FakeRetriever:
        def __init__(self):
            self.query = ""

        def retrieve(self, query, *, title, params, lookback_days, limit, report_period=None):
            self.query = query
            assert title == "人工智能"
            assert params == {"param": "人工智能"}
            assert lookback_days == 7
            assert report_period is not None
            assert report_period.start_date
            assert report_period.end_date
            return [
                EvidenceSnippet(
                    source="ingestion:news",
                    title="AI 新闻",
                    content="人工智能产业链本周出现多条政策和产品进展。",
                    published_at="2026-06-01",
                    url="https://example.test/ai",
                )
            ]

    class FakeGateway:
        def __init__(self):
            self.messages = []
            self.task = None

        def chat(self, messages, model=None, temperature=0.7, max_tokens=None, task=None, **kwargs):
            self.messages = messages
            self.task = task
            assert task in {"reporting", "default"}
            assert "人工智能产业链本周出现多条政策和产品进展" in messages[-1]["content"]
            return ModelResponse(
                content="我们根据提供的evidence撰写。可以写：人工智能板块本周围绕政策和产品进展延续活跃。",
                model_name="deepseek-chat",
                provider="deepseek",
                tokens_used=123,
                latency_ms=456,
            )

    retriever = FakeRetriever()
    gateway = FakeGateway()
    service = ReportProjectGenerationService(retriever=retriever, model_gateway=gateway)
    result = service.generate_placeholders(
        project=project,
        report_config={
            "placeholders": {
                "人工智能": {
                    "title": "人工智能",
                    "type": "prompt",
                    "prompt_template": "人工智能",
                    "max_words": 120,
                    "params": {"param": "人工智能"},
                }
            }
        },
        prompt_templates_source=(
            "## 人工智能\n\n"
            "```text\n"
            "检索 Query：请检索本周人工智能相关新闻\n\n"
            "写作要求：控制在 100 字以内，不输出投资建议。\n"
            "```"
        ),
    )

    assert retriever.query == "请检索本周人工智能相关新闻"
    assert gateway.task in {"reporting", "default"}
    assert result.placeholders["人工智能"] == "人工智能板块本周围绕政策和产品进展延续活跃。"
    assert result.sections[0].provider == "deepseek"
    assert result.sections[0].evidence_count == 1


def test_generation_service_renders_structured_generation_constraints(tmp_path: Path):
    """字数、新闻条数和禁用规则应由 report_config 参数进入最终模型消息。"""
    project_dir = tmp_path / "华安ETF周报"
    (project_dir / "templates").mkdir(parents=True)
    (project_dir / "data").mkdir()
    (project_dir / "config").mkdir()
    (project_dir / "generated").mkdir()
    (project_dir / "runs").mkdir()
    write_minimal_docx(project_dir / "templates" / "report_template.docx", "{{A股市场回顾}}")
    write_minimal_xlsx(project_dir / "data" / "data.xlsx")
    (project_dir / "config" / "report_config.yaml").write_text(
        "placeholders: {}\n", encoding="utf-8"
    )
    (project_dir / "project.yaml").write_text(
        "\n".join(
            [
                "name: 华安ETF周报",
                "active_word_template: templates/report_template.docx",
                "active_excel_workbook: data/data.xlsx",
                "report_config: config/report_config.yaml",
                "output_dir: generated",
                "run_log_dir: runs",
            ]
        ),
        encoding="utf-8",
    )
    project = ReportProjectManager(projects_root=tmp_path).get_project("华安ETF周报")

    class FakeRetriever:
        def retrieve(self, query, **kwargs):
            return [
                EvidenceSnippet(
                    source="ingestion:news",
                    title="市场热点",
                    content="CPO、算力、先进封装、机器人和新能源方向均有新闻事实。",
                    published_at="2026-06-05",
                )
            ]

    class FakeGateway:
        def chat(self, messages, model=None, temperature=0.7, max_tokens=None, task=None, **kwargs):
            message = messages[-1]["content"]
            assert "生成约束：" in message
            assert "写作参数：" in message
            assert "目标字数：约 100 字" in message
            assert "最大字数：不超过 150 字" in message
            assert "至少使用 5 条 evidence/news 信息" in message
            assert "不得使用 Wind 数据" in message
            assert "不得使用日度数据" in message
            assert "使用数据或数值时必须说明来源" in message
            assert "生成一段正文，不输出换行符" in message
            assert "禁止出现这些短语：根据文件、据报道、数据显示" in message
            assert "禁止提及这些实体类别：指数名称、公司名称、证券机构、个股名称、ETF名称" in message
            assert "控制在 100-150 字" not in message
            return ModelResponse(
                content="本周市场热点依次为 CPO、算力、先进封装，板块呈现快速轮动特征。",
                model_name="deepseek-chat",
                provider="deepseek",
                tokens_used=88,
                latency_ms=120,
            )

    service = ReportProjectGenerationService(retriever=FakeRetriever(), model_gateway=FakeGateway())

    result = service.generate_placeholders(
        project=project,
        report_config={
            "placeholders": {
                "A股市场回顾": {
                    "title": "A股市场回顾",
                    "type": "prompt",
                    "prompt_template": "A股市场回顾",
                    "target_words": 100,
                    "max_words": 150,
                    "min_news_count": 5,
                    "validators": {
                        "forbid_wind_data": True,
                        "forbid_daily_data": True,
                        "require_source_for_numbers": True,
                        "no_newline": True,
                        "forbidden_phrases": ["根据文件", "据报道", "数据显示"],
                        "forbid_entities": [
                            "指数名称",
                            "公司名称",
                            "证券机构",
                            "个股名称",
                            "ETF名称",
                        ],
                    },
                }
            }
        },
        prompt_templates_source=(
            "## A股市场回顾\n\n"
            "```text\n"
            "检索 Query：请基于上传的全部新闻内容，找出本周市场热点板块和概念。\n\n"
            "写作格式：本周市场热点依次为【列出市场热点板块/概念】。\n"
            "```"
        ),
    )

    assert "本周市场热点依次为 CPO、算力、先进封装" in result.placeholders["A股市场回顾"]


def test_generation_service_fills_report_period_placeholders(tmp_path: Path):
    """报告周期占位符应由生成上下文自动填充，不进入 LLM。"""
    project_dir = tmp_path / "华安ETF周报"
    (project_dir / "templates").mkdir(parents=True)
    (project_dir / "data").mkdir()
    (project_dir / "config").mkdir()
    (project_dir / "generated").mkdir()
    (project_dir / "runs").mkdir()
    write_minimal_docx(
        project_dir / "templates" / "report_template.docx",
        "{{开始日期}} {{结束日期}}",
    )
    write_minimal_xlsx(project_dir / "data" / "data.xlsx")
    (project_dir / "config" / "report_config.yaml").write_text(
        "placeholders: {}\n", encoding="utf-8"
    )
    (project_dir / "project.yaml").write_text(
        "\n".join(
            [
                "name: 华安ETF周报",
                "active_word_template: templates/report_template.docx",
                "active_excel_workbook: data/data.xlsx",
                "report_config: config/report_config.yaml",
                "output_dir: generated",
                "run_log_dir: runs",
            ]
        ),
        encoding="utf-8",
    )
    project = ReportProjectManager(projects_root=tmp_path).get_project("华安ETF周报")

    class FailRetriever:
        def retrieve(self, **kwargs):
            raise AssertionError("report_period placeholders should not retrieve evidence")

    class FailGateway:
        def chat(self, **kwargs):
            raise AssertionError("report_period placeholders should not call model")

    service = ReportProjectGenerationService(retriever=FailRetriever(), model_gateway=FailGateway())

    result = service.generate_placeholders(
        project=project,
        report_config={
            "placeholders": {
                "开始日期": {"title": "开始日期", "type": "report_period", "field": "start_date"},
                "结束日期": {"title": "结束日期", "type": "report_period", "field": "end_date"},
            }
        },
        prompt_templates_source="",
        report_date="2026-06-05",
    )

    assert result.placeholders == {
        "开始日期": "2026-06-01",
        "结束日期": "2026-06-05",
    }


def test_generation_service_handles_output_shape_placeholder_protocol(tmp_path: Path):
    """通用占位符协议应让短字段/固定文案确定性填充，表格和图表不进入 LLM。"""
    project_dir = tmp_path / "通用模板"
    (project_dir / "templates").mkdir(parents=True)
    (project_dir / "data").mkdir()
    (project_dir / "config").mkdir()
    (project_dir / "generated").mkdir()
    (project_dir / "runs").mkdir()
    write_minimal_docx(
        project_dir / "templates" / "report_template.docx",
        "{{开始日期}} {{免责声明}} {{chart_nav}} {{calendar_table}}",
    )
    write_minimal_xlsx(project_dir / "data" / "data.xlsx")
    (project_dir / "config" / "report_config.yaml").write_text(
        "placeholders: {}\n", encoding="utf-8"
    )
    (project_dir / "project.yaml").write_text(
        "\n".join(
            [
                "name: 通用模板",
                "active_word_template: templates/report_template.docx",
                "active_excel_workbook: data/data.xlsx",
                "report_config: config/report_config.yaml",
                "output_dir: generated",
                "run_log_dir: runs",
            ]
        ),
        encoding="utf-8",
    )
    project = ReportProjectManager(projects_root=tmp_path).get_project("通用模板")

    class FailRetriever:
        def retrieve(self, **kwargs):
            raise AssertionError("deterministic placeholders should not retrieve evidence")

    class FailGateway:
        def chat(self, **kwargs):
            raise AssertionError("deterministic placeholders should not call model")

    service = ReportProjectGenerationService(retriever=FailRetriever(), model_gateway=FailGateway())

    result = service.generate_placeholders(
        project=project,
        report_config={
            "placeholders": {
                "开始日期": {
                    "title": "开始日期",
                    "type": "field",
                    "format": "date",
                    "source": {"kind": "report_period", "field": "start_date"},
                },
                "免责声明": {
                    "title": "免责声明",
                    "type": "static_text",
                    "value": "本报告仅供参考。",
                },
                "chart_nav": {
                    "title": "净值走势图",
                    "type": "chart",
                    "source": {"kind": "excel_chart", "workbook": "data.xlsx"},
                },
                "calendar_table": {
                    "title": "全球投资日历",
                    "type": "table",
                    "source": {"kind": "excel_range", "workbook": "data.xlsx"},
                },
            }
        },
        prompt_templates_source="",
        report_date="2026-06-05",
    )

    assert result.placeholders == {
        "开始日期": "2026-06-01",
        "免责声明": "本报告仅供参考。",
    }
    assert result.warnings == []
    assert result.sections == []


def test_generation_service_builds_composite_market_review_from_excel_and_evidence(
    tmp_path: Path,
):
    """A股市场回顾应先用 Excel 真实数据生成固定句，再用 evidence 生成热点句。"""
    project_dir = tmp_path / "华安ETF周报"
    (project_dir / "templates").mkdir(parents=True)
    (project_dir / "data").mkdir()
    (project_dir / "config").mkdir()
    (project_dir / "generated").mkdir()
    (project_dir / "runs").mkdir()
    write_minimal_docx(project_dir / "templates" / "report_template.docx", "{{A股市场回顾}}")
    write_market_review_xlsx(project_dir / "data" / "周报数据.xlsx")
    (project_dir / "config" / "report_config.yaml").write_text(
        "placeholders: {}\n", encoding="utf-8"
    )
    (project_dir / "project.yaml").write_text(
        "\n".join(
            [
                "name: 华安ETF周报",
                "active_word_template: templates/report_template.docx",
                "active_excel_workbook: data/周报数据.xlsx",
                "report_config: config/report_config.yaml",
                "output_dir: generated",
                "run_log_dir: runs",
            ]
        ),
        encoding="utf-8",
    )
    project = ReportProjectManager(projects_root=tmp_path).get_project("华安ETF周报")

    class FakeRetriever:
        def retrieve(self, query, *, title, params, lookback_days, limit, report_period=None):
            assert "市场热点" in query
            assert title == "A股市场回顾"
            assert report_period is not None
            return [
                EvidenceSnippet(
                    source="ingestion:news",
                    title="CPO 新闻",
                    content="CPO、算力租赁和先进封装本周活跃，AI 光模块需求受到关注。",
                    published_at="2026-06-03",
                )
            ]

    class FakeGateway:
        def chat(self, messages, model=None, temperature=0.7, max_tokens=None, task=None, **kwargs):
            prompt = messages[-1]["content"]
            assert "已完成的市场数据开头" in prompt
            assert "续写要求：" in prompt
            assert "请只输出续写正文" in prompt
            assert "沪深300" in prompt
            return ModelResponse(
                content="本周市场热点依次为 CPO、算力租赁、先进封装，板块呈现快速轮动特征。",
                model_name="deepseek-chat",
                provider="deepseek",
                tokens_used=66,
                latency_ms=100,
            )

    service = ReportProjectGenerationService(retriever=FakeRetriever(), model_gateway=FakeGateway())

    result = service.generate_placeholders(
        project=project,
        report_config={
            "placeholders": {
                "A股市场回顾": {
                    "title": "A股市场回顾",
                    "type": "composite_market_review",
                    "prompt_template": "A股市场回顾",
                    "data_source": {
                        "workbook": "周报数据.xlsx",
                        "domestic_sheet": "错误Sheet",
                        "turnover_sheet": "错误Sheet",
                    },
                    "components": [
                        {
                            "name": "市场表现与交易面",
                            "type": "data_template",
                            "source": "excel",
                            "fields": {
                                "market_trend": {
                                    "workbook": "周报数据.xlsx",
                                    "sheet": "国内",
                                    "range": "B2:C6",
                                },
                                "index_performance": {
                                    "workbook": "周报数据.xlsx",
                                    "sheet": "国内",
                                    "range": "B2:C6",
                                },
                                "avg_turnover": {
                                    "workbook": "周报数据.xlsx",
                                    "sheet": "市场成交",
                                    "cell": "B2",
                                },
                                "turnover_trend": {
                                    "workbook": "周报数据.xlsx",
                                    "sheet": "市场成交",
                                    "range": "B2:C2",
                                },
                            },
                        }
                    ],
                    "max_words": 180,
                }
            }
        },
        prompt_templates_source=(
            "## A股市场回顾\n"
            "检索 Query：请基于上传的全部新闻内容，找出本周所有的市场热点板块和概念。\n\n"
            "写作要求：概括本周 A 股市场热点、板块轮动和风格变化。"
        ),
        report_date="2026-06-05",
    )

    content = result.placeholders["A股市场回顾"]
    assert "沪深300涨0.86%" in content
    assert "中证500涨0.41%" in content
    assert "中证1000跌0.04%" in content
    assert "创业板50跌0.42%" in content
    assert "科创50涨2.13%" in content
    assert "本周日均成交额在2.55万亿左右" in content
    assert "较上周放大" in content
    assert "本周市场热点依次为 CPO、算力租赁、先进封装" in content
    assert result.sections[0].prompt_template == "A股市场回顾"
    assert result.sections[0].evidence_count == 1


def test_generation_service_builds_gold_and_oil_reviews_from_excel(tmp_path: Path):
    """黄金/原油市场回顾应直接读取周报数据 Excel 生成固定文本。"""
    project_dir = tmp_path / "华安ETF周报"
    (project_dir / "templates").mkdir(parents=True)
    (project_dir / "data").mkdir()
    (project_dir / "config").mkdir()
    (project_dir / "generated").mkdir()
    (project_dir / "runs").mkdir()
    write_minimal_docx(project_dir / "templates" / "report_template.docx", "{{黄金市场回顾}}{{原油市场回顾}}")
    write_market_review_xlsx(project_dir / "data" / "周报数据.xlsx")
    (project_dir / "config" / "report_config.yaml").write_text(
        "placeholders: {}\n", encoding="utf-8"
    )
    (project_dir / "project.yaml").write_text(
        "\n".join(
            [
                "name: 华安ETF周报",
                "active_word_template: templates/report_template.docx",
                "active_excel_workbook: data/周报数据.xlsx",
                "report_config: config/report_config.yaml",
                "output_dir: generated",
                "run_log_dir: runs",
            ]
        ),
        encoding="utf-8",
    )
    project = ReportProjectManager(projects_root=tmp_path).get_project("华安ETF周报")

    class NoRetriever:
        def retrieve(self, *args, **kwargs):
            raise AssertionError("Excel fixed market reviews should not retrieve evidence")

    class NoGateway:
        def chat(self, *args, **kwargs):
            raise AssertionError("Excel fixed market reviews should not call model")

    service = ReportProjectGenerationService(retriever=NoRetriever(), model_gateway=NoGateway())

    result = service.generate_placeholders(
        project=project,
        report_config={
            "placeholders": {
                "黄金市场回顾": {
                    "title": "黄金市场回顾",
                    "type": "excel_commodity_market_review",
                    "prompt_template": "黄金市场回顾",
                    "data_source": {
                        "kind": "gold",
                        "workbook": "周报数据.xlsx",
                        "sheet": "黄金",
                    },
                },
                "原油市场回顾": {
                    "title": "原油市场回顾",
                    "type": "excel_commodity_market_review",
                    "prompt_template": "原油市场回顾",
                    "data_source": {
                        "kind": "oil",
                        "workbook": "周报数据.xlsx",
                        "sheet": "石油",
                    },
                },
            }
        },
        prompt_templates_source="",
    )

    assert result.placeholders["黄金市场回顾"] == (
        "截止本周，伦敦现货黄金收于4704.74美元/盎司（周环比-2.74%），" "国内AU9999黄金收于1033.25元/克（周环比-1.75%）。"
    )
    assert result.placeholders["原油市场回顾"] == (
        "截至本周，布伦特原油期货周均价为106.01美元/桶，较上周五涨13.59美元/桶；"
        "WTI原油期货周均价为97.0美元/桶，较上周五涨13.0美元/桶。"
        "和周初相比主要油品涨跌幅分别为：布伦特原油（17.29%）、WTI原油（17.45%）。"
    )
    assert [section.provider for section in result.sections] == ["excel", "excel"]


def test_filter_and_rank_evidence_applies_must_any_exclude_and_scores():
    """第一阶段检索配置应过滤明显无关 evidence，并按关键词相关度排序。"""
    snippets = [
        EvidenceSnippet(
            source="ingestion:cls",
            title="世贸组织：全球货物贸易保持韧性",
            content="人工智能相关电子元件需求上升。",
            published_at="2026-06-05",
        ),
        EvidenceSnippet(
            source="ingestion:zq",
            title="商业航天卫星应用政策落地",
            content="商业航天、卫星互联网和运载火箭产业化加速。",
            published_at="2026-06-05",
        ),
        EvidenceSnippet(
            source="ingestion:zq",
            title="私募基金监管新规",
            content="私募基金风险防范。",
            published_at="2026-06-05",
        ),
    ]

    ranked = filter_and_rank_evidence(
        snippets,
        RetrievalConfig(
            must_any=["商业航天", "卫星", "火箭"],
            exclude=["私募基金"],
            min_keyword_score=1.0,
        ),
    )

    assert [item.title for item in ranked] == ["商业航天卫星应用政策落地"]
    assert ranked[0].keyword_score >= 3
    assert ranked[0].matched_terms == ["商业航天", "卫星", "火箭"]


def test_build_retrieval_config_reads_nested_query_terms():
    """工作台生成的 retrieval.query_terms 配置应被生成器正确读取。"""
    config = {
        "retrieval": {
            "mode": "hybrid",
            "top_k": 6,
            "candidate_k": 30,
            "min_keyword_score": 1,
            "fusion": {
                "method": "rrf",
                "keyword_weight": 0.65,
                "semantic_weight": 0.35,
                "rrf_k": 50,
                "semantic_candidate_k": 80,
            },
            "rerank": {
                "enabled": True,
                "provider": "bge-reranker",
                "model": "BAAI/bge-reranker-large",
                "top_n": 12,
                "min_score": 0.4,
            },
            "query_terms": {
                "must_any": ["航天", "商业航天", "卫星"],
                "exclude": ["私募基金"],
            },
        }
    }

    retrieval_config = build_retrieval_config(config, default_top_k=8)

    assert retrieval_config.mode == "hybrid"
    assert retrieval_config.top_k == 6
    assert retrieval_config.candidate_k == 30
    assert retrieval_config.must_any == ["航天", "商业航天", "卫星"]
    assert retrieval_config.exclude == ["私募基金"]
    assert retrieval_config.min_keyword_score == 1
    assert retrieval_config.fusion_method == "rrf"
    assert retrieval_config.keyword_weight == 0.65
    assert retrieval_config.semantic_weight == 0.35
    assert retrieval_config.rrf_k == 50
    assert retrieval_config.semantic_candidate_k == 80
    assert retrieval_config.rerank_enabled is True
    assert retrieval_config.rerank_provider == "bge-reranker"
    assert retrieval_config.rerank_model == "BAAI/bge-reranker-large"
    assert retrieval_config.rerank_top_n == 12
    assert retrieval_config.min_rerank_score == 0.4


def test_build_retrieval_config_prefers_flat_keywords():
    """新工作台配置使用 retrieval.keywords，生成器应直接读取。"""
    config = {
        "retrieval": {
            "keywords": ["电力设备", "新能源", "光伏"],
            "query_terms": {"must_any": ["旧关键词"]},
        }
    }

    retrieval_config = build_retrieval_config(config, default_top_k=8)

    assert retrieval_config.must_any == ["电力设备", "新能源", "光伏"]


def test_hybrid_filter_and_rank_fuses_keyword_and_semantic_scores():
    """Hybrid 模式应记录语义分、融合分和排名，便于调试 evidence 来源。"""
    snippets = [
        EvidenceSnippet(
            source="ingestion:zq",
            title="商业航天卫星发射提速",
            content="商业航天和卫星互联网订单增加。",
            published_at="2026-06-06",
        ),
        EvidenceSnippet(
            source="ingestion:cls",
            title="海南打造火箭链和航天产业体系",
            content="火箭链、卫星链和航天+产业体系加快建设。",
            published_at="2026-06-05",
        ),
        EvidenceSnippet(
            source="ingestion:cls",
            title="消费电子新品发布",
            content="手机和耳机新品发布。",
            published_at="2026-06-05",
        ),
    ]

    ranked = filter_and_rank_evidence(
        snippets,
        RetrievalConfig(
            mode="hybrid",
            must_any=["商业航天", "卫星", "火箭", "航天"],
            semantic_weight=0.4,
            keyword_weight=0.6,
        ),
        query="商业航天 卫星 火箭",
    )

    assert [item.title for item in ranked] == [
        "商业航天卫星发射提速",
        "海南打造火箭链和航天产业体系",
    ]
    assert ranked[0].retrieval_rank == 1
    assert ranked[1].retrieval_rank == 2
    assert ranked[0].retrieval_score is not None
    assert ranked[0].semantic_score is not None
    assert ranked[0].retrieval_method == "hybrid_rrf"


def test_generation_service_reranks_evidence_with_local_bge(tmp_path: Path, monkeypatch):
    """启用本地 rerank 后，应先取更多候选，再按本地模型分数选入最终 prompt。"""
    project_dir = tmp_path / "华安ETF周报"
    (project_dir / "templates").mkdir(parents=True)
    (project_dir / "data").mkdir()
    (project_dir / "config").mkdir()
    (project_dir / "generated").mkdir()
    (project_dir / "runs").mkdir()
    write_minimal_docx(project_dir / "templates" / "report_template.docx", "{{航天}}")
    write_minimal_xlsx(project_dir / "data" / "data.xlsx")
    (project_dir / "config" / "report_config.yaml").write_text(
        "placeholders: {}\n", encoding="utf-8"
    )
    (project_dir / "project.yaml").write_text(
        "\n".join(
            [
                "name: 华安ETF周报",
                "active_word_template: templates/report_template.docx",
                "active_excel_workbook: data/data.xlsx",
                "report_config: config/report_config.yaml",
                "output_dir: generated",
                "run_log_dir: runs",
            ]
        ),
        encoding="utf-8",
    )
    project = ReportProjectManager(projects_root=tmp_path).get_project("华安ETF周报")

    class FakeRetriever:
        def __init__(self):
            self.limit = None

        def retrieve(
            self,
            query,
            *,
            title,
            params,
            lookback_days,
            limit,
            report_period=None,
            retrieval_config=None,
        ):
            self.limit = limit
            return [
                EvidenceSnippet(source="test", title="弱相关", content="行业泛泛而谈"),
                EvidenceSnippet(source="test", title="强相关", content="商业航天卫星火箭发射提速"),
                EvidenceSnippet(source="test", title="一般相关", content="航天产业政策更新"),
            ]

    class FakeGateway:
        def __init__(self):
            self.calls = []

        def chat(self, messages, model=None, temperature=0.7, max_tokens=None, task=None, **kwargs):
            self.calls.append(messages)
            return ModelResponse(
                content="重排后生成正文",
                model_name="deepseek-chat",
                provider="deepseek",
                tokens_used=50,
                latency_ms=1.0,
            )

    retriever = FakeRetriever()
    gateway = FakeGateway()
    service = ReportProjectGenerationService(retriever=retriever, model_gateway=gateway)

    def fake_local_rerank(*, title, query, evidence, retrieval_config, final_limit):
        assert retrieval_config.rerank_provider == "bge-reranker"
        assert retrieval_config.rerank_model == "BAAI/bge-reranker-large"
        return [
            EvidenceSnippet(
                source=evidence[1].source,
                title=evidence[1].title,
                content=evidence[1].content,
                rerank_score=0.96,
                rerank_rank=1,
                rerank_reason="本地 reranker 相关",
            )
        ][:final_limit]

    monkeypatch.setattr(
        "reporting.projects.generation.rerank_evidence_with_local_model",
        fake_local_rerank,
    )

    result = service.generate_placeholders(
        project=project,
        report_config={
            "placeholders": {
                "航天": {
                    "title": "航天",
                    "prompt_template": "航天",
                    "evidence_limit": 1,
                    "retrieval": {
                        "mode": "hybrid",
                        "top_k": 1,
                        "rerank": {
                            "enabled": True,
                            "provider": "bge-reranker",
                            "model": "BAAI/bge-reranker-large",
                            "top_n": 3,
                            "min_score": 0.35,
                        },
                    },
                }
            }
        },
        prompt_templates_source="## 航天\n检索 Query：商业航天 卫星 火箭\n写作要求：简洁",
        report_date="2026-06-05",
    )

    assert retriever.limit == 3
    assert result.placeholders["航天"] == "重排后生成正文"
    assert result.sections[0].evidence_count == 1
    assert result.sections[0].evidence[0].title == "强相关"
    assert result.sections[0].evidence[0].rerank_rank == 1
    assert result.sections[0].evidence[0].rerank_score == 0.96
    assert result.sections[0].evidence[0].rerank_reason == "本地 reranker 相关"
    assert len(gateway.calls) == 1
    assert "证据候选" not in gateway.calls[0][-1]["content"]


def test_generation_service_generates_independent_prompt_sections_concurrently(tmp_path: Path):
    """多个普通 prompt section 应按有界并发生成，并保持输出顺序稳定。"""
    project_dir = tmp_path / "华安ETF周报"
    (project_dir / "templates").mkdir(parents=True)
    (project_dir / "data").mkdir()
    (project_dir / "config").mkdir()
    (project_dir / "generated").mkdir()
    (project_dir / "runs").mkdir()
    write_minimal_docx(project_dir / "templates" / "report_template.docx", "{{人工智能}}")
    write_minimal_xlsx(project_dir / "data" / "data.xlsx")
    (project_dir / "config" / "report_config.yaml").write_text(
        "placeholders: {}\n", encoding="utf-8"
    )
    (project_dir / "project.yaml").write_text(
        "\n".join(
            [
                "name: 华安ETF周报",
                "active_word_template: templates/report_template.docx",
                "active_excel_workbook: data/data.xlsx",
                "report_config: config/report_config.yaml",
                "output_dir: generated",
                "run_log_dir: runs",
            ]
        ),
        encoding="utf-8",
    )
    project = ReportProjectManager(projects_root=tmp_path).get_project("华安ETF周报")

    class FakeRetriever:
        def retrieve(self, query, *, title, params, lookback_days, limit, report_period=None):
            return [EvidenceSnippet(source="test", title=title, content=f"{title} evidence")]

    class ConcurrentGateway:
        def __init__(self):
            self.active = 0
            self.max_active = 0
            self.lock = threading.Lock()

        def chat(self, messages, model=None, temperature=0.7, max_tokens=None, task=None, **kwargs):
            with self.lock:
                self.active += 1
                self.max_active = max(self.max_active, self.active)
            try:
                time.sleep(0.05)
                title = messages[-1]["content"].split("| ", 1)[1].split("\n", 1)[0]
                return ModelResponse(
                    content=f"{title}正文",
                    model_name="deepseek-chat",
                    provider="deepseek",
                    tokens_used=10,
                    latency_ms=50,
                )
            finally:
                with self.lock:
                    self.active -= 1

    gateway = ConcurrentGateway()
    service = ReportProjectGenerationService(
        retriever=FakeRetriever(),
        model_gateway=gateway,
        max_parallel_sections=4,
    )
    placeholders = {
        f"段落{i}": {
            "title": f"段落{i}",
            "type": "prompt",
            "prompt_template": f"段落{i}",
        }
        for i in range(1, 5)
    }

    result = service.generate_placeholders(
        project=project,
        report_config={"placeholders": placeholders},
        prompt_templates_source="\n\n".join(
            f"## 段落{i}\n检索 Query：段落{i}\n\n写作要求：短句" for i in range(1, 5)
        ),
    )

    assert gateway.max_active > 1
    assert list(result.placeholders) == ["段落1", "段落2", "段落3", "段落4"]
    assert [section.placeholder for section in result.sections] == [
        "段落1",
        "段落2",
        "段落3",
        "段落4",
    ]


def test_render_report_project_generates_from_config_and_writes_generation_log(
    tmp_path: Path, monkeypatch
):
    """render 接口应默认从配置生成占位符并记录证据/模型信息。"""
    project_dir = tmp_path / "华安ETF周报"
    (project_dir / "templates").mkdir(parents=True)
    (project_dir / "data").mkdir()
    (project_dir / "config").mkdir()
    (project_dir / "generated").mkdir()
    (project_dir / "runs").mkdir()
    (project_dir / "templates" / "report_template.docx").write_bytes(b"docx")
    (project_dir / "data" / "data.xlsx").write_bytes(b"xlsx")
    write_global_calendar_xlsx(project_dir / "data" / "全球经济日历.xlsx")
    (project_dir / "config" / "report_config.yaml").write_text(
        "\n".join(
            [
                "placeholders:",
                "  人工智能:",
                "    title: 人工智能",
                "    prompt_template: 人工智能",
                "tables:",
                "  global_investment_calendar:",
                "    title: 下周全球投资日历",
                "    enabled: true",
                "    placeholder: 下周全球投资日历",
                "    workbook: 全球经济日历.xlsx",
                "    sheet: 经济数据",
                "    columns:",
                "      - 日期",
                "      - 国家/地区",
                "      - 指标名称",
                "    filter:",
                "      column: 重要性",
                "      equals: 重要",
            ]
        ),
        encoding="utf-8",
    )
    (project_dir / "config" / "prompt_templates.md").write_text(
        "## 人工智能\n检索 Query：AI\n\n写作要求：周报口吻",
        encoding="utf-8",
    )
    (project_dir / "project.yaml").write_text(
        "\n".join(
            [
                "name: 华安ETF周报",
                "active_word_template: templates/report_template.docx",
                "active_excel_workbook: data/data.xlsx",
                "report_config: config/report_config.yaml",
                "prompt_templates: config/prompt_templates.md",
                "output_dir: generated",
                "run_log_dir: runs",
            ]
        ),
        encoding="utf-8",
    )

    import app.api.routes.report_projects as report_projects_route

    monkeypatch.setattr(
        report_projects_route,
        "report_project_manager",
        ReportProjectManager(projects_root=tmp_path),
    )

    class FakeGenerationService:
        def generate_placeholders(self, **kwargs):
            assert kwargs["project"].name == "华安ETF周报"
            assert kwargs["manual_placeholders"] == {}
            assert kwargs["lookback_days"] == 7
            assert kwargs["report_period"].start_date == "2026-06-01"
            assert kwargs["report_period"].end_date == "2026-06-05"
            return ReportGenerationResult(
                placeholders={"人工智能": "AI 生成段落"},
                sections=[
                    GeneratedSectionInfo(
                        placeholder="人工智能",
                        title="人工智能",
                        prompt_template="人工智能",
                        retrieval_query="AI",
                        evidence_count=2,
                        model_name="deepseek-chat",
                        provider="deepseek",
                        tokens_used=88,
                        evidence=[
                            EvidenceSnippet(
                                source="ingestion:cls",
                                title="AI 新闻",
                                content="AI 产业链本周有新增证据。",
                                published_at="2026-06-05",
                                keyword_score=3.0,
                                semantic_score=0.8,
                                retrieval_score=0.02,
                                retrieval_rank=1,
                                retrieval_method="hybrid_rrf",
                                rerank_score=90.0,
                                rerank_rank=1,
                                rerank_reason="相关",
                                matched_terms=["AI"],
                            )
                        ],
                    )
                ],
            )

    monkeypatch.setattr(
        report_projects_route,
        "_get_report_generation_service",
        lambda: FakeGenerationService(),
    )

    class FakeChartService:
        def generate_and_embed(self, **kwargs):
            return [
                GeneratedChartInfo(
                    chart_id="industry_weekly_performance",
                    title="申万一级行业周涨跌幅",
                    workbook="周报图表.xlsx",
                    source_chart="xl/charts/chart2.xml",
                    replace_kind="chart_to_image",
                    point_count=31,
                    warnings=["图表缓存数值全为 0，请确认 Excel/Wind 已刷新并保存"],
                )
            ]

    monkeypatch.setattr(
        report_projects_route,
        "_get_report_chart_service",
        lambda: FakeChartService(),
    )

    captured = {}

    def fake_save_from_template(output_path, template_path, sections, placeholders, **kwargs):
        captured["placeholders"] = placeholders
        captured["tables"] = kwargs["tables"]
        output_path.write_bytes(b"rendered")

    with patch(
        "reporting.projections.word.WordProjection.save_from_template",
        side_effect=fake_save_from_template,
    ):
        response = client.post(
            "/api/report-projects/华安ETF周报/render",
            json={"report_date": "2026-06-05"},
        )

    assert response.status_code == 200
    data = response.json()
    assert captured["placeholders"] == {"人工智能": "AI 生成段落"}
    assert captured["tables"][0].title == "下周全球投资日历"
    assert captured["tables"][0].headers == ["日期", "国家/地区", "指标名称"]
    assert captured["tables"][0].rows == [
        ["2026-06-15", "美国", "6月纽约联储制造业指数"],
        ["2026-06-16", "欧盟", "5月欧元区CPI:同比"],
    ]
    assert data["generated_placeholder_count"] == 1
    assert data["evidence_count"] == 2
    assert data["run_log_url"].startswith("/api/report-projects/华安ETF周报/runs/")
    assert "图表缓存数值全为 0" in data["warnings"][0]
    run_files = list((project_dir / "runs").glob("*.json"))
    assert len(run_files) == 1
    run_record = json.loads(run_files[0].read_text(encoding="utf-8"))
    assert run_record["generation"]["sections"][0]["provider"] == "deepseek"
    assert run_record["generation"]["sections"][0]["retrieval_query"] == "AI"
    assert run_record["generation"]["sections"][0]["evidence"][0] == {
        "source": "ingestion:cls",
        "title": "AI 新闻",
        "content": "AI 产业链本周有新增证据。",
        "published_at": "2026-06-05",
        "url": None,
        "keyword_score": 3.0,
        "semantic_score": 0.8,
        "retrieval_score": 0.02,
        "retrieval_rank": 1,
        "retrieval_method": "hybrid_rrf",
        "rerank_score": 90.0,
        "rerank_rank": 1,
        "rerank_reason": "相关",
        "matched_terms": ["AI"],
    }
    assert run_record["charts"][0]["chart_id"] == "industry_weekly_performance"
    assert run_record["tables"][0] == {
        "table_id": "global_investment_calendar",
        "title": "下周全球投资日历",
        "workbook": "全球经济日历.xlsx",
        "sheet": "经济数据",
        "row_count": 2,
        "warnings": [],
    }
    assert run_record["report_period"] == {
        "start_date": "2026-06-01",
        "end_date": "2026-06-05",
    }

    run_response = client.get(data["run_log_url"])
    assert run_response.status_code == 200
    assert run_response.json()["generation"]["sections"][0]["evidence"][0]["title"] == "AI 新闻"


def _write_preview_project(tmp_path: Path) -> Path:
    project_dir = tmp_path / "华安ETF周报"
    (project_dir / "templates").mkdir(parents=True)
    (project_dir / "data").mkdir()
    (project_dir / "config").mkdir()
    (project_dir / "generated").mkdir()
    (project_dir / "runs").mkdir()
    write_minimal_docx(project_dir / "templates" / "report_template.docx", "{{人工智能}}")
    write_minimal_xlsx(project_dir / "data" / "data.xlsx")
    (project_dir / "config" / "report_config.yaml").write_text("sections: []\n", encoding="utf-8")
    (project_dir / "generated" / "preview.docx").write_bytes(
        (project_dir / "templates" / "report_template.docx").read_bytes()
    )
    (project_dir / "project.yaml").write_text(
        "\n".join(
            [
                "name: 华安ETF周报",
                "active_word_template: templates/report_template.docx",
                "active_excel_workbook: data/data.xlsx",
                "report_config: config/report_config.yaml",
                "output_dir: generated",
                "run_log_dir: runs",
            ]
        ),
        encoding="utf-8",
    )
    return project_dir


def test_preview_report_project_file_prefers_word_pdf_preview(tmp_path: Path, monkeypatch):
    """生成后的 Word 预览应优先使用 Word 导出的 PDF，保留页眉、配色和图表等原始视觉。"""
    _write_preview_project(tmp_path)

    import app.api.routes.report_projects as report_projects_route

    monkeypatch.setattr(
        report_projects_route,
        "report_project_manager",
        ReportProjectManager(projects_root=tmp_path),
    )
    monkeypatch.setattr(
        report_projects_route,
        "_build_word_pdf_preview",
        lambda path: b"%PDF-fake-bytes",
        raising=False,
    )
    monkeypatch.setattr(
        report_projects_route,
        "_render_pdf_preview_pages",
        lambda pdf_bytes: [
            {
                "src": "data:image/png;base64,cGFnZTE=",
                "width": 612,
                "height": 792,
                "label": "第 1 页",
            },
            {
                "src": "data:image/png;base64,cGFnZTI=",
                "width": 612,
                "height": 792,
                "label": "第 2 页",
            },
        ],
        raising=False,
    )
    monkeypatch.setattr(
        report_projects_route,
        "_build_quicklook_preview_image",
        lambda path: b"fake-png-bytes",
        raising=False,
    )

    response = client.get("/api/report-projects/华安ETF周报/preview/preview.docx")

    assert response.status_code == 200
    assert "charset=utf-8" in response.headers["content-type"]
    assert response.headers["cache-control"] == "no-store, max-age=0"
    assert '<meta charset="utf-8">' in response.text
    assert "docx-word-page-preview" in response.text
    assert "docx-preview-stage" in response.text
    assert "data-preview-zoom-in" in response.text
    assert "data-preview-zoom-out" in response.text
    assert "data-preview-zoom-input" in response.text
    assert "docx-preview-tool-icon docx-preview-icon-sidebar" in response.text
    assert "docx-preview-tool-icon docx-preview-icon-minus" in response.text
    assert "docx-preview-tool-icon docx-preview-icon-plus" in response.text
    assert "docx-preview-tool-icon docx-preview-icon-fit-width" in response.text
    assert "docx-preview-tool-icon docx-preview-icon-fit-page" in response.text
    assert ">▤</button>" not in response.text
    assert ">-</button>" not in response.text
    assert ">+</button>" not in response.text
    assert ">↔</button>" not in response.text
    assert ">⤢</button>" not in response.text
    assert "data-preview-zoom-preset" not in response.text
    assert ">50%</button>" not in response.text
    assert ">75%</button>" not in response.text
    assert ">100%</button>" not in response.text
    assert ">125%</button>" not in response.text
    assert "data-preview-fit-width" in response.text
    assert "data-preview-fit-page" in response.text
    assert "data-preview-page-input" in response.text
    assert "data-preview-page-total" in response.text
    assert "data-preview-thumbnails-toggle" in response.text
    assert "data-preview-thumbnails" in response.text
    assert ".docx-preview-stage{grid-row:2;grid-column:2;" in response.text
    assert "localStorage.getItem('af-color-scheme')" in response.text
    assert "parent.document.documentElement" in response.text
    assert "MutationObserver" in response.text
    assert "attributeFilter:['data-theme','data-color-scheme']" in response.text
    assert '[data-color-scheme="obsidian"]' in response.text
    assert "--preview-accent:var(--accent)" in response.text
    assert "--preview-control-bg:" in response.text
    assert "--preview-control-active-bg:" in response.text
    assert (
        "data-preview-thumbnails-toggle][aria-pressed='true']{" "background:var(--preview-accent)"
    ) in response.text
    assert "background:#f5f5f7" not in response.text
    assert "background:#34343a" not in response.text
    assert "#0a84ff" not in response.text
    assert "rgba(10,132,255" not in response.text
    assert "AccentColor" not in response.text
    assert "function loadThumbnails()" in response.text
    assert "localStorage" in response.text
    assert "wheel" in response.text
    assert "data-preview-loading" in response.text
    assert "data-preview-error" in response.text
    assert 'data-preview-layout="double"' in response.text
    assert 'data-preview-page-count="2"' in response.text
    assert 'data-layout="double" data-thumbnails="open"' in response.text
    assert 'aria-pressed="true" data-preview-thumbnails-toggle' in response.text
    assert 'aria-pressed="true" data-preview-layout="double"' in response.text
    assert 'aria-disabled="true" disabled data-preview-layout="double"' not in response.text
    assert "docx-preview-layout-switch" in response.text
    assert "docx-preview-layout-icon docx-preview-layout-single" in response.text
    assert "docx-preview-layout-icon docx-preview-layout-double" in response.text
    assert ">▯</button>" not in response.text
    assert ">▯▯</button>" not in response.text
    assert "docx-preview-file" not in response.text
    assert "data:image/png;base64,cGFnZTE=" in response.text
    assert "data:image/png;base64,cGFnZTI=" in response.text
    assert "Word 原版预览" not in response.text
    assert "docx-native-preview-workspace" not in response.text


def test_word_pdf_preview_render_failure_still_uses_owned_viewer(monkeypatch):
    """PDF 渲染失败时也不能退回浏览器 PDF 插件或旧标题。"""
    import app.api.routes.report_projects as report_projects_route

    monkeypatch.setattr(
        report_projects_route,
        "_render_pdf_preview_pages",
        lambda pdf_bytes: [],
        raising=False,
    )

    html = report_projects_route._word_pdf_preview_html("preview.docx", b"%PDF-fake")

    assert "docx-word-page-preview" in html
    assert "data-preview-zoom-in" in html
    assert "data-preview-zoom-input" in html
    assert 'aria-label="缩放百分比"' in html
    assert "data-preview-fit-width" in html
    assert "data-preview-fit-page" in html
    assert "data-preview-page-input" in html
    assert "data-preview-thumbnails-toggle" in html
    assert "docx-preview-tool-icon docx-preview-icon-sidebar" in html
    assert "docx-preview-tool-icon docx-preview-icon-minus" in html
    assert "docx-preview-tool-icon docx-preview-icon-plus" in html
    assert "docx-preview-tool-icon docx-preview-icon-fit-width" in html
    assert "docx-preview-tool-icon docx-preview-icon-fit-page" in html
    assert ">▤</button>" not in html
    assert ">-</button>" not in html
    assert ">+</button>" not in html
    assert ">↔</button>" not in html
    assert ">⤢</button>" not in html
    assert "data-preview-zoom-preset" not in html
    assert ".docx-preview-stage{grid-row:2;grid-column:2;" in html
    assert "function loadThumbnails()" in html
    assert 'data-preview-layout="double"' in html
    assert "docx-preview-layout-switch" in html
    assert "docx-preview-layout-icon docx-preview-layout-single" in html
    assert "docx-preview-layout-icon docx-preview-layout-double" in html
    assert ">▯</button>" not in html
    assert ">▯▯</button>" not in html
    assert "docx-preview-file" not in html
    assert "Word 原版预览" not in html
    assert "application/pdf" not in html


def test_word_page_preview_can_use_cached_page_asset_urls():
    """性能优化下，预览 HTML 应引用缓存图片 URL，而不是把所有页面 base64 塞进 HTML。"""
    import app.api.routes.report_projects as report_projects_route

    html = report_projects_route._word_page_preview_html(
        "preview.docx",
        [
            {
                "src": "/api/report-projects/华安ETF周报/preview-assets/preview.docx/page-001.png",
                "width": 612,
                "height": 792,
                "label": "第 1 页",
            }
        ],
    )

    assert (
        'data-src="/api/report-projects/华安ETF周报/preview-assets/preview.docx/page-001.png"' in html
    )
    assert 'loading="lazy"' in html
    assert "data:image/png;base64" not in html


def test_rendered_word_page_asset_urls_include_pdf_cache_buster(tmp_path: Path):
    """页图片 URL 应带 PDF hash，避免 WebView 复用旧 LibreOffice 图片缓存。"""
    import app.api.routes.report_projects as report_projects_route

    docx_path = tmp_path / "preview.docx"
    docx_path.write_bytes(b"docx")
    pdf_bytes = b"%PDF-word-preview"
    pdf_hash = hashlib.sha256(pdf_bytes).hexdigest()
    asset_dir = report_projects_route._word_preview_page_asset_dir(docx_path)
    asset_dir.mkdir(parents=True)
    (asset_dir / "page-001.png").write_bytes(b"png")
    (asset_dir / "manifest.json").write_text(
        json.dumps(
            {
                "source": {
                    "mtime_ns": docx_path.stat().st_mtime_ns,
                    "size": docx_path.stat().st_size,
                    "version": report_projects_route.WORD_PREVIEW_LAYOUT_VERSION,
                    "pdf_sha256": pdf_hash,
                },
                "pages": [
                    {
                        "asset": "page-001.png",
                        "width": 1224,
                        "height": 1584,
                        "label": "第 1 页",
                    }
                ],
            }
        ),
        encoding="utf-8",
    )

    pages = report_projects_route._render_pdf_preview_page_assets(
        pdf_bytes,
        docx_path,
        "/api/report-projects/华安ETF周报/preview-assets/preview.docx",
    )

    assert pages[0]["src"].endswith(f"page-001.png?v={pdf_hash[:16]}")


def test_preview_report_project_file_prefers_cached_page_asset_urls(tmp_path: Path, monkeypatch):
    """真实预览接口应优先返回按需加载的页图片 URL，降低首屏 HTML 体积。"""
    _write_preview_project(tmp_path)

    import app.api.routes.report_projects as report_projects_route

    monkeypatch.setattr(
        report_projects_route,
        "report_project_manager",
        ReportProjectManager(projects_root=tmp_path),
    )
    monkeypatch.setattr(
        report_projects_route,
        "_build_word_pdf_preview",
        lambda path: b"%PDF-fake-bytes",
        raising=False,
    )
    monkeypatch.setattr(
        report_projects_route,
        "_render_pdf_preview_page_assets",
        lambda pdf_bytes, source_path, asset_base_url: [
            {
                "src": f"{asset_base_url}/page-001.png",
                "width": 612,
                "height": 792,
                "label": "第 1 页",
            },
            {
                "src": f"{asset_base_url}/page-002.png",
                "width": 612,
                "height": 792,
                "label": "第 2 页",
            },
        ],
        raising=False,
    )

    response = client.get("/api/report-projects/华安ETF周报/preview/preview.docx")

    assert response.status_code == 200
    assert (
        'data-src="/api/report-projects/华安ETF周报/preview-assets/preview.docx/page-001.png"'
        in response.text
    )
    assert (
        'data-src="/api/report-projects/华安ETF周报/preview-assets/preview.docx/page-002.png"'
        in response.text
    )
    assert "data:image/png;base64" not in response.text


def test_build_word_pdf_preview_uses_cached_word_pdf_without_launching_converters(
    tmp_path: Path, monkeypatch
):
    """已有 Word PDF 缓存时，应直接读取，避免重复启动 Word 或 LibreOffice。"""
    import app.api.routes.report_projects as report_projects_route

    docx_path = tmp_path / "preview.docx"
    docx_path.write_bytes(b"docx")
    cache_path = report_projects_route._word_pdf_preview_cache_path(docx_path, engine="word")
    cache_path.parent.mkdir(parents=True)
    cache_path.write_bytes(b"%PDF-cached-word")

    def fail_run(*args, **kwargs):
        raise AssertionError("converter should not be launched for cached previews")

    monkeypatch.setattr(report_projects_route.subprocess, "run", fail_run)
    monkeypatch.setattr(report_projects_route, "_find_soffice_command", lambda: None)

    assert report_projects_route._build_word_pdf_preview(docx_path) == b"%PDF-cached-word"


def test_build_word_pdf_preview_prefers_microsoft_word_export(tmp_path: Path, monkeypatch):
    """macOS 有 Word 时，应优先用 Word 导出，保留原版分页和字体。"""
    import app.api.routes.report_projects as report_projects_route

    docx_path = tmp_path / "preview.docx"
    docx_path.write_bytes(b"docx")
    commands = []

    def fake_run(command, input, check, capture_output, text, timeout):
        commands.append((command, input))
        output_path = Path(re.search(r'set outputPath to "([^"]+)"', input).group(1))
        output_path.write_bytes(b"%PDF-from-word")

        class Result:
            returncode = 0
            stderr = ""

        return Result()

    monkeypatch.setattr(
        report_projects_route,
        "_find_microsoft_word_app",
        lambda: Path("/Applications/Microsoft Word.app"),
    )
    monkeypatch.setattr(report_projects_route, "_find_soffice_command", lambda: "/usr/bin/soffice")
    monkeypatch.setattr(report_projects_route.subprocess, "run", fake_run)

    pdf_bytes = report_projects_route._build_word_pdf_preview(docx_path)

    assert pdf_bytes == b"%PDF-from-word"
    assert commands
    assert commands[0][0] == ["osascript"]
    assert "com.microsoft.Word" in commands[0][1]
    assert (
        report_projects_route._word_pdf_preview_cache_path(docx_path, engine="word").read_bytes()
        == b"%PDF-from-word"
    )


def test_build_word_pdf_preview_uses_local_soffice_converter(tmp_path: Path, monkeypatch):
    """Word 不可用但有 LibreOffice/soffice 时，应本地离线生成 PDF 缓存。"""
    import app.api.routes.report_projects as report_projects_route

    docx_path = tmp_path / "preview.docx"
    docx_path.write_bytes(b"docx")
    commands = []

    def fake_run(command, check, capture_output, text, timeout):
        commands.append(command)
        outdir = Path(command[command.index("--outdir") + 1])
        (outdir / "preview.pdf").write_bytes(b"%PDF-from-soffice")

        class Result:
            returncode = 0
            stderr = ""

        return Result()

    monkeypatch.setattr(report_projects_route, "_find_microsoft_word_app", lambda: None)
    monkeypatch.setattr(report_projects_route, "_find_soffice_command", lambda: "/usr/bin/soffice")
    monkeypatch.setattr(report_projects_route.subprocess, "run", fake_run)

    pdf_bytes = report_projects_route._build_word_pdf_preview(docx_path)

    assert pdf_bytes == b"%PDF-from-soffice"
    assert commands
    assert commands[0][0] == "/usr/bin/soffice"
    assert "--headless" in commands[0]
    assert "--convert-to" in commands[0]
    assert "pdf:writer_pdf_Export" in commands[0]
    assert (
        report_projects_route._word_pdf_preview_cache_path(docx_path, engine="soffice").read_bytes()
        == b"%PDF-from-soffice"
    )


def _write_preview_docx_with_anchored_chart(tmp_path: Path, *, chart_in_caption: bool) -> Path:
    """构造最小 DOCX，避免预览规则测试依赖本机生成的历史报告。"""
    anchor = """
        <wp:anchor distT="0" distB="0" distL="0" distR="0">
          <wp:extent cx="1000" cy="1000"/>
          <wp:docPr id="1" name="Chart 1"/>
          <a:graphic><a:graphicData><c:chart r:id="rId1"/></a:graphicData></a:graphic>
        </wp:anchor>
    """
    drawing = f"<w:r><w:drawing>{anchor}</w:drawing></w:r>"
    caption = "图1：申万一级各板块表现"
    if chart_in_caption:
        body = f"<w:p><w:r><w:t>{caption}</w:t></w:r>{drawing}</w:p>"
    else:
        body = f"<w:p>{drawing}</w:p><w:p><w:r><w:t>{caption}</w:t></w:r></w:p>"
    document_xml = f"""<?xml version="1.0" encoding="UTF-8"?>
    <w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main"
        xmlns:wp="http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
        xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"
        xmlns:c="http://schemas.openxmlformats.org/drawingml/2006/chart"
        xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
      <w:body>{body}<w:sectPr/></w:body>
    </w:document>"""
    source = tmp_path / "anchored-chart.docx"
    with zipfile.ZipFile(source, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", document_xml)
    return source


def test_word_preview_normalizes_anchored_chart_below_caption(tmp_path: Path):
    """历史报告里的浮动图表应在预览副本中移到“图1”标题下方。"""
    import app.api.routes.report_projects as report_projects_route

    source = _write_preview_docx_with_anchored_chart(tmp_path, chart_in_caption=False)
    normalized = tmp_path / source.name

    assert report_projects_route._write_docx_with_inline_preview_charts(source, normalized)

    ns = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
        "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    }
    with zipfile.ZipFile(normalized) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    body = root.find("w:body", ns)
    assert body is not None
    paragraphs = [child for child in list(body) if child.tag.endswith("}p")]

    caption_index = None
    chart_index = None
    for index, paragraph in enumerate(paragraphs):
        text = "".join(node.text or "" for node in paragraph.findall(".//w:t", ns)).strip()
        if text == "图1：申万一级各板块表现":
            caption_index = index
        if paragraph.find(".//c:chart", ns) is not None:
            chart_index = index
            assert paragraph.find(".//wp:anchor", ns) is None
            assert paragraph.find(".//wp:inline", ns) is not None
            break

    assert caption_index is not None
    assert chart_index == caption_index + 1


def test_word_preview_normalizes_chart_embedded_in_caption_paragraph(tmp_path: Path):
    """标题段落自身携带浮动图时，也应拆成标题在前、图表在后。"""
    import app.api.routes.report_projects as report_projects_route

    source = _write_preview_docx_with_anchored_chart(tmp_path, chart_in_caption=True)
    normalized = tmp_path / source.name

    assert report_projects_route._write_docx_with_inline_preview_charts(source, normalized)

    ns = {
        "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main",
        "wp": "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing",
        "c": "http://schemas.openxmlformats.org/drawingml/2006/chart",
    }
    with zipfile.ZipFile(normalized) as archive:
        root = ET.fromstring(archive.read("word/document.xml"))
    body = root.find("w:body", ns)
    assert body is not None
    paragraphs = [child for child in list(body) if child.tag.endswith("}p")]

    caption_index = None
    chart_index = None
    for index, paragraph in enumerate(paragraphs):
        text = "".join(node.text or "" for node in paragraph.findall(".//w:t", ns)).strip()
        if text == "图1：申万一级各板块表现":
            caption_index = index
            assert paragraph.find(".//c:chart", ns) is None
        if paragraph.find(".//c:chart", ns) is not None:
            chart_index = index
            assert paragraph.find(".//wp:anchor", ns) is None
            assert paragraph.find(".//wp:inline", ns) is not None
            break

    assert caption_index is not None
    assert chart_index == caption_index + 1


def test_build_word_pdf_preview_uses_normalized_copy_for_soffice(tmp_path: Path, monkeypatch):
    """LibreOffice 转换应读取预览规范化副本，避免直接解释 Word 浮动锚点。"""
    import app.api.routes.report_projects as report_projects_route

    source = _write_preview_docx_with_anchored_chart(tmp_path, chart_in_caption=False)
    docx_path = tmp_path / source.name
    docx_path.write_bytes(source.read_bytes())
    commands = []

    def fake_run(command, check, capture_output, text, timeout):
        commands.append(command)
        converted_docx = Path(command[-1])
        outdir = Path(command[command.index("--outdir") + 1])
        (outdir / f"{converted_docx.stem}.pdf").write_bytes(b"%PDF-from-normalized-copy")

        class Result:
            returncode = 0
            stderr = ""

        return Result()

    monkeypatch.setattr(report_projects_route, "_find_microsoft_word_app", lambda: None)
    monkeypatch.setattr(report_projects_route, "_find_soffice_command", lambda: "/usr/bin/soffice")
    monkeypatch.setattr(report_projects_route.subprocess, "run", fake_run)

    pdf_bytes = report_projects_route._build_word_pdf_preview(docx_path)

    assert pdf_bytes == b"%PDF-from-normalized-copy"
    assert commands
    assert Path(commands[0][-1]).parent != docx_path.parent
    assert report_projects_route.WORD_PREVIEW_LAYOUT_VERSION in str(
        report_projects_route._word_pdf_preview_cache_path(docx_path)
    )
    assert report_projects_route.WORD_PREVIEW_LAYOUT_VERSION in str(
        report_projects_route._word_preview_page_asset_dir(docx_path)
    )


def test_preview_report_project_file_falls_back_to_system_quicklook_image(
    tmp_path: Path, monkeypatch
):
    """Microsoft Word 不可用时，应退到系统 Quick Look 图。"""
    _write_preview_project(tmp_path)

    import app.api.routes.report_projects as report_projects_route

    monkeypatch.setattr(
        report_projects_route,
        "report_project_manager",
        ReportProjectManager(projects_root=tmp_path),
    )
    monkeypatch.setattr(
        report_projects_route,
        "_build_word_pdf_preview",
        lambda path: None,
        raising=False,
    )
    monkeypatch.setattr(
        report_projects_route,
        "_build_quicklook_preview_image",
        lambda path: b"fake-png-bytes",
        raising=False,
    )

    response = client.get("/api/report-projects/华安ETF周报/preview/preview.docx")

    assert response.status_code == 200
    assert "charset=utf-8" in response.headers["content-type"]
    assert '<meta charset="utf-8">' in response.text
    assert "docx-word-page-preview" in response.text
    assert "docx-preview-stage" in response.text
    assert "docx-native-preview-fit-page" in response.text
    assert "max-height:calc(100vh - 76px)" in response.text
    assert "object-fit:contain" in response.text
    assert "data-preview-zoom-in" in response.text
    assert "data-preview-zoom-out" in response.text
    assert 'data-preview-layout="double"' in response.text
    assert 'data-preview-page-count="1"' in response.text
    assert 'aria-disabled="true" disabled data-preview-layout="double"' in response.text
    assert "data:image/png;base64,ZmFrZS1wbmctYnl0ZXM=" in response.text
    assert "系统缩略预览" in response.text


def test_quicklook_preview_uses_actual_image_dimensions():
    """缩略图适配应使用图片真实尺寸，避免一页看不完整。"""
    import io

    from PIL import Image

    import app.api.routes.report_projects as report_projects_route

    image_io = io.BytesIO()
    Image.new("RGB", (640, 900), "white").save(image_io, format="PNG")

    html = report_projects_route._quicklook_preview_html("preview.docx", image_io.getvalue())

    assert 'data-page-width="640"' in html
    assert 'data-page-height="900"' in html
    assert "data-preview-fit" in html


def test_preview_report_project_file_does_not_launch_word_export_without_cache(
    tmp_path: Path, monkeypatch
):
    """打开预览不应主动启动 Word；无缓存时直接退到系统缩略预览。"""
    _write_preview_project(tmp_path)

    import app.api.routes.report_projects as report_projects_route

    monkeypatch.setattr(
        report_projects_route,
        "report_project_manager",
        ReportProjectManager(projects_root=tmp_path),
    )
    monkeypatch.setattr(
        report_projects_route,
        "_build_word_pdf_preview",
        lambda path: None,
        raising=False,
    )

    monkeypatch.setattr(
        report_projects_route,
        "_find_soffice_command",
        lambda: None,
        raising=False,
    )
    monkeypatch.setattr(
        report_projects_route,
        "_build_quicklook_preview_image",
        lambda path: b"fake-png-bytes",
        raising=False,
    )

    response = client.get("/api/report-projects/华安ETF周报/preview/preview.docx")

    assert response.status_code == 200
    assert "系统缩略预览" in response.text


def test_preview_report_project_file_falls_back_to_docx_html(tmp_path: Path, monkeypatch):
    """Quick Look 不可用时，仍应能返回轻量 HTML 预览。"""
    _write_preview_project(tmp_path)

    import app.api.routes.report_projects as report_projects_route

    monkeypatch.setattr(
        report_projects_route,
        "report_project_manager",
        ReportProjectManager(projects_root=tmp_path),
    )
    monkeypatch.setattr(
        report_projects_route,
        "_build_word_pdf_preview",
        lambda path: None,
        raising=False,
    )
    monkeypatch.setattr(
        report_projects_route,
        "_build_quicklook_preview_image",
        lambda path: None,
        raising=False,
    )

    response = client.get("/api/report-projects/华安ETF周报/preview/preview.docx")

    assert response.status_code == 200
    assert "docx-preview-workspace" in response.text
    assert "docx-preview-page" in response.text
    assert "width: 210mm" in response.text
    assert "min-height: 297mm" in response.text
    assert "box-shadow" in response.text
    assert "{{人工智能}}" in response.text


def test_open_report_project_generated_folder_uses_platform_file_manager(
    tmp_path: Path, monkeypatch
):
    """桌面端应能从报告页面打开项目 generated 文件夹。"""
    project_dir = tmp_path / "华安ETF周报"
    (project_dir / "templates").mkdir(parents=True)
    (project_dir / "data").mkdir()
    (project_dir / "config").mkdir()
    (project_dir / "generated").mkdir()
    (project_dir / "runs").mkdir()
    write_minimal_docx(project_dir / "templates" / "report_template.docx", "{{人工智能}}")
    write_minimal_xlsx(project_dir / "data" / "data.xlsx")
    (project_dir / "config" / "report_config.yaml").write_text("sections: []\n", encoding="utf-8")
    (project_dir / "project.yaml").write_text(
        "\n".join(
            [
                "name: 华安ETF周报",
                "active_word_template: templates/report_template.docx",
                "active_excel_workbook: data/data.xlsx",
                "report_config: config/report_config.yaml",
                "output_dir: generated",
                "run_log_dir: runs",
            ]
        ),
        encoding="utf-8",
    )

    import app.api.routes.report_projects as report_projects_route

    monkeypatch.setattr(
        report_projects_route,
        "report_project_manager",
        ReportProjectManager(projects_root=tmp_path),
    )
    monkeypatch.setattr(report_projects_route.sys, "platform", "darwin")
    opened = {}

    def fake_popen(command):
        opened["command"] = command

    monkeypatch.setattr(report_projects_route.subprocess, "Popen", fake_popen)

    response = client.post("/api/report-projects/华安ETF周报/open-folder")

    assert response.status_code == 200
    assert response.json() == {
        "success": True,
        "folder_path": str(project_dir / "generated"),
    }
    assert opened == {"command": ["open", str(project_dir / "generated")]}


def test_open_report_project_generated_folder_reveals_selected_file_on_macos(
    tmp_path: Path, monkeypatch
):
    """选中某个生成报告时，打开文件夹应在 Finder 中定位这个文件。"""
    project_dir = tmp_path / "华安ETF周报"
    (project_dir / "templates").mkdir(parents=True)
    (project_dir / "data").mkdir()
    (project_dir / "config").mkdir()
    (project_dir / "generated").mkdir()
    (project_dir / "runs").mkdir()
    write_minimal_docx(project_dir / "templates" / "report_template.docx", "{{人工智能}}")
    write_minimal_docx(project_dir / "generated" / "20260605_华安ETF周报.docx", "report")
    write_minimal_xlsx(project_dir / "data" / "data.xlsx")
    (project_dir / "config" / "report_config.yaml").write_text("sections: []\n", encoding="utf-8")
    (project_dir / "project.yaml").write_text(
        "\n".join(
            [
                "name: 华安ETF周报",
                "active_word_template: templates/report_template.docx",
                "active_excel_workbook: data/data.xlsx",
                "report_config: config/report_config.yaml",
                "output_dir: generated",
                "run_log_dir: runs",
            ]
        ),
        encoding="utf-8",
    )

    import app.api.routes.report_projects as report_projects_route

    monkeypatch.setattr(
        report_projects_route,
        "report_project_manager",
        ReportProjectManager(projects_root=tmp_path),
    )
    monkeypatch.setattr(report_projects_route.sys, "platform", "darwin")
    opened = {}

    def fake_popen(command):
        opened["command"] = command

    monkeypatch.setattr(report_projects_route.subprocess, "Popen", fake_popen)

    response = client.post(
        "/api/report-projects/华安ETF周报/open-folder",
        params={"file_name": "20260605_华安ETF周报.docx"},
    )

    selected_file = project_dir / "generated" / "20260605_华安ETF周报.docx"
    assert response.status_code == 200
    assert opened == {"command": ["open", "-R", str(selected_file)]}

    missing_response = client.post(
        "/api/report-projects/华安ETF周报/open-folder",
        params={"file_name": "missing.docx"},
    )

    assert missing_response.status_code == 404
    assert "Generated report not found" in missing_response.json()["detail"]


def test_upload_report_project_package_creates_project_folder(tmp_path: Path, monkeypatch):
    """上传项目包应创建 report_projects 子目录及 project.yaml。"""
    import app.api.routes.report_projects as report_projects_route

    monkeypatch.setattr(
        report_projects_route,
        "report_project_manager",
        ReportProjectManager(projects_root=tmp_path),
    )

    response = client.post(
        "/api/report-projects/upload",
        data={"project_name": "新周报"},
        files=[
            (
                "word_template",
                (
                    "report_template.docx",
                    b"docx",
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ),
            ),
            (
                "excel_workbook",
                (
                    "data.xlsx",
                    b"xlsx",
                    "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ),
            ),
            ("report_config", ("report_config.yaml", b"sections: []\n", "text/yaml")),
            ("prompt_templates", ("prompt_templates.md", b"prompt", "text/markdown")),
            ("data_files", ("domestic.json", b"{}", "application/json")),
        ],
    )

    assert response.status_code == 200
    data = response.json()
    assert data["name"] == "新周报"
    project_dir = tmp_path / "新周报"
    assert (project_dir / "templates" / "report_template.docx").read_bytes() == b"docx"
    assert (project_dir / "data" / "data.xlsx").read_bytes() == b"xlsx"
    assert (project_dir / "config" / "report_config.yaml").read_text(encoding="utf-8")
    assert (project_dir / "config" / "prompt_templates.md").read_bytes() == b"prompt"
    assert (project_dir / "data" / "domestic.json").read_bytes() == b"{}"
    assert "prompt_templates: config/prompt_templates.md" in (
        project_dir / "project.yaml"
    ).read_text(encoding="utf-8")


def test_upload_report_project_allows_word_only_package(tmp_path: Path, monkeypatch):
    """只上传 Word 模板也应创建项目，其它资产后续可补充。"""
    import app.api.routes.report_projects as report_projects_route

    monkeypatch.setattr(
        report_projects_route,
        "report_project_manager",
        ReportProjectManager(projects_root=tmp_path),
    )

    response = client.post(
        "/api/report-projects/upload",
        data={"project_name": "仅Word周报"},
        files=[
            (
                "word_template",
                (
                    "report_template.docx",
                    b"docx",
                    "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ),
            ),
        ],
    )

    assert response.status_code == 200
    data = response.json()
    assert data["name"] == "仅Word周报"
    assert data["word_template_filename"] == "report_template.docx"
    assert data["excel_workbook_filename"] == ""
    assert data["report_config_filename"] == "report_config.yaml"

    project_dir = tmp_path / "仅Word周报"
    assert (project_dir / "templates" / "report_template.docx").read_bytes() == b"docx"
    section_source = (project_dir / "config" / "report_config.yaml").read_text(encoding="utf-8")
    assert "placeholders:" in section_source
    assert "sections: []" in section_source

    project_yaml = (project_dir / "project.yaml").read_text(encoding="utf-8")
    assert "active_word_template: templates/report_template.docx" in project_yaml
    assert "report_config: config/report_config.yaml" in project_yaml
    assert "active_excel_workbook" not in project_yaml


def test_upload_report_project_creates_ppt_project_package(tmp_path: Path, monkeypatch):
    """上传 PPT 模板时应创建 PPT 项目而不是 Word 项目。"""
    import app.api.routes.report_projects as report_projects_route

    monkeypatch.setattr(
        report_projects_route,
        "report_project_manager",
        ReportProjectManager(projects_root=tmp_path),
    )

    response = client.post(
        "/api/report-projects/upload",
        data={"project_name": "静态PPT", "project_type": "ppt"},
        files=[
            (
                "ppt_template",
                (
                    "report_template.pptx",
                    b"pptx",
                    "application/vnd.openxmlformats-officedocument.presentationml.presentation",
                ),
            ),
        ],
    )

    assert response.status_code == 200
    data = response.json()
    assert data["name"] == "静态PPT"
    assert data["project_type"] == "ppt"
    assert data["template_filename"] == "report_template.pptx"
    assert data["ppt_template_filename"] == "report_template.pptx"
    assert data["word_template_filename"] == ""

    project_dir = tmp_path / "静态PPT"
    assert (project_dir / "templates" / "report_template.pptx").read_bytes() == b"pptx"
    section_source = (project_dir / "config" / "report_config.yaml").read_text(encoding="utf-8")
    assert "placeholders:" in section_source
    project_yaml = (project_dir / "project.yaml").read_text(encoding="utf-8")
    assert "project_type: ppt" in project_yaml
    assert "active_ppt_template: templates/report_template.pptx" in project_yaml
    assert "active_word_template" not in project_yaml


def test_rename_report_project_updates_folder_and_yaml(tmp_path: Path, monkeypatch):
    """报告项目改名应同步移动目录并更新 project.yaml。"""
    project_dir = tmp_path / "旧周报"
    (project_dir / "templates").mkdir(parents=True)
    (project_dir / "data").mkdir()
    (project_dir / "config").mkdir()
    (project_dir / "generated").mkdir()
    (project_dir / "runs").mkdir()
    (project_dir / "templates" / "report_template.docx").write_bytes(b"docx")
    (project_dir / "data" / "data.xlsx").write_bytes(b"xlsx")
    (project_dir / "config" / "report_config.yaml").write_text("sections: []\n", encoding="utf-8")
    (project_dir / "project.yaml").write_text(
        "\n".join(
            [
                "name: 旧周报",
                "active_word_template: templates/report_template.docx",
                "active_excel_workbook: data/data.xlsx",
                "report_config: config/report_config.yaml",
                "output_dir: generated",
                "run_log_dir: runs",
            ]
        ),
        encoding="utf-8",
    )

    import app.api.routes.report_projects as report_projects_route

    monkeypatch.setattr(
        report_projects_route,
        "report_project_manager",
        ReportProjectManager(projects_root=tmp_path),
    )

    response = client.patch("/api/report-projects/旧周报", json={"project_name": "新周报"})

    assert response.status_code == 200
    assert response.json()["name"] == "新周报"
    assert not project_dir.exists()
    new_project_dir = tmp_path / "新周报"
    assert (new_project_dir / "templates" / "report_template.docx").exists()
    assert "name: 新周报" in (new_project_dir / "project.yaml").read_text(encoding="utf-8")
