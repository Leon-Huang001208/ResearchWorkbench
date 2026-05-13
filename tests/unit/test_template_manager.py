"""
测试 TemplateManager
"""
from pathlib import Path
from unittest.mock import MagicMock, patch

import pytest
import yaml

from core.contracts import RetrievalProfileType, SectionSpec, TemplateConfig
from reporting.templates.template_manager import TemplateManager


@pytest.fixture
def temp_template_dir(tmp_path: Path) -> Path:
    """临时模板目录"""
    return tmp_path / "templates"


@pytest.fixture
def template_manager(temp_template_dir: Path) -> TemplateManager:
    """TemplateManager 实例"""
    return TemplateManager(templates_dir=temp_template_dir)


@pytest.fixture
def sample_template_config() -> TemplateConfig:
    """示例模板配置"""
    sections = [
        SectionSpec(
            key="executive_summary",
            title="执行摘要",
            target_words=200,
            placeholder="executive_summary",
        ),
        SectionSpec(
            key="market_analysis",
            title="市场分析",
            target_words=300,
            placeholder="market_analysis",
        ),
    ]
    return TemplateConfig(
        name="test_template",
        description="测试模板",
        version="1.0",
        sections=sections,
        default_retrieval_profile=RetrievalProfileType.WEEKLY_REPORT,
    )


class TestTemplateManager:
    """测试 TemplateManager"""

    def test_initialization_creates_directories(
        self, template_manager: TemplateManager
    ):
        """测试初始化创建目录结构"""
        assert template_manager.yaml_dir.exists()
        assert template_manager.docx_dir.exists()
        assert template_manager.pptx_dir.exists()
        assert template_manager.excel_dir.exists()

    def test_save_and_load_template(
        self, template_manager: TemplateManager, sample_template_config: TemplateConfig
    ):
        """测试保存和加载模板"""
        # 保存模板
        template_manager.save_template(sample_template_config)

        # 验证文件已创建
        yaml_path = template_manager.yaml_dir / "test_template.yaml"
        assert yaml_path.exists()

        # 加载模板
        loaded = template_manager.load_template("test_template")
        assert loaded.name == "test_template"
        assert loaded.description == "测试模板"
        assert len(loaded.sections) == 2

    def test_list_templates(
        self, template_manager: TemplateManager, sample_template_config: TemplateConfig
    ):
        """测试列出模板"""
        # 初始为空
        assert template_manager.list_templates() == []

        # 保存模板
        template_manager.save_template(sample_template_config)

        # 列出模板
        templates = template_manager.list_templates()
        assert len(templates) == 1
        assert "test_template" in templates

    def test_delete_template(
        self, template_manager: TemplateManager, sample_template_config: TemplateConfig
    ):
        """测试删除模板"""
        template_manager.save_template(sample_template_config)
        assert "test_template" in template_manager.list_templates()

        template_manager.delete_template("test_template")
        assert "test_template" not in template_manager.list_templates()

    def test_save_template_file(
        self, template_manager: TemplateManager, sample_template_config: TemplateConfig
    ):
        """测试保存模板文件"""
        # 先保存 YAML 配置
        template_manager.save_template(sample_template_config)

        # 保存 DOCX 模板文件
        test_content = b"fake docx content"
        saved_path = template_manager.save_template_file(
            "test_template", "docx", test_content
        )

        assert saved_path.exists()
        assert saved_path.read_bytes() == test_content

    def test_get_template_file_path(
        self, template_manager: TemplateManager, sample_template_config: TemplateConfig
    ):
        """测试获取模板文件路径"""
        # 先保存 YAML 配置
        template_manager.save_template(sample_template_config)

        # 不存在时返回 None
        assert template_manager.get_template_file_path("test_template", "docx") is None

        # 保存文件后返回路径
        test_content = b"fake docx content"
        template_manager.save_template_file("test_template", "docx", test_content)

        path = template_manager.get_template_file_path("test_template", "docx")
        assert path is not None
        assert path.exists()

    def test_delete_template_file(
        self, template_manager: TemplateManager, sample_template_config: TemplateConfig
    ):
        """测试删除模板文件"""
        template_manager.save_template(sample_template_config)
        template_manager.save_template_file("test_template", "docx", b"test")

        assert template_manager.get_template_file_path("test_template", "docx") is not None

        deleted = template_manager.delete_template_file("test_template", "docx")
        assert deleted is True
        assert template_manager.get_template_file_path("test_template", "docx") is None

    def test_extract_placeholders_from_text(self, template_manager: TemplateManager):
        """测试从文本中提取占位符"""
        text = "Hello {{name}}, welcome to {company}!"
        placeholders = template_manager._extract_placeholders_from_text(text)

        assert placeholders == {"name", "company"}

    def test_extract_placeholders_no_overlap(self, template_manager: TemplateManager):
        """测试占位符不重叠提取"""
        text = "{{name}} {name}"
        placeholders = template_manager._extract_placeholders_from_text(text)

        # 应该只提取一次
        assert placeholders == {"name"}

    def test_list_all_template_files(
        self, template_manager: TemplateManager, sample_template_config: TemplateConfig
    ):
        """测试列出所有模板文件"""
        template_manager.save_template(sample_template_config)
        template_manager.save_template_file("test_template", "docx", b"docx")
        template_manager.save_template_file("test_template", "excel", b"excel")

        all_files = template_manager.list_all_template_files()

        assert "test_template" in all_files
        assert all_files["test_template"]["docx"] is not None
        assert all_files["test_template"]["pptx"] is None
        assert all_files["test_template"]["excel"] is not None

    def test_save_template_file_updates_config(
        self, template_manager: TemplateManager, sample_template_config: TemplateConfig
    ):
        """测试保存模板文件更新配置"""
        template_manager.save_template(sample_template_config)

        # 保存前没有 word_template_path
        config_before = template_manager.load_template("test_template")
        assert config_before.word_template_path is None

        # 保存 DOCX 文件
        template_manager.save_template_file("test_template", "docx", b"test")

        # 保存后应该更新了 word_template_path
        config_after = template_manager.load_template("test_template")
        assert config_after.word_template_path is not None
        assert "test_template_template.docx" in config_after.word_template_path

    def test_save_template_file_overwrite(
        self, template_manager: TemplateManager, sample_template_config: TemplateConfig
    ):
        """测试覆盖保存模板文件"""
        template_manager.save_template(sample_template_config)
        template_manager.save_template_file("test_template", "docx", b"v1")

        # 再次保存（不覆盖，应该失败）
        with pytest.raises(FileExistsError):
            template_manager.save_template_file("test_template", "docx", b"v2")

        # 覆盖保存
        saved = template_manager.save_template_file(
            "test_template", "docx", b"v2", overwrite=True
        )
        assert saved.read_bytes() == b"v2"

    def test_create_weekly_report_template(self, template_manager: TemplateManager):
        """测试创建周报模板"""
        template = template_manager.create_weekly_report_template()

        assert template.name == "weekly_report"
        assert len(template.sections) == 5
        assert template.sections[0].key == "market_summary"
        assert template.sections[0].placeholder == "market_summary_placeholder"


class TestPlaceholderDiscovery:
    """测试占位符发现（需要 python-docx 或 python-pptx）"""

    @pytest.fixture
    def docx_available(self) -> bool:
        """检查 python-docx 是否可用"""
        try:
            import docx  # noqa: F401
            return True
        except ImportError:
            return False

    @pytest.fixture
    def pptx_available(self) -> bool:
        """检查 python-pptx 是否可用"""
        try:
            import pptx  # noqa: F401
            return True
        except ImportError:
            return False

    @pytest.mark.skipif(not pytest.importorskip("docx"), reason="python-docx not available")
    def test_discover_placeholders_from_docx(
        self, template_manager: TemplateManager, sample_template_config: TemplateConfig
    ):
        """测试从 DOCX 发现占位符"""
        import docx
        from io import BytesIO

        # 创建一个简单的测试文档
        doc = docx.Document()
        doc.add_paragraph("Hello {{name}}!")
        doc.add_paragraph("Welcome to {company}.")

        # 保存到内存
        doc_stream = BytesIO()
        doc.save(doc_stream)
        doc_content = doc_stream.getvalue()

        # 保存模板
        template_manager.save_template(sample_template_config)
        template_manager.save_template_file("test_template", "docx", doc_content)

        # 发现占位符
        placeholders = template_manager.discover_placeholders_from_docx("test_template")

        assert placeholders == {"name", "company"}

    @pytest.mark.skipif(not pytest.importorskip("pptx"), reason="python-pptx not available")
    def test_discover_placeholders_from_pptx(
        self, template_manager: TemplateManager, sample_template_config: TemplateConfig
    ):
        """测试从 PPTX 发现占位符"""
        import pptx
        from pptx import Presentation
        from io import BytesIO

        # 创建一个简单的测试演示文稿
        prs = Presentation()

        # Title slide with placeholder
        slide_layout = prs.slide_layouts[0]
        slide = prs.slides.add_slide(slide_layout)
        slide.shapes.title.text = "Hello {{name}}!"

        # Content slide
        slide_layout = prs.slide_layouts[1]
        slide = prs.slides.add_slide(slide_layout)
        slide.placeholders[1].text = "Welcome to {company}."

        # 保存到内存
        prs_stream = BytesIO()
        prs.save(prs_stream)
        prs_content = prs_stream.getvalue()

        # 保存模板
        template_manager.save_template(sample_template_config)
        template_manager.save_template_file("test_template", "pptx", prs_content)

        # 发现占位符
        placeholders = template_manager.discover_placeholders_from_pptx("test_template")

        assert placeholders == {"name", "company"}

    def test_discover_placeholders_from_pptx_raises_when_not_found(
        self, template_manager: TemplateManager, sample_template_config: TemplateConfig
    ):
        """测试 PPTX 文件不存在时抛出异常"""
        template_manager.save_template(sample_template_config)

        with pytest.raises(FileNotFoundError):
            template_manager.discover_placeholders_from_pptx("test_template")
