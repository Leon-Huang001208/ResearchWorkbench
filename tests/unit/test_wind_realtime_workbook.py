"""Wind realtime workbook generation and snapshot parsing tests."""


def test_build_realtime_workbook_writes_view_ranges(tmp_path):
    from openpyxl import load_workbook

    from services.wind_realtime_workbook import build_realtime_workbook

    catalog_path = tmp_path / "wind_index_catalog.csv"
    workbook_path = tmp_path / "Research Workbench_Wind_Realtime.xlsx"
    catalog_path.write_text(
        "wind_code,name,family,category,is_active,priority,is_concept,view_key,view_label\n"
        "CI005001.WI,石油石化,citic_l1,中信一级行业,true,100,false,citic_l1,中信一级\n"
        "CI005002.WI,煤炭,citic_l1,中信一级行业,true,99,false,citic_l1,中信一级\n"
        "801010.SI,农林牧渔,sw_l1,申万一级行业,true,90,false,sw_l1,申万一级\n",
        encoding="utf-8",
    )

    build_realtime_workbook(catalog_path, workbook_path)

    workbook = load_workbook(workbook_path, data_only=False, read_only=True)
    assert "ViewRanges" in workbook.sheetnames
    rows = list(workbook["ViewRanges"].iter_rows(values_only=True))
    assert rows[0] == (
        "view_key",
        "view_label",
        "snapshot_start_row",
        "row_count",
        "updated_at",
    )
    assert rows[1][:4] == ("citic_l1", "中信一级", 2, 2)
    assert rows[2][:4] == ("sw_l1", "申万一级", 4, 1)
    assert "RealtimeRaw" in workbook.sheetnames
    assert "Snapshot" in workbook.sheetnames
    assert workbook["RealtimeRaw"]["F2"].value == (
        '=wss("CI005001.WI,CI005002.WI","sec_name,rt_last,rt_pct_chg",' '"cols=3;rows=2")'
    )
    assert workbook["RealtimeRaw"]["F4"].value == (
        '=wss("801010.SI","sec_name,rt_last,rt_pct_chg","cols=3;rows=1")'
    )


def test_build_realtime_workbook_splits_large_view_wss_formulas(tmp_path, monkeypatch):
    from openpyxl import load_workbook

    from services import wind_realtime_workbook as module

    catalog_path = tmp_path / "wind_index_catalog.csv"
    workbook_path = tmp_path / "Research Workbench_Wind_Realtime.xlsx"
    catalog_path.write_text(
        "wind_code,name,family,category,is_active,priority,is_concept,view_key,view_label\n"
        "884001.WI,概念一,wind_concept,热门概念,true,100,true,wind_hot_concept,Wind热门概念\n"
        "884002.WI,概念二,wind_concept,热门概念,true,99,true,wind_hot_concept,Wind热门概念\n"
        "884003.WI,概念三,wind_concept,热门概念,true,98,true,wind_hot_concept,Wind热门概念\n"
        "801010.SI,农林牧渔,sw_l1,申万一级行业,true,90,false,sw_l1,申万一级\n",
        encoding="utf-8",
    )
    monkeypatch.setattr(module, "WIND_WSS_FORMULA_BATCH_SIZE", 2)

    module.build_realtime_workbook(catalog_path, workbook_path)

    workbook = load_workbook(workbook_path, data_only=False, read_only=True)
    assert workbook["ViewRanges"]["A2"].value == "wind_hot_concept"
    assert workbook["ViewRanges"]["C2"].value == 2
    assert workbook["ViewRanges"]["D2"].value == 3
    assert workbook["RealtimeRaw"]["F2"].value == (
        '=wss("884001.WI,884002.WI","sec_name,rt_last,rt_pct_chg",' '"cols=3;rows=2")'
    )
    assert workbook["RealtimeRaw"]["F4"].value == (
        '=wss("884003.WI","sec_name,rt_last,rt_pct_chg","cols=3;rows=1")'
    )
    assert workbook["RealtimeRaw"]["F5"].value == (
        '=wss("801010.SI","sec_name,rt_last,rt_pct_chg","cols=3;rows=1")'
    )

    batch_rows = module._load_batch_formula_rows(workbook_path)
    assert [row_number for row_number, _formula in batch_rows] == [2, 4, 5]


def test_payload_from_snapshot_filters_inactive_catalog_rows(monkeypatch):
    from datetime import UTC, datetime

    from services import wind_realtime_workbook as module

    now = datetime(2026, 7, 1, 9, 30, tzinfo=UTC)
    monkeypatch.setattr(
        module,
        "_active_catalog_codes_for_view",
        lambda view_key: {"8841892.WI"},
    )
    snapshot = module.WorkbookSnapshot(
        rows=(
            module.WorkbookSnapshotRow(
                view_key="wind_hot_concept",
                view_label="Wind热门概念",
                wind_code="884833.WI",
                name="折叠屏指数",
                last=100.0,
                pct_change=7.27,
                is_concept=True,
                source="wind",
                updated_at=now,
                status="ok",
            ),
            module.WorkbookSnapshotRow(
                view_key="wind_hot_concept",
                view_label="Wind热门概念",
                wind_code="8841892.WI",
                name="光芯片指数",
                last=100.0,
                pct_change=7.99,
                is_concept=True,
                source="wind",
                updated_at=now,
                status="ok",
            ),
        ),
        status="ok",
        is_stale=False,
        updated_at=now,
        error_count=0,
        message="",
    )

    payload = module.payload_from_snapshot(
        snapshot,
        view_key="wind_hot_concept",
        limit=30,
        view_label="Wind热门概念",
    )

    assert [item["name"] for item in payload["up"]] == ["光芯片"]
    assert payload["catalog_filtered_count"] == 1


def test_parse_snapshot_rows_skips_wind_fetching_placeholders():
    from datetime import UTC, datetime

    from services import wind_realtime_workbook as module

    snapshot = module.parse_snapshot_rows(
        [
            {
                "view_key": "wind_hot_concept",
                "view_label": "Wind热门概念",
                "wind_code": "8841924.WI",
                "name": "Fetching...",
                "last": 100.0,
                "pct_change": -2.5,
                "is_concept": True,
                "source": "wind",
                "updated_at": "2026-07-02T09:30:00+08:00",
                "status": "ok",
            },
            {
                "view_key": "wind_hot_concept",
                "view_label": "Wind热门概念",
                "wind_code": "8841892.WI",
                "name": "光芯片指数",
                "last": 100.0,
                "pct_change": 7.99,
                "is_concept": True,
                "source": "wind",
                "updated_at": "2026-07-02T09:30:00+08:00",
                "status": "ok",
            },
        ],
        now=datetime(2026, 7, 2, 1, 30, tzinfo=UTC),
    )

    assert [row.name for row in snapshot.rows] == ["光芯片指数"]


def test_resolve_workbook_path_explicit_overrides_env(monkeypatch, tmp_path):
    """显式参数优先于环境变量。"""
    from services import wind_realtime_workbook as module

    explicit = tmp_path / "explicit.xlsx"
    monkeypatch.setenv(module.WIND_WORKBOOK_PATH_ENV, str(tmp_path / "env.xlsx"))
    assert module.resolve_workbook_path(explicit) == explicit.resolve()


def test_resolve_workbook_path_env_override(monkeypatch, tmp_path):
    """无显式参数时，环境变量覆盖默认路径。"""
    from services import wind_realtime_workbook as module

    env_path = tmp_path / "env.xlsx"
    monkeypatch.setenv(module.WIND_WORKBOOK_PATH_ENV, str(env_path))
    assert module.resolve_workbook_path(None) == env_path.resolve()


def test_resolve_workbook_path_defaults_when_no_env(monkeypatch):
    """无显式参数且无环境变量时，回退到平台默认路径。"""
    from services import wind_realtime_workbook as module

    monkeypatch.delenv(module.WIND_WORKBOOK_PATH_ENV, raising=False)
    result = module.resolve_workbook_path(None)
    assert result.name == "Research Workbench_Wind_Realtime.xlsx"
    assert result == module.DEFAULT_WORKBOOK_PATH
