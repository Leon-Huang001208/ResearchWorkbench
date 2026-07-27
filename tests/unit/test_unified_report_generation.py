"""Unified report configuration generation and rendering regression tests."""

from __future__ import annotations

from pathlib import Path

import pytest

from reporting.projects.generation import ReportGenerationResult
from reporting.projects.plan import compile_report_plan
from reporting.projects.project_manager import ReportProject
from reporting.projects.run import ReportProjectRunRequest, ReportProjectRunService
from reporting.projects.unified_config import (
    UnifiedReportConfigError,
    parse_unified_report_config,
)


def test_compile_plan_uses_unified_config_and_markdown_prompt():
    """Preflight consumes the parsed contract while prompts remain Markdown-owned."""
    config = parse_unified_report_config(
        {
            "defaults": {"retrieval": {"top_k": 6}},
            "placeholders": {
                "市场回顾": {
                    "type": "paragraph",
                    "title": "市场回顾",
                    "prompt_template": "market_review",
                    "retrieval": {"keywords": ["A股"]},
                }
            },
        }
    )

    plan = compile_report_plan(
        config,
        "## market_review\n检索 Query：A股\n\n写作要求：仅依据证据撰写。",
        report_date="2026-07-26",
    )

    assert plan.ready is True
    assert plan.placeholders[0].placeholder == "市场回顾"
    assert plan.placeholders[0].prompt_template == "market_review"
    assert plan.placeholders[0].retrieval_config is not None
    assert plan.placeholders[0].retrieval_config.must_any == ["A股"]


def test_run_uses_unified_render_dispatch_for_rich_text(tmp_path: Path):
    """The single Word path applies unified rich-text rendering after generation."""
    from docx import Document

    project_dir = tmp_path / "weekly"
    template_dir = project_dir / "templates"
    output_dir = project_dir / "generated"
    run_log_dir = project_dir / "runs"
    template_dir.mkdir(parents=True)
    output_dir.mkdir()
    run_log_dir.mkdir()
    template_path = template_dir / "weekly.docx"
    document = Document()
    document.add_paragraph("{{市场回顾}}")
    document.save(template_path)

    project = ReportProject(
        name="周报",
        slug="weekly",
        project_dir=project_dir,
        word_template_path=template_path,
        excel_workbook_path=project_dir / "data.xlsx",
        section_config_path=project_dir / "config" / "report_config.yaml",
        output_dir=output_dir,
        run_log_dir=run_log_dir,
    )
    raw_config = {
        "placeholders": {
            "市场回顾": {
                "type": "paragraph",
                "title": "市场回顾",
                "prompt_template": "market_review",
                "retrieval": {"keywords": ["A股"]},
                "rendering": {
                    "runs": [
                        {"text": "市场：", "bold": True},
                        {"is_dynamic": True},
                    ]
                },
            }
        }
    }
    markdown_prompt = "## market_review\n检索 Query：A股\n\n写作要求：仅依据证据撰写。"

    class GenerationService:
        def generate_placeholders(self, **kwargs):
            assert (
                kwargs["section_config"]
                == parse_unified_report_config(raw_config).to_generation_dict()
            )
            assert kwargs["prompt_templates_source"] == markdown_prompt
            return ReportGenerationResult(
                placeholders={"市场回顾": "本周市场震荡上行。"},
                sections=[],
                warnings=[],
            )

    class ChartService:
        def generate_and_embed(self, **kwargs):
            return []

    result = ReportProjectRunService(
        generation_service=GenerationService(),
        chart_service=ChartService(),
        table_builder=lambda **_: ([], []),
    ).execute(
        project=project,
        section_config=raw_config,
        prompt_templates_source=markdown_prompt,
        request=ReportProjectRunRequest(report_date="2026-07-26"),
    )

    rendered = Document(result.output_path)
    runs = rendered.paragraphs[0].runs
    assert [run.text for run in runs] == ["市场：", "本周市场震荡上行。"]
    assert runs[0].bold is True


def test_run_rejects_unmigrated_v2_without_version_fallback(tmp_path: Path):
    """Runtime rejects legacy V2 input instead of selecting a version-specific path."""
    project = ReportProject(
        name="周报",
        slug="weekly",
        project_dir=tmp_path,
        word_template_path=tmp_path / "weekly.docx",
        excel_workbook_path=tmp_path / "data.xlsx",
        section_config_path=tmp_path / "config" / "report_config.yaml",
        output_dir=tmp_path,
        run_log_dir=tmp_path,
    )

    with pytest.raises(UnifiedReportConfigError, match="迁移"):
        ReportProjectRunService().execute(
            project=project,
            section_config={"meta": {}, "template": {}, "placeholders": {}},
            prompt_templates_source="",
            request=ReportProjectRunRequest(generate_from_config=False),
        )


def test_unified_render_dispatch_replaces_rich_tokens_in_body_tables_headers_and_footers(
    tmp_path: Path,
):
    """Rich unified rendering uses one dispatcher across every Word paragraph location."""
    from docx import Document

    project_dir = tmp_path / "weekly"
    template_dir = project_dir / "templates"
    output_dir = project_dir / "generated"
    run_log_dir = project_dir / "runs"
    template_dir.mkdir(parents=True)
    output_dir.mkdir()
    run_log_dir.mkdir()
    template_path = template_dir / "weekly.docx"
    document = Document()
    document.add_paragraph("{{body}}")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "{{table}}"
    document.sections[0].header.paragraphs[0].text = "{{header}}"
    document.sections[0].footer.paragraphs[0].text = "{{footer}}"
    document.save(template_path)

    project = ReportProject(
        name="周报",
        slug="weekly",
        project_dir=project_dir,
        word_template_path=template_path,
        excel_workbook_path=project_dir / "data.xlsx",
        section_config_path=project_dir / "config" / "report_config.yaml",
        output_dir=output_dir,
        run_log_dir=run_log_dir,
    )
    raw_config = {
        "placeholders": {
            key: {
                "type": "paragraph",
                "prompt_template": key,
                "rendering": {"runs": [{"text": f"{key}: ", "bold": True}, {"is_dynamic": True}]},
            }
            for key in ("body", "table", "header", "footer")
        }
    }

    class GenerationService:
        def generate_placeholders(self, **_):
            return ReportGenerationResult(
                placeholders={key: f"{key}-text" for key in raw_config["placeholders"]},
                sections=[],
                warnings=[],
            )

    class ChartService:
        def generate_and_embed(self, **_):
            return []

    result = ReportProjectRunService(
        generation_service=GenerationService(),
        chart_service=ChartService(),
        table_builder=lambda **_: ([], []),
    ).execute(
        project=project,
        section_config=raw_config,
        prompt_templates_source="\n".join(f"## {key}" for key in raw_config["placeholders"]),
        request=ReportProjectRunRequest(),
    )

    rendered = Document(result.output_path)
    assert rendered.paragraphs[0].text == "body: body-text"
    assert rendered.tables[0].cell(0, 0).text == "table: table-text"
    assert rendered.sections[0].header.paragraphs[0].text == "header: header-text"
    assert rendered.sections[0].footer.paragraphs[0].text == "footer: footer-text"


def test_ppt_run_uses_normalized_unified_config(tmp_path: Path):
    """PPT generation receives the same normalized placeholder-only model as Word."""
    template_path = tmp_path / "weekly.pptx"
    template_path.write_bytes(b"template")
    project = ReportProject(
        name="周报",
        slug="weekly",
        project_dir=tmp_path,
        word_template_path=tmp_path / "weekly.docx",
        excel_workbook_path=tmp_path / "data.xlsx",
        section_config_path=tmp_path / "config" / "report_config.yaml",
        output_dir=tmp_path,
        run_log_dir=tmp_path,
        project_type="ppt",
        ppt_template_path=template_path,
    )
    raw_config = {"placeholders": {"title": {"type": "static_text", "value": "周报标题"}}}
    observed: dict[str, object] = {}

    class GenerationService:
        def generate_placeholders(self, **kwargs):
            observed["config"] = kwargs["section_config"]
            return ReportGenerationResult(placeholders={"title": "周报标题"}, sections=[], warnings=[])

    class PptProjection:
        def save_from_template(self, **kwargs):
            observed["placeholders"] = kwargs["placeholders"]
            kwargs["output_path"].write_bytes(b"rendered")
            return type("ProjectionResult", (), {"missing_placeholders": [], "replaced_count": 1})()

    result = ReportProjectRunService(
        generation_service=GenerationService(),
        ppt_projection_factory=PptProjection,
    ).execute(
        project=project,
        section_config=raw_config,
        prompt_templates_source="",
        request=ReportProjectRunRequest(),
    )

    assert observed["config"] == parse_unified_report_config(raw_config).to_generation_dict()
    assert observed["placeholders"] == {"title": "周报标题"}
    assert result.output_path.read_bytes() == b"rendered"


def test_ppt_rejects_unified_rich_rendering_before_projection(tmp_path: Path):
    """PPT refuses unsupported rich rendering instead of silently bypassing it."""
    template_path = tmp_path / "weekly.pptx"
    template_path.write_bytes(b"template")
    project = ReportProject(
        name="周报",
        slug="weekly",
        project_dir=tmp_path,
        word_template_path=tmp_path / "weekly.docx",
        excel_workbook_path=tmp_path / "data.xlsx",
        section_config_path=tmp_path / "config" / "report_config.yaml",
        output_dir=tmp_path,
        run_log_dir=tmp_path,
        project_type="ppt",
        ppt_template_path=template_path,
    )

    with pytest.raises(UnifiedReportConfigError, match="PPT.*rendering.*正文"):
        ReportProjectRunService().execute(
            project=project,
            section_config={
                "placeholders": {"正文": {"rendering": {"runs": [{"is_dynamic": True}]}}}
            },
            prompt_templates_source="",
            request=ReportProjectRunRequest(generate_from_config=False),
        )


def test_runtime_no_longer_exposes_v2_renderer_or_generation_adapter():
    """Legacy V2 objects are not public runtime routes after the contract cutover."""
    import reporting.rendering.template_renderer as template_renderer
    from reporting.projects.generation import ReportProjectGenerationService

    assert not hasattr(template_renderer, "ConfigDrivenTemplateRenderer")
    assert not hasattr(ReportProjectGenerationService, "generate_placeholder_content")
