"""Wind index structure probe workbook tests."""

from datetime import datetime, timezone


def test_build_probe_workbook_writes_formula_catalog_and_results(tmp_path):
    from openpyxl import load_workbook

    from services.wind_index_structure_probe import build_index_structure_probe_workbook

    workbook_path = tmp_path / "AlphaFoundry_Wind_Index_Structure_Probe.xlsx"
    build_index_structure_probe_workbook(
        workbook_path=workbook_path,
        trade_date="2026-06-24",
        index_codes=["000300.SH"],
        etf_codes=["510300.SH"],
    )

    workbook = load_workbook(workbook_path, data_only=False, read_only=True)

    assert workbook.sheetnames == [
        "README",
        "Config",
        "FormulaCatalog",
        "ProbeTargets",
        "ProbeResults",
        "Health",
        "FormulaLog",
    ]
    assert workbook["Config"]["B2"].value == "2026-06-24"

    formula_rows = list(workbook["FormulaCatalog"].iter_rows(values_only=True))
    assert formula_rows[0] == (
        "formula_key",
        "domain",
        "label",
        "template",
        "expected_shape",
        "notes",
    )
    formula_keys = {row[0] for row in formula_rows[1:]}
    assert "index_name_wss" in formula_keys
    assert "index_constituent_wset" in formula_keys
    assert "etf_aum_wss" in formula_keys

    result_rows = list(workbook["ProbeResults"].iter_rows(values_only=True))
    assert result_rows[0] == (
        "probe_id",
        "domain",
        "target_code",
        "formula_key",
        "label",
        "formula_text",
        "value",
        "status",
        "updated_at",
        "notes",
    )
    first_result = result_rows[1]
    assert first_result[:5] == (
        "index:000300.SH:index_name_wss",
        "index",
        "000300.SH",
        "index_name_wss",
        "指数名称 WSS",
    )
    assert first_result[5] == '=@wss("000300.SH","sec_name")'
    assert first_result[6] == '=@wss("000300.SH","sec_name")'
    assert first_result[7].startswith("=IF(OR(ISERROR(G2)")


def test_parse_probe_result_rows_keeps_ok_values_and_counts_errors():
    from services.wind_index_structure_probe import parse_probe_result_rows

    snapshot = parse_probe_result_rows(
        [
            {
                "probe_id": "index:000300.SH:index_name_wss",
                "domain": "index",
                "target_code": "000300.SH",
                "formula_key": "index_name_wss",
                "label": "指数名称 WSS",
                "formula_text": '=@wss("000300.SH","sec_name")',
                "value": "沪深300",
                "status": "ok",
                "updated_at": "2026-06-24T15:10:00+08:00",
                "notes": "",
            },
            {
                "probe_id": "index:000300.SH:index_constituent_wset",
                "domain": "index",
                "target_code": "000300.SH",
                "formula_key": "index_constituent_wset",
                "label": "指数成分 WSET",
                "formula_text": "bad",
                "value": "#VALUE!",
                "status": "formula_error",
                "updated_at": "2026-06-24T15:10:00+08:00",
                "notes": "",
            },
        ],
        now=datetime(2026, 6, 24, 15, 11, tzinfo=timezone.utc),
    )

    assert snapshot.status == "ok"
    assert snapshot.error_count == 1
    assert len(snapshot.rows) == 1
    assert snapshot.rows[0].target_code == "000300.SH"
    assert snapshot.rows[0].value == "沪深300"


def test_parse_probe_result_rows_treats_wind_failure_text_as_error():
    from services.wind_index_structure_probe import parse_probe_result_rows

    snapshot = parse_probe_result_rows(
        [
            {
                "probe_id": "etf:510300.SH:etf_aum_wss",
                "domain": "etf",
                "target_code": "510300.SH",
                "formula_key": "etf_aum_wss",
                "label": "ETF 规模 WSS",
                "formula_text": '=@wss("510300.SH","fund_assetnetvalue")',
                "value": "无法读取数据！",
                "status": "ok",
                "updated_at": "2026-06-24T15:10:00+08:00",
                "notes": "",
            }
        ]
    )

    assert snapshot.status == "all_failed"
    assert snapshot.error_count == 1
    assert snapshot.rows == ()
