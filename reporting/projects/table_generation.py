"""Deterministic table generation for report projects.

This module converts project-owned Excel workbooks into Word ``TableSpec``
objects. It is intentionally separate from LLM section generation: tables such
as economic calendars should be copied from refreshed workbooks, not generated.
"""
from __future__ import annotations

from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Dict, List, Tuple

from core.contracts import TableSpec
from core.observability import get_logger
from reporting.projects.project_manager import ReportProject

logger = get_logger(__name__)


@dataclass(frozen=True)
class GeneratedTableInfo:
    """Run-log metadata for one deterministic project table."""

    table_id: str
    title: str
    workbook: str
    sheet: str
    row_count: int
    warnings: List[str] = field(default_factory=list)

    def to_dict(self) -> Dict[str, Any]:
        """Return a JSON-serializable run-log record."""
        return {
            "table_id": self.table_id,
            "title": self.title,
            "workbook": self.workbook,
            "sheet": self.sheet,
            "row_count": self.row_count,
            "warnings": list(self.warnings),
        }


def build_project_tables(
    *,
    project: ReportProject,
    section_config: Dict[str, Any],
) -> Tuple[List[TableSpec], List[Dict[str, Any]]]:
    """Build configured Word tables from project Excel assets."""
    table_configs = section_config.get("tables")
    if not isinstance(table_configs, dict):
        return [], []

    tables: List[TableSpec] = []
    infos: List[Dict[str, Any]] = []
    for table_id, raw_config in table_configs.items():
        if not isinstance(raw_config, dict) or raw_config.get("enabled") is False:
            continue
        table, info = _build_one_table(project, str(table_id), raw_config)
        if table:
            tables.append(table)
        infos.append(info.to_dict())
    return tables, infos


def _build_one_table(
    project: ReportProject,
    table_id: str,
    config: Dict[str, Any],
) -> Tuple[TableSpec | None, GeneratedTableInfo]:
    workbook_name = str(config.get("workbook") or "")
    sheet_name = str(config.get("sheet") or "")
    title = str(config.get("title") or table_id)
    warnings: List[str] = []
    try:
        workbook_path = _resolve_project_workbook(project, workbook_name)
        headers, rows = _read_excel_table(workbook_path, config)
        table = TableSpec(
            table_id=table_id,
            title=title,
            headers=headers,
            rows=rows,
            placeholder=str(config.get("placeholder") or title),
        )
        if not rows:
            warnings.append(f"表格 {title} 未读取到数据行")
        return table, GeneratedTableInfo(
            table_id=table_id,
            title=title,
            workbook=workbook_name,
            sheet=sheet_name,
            row_count=len(rows),
            warnings=warnings,
        )
    except Exception as exc:
        warning = f"表格 {title} 生成失败: {exc}"
        logger.warning(
            "Failed to build report project table",
            project=project.name,
            table_id=table_id,
            error=str(exc),
        )
        warnings.append(warning)
        return None, GeneratedTableInfo(
            table_id=table_id,
            title=title,
            workbook=workbook_name,
            sheet=sheet_name,
            row_count=0,
            warnings=warnings,
        )


def _resolve_project_workbook(project: ReportProject, workbook_name: str) -> Path:
    if not workbook_name:
        return project.excel_workbook_path
    candidates = [
        project.project_dir / "data" / workbook_name,
        project.project_dir / workbook_name,
        Path(workbook_name),
    ]
    for candidate in candidates:
        if candidate.exists():
            return candidate
    raise FileNotFoundError(f"Workbook not found: {workbook_name}")


def _read_excel_table(
    workbook_path: Path, config: Dict[str, Any]
) -> Tuple[List[str], List[List[Any]]]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - dependency exists in project env
        raise RuntimeError("openpyxl is required for table generation") from exc

    sheet_name = str(config.get("sheet") or "")
    columns = [str(item) for item in config.get("columns") or []]
    if not columns:
        raise ValueError("table columns are required")

    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"Worksheet not found: {sheet_name}")

    sheet = workbook[sheet_name]
    header_row = [
        str(value).strip() if value is not None else ""
        for value in next(sheet.iter_rows(min_row=1, max_row=1, values_only=True))
    ]
    column_indexes = [_require_column(header_row, column) for column in columns]
    filter_config = config.get("filter") if isinstance(config.get("filter"), dict) else {}
    filter_index = None
    filter_value = None
    if filter_config:
        filter_index = _require_column(header_row, str(filter_config.get("column") or ""))
        filter_value = filter_config.get("equals")

    max_rows = int(config.get("max_rows") or 0)
    rows: List[List[Any]] = []
    for excel_row in sheet.iter_rows(min_row=2, values_only=True):
        if filter_index is not None and excel_row[filter_index] != filter_value:
            continue
        row = [_format_cell_value(excel_row[index]) for index in column_indexes]
        if all(value in ("", None) for value in row):
            continue
        rows.append(row)
        if max_rows and len(rows) >= max_rows:
            break
    return columns, rows


def _require_column(headers: List[str], column: str) -> int:
    if column not in headers:
        raise ValueError(f"Column not found: {column}")
    return headers.index(column)


def _format_cell_value(value: Any) -> Any:
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value
