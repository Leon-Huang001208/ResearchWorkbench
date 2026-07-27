"""Unified optional rendering for report project Word output.

The report runtime accepts :class:`UnifiedReportConfig` only.  Flat text,
tables, and project charts are projected first; this module then dispatches
the unified rendering block for rich text, visibility, and chart grids.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any, Dict, Iterable

from core.contracts.reporting import ChartGridSpec, RichTextSpec
from core.observability import get_logger
from reporting.projects.unified_config import UnifiedReportConfig

logger = get_logger(__name__)

try:
    from docx import Document as DocxDocument

    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


class UnifiedTemplateRenderer:
    """Dispatch unified rendering settings across a projected Word document."""

    def __init__(
        self,
        *,
        config: UnifiedReportConfig,
        chart_service: Any = None,
        project_dir: Path | None = None,
    ) -> None:
        if not DOCX_AVAILABLE:
            raise ImportError("python-docx is required for unified Word rendering")
        self._config = config
        self._project_dir = project_dir or Path.cwd()
        from reporting.rendering.chart_grid_injector import ChartGridInjector
        from reporting.rendering.rich_text_injector import RichTextInjector

        self._chart_grid_injector = ChartGridInjector(chart_service=chart_service)
        self._rich_text_injector = RichTextInjector()

    def rendering_placeholders(self) -> set[str]:
        """Return placeholders withheld from the flat projection step."""
        return {
            key
            for key, placeholder in self._config.placeholders.items()
            if placeholder.rendering.configured
        }

    def render(
        self,
        output_path: Path,
        *,
        generated_texts: Dict[str, str],
        context: Dict[str, Any] | None = None,
    ) -> Path:
        """Apply rendering settings to body, table, header, and footer tokens."""
        document = DocxDocument(str(output_path))
        render_context = {**generated_texts, **(context or {})}
        for key, placeholder in self._config.placeholders.items():
            if not placeholder.rendering.configured:
                continue
            try:
                self._dispatch_placeholder(
                    document, key, placeholder, generated_texts, render_context
                )
            except Exception:
                logger.exception("Unified placeholder rendering failed", extra={"key": key})
        document.save(str(output_path))
        logger.info(
            "Unified template rendering complete",
            extra={
                "output": str(output_path),
                "placeholder_count": len(self.rendering_placeholders()),
            },
        )
        return output_path

    def _dispatch_placeholder(
        self,
        document: Any,
        key: str,
        placeholder: Any,
        generated_texts: Dict[str, str],
        context: Dict[str, Any],
    ) -> None:
        rendering = placeholder.rendering
        if not self._is_visible(rendering.visible_if, context):
            self._remove_placeholder(document, key)
            return
        if rendering.chart_grid is not None:
            self._chart_grid_injector.inject_at_placeholder(
                doc=document,
                placeholder_key=key,
                spec=ChartGridSpec(**dict(rendering.chart_grid)),
                project_dir=self._project_dir,
            )
            return

        paragraph = self._find_placeholder(document, key)
        if paragraph is None:
            logger.warning("Unified rendering placeholder not found", extra={"key": key})
            return
        if rendering.paragraph_style:
            try:
                paragraph.style = rendering.paragraph_style
            except Exception:
                logger.warning(
                    "Unified rendering paragraph style unavailable",
                    extra={"key": key, "style": rendering.paragraph_style},
                )
        self._rich_text_injector.inject(
            paragraph=paragraph,
            placeholder_key=key,
            spec=RichTextSpec(**rendering.to_mapping()),
            generated_text=str(generated_texts.get(key) or ""),
        )

    @staticmethod
    def _is_visible(expression: str | None, context: Dict[str, Any]) -> bool:
        if not expression:
            return True
        try:
            import jinja2

            result = jinja2.Template(expression).render(**context).strip().lower()
            return result in {"1", "true", "yes", "on", "是"}
        except Exception:
            logger.exception("Unified rendering visibility evaluation failed")
            return True

    def _remove_placeholder(self, document: Any, key: str) -> None:
        paragraph = self._find_placeholder(document, key)
        if paragraph is not None:
            paragraph._element.getparent().remove(paragraph._element)

    def _find_placeholder(self, document: Any, key: str) -> Any:
        patterns = (f"{{{{{key}}}}}", f"{{{key}}}")
        return next(
            (
                paragraph
                for paragraph in self._iter_document_paragraphs(document)
                if any(pattern in (paragraph.text or "") for pattern in patterns)
            ),
            None,
        )

    def _iter_document_paragraphs(self, document: Any) -> Iterable[Any]:
        yield from self._iter_container_paragraphs(document)
        for section in document.sections:
            yield from self._iter_container_paragraphs(section.header)
            yield from self._iter_container_paragraphs(section.footer)

    def _iter_container_paragraphs(self, container: Any) -> Iterable[Any]:
        yield from container.paragraphs
        for table in container.tables:
            yield from self._iter_table_paragraphs(table)

    def _iter_table_paragraphs(self, table: Any) -> Iterable[Any]:
        for row in table.rows:
            for cell in row.cells:
                yield from self._iter_container_paragraphs(cell)
