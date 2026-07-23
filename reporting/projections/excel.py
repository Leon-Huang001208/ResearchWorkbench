"""
Excel 文档投影 - 将报告输出为 Excel 格式.

Supports data tables, chart generation, and template-based Excel.
"""

from datetime import datetime
from io import BytesIO
from pathlib import Path
from typing import Any, Dict, List, Optional

try:
    import openpyxl as _openpyxl_mod
except ImportError:
    _openpyxl_mod = None

try:
    from PIL import Image as _PIL_Image
except ImportError:
    _PIL_Image = None

from core.contracts import ChartSpec, TableSpec
from core.observability import get_logger

logger = get_logger(__name__)


class ExcelProjection:
    """Excel 报告投影.

    Supports creating Excel reports with tables and charts.
    """

    def __init__(self):
        self._has_openpyxl = self._check_openpyxl()
        self._has_pil = self._check_pil()

    def _check_openpyxl(self) -> bool:
        """检查 openpyxl 是否可用."""
        return _openpyxl_mod is not None

    def _check_pil(self) -> bool:
        """检查 PIL/Pillow 是否可用."""
        return _PIL_Image is not None

    def save(
        self,
        output_path: Path | str,
        title: str,
        tables: Optional[List[TableSpec]] = None,
        charts: Optional[List[ChartSpec]] = None,
        metadata: Optional[Dict[str, Any]] = None,
    ):
        """保存为 Excel 文档.

        Args:
            output_path: Output file path.
            title: Report title.
            tables: List of table specifications.
            charts: List of chart specifications.
            metadata: Additional metadata.
        """
        if not self._has_openpyxl:
            raise ImportError(
                "openpyxl is required for Excel output. "
                "Please install it with: pip install openpyxl"
            )

        from openpyxl import Workbook
        from openpyxl.styles import Font

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = "报告"

        # Title
        ws["A1"] = title
        ws["A1"].font = Font(size=16, bold=True)
        ws.merge_cells("A1:D1")

        # Metadata
        current_row = 3
        if metadata:
            for key, value in metadata.items():
                ws[f"A{current_row}"] = f"{key}:"
                ws[f"A{current_row}"].font = Font(bold=True)
                ws[f"B{current_row}"] = str(value)
                ws.merge_cells(f"B{current_row}:D{current_row}")
                current_row += 1

        # Generation time
        ws[f"A{current_row}"] = "生成时间:"
        ws[f"A{current_row}"].font = Font(bold=True)
        ws[f"B{current_row}"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
        ws.merge_cells(f"B{current_row}:D{current_row}")
        current_row += 2

        # Add tables
        if tables:
            for table_spec in tables:
                current_row = self._add_table_to_worksheet(
                    ws,
                    table_spec,
                    current_row,
                )
                current_row += 2

        # Add charts - note: openpyxl has limited chart support
        if charts:
            logger.warning("Chart generation requires additional libraries, skipping for now")

        wb.save(str(output_path))
        logger.info(f"Saved Excel report to: {output_path}")

    def _add_table_to_worksheet(
        self,
        ws: Any,
        table_spec: TableSpec,
        start_row: int,
    ) -> int:
        """向工作表添加表格.

        Args:
            ws: Worksheet object.
            table_spec: Table specification.
            start_row: Starting row number.

        Returns:
            Next row number after the table.
        """
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

        thin_border = Border(
            left=Side(style="thin"),
            right=Side(style="thin"),
            top=Side(style="thin"),
            bottom=Side(style="thin"),
        )

        header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")

        # Table title
        ws.cell(row=start_row, column=1, value=table_spec.title)
        ws.cell(row=start_row, column=1).font = Font(size=12, bold=True)
        start_row += 1

        if not table_spec.rows:
            return start_row + 1

        # Header row
        if table_spec.headers:
            for col_idx, header in enumerate(table_spec.headers, 1):
                cell = ws.cell(row=start_row, column=col_idx, value=header)
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = header_fill
                cell.alignment = Alignment(horizontal="center")
                cell.border = thin_border
            start_row += 1

        # Data rows
        for row_data in table_spec.rows:
            for col_idx, cell_data in enumerate(row_data, 1):
                cell = ws.cell(row=start_row, column=col_idx, value=cell_data)
                cell.border = thin_border
            start_row += 1

        return start_row

    def save_dataframe(
        self,
        output_path: Path | str,
        df: Any,
        sheet_name: str = "数据",
        title: Optional[str] = None,
    ):
        """保存 DataFrame 到 Excel.

        Args:
            output_path: Output file path.
            df: Pandas DataFrame.
            sheet_name: Name of the sheet.
            title: Optional title for the sheet.
        """
        if not self._has_openpyxl:
            raise ImportError(
                "openpyxl is required for Excel output. "
                "Please install it with: pip install openpyxl"
            )

        from openpyxl import Workbook
        from openpyxl.styles import Font
        from openpyxl.utils.dataframe import dataframe_to_rows

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        wb = Workbook()
        ws = wb.active
        assert ws is not None
        ws.title = sheet_name

        current_row = 1

        if title:
            ws.cell(row=current_row, column=1, value=title)
            ws.cell(row=current_row, column=1).font = Font(size=14, bold=True)
            current_row += 2

        for r_idx, row in enumerate(dataframe_to_rows(df, index=True, header=True), current_row):
            for c_idx, value in enumerate(row, 1):
                ws.cell(row=r_idx, column=c_idx, value=value)

        wb.save(str(output_path))
        logger.info(f"Saved DataFrame to Excel: {output_path}")

    def generate_chart_image(
        self,
        chart_spec: ChartSpec,
        data: List[List[Any]],
        width: int = 600,
        height: int = 400,
    ) -> Optional[bytes]:
        """生成图表图片.

        Note: Requires matplotlib for actual chart generation.

        Args:
            chart_spec: Chart specification.
            data: Chart data.
            width: Image width.
            height: Image height.

        Returns:
            Image bytes if successful, None otherwise.
        """
        try:
            import matplotlib

            matplotlib.use("Agg")
            import matplotlib.pyplot as plt

            fig, ax = plt.subplots(figsize=(width / 100, height / 100))

            chart_type = chart_spec.chart_type
            values = data[1:] if len(data) > 1 else []

            if chart_type == "bar" and values:
                x = list(range(len(values)))
                y = [float(v[1]) if len(v) > 1 else 0.0 for v in values]
                ax.bar(x, y)
                ax.set_xticks(x)
                ax.set_xticklabels([str(v[0]) if len(v) > 0 else "" for v in values])

            elif chart_type == "line" and values:
                x = list(range(len(values)))
                y = [float(v[1]) if len(v) > 1 else 0.0 for v in values]
                ax.plot(x, y, marker="o")
                ax.set_xticks(x)
                ax.set_xticklabels([str(v[0]) if len(v) > 0 else "" for v in values])

            elif chart_type == "pie" and values:
                labels = [str(v[0]) if len(v) > 0 else "" for v in values]
                sizes = [float(v[1]) if len(v) > 1 else 0.0 for v in values]
                ax.pie(sizes, labels=labels, autopct="%1.1f%%")

            elif chart_type == "scatter" and values:
                scatter_x = [float(v[0]) if len(v) > 0 else 0.0 for v in values]
                scatter_y = [float(v[1]) if len(v) > 1 else 0.0 for v in values]
                ax.scatter(scatter_x, scatter_y)

            ax.set_title(chart_spec.title)

            # Save to bytes
            img_buffer = BytesIO()
            plt.tight_layout()
            fig.savefig(img_buffer, format="png", dpi=100)
            plt.close(fig)
            img_buffer.seek(0)

            return img_buffer.getvalue()

        except ImportError as e:
            logger.warning(f"Matplotlib not available: {e}, cannot generate chart image")
            return None
        except Exception as e:
            logger.error(f"Failed to generate chart image: {e}", exc_info=True)
            return None

    def save_from_template(
        self,
        output_path: Path | str,
        template_path: Path | str,
        data_sheets: Optional[Dict[str, List[List[Any]]]] = None,
        chart_specs: Optional[List[ChartSpec]] = None,
        generate_chart_images: bool = False,
    ) -> Dict[str, bytes]:
        """从模板保存 Excel 文档，并可选生成图表图片.

        Args:
            output_path: Output file path.
            template_path: Path to Excel template.
            data_sheets: Dict of sheet name to data.
            chart_specs: List of chart specifications.
            generate_chart_images: Whether to generate chart images for Word embedding.

        Returns:
            Dict of chart ID to image bytes.
        """
        if not self._has_openpyxl:
            raise ImportError(
                "openpyxl is required for Excel output. "
                "Please install it with: pip install openpyxl"
            )

        from openpyxl import load_workbook

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        template_path = Path(template_path)
        if not template_path.exists():
            raise FileNotFoundError(f"Template not found: {template_path}")

        wb = load_workbook(str(template_path), data_only=True)

        # 填充数据到各个工作表
        if data_sheets:
            for sheet_name, data in data_sheets.items():
                if sheet_name in wb.sheetnames:
                    ws = wb[sheet_name]
                else:
                    ws = wb.create_sheet(sheet_name)

                # 清空原有数据
                for row in ws.iter_rows():
                    for cell in row:
                        cell.value = None

                # 写入新数据
                for row_idx, row_data in enumerate(data, 1):
                    for col_idx, cell_data in enumerate(row_data, 1):
                        ws.cell(row=row_idx, column=col_idx, value=cell_data)

        # 保存 Excel 文件
        wb.save(str(output_path))
        logger.info(f"Saved Excel report from template to: {output_path}")

        # 生成图表图片（如果需要）
        chart_images = {}
        if generate_chart_images and chart_specs:
            for chart_spec in chart_specs:
                # 查找图表数据所在的工作表和范围
                if chart_spec.data_range:
                    try:
                        sheet_name, range_str = chart_spec.data_range.split("!", 1)
                        if sheet_name in wb.sheetnames:
                            ws = wb[sheet_name]
                            # 读取范围数据
                            data = []
                            for row in ws[range_str]:
                                row_data = []
                                for cell in row:
                                    row_data.append(cell.value)
                                data.append(row_data)

                            # 生成图表图片
                            img_bytes = self.generate_chart_image(chart_spec, data)
                            if img_bytes:
                                chart_images[chart_spec.chart_id] = img_bytes
                                logger.info(f"Generated chart image: {chart_spec.chart_id}")
                    except Exception as e:
                        logger.error(
                            f"Failed to generate chart {chart_spec.chart_id}: {e}", exc_info=True
                        )

        return chart_images

    def generate_charts_from_data(
        self,
        chart_specs: List[ChartSpec],
        data_sheets: Dict[str, List[List[Any]]],
    ) -> Dict[str, bytes]:
        """从数据直接生成所有图表图片，不需要Excel模板.

        Args:
            chart_specs: List of chart specifications.
            data_sheets: Dict of sheet name to data.

        Returns:
            Dict of chart ID to image bytes.
        """
        chart_images = {}

        for chart_spec in chart_specs:
            try:
                # 解析数据范围
                if chart_spec.data_range:
                    sheet_name, range_str = chart_spec.data_range.split("!", 1)
                    if sheet_name in data_sheets:
                        data = data_sheets[sheet_name]
                        # 简单的范围解析，只支持A1:B5格式
                        import re

                        range_match = re.match(r"([A-Z]+)(\d+):([A-Z]+)(\d+)", range_str)
                        if range_match:
                            start_col, start_row, end_col, end_row = range_match.groups()
                            start_row_idx = int(start_row) - 1
                            end_row_idx = int(end_row)

                            # 提取数据
                            chart_data = data[start_row_idx:end_row_idx]
                            if chart_data:
                                img_bytes = self.generate_chart_image(chart_spec, chart_data)
                                if img_bytes:
                                    chart_images[chart_spec.chart_id] = img_bytes
                                    logger.info(f"Generated chart image: {chart_spec.chart_id}")
            except Exception as e:
                logger.error(f"Failed to generate chart {chart_spec.chart_id}: {e}", exc_info=True)

        return chart_images
