"""Word 文档投影 - 将报告输出为 Word 格式.

Supports both simple Word generation and template-based generation
with placeholder replacement.
"""
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any, Dict, List, Optional

try:
    import docx as _docx_mod  # noqa: F401 – re-exported on demand
except ImportError:
    _docx_mod = None  # type: ignore[assignment]

from core.contracts import SectionOutput, TableSpec
from core.observability import get_logger

logger = get_logger(__name__)


class WordProjection:
    """Word 报告投影.

    Supports both simple Word generation and template-based generation
    with placeholder replacement.
    """

    def __init__(self):
        self._has_docx = self._check_docx()

    def _check_docx(self) -> bool:
        """检查 python-docx 是否可用."""
        return _docx_mod is not None

    def save(
        self,
        output_path: Path | str,
        title: str,
        sections: list[SectionOutput],
        metadata: dict[str, Any] | None = None,
    ):
        """保存为 Word 文档.

        Args:
            output_path: Output file path.
            title: Report title.
            sections: List of section outputs.
            metadata: Additional metadata.
        """
        if not self._has_docx:
            raise ImportError(
                "python-docx is required for Word output. "
                "Please install it with: pip install python-docx"
            )

        from docx import Document
        from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
        from docx.shared import Pt

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        doc = Document()

        # 标题
        title_para = doc.add_heading(title, level=0)
        title_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        # 元数据
        if metadata:
            info_para = doc.add_paragraph()
            for key, value in metadata.items():
                run = info_para.add_run(f"{key}: {value}\n")
                run.font.size = Pt(10)
                run.font.italic = True

        # 生成时间
        time_para = doc.add_paragraph()
        time_run = time_para.add_run(f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        time_run.font.size = Pt(9)
        time_run.font.italic = True

        doc.add_paragraph()

        # 段落内容
        for section in sections:
            # 段落标题
            doc.add_heading(section.key.replace("_", " ").title(), level=2)

            # 段落内容
            content_para = doc.add_paragraph(section.content)
            content_para.paragraph_format.line_spacing = 1.5

            # 证据引用
            if section.evidence_refs:
                ref_para = doc.add_paragraph()
                ref_run = ref_para.add_run("参考文献:")
                ref_run.bold = True
                for ref in section.evidence_refs:
                    doc.add_paragraph(f"[{ref}] 来源待补充", style="List Bullet")

            # 警告
            if section.warnings:
                warning_para = doc.add_paragraph()
                warning_run = warning_para.add_run("⚠️ 警告:")
                warning_run.bold = True
                from docx.shared import RGBColor

                warning_run.font.color.rgb = RGBColor(255, 0, 0)
                for warning in section.warnings:
                    doc.add_paragraph(warning, style="List Bullet")

            doc.add_paragraph()

        doc.save(str(output_path))
        logger.info(f"Saved Word report to: {output_path}")

    def save_review_version(
        self,
        output_path: Path | str,
        title: str,
        sections: list[SectionOutput],
        metadata: dict[str, Any] | None = None,
        include_fact_cards: bool = True,
        include_validation: bool = True,
        include_evidence: bool = True,
    ):
        """保存审稿友好版本的 Word 文档.

        这个版本包含更多细节，方便审稿人检查内容质量：
        - 每个段落展示完整信息
        - 可选择包含事实卡
        - 可选择包含校验结果
        - 可选择包含证据来源

        Args:
            output_path: Output file path.
            title: Report title.
            sections: List of section outputs.
            metadata: Additional metadata.
            include_fact_cards: Whether to include fact cards in the output.
            include_validation: Whether to include validation results.
            include_evidence: Whether to include evidence references.
        """
        if not self._has_docx:
            raise ImportError(
                "python-docx is required for Word output. "
                "Please install it with: pip install python-docx"
            )

        from docx import Document
        from docx.enum.section import WD_SECTION
        from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
        from docx.shared import Pt, RGBColor

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        doc = Document()

        # 标题
        title_para = doc.add_heading(f"【审稿版】{title}", level=0)
        title_para.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER

        # 元数据
        if metadata:
            info_para = doc.add_paragraph()
            for key, value in metadata.items():
                run = info_para.add_run(f"{key}: {value}\n")
                run.font.size = Pt(10)
                run.font.italic = True

        # 生成时间
        time_para = doc.add_paragraph()
        time_run = time_para.add_run(f"生成时间: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}")
        time_run.font.size = Pt(9)
        time_run.font.italic = True

        doc.add_paragraph()

        # 内容概览
        overview_heading = doc.add_heading("内容概览", level=1)
        overview_heading.runs[0].font.color.rgb = RGBColor(0, 51, 102)

        overview_table = doc.add_table(rows=len(sections) + 1, cols=4)
        overview_table.style = "Table Grid"
        overview_table.rows[0].cells[0].text = "段落ID"
        overview_table.rows[0].cells[1].text = "段落标题"
        overview_table.rows[0].cells[2].text = "字数"
        overview_table.rows[0].cells[3].text = "校验状态"

        for idx, section in enumerate(sections):
            row = overview_table.rows[idx + 1]
            row.cells[0].text = section.key
            row.cells[1].text = section.title
            if section.validation_results:
                row.cells[2].text = str(section.validation_results.word_count)
                status = "✅ 通过" if section.validation_results.overall_passed else "❌ 未通过"
                row.cells[3].text = status
            else:
                row.cells[2].text = str(self._count_words(section.content))
                row.cells[3].text = "⚠️ 未校验"

        doc.add_paragraph()

        # 段落详情
        details_heading = doc.add_heading("段落详情", level=1)
        details_heading.runs[0].font.color.rgb = RGBColor(0, 51, 102)
        doc.add_paragraph()

        for section_idx, section in enumerate(sections):
            # 段落标题
            section_heading = doc.add_heading(
                f"{section_idx + 1}. {section.title} (ID: {section.key})", level=2
            )
            section_heading.runs[0].font.color.rgb = RGBColor(0, 102, 204)

            # 段落内容
            content_heading = doc.add_heading("生成内容", level=3)
            content_heading.runs[0].font.size = Pt(12)

            content_para = doc.add_paragraph(section.content)
            content_para.paragraph_format.line_spacing = 1.5
            content_para.paragraph_format.left_indent = Pt(20)

            # 校验结果
            if include_validation and section.validation_results:
                validation_heading = doc.add_heading("校验结果", level=3)
                validation_heading.runs[0].font.size = Pt(12)

                validation_table = doc.add_table(
                    rows=len(section.validation_results.results) + 1, cols=3
                )
                validation_table.style = "Table Grid"
                validation_table.rows[0].cells[0].text = "检查项"
                validation_table.rows[0].cells[1].text = "状态"
                validation_table.rows[0].cells[2].text = "说明"

                for r_idx, result in enumerate(section.validation_results.results):
                    row = validation_table.rows[r_idx + 1]
                    row.cells[0].text = result.check_name
                    status = (
                        "✅ 通过"
                        if result.passed
                        else "❌ 失败"
                        if result.severity == "error"
                        else "⚠️ 警告"
                    )
                    row.cells[1].text = status
                    row.cells[2].text = result.message

                    # 设置颜色
                    if not result.passed:
                        for cell in row.cells:
                            for paragraph in cell.paragraphs:
                                for run in paragraph.runs:
                                    if result.severity == "error":
                                        run.font.color.rgb = RGBColor(255, 0, 0)
                                    else:
                                        run.font.color.rgb = RGBColor(255, 153, 0)

            # 事实卡
            if include_fact_cards and section.fact_card:
                fact_card_heading = doc.add_heading("事实卡", level=3)
                fact_card_heading.runs[0].font.size = Pt(12)

                fact_card = section.fact_card
                fact_content = []

                if fact_card.key_changes:
                    fact_content.append("关键变化：")
                    fact_content.extend([f"- {c}" for c in fact_card.key_changes])
                    fact_content.append("")

                if fact_card.drivers:
                    fact_content.append("驱动因素：")
                    fact_content.extend([f"- {d}" for d in fact_card.drivers])
                    fact_content.append("")

                if fact_card.impacts:
                    fact_content.append("影响分析：")
                    fact_content.extend([f"- {i}" for i in fact_card.impacts])
                    fact_content.append("")

                if fact_card.risks:
                    fact_content.append("风险提示：")
                    fact_content.extend([f"- {r}" for r in fact_card.risks])
                    fact_content.append("")

                if fact_card.watch_points:
                    fact_content.append("观察重点：")
                    fact_content.extend([f"- {w}" for w in fact_card.watch_points])
                    fact_content.append("")

                if fact_card.source_refs:
                    fact_content.append("来源引用：")
                    fact_content.extend([f"- {s}" for s in fact_card.source_refs])

                fact_para = doc.add_paragraph("\n".join(fact_content))
                fact_para.paragraph_format.left_indent = Pt(20)
                fact_para.paragraph_format.line_spacing = 1.2

            # 警告
            if section.warnings:
                warning_heading = doc.add_heading("警告信息", level=3)
                warning_heading.runs[0].font.size = Pt(12)
                warning_heading.runs[0].font.color.rgb = RGBColor(255, 0, 0)

                for warning in section.warnings:
                    warn_para = doc.add_paragraph(f"⚠️ {warning}", style="List Bullet")
                    warn_para.paragraph_format.left_indent = Pt(20)
                    for run in warn_para.runs:
                        run.font.color.rgb = RGBColor(255, 0, 0)

            # 证据引用
            if include_evidence and section.evidence_refs:
                evidence_heading = doc.add_heading("证据引用", level=3)
                evidence_heading.runs[0].font.size = Pt(12)

                ref_para = doc.add_paragraph("本段落使用了以下证据：")
                ref_para.paragraph_format.left_indent = Pt(20)
                for ref in section.evidence_refs:
                    doc.add_paragraph(f"[{ref}] 来源待补充", style="List Bullet")

            # 分页
            if section_idx < len(sections) - 1:
                doc.add_section(WD_SECTION.NEW_PAGE)

        doc.save(str(output_path))
        logger.info(f"Saved review version Word report to: {output_path}")

    def _count_words(self, content: str) -> int:
        """计算中文字数和英文单词数总和."""
        import re

        # Count Chinese characters
        chinese_chars = len(re.findall(r"[一-鿿]", content))
        # Count English words
        english_words = len(re.findall(r"\b[a-zA-Z]+\b", content))
        return chinese_chars + english_words

    def save_from_template(
        self,
        output_path: Path | str,
        template_path: Path | str,
        sections: List[SectionOutput],
        placeholders: Optional[Dict[str, str]] = None,
        tables: Optional[List[TableSpec]] = None,
        chart_images: Optional[Dict[str, bytes]] = None,
    ):
        """从模板保存 Word 文档.

        Replaces placeholders in a Word template with content.

        Args:
            output_path: Output file path.
            template_path: Path to Word template.
            sections: List of section outputs.
            placeholders: Additional placeholder mappings.
            tables: List of table specifications.
            chart_images: Dict of chart ID to image bytes.
        """
        if not self._has_docx:
            raise ImportError(
                "python-docx is required for Word output. "
                "Please install it with: pip install python-docx"
            )

        from docx import Document

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)

        template_path = Path(template_path)
        if not template_path.exists():
            raise FileNotFoundError(f"Template not found: {template_path}")

        doc = Document(str(template_path))

        # Build placeholder map
        placeholder_map = dict(placeholders or {})

        # Add section content to placeholder map - support both key and text_<key> formats
        for section in sections:
            placeholder_map[section.key] = section.content
            placeholder_map[f"text_{section.key}"] = section.content

        # Replace placeholders in paragraphs
        self._replace_placeholders_in_document(doc, placeholder_map)

        # Add tables
        if tables:
            self._add_tables_to_document(doc, tables)

        # Add chart images
        if chart_images:
            self._add_chart_images_to_document(doc, chart_images)

        doc.save(str(output_path))
        logger.info(f"Saved Word report from template to: {output_path}")

    def _replace_placeholders_in_document(
        self,
        doc: Any,
        placeholder_map: Dict[str, str],
    ):
        """在文档中替换占位符.

        Args:
            doc: Word Document object.
            placeholder_map: Mapping of placeholder to replacement text.
        """
        # Replace in paragraphs
        for paragraph in doc.paragraphs:
            self._replace_placeholders_in_paragraph(paragraph, placeholder_map)

        # Replace in tables
        for table in doc.tables:
            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        self._replace_placeholders_in_paragraph(paragraph, placeholder_map)

        # Replace in headers and footers
        for section in doc.sections:
            header = section.header
            for paragraph in header.paragraphs:
                self._replace_placeholders_in_paragraph(paragraph, placeholder_map)

            footer = section.footer
            for paragraph in footer.paragraphs:
                self._replace_placeholders_in_paragraph(paragraph, placeholder_map)

    def _replace_placeholders_in_paragraph(
        self,
        paragraph: Any,
        placeholder_map: Dict[str, str],
    ):
        """在段落中替换占位符.

        Args:
            paragraph: Word Paragraph object.
            placeholder_map: Mapping of placeholder to replacement text.
        """
        replacement_items = sorted(
            placeholder_map.items(),
            key=lambda item: len(str(item[0])),
            reverse=True,
        )
        for placeholder, replacement in replacement_items:
            # Support both {{placeholder}} and {placeholder} formats inline.
            # Bare placeholders are replaced only when the whole paragraph is
            # the placeholder, otherwise common words like "黄金" can corrupt
            # generated content after a longer placeholder was replaced.
            braced_patterns = [
                f"{{{{{placeholder}}}}}",
                f"{{{placeholder}}}",
            ]
            bare_pattern = str(placeholder)
            allow_bare_pattern = self._allows_bare_placeholder_match(bare_pattern)
            patterns = braced_patterns + ([bare_pattern] if allow_bare_pattern else [])

            for pattern in patterns:
                if pattern == bare_pattern and paragraph.text.strip() != bare_pattern.strip():
                    continue
                if pattern in paragraph.text:
                    # Simple replacement - replace entire text
                    if paragraph.text.strip() == pattern.strip():
                        # Clear existing runs and add new one
                        for run in paragraph.runs:
                            run.clear()
                        if paragraph.runs:
                            paragraph.runs[0].text = replacement
                        else:
                            paragraph.add_run(replacement)
                    else:
                        # Replace within text
                        full_text = paragraph.text
                        new_text = full_text.replace(pattern, replacement)
                        if new_text != full_text:
                            # Recreate the paragraph
                            runs = list(paragraph.runs)
                            # Save formatting from first run
                            first_run = runs[0] if runs else None

                            # Clear existing runs
                            for run in runs:
                                run.clear()

                            # Add new text with saved formatting
                            if paragraph.runs:
                                paragraph.runs[0].text = new_text
                            else:
                                new_run = paragraph.add_run(new_text)
                                if first_run:
                                    new_run.bold = first_run.bold
                                    new_run.italic = first_run.italic
                                    new_run.underline = first_run.underline
                                    if first_run.font.size:
                                        new_run.font.size = first_run.font.size
                                    if first_run.font.color and first_run.font.color.rgb:
                                        new_run.font.color.rgb = first_run.font.color.rgb

    @staticmethod
    def _allows_bare_placeholder_match(placeholder: str) -> bool:
        """Return whether a non-braced placeholder can be matched safely.

        Project report templates should use ``{{...}}`` for Chinese section
        placeholders. Bare matching is kept only for legacy ASCII template
        fields such as ``title`` or ``text_summary``; matching common Chinese
        words like "美国" or "黄金" corrupts static tables and product names.
        """
        import re

        return bool(re.fullmatch(r"[A-Za-z_][A-Za-z0-9_:-]*", placeholder))

    def _add_tables_to_document(
        self,
        doc: Any,
        tables: List[TableSpec],
    ):
        """向文档添加表格.

        Args:
            doc: Word Document object.
            tables: List of table specifications.
        """
        for table_spec in tables:
            # Find placeholder or add at end
            if table_spec.placeholder:
                inserted = self._insert_table_at_placeholder(doc, table_spec)
                if inserted:
                    continue
            if table_spec.title:
                replaced = self._replace_table_after_title(doc, table_spec)
                if replaced:
                    continue

            # Add at end as fallback
            doc.add_heading(table_spec.title, level=3)
            self._create_word_table(doc, table_spec)
            doc.add_paragraph()

    def _insert_table_at_placeholder(
        self,
        doc: Any,
        table_spec: TableSpec,
    ) -> bool:
        """在占位符处插入表格.

        Args:
            doc: Word Document object.
            table_spec: Table specification.

        Returns:
            True if inserted successfully.
        """
        # Use normalized placeholder name
        placeholder = table_spec.normalized_placeholder
        patterns = [f"{{{{{placeholder}}}}}", f"{{{placeholder}}}"]
        if self._allows_bare_placeholder_match(placeholder):
            patterns.append(placeholder)

        for para_idx, paragraph in enumerate(doc.paragraphs):
            for pattern in patterns:
                if pattern in paragraph.text:
                    # Found placeholder
                    table = self._create_word_table(doc, table_spec)
                    paragraph._p.addnext(table._tbl)
                    paragraph._element.getparent().remove(paragraph._element)
                    return True

        return False

    def _replace_table_after_title(self, doc: Any, table_spec: TableSpec) -> bool:
        """Replace the first table after a title paragraph with generated rows."""
        title = table_spec.title.strip()
        if not title:
            return False

        body = doc.element.body
        children = list(body)
        for index, child in enumerate(children):
            if child.tag.rsplit("}", 1)[-1] != "p":
                continue
            paragraph_text = "".join(child.itertext()).strip()
            if title not in paragraph_text:
                continue

            table = self._create_word_table(doc, table_spec)
            for next_child in children[index + 1 :]:
                if next_child.tag.rsplit("}", 1)[-1] == "tbl":
                    self._copy_table_format(next_child, table._tbl)
                    body.replace(next_child, table._tbl)
                    return True
                if next_child.tag.rsplit("}", 1)[-1] == "p" and "".join(next_child.itertext()).strip():
                    child.addnext(table._tbl)
                    return True
            child.addnext(table._tbl)
            return True
        return False

    def _create_word_table(self, doc: Any, table_spec: TableSpec) -> Any:
        """Create and populate a Word table for a table spec."""
        num_rows = len(table_spec.rows) + (1 if table_spec.headers else 0)
        num_cols = (
            len(table_spec.headers)
            if table_spec.headers
            else (len(table_spec.rows[0]) if table_spec.rows else 1)
        )
        table = doc.add_table(rows=max(num_rows, 1), cols=max(num_cols, 1))
        table.style = "Table Grid"

        row_offset = 0
        if table_spec.headers:
            for idx, header in enumerate(table_spec.headers):
                table.rows[0].cells[idx].text = str(header)
            row_offset = 1

        for row_idx, row_data in enumerate(table_spec.rows):
            for col_idx, cell_data in enumerate(row_data[:num_cols]):
                table.rows[row_idx + row_offset].cells[col_idx].text = str(cell_data)
        return table

    def _copy_table_format(self, source_tbl: Any, target_tbl: Any) -> None:
        """Copy reusable table layout/style XML from an existing template table."""
        self._replace_child_by_local_name(
            target_tbl,
            "tblPr",
            self._child_by_local_name(source_tbl, "tblPr"),
        )
        self._replace_child_by_local_name(
            target_tbl,
            "tblGrid",
            self._child_by_local_name(source_tbl, "tblGrid"),
        )

        source_rows = [child for child in source_tbl if child.tag.rsplit("}", 1)[-1] == "tr"]
        target_rows = [child for child in target_tbl if child.tag.rsplit("}", 1)[-1] == "tr"]
        if not source_rows or not target_rows:
            return

        header_source = source_rows[0]
        data_source = source_rows[1] if len(source_rows) > 1 else source_rows[0]
        for index, target_row in enumerate(target_rows):
            source_row = header_source if index == 0 else data_source
            self._replace_child_by_local_name(
                target_row,
                "trPr",
                self._child_by_local_name(source_row, "trPr"),
            )
            source_cells = [
                child for child in source_row if child.tag.rsplit("}", 1)[-1] == "tc"
            ]
            target_cells = [
                child for child in target_row if child.tag.rsplit("}", 1)[-1] == "tc"
            ]
            for cell_index, target_cell in enumerate(target_cells):
                if not source_cells:
                    break
                source_cell = source_cells[min(cell_index, len(source_cells) - 1)]
                self._replace_child_by_local_name(
                    target_cell,
                    "tcPr",
                    self._child_by_local_name(source_cell, "tcPr"),
                )

    @staticmethod
    def _child_by_local_name(element: Any, local_name: str) -> Any | None:
        for child in element:
            if child.tag.rsplit("}", 1)[-1] == local_name:
                return child
        return None

    @staticmethod
    def _replace_child_by_local_name(element: Any, local_name: str, replacement: Any | None) -> None:
        if replacement is None:
            return
        for index, child in enumerate(list(element)):
            if child.tag.rsplit("}", 1)[-1] == local_name:
                element.remove(child)
                element.insert(index, deepcopy(replacement))
                return
        element.insert(0, deepcopy(replacement))

    def _add_chart_images_to_document(
        self,
        doc: Any,
        chart_images: Dict[str, bytes],
    ):
        """向文档添加图表图片.

        Args:
            doc: Word Document object.
            chart_images: Dict of chart ID to image bytes.
        """
        from io import BytesIO

        from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
        from docx.shared import Inches

        for chart_id, image_bytes in chart_images.items():
            # Look for placeholder - support both chart_id and chart_{chart_id} formats
            placeholder_found = False
            patterns = [
                f"{{{{{chart_id}}}}}",
                f"{{{chart_id}}}",
                chart_id,
                f"{{{{chart_{chart_id}}}}}",
                f"{{chart_{chart_id}}}",
                f"chart_{chart_id}",
            ]

            for paragraph in doc.paragraphs:
                for pattern in patterns:
                    if pattern in paragraph.text:
                        # Replace placeholder with image
                        paragraph.clear()
                        image_stream = BytesIO(image_bytes)
                        run = paragraph.add_run()
                        run.add_picture(image_stream, width=Inches(4.5))
                        # 居中显示
                        paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                        placeholder_found = True
                        break
                if placeholder_found:
                    break

            # 也在表格单元格中查找占位符
            if not placeholder_found:
                for table in doc.tables:
                    for row in table.rows:
                        for cell in row.cells:
                            for paragraph in cell.paragraphs:
                                for pattern in patterns:
                                    if pattern in paragraph.text:
                                        paragraph.clear()
                                        image_stream = BytesIO(image_bytes)
                                        run = paragraph.add_run()
                                        run.add_picture(image_stream, width=Inches(4.5))
                                        paragraph.alignment = WD_PARAGRAPH_ALIGNMENT.CENTER
                                        placeholder_found = True
                                        break
                            if placeholder_found:
                                break
                        if placeholder_found:
                            break

            if not placeholder_found:
                # Add at end
                doc.add_paragraph(f"Chart: {chart_id}")
                image_stream = BytesIO(image_bytes)
                doc.add_picture(image_stream, width=Inches(4.5))
                doc.add_paragraph()
